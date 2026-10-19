import json
from datetime import timedelta
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.management import call_command
from django.core.exceptions import ValidationError
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.db.models import Count, Max, Q, Sum
from django.test import TestCase, TransactionTestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework import status
from rest_framework.authtoken.models import Token
from rest_framework.test import APIClient

from .seeders import ensure_initial_reference_data
from .seed_data import INITIAL_TOURIST_RECORDS, REFERENCE_TABLES
from .models import (
    ACTION_CREATE,
    ACTION_UPDATE,
    BOOKING_STATUS_ARRIVED,
    MODULE_TOURISM,
    ActiveOption,
    ActivityLog,
    Barangay,
    BoatCapacityFare,
    BoatType,
    Country,
    EmergencyContact,
    FeedbackEntry,
    HouseholdSanitationRecord,
    Itinerary,
    Notification,
    NOTIFICATION_AUDIENCE_PUBLIC,
    NOTIFICATION_MODULE_GENERAL,
    NOTIFICATION_MODULE_SANITATION,
    NOTIFICATION_MODULE_TOURISM,
    NOTIFICATION_TYPE_PUBLIC_ADVISORY,
    ROLE_ADMIN,
    ROLE_ESTABLISHMENT,
    ROLE_SANITATION,
    ROLE_TOURIST,
    ROLE_TOURISM,
    Province,
    Region,
    Resort,
    SanitaryBusinessType,
    SanitaryComplaint,
    SanitaryEstablishment,
    SanitaryInspection,
    TouristRecord,
    TravelMode,
    UserProfile,
    VisitPurpose,
)


class AuthApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_login_returns_token_and_role(self):
        user = User.objects.create_user(
            username="tourism_admin",
            password="Tourism@123",
        )
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)

        response = self.client.post(
            "/api/auth/login/",
            {"username": "tourism_admin", "password": "Tourism@123"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("token", response.json())
        self.assertEqual(response.json()["user"]["profile"]["role"], ROLE_TOURISM)

    def test_sanitation_user_cannot_access_tourism_endpoint(self):
        user = User.objects.create_user(
            username="sanitation_admin",
            password="Sanitation@123",
        )
        UserProfile.objects.create(user=user, role=ROLE_SANITATION)
        token = Token.objects.create(user=user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

        response = self.client.get("/api/booking-management/")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_public_tourist_registration_assigns_tourist_role(self):
        payload = {
            "full_name": "Maria Tourist",
            "email": "maria.tourist@example.com",
            "password": "SecurePassword@123",
            "contact_number": "09171234567",
        }
        response = self.client.post("/api/auth/register/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["user"]["profile"]["role"], ROLE_TOURIST)
        self.assertEqual(response.json()["user"]["profile"]["role_label"], "Tourist")

        created_user = User.objects.get(username="maria.tourist@example.com")
        self.assertEqual(created_user.profile.role, ROLE_TOURIST)

    def test_tourist_cannot_access_staff_tourism_endpoints(self):
        tourist_user = User.objects.create_user(
            username="registered_tourist",
            password="Password@123",
        )
        UserProfile.objects.create(user=tourist_user, role=ROLE_TOURIST)
        token = Token.objects.create(user=tourist_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

        # 1. Staff Booking Management Web API
        resp_booking = self.client.get("/api/booking-management/")
        self.assertEqual(resp_booking.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(resp_booking.json()["detail"], "You do not have access to this module.")

        # 2. Staff-only Mobile Tourism Record Lookup
        resp_lookup = self.client.get("/api/mobile/tourism/records/lookup/?query=SURV-2026-00001")
        self.assertEqual(resp_lookup.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(resp_lookup.json()["detail"], "Only tourism staff and administrators can look up visitor records.")

    def test_tourism_staff_retains_access_to_staff_tourism_endpoints(self):
        staff_user = User.objects.create_user(
            username="tourism_staff_member",
            password="Password@123",
        )
        UserProfile.objects.create(user=staff_user, role=ROLE_TOURISM)
        token = Token.objects.create(user=staff_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

        # 1. Staff Booking Management Web API
        resp_booking = self.client.get("/api/booking-management/")
        self.assertEqual(resp_booking.status_code, status.HTTP_200_OK)

        # 2. Staff-only Mobile Tourism Record Lookup (access granted through permission boundary, returns 404 for missing query)
        resp_lookup = self.client.get("/api/mobile/tourism/records/lookup/?query=NONEXISTENT_ID")
        self.assertNotEqual(resp_lookup.status_code, status.HTTP_403_FORBIDDEN)


class RemovedEstablishmentRegistrationTests(TestCase):
    """The establishment username/password registration is gone; owners use
    the Establishment Portal with the code on their Owner's Slip."""

    def test_both_registration_routes_are_404_and_create_nothing(self):
        users_before = User.objects.count()
        for url in (
            "/api/auth/register-establishment/",
            "/api/mobile/sanitation/register-establishment/",
        ):
            response = APIClient().post(
                url,
                {"username": "would_be_owner", "password": "Owner@12345", "permit_number": "SP-1"},
                format="json",
            )
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND, url)
        self.assertEqual(User.objects.count(), users_before)

    def test_the_shared_login_still_signs_in_an_existing_establishment_account(self):
        # Existing establishment users, their rows and tokens are left alone.
        user = User.objects.create_user(username="old_owner", password="Owner@12345")
        UserProfile.objects.create(user=user, role=ROLE_ESTABLISHMENT)

        response = APIClient().post(
            "/api/auth/login/", {"username": "old_owner", "password": "Owner@12345"}, format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.assertEqual(response.json()["user"]["profile"]["role"], ROLE_ESTABLISHMENT)


def ensure_test_barangays():
    """The 40 official Mauban barangays, as seeded into api_barangay."""
    from .seed_data import MAUBAN_BARANGAYS

    for index, name in enumerate(MAUBAN_BARANGAYS, start=1):
        Barangay.objects.get_or_create(
            name=name, defaults={"display_order": index, "is_active": True}
        )


def ensure_test_reference_tables():
    if not Country.objects.exists():
        Country.objects.bulk_create([Country(**c) for c in REFERENCE_TABLES["countries"]], ignore_conflicts=True)
    if not Region.objects.exists():
        Region.objects.bulk_create([Region(**r) for r in REFERENCE_TABLES["regions"]], ignore_conflicts=True)
    if not Province.objects.exists():
        reg_calabarzon = Region.objects.filter(code="04").first() or Region.objects.first()
        reg_ncr = Region.objects.filter(code="13").first() or reg_calabarzon
        provinces = []
        for p in REFERENCE_TABLES["provinces"]:
            reg = reg_ncr if p["name"] == "Metro Manila" else reg_calabarzon
            provinces.append(Province(id=p["id"], name=p["name"], region=reg))
        Province.objects.bulk_create(provinces, ignore_conflicts=True)
    if not Itinerary.objects.exists():
        Itinerary.objects.bulk_create([Itinerary(**it) for it in REFERENCE_TABLES["itineraries"]], ignore_conflicts=True)
    if not TravelMode.objects.exists():
        TravelMode.objects.bulk_create([TravelMode(**tm) for tm in REFERENCE_TABLES["travel_modes"]], ignore_conflicts=True)
    if not BoatType.objects.exists():
        BoatType.objects.bulk_create([BoatType(**bt) for bt in REFERENCE_TABLES["boat_types"]], ignore_conflicts=True)
    if not VisitPurpose.objects.exists():
        VisitPurpose.objects.bulk_create([VisitPurpose(**vp) for vp in REFERENCE_TABLES["visit_purposes"]], ignore_conflicts=True)
    if not Resort.objects.exists():
        Resort.objects.bulk_create([Resort(**res) for res in REFERENCE_TABLES["resorts"]], ignore_conflicts=True)


class SanitaryEstablishmentRecordsApiTests(TestCase):
    """The server contract the Establishment Records Register/Edit form relies on."""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="records_sanitation", password="Password@123"
        )
        UserProfile.objects.create(user=self.user, role=ROLE_SANITATION)
        token = Token.objects.create(user=self.user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

        self.market_stall = SanitaryBusinessType.objects.create(
            name="Public Market Stall", inspection_frequency="monthly"
        )
        self.karaoke = SanitaryBusinessType.objects.create(
            name="Karaoke / Video Bar / CSW", inspection_frequency="monthly"
        )

    def create_permitted_establishment(self):
        return SanitaryEstablishment.objects.create(
            business_name="Existing Bakery",
            owner_name="Juan Dela Cruz",
            business_type=self.market_stall,
            permit_size="large",
            barangay="Daungan",
            address="12 Rizal St",
            contact_number="09171234567",
            has_permit=True,
            permit_number="LG-2026-007",
            permit_issued_date="2026-02-01",
            permit_expiry_date="2026-12-31",
            compliance_status="for_completion",
            permit_status="conditional",
            latitude=14.191234,
            longitude=121.735678,
            remarks="Existing remarks",
        )

    def test_registration_payload_creates_establishment_with_no_permit_on_record(self):
        # Exactly what the Register New Establishment form sends.
        response = self.client.post(
            "/api/sanitation/establishments/",
            {
                "business_name": "New Sari-Sari Store",
                "owner_name": "Ana Reyes",
                "business_type": self.market_stall.id,
                "barangay": "Daungan",
                "address": "1 Test St",
                "contact_number": "09170000000",
                "latitude": 14.188,
                "longitude": 121.732,
                "has_permit": False,
                "permit_number": "",
                "permit_issued_date": None,
                "permit_expiry_date": None,
                "compliance_status": "no_permit",
                "permit_status": "no_permit",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        created = SanitaryEstablishment.objects.get(pk=response.data["id"])
        self.assertFalse(created.has_permit)
        self.assertEqual(created.permit_number, "")
        self.assertIsNone(created.permit_issued_date)
        self.assertIsNone(created.permit_expiry_date)
        self.assertEqual(created.compliance_status, "no_permit")
        self.assertEqual(created.permit_status, "no_permit")
        self.assertEqual(created.business_type_id, self.market_stall.id)
        # Not asked for at registration: the model defaults apply.
        self.assertEqual(created.permit_size, "sp")
        self.assertEqual(created.remarks, "")
        self.assertEqual(created.latitude, 14.188)
        self.assertEqual(created.longitude, 121.732)
        # The internal record id is not a permit number.
        self.assertNotEqual(str(created.pk), created.permit_number)

    def test_patch_of_profile_fields_preserves_permit_location_and_coverage(self):
        establishment = self.create_permitted_establishment()

        response = self.client.patch(
            f"/api/sanitation/establishments/{establishment.id}/",
            {"contact_number": "09990000000", "address": "99 New Address"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        establishment.refresh_from_db()
        self.assertEqual(establishment.contact_number, "09990000000")
        self.assertEqual(establishment.address, "99 New Address")
        self.assertTrue(establishment.has_permit)
        self.assertEqual(establishment.permit_number, "LG-2026-007")
        self.assertEqual(str(establishment.permit_issued_date), "2026-02-01")
        self.assertEqual(str(establishment.permit_expiry_date), "2026-12-31")
        self.assertEqual(establishment.compliance_status, "for_completion")
        self.assertEqual(establishment.permit_status, "conditional")
        self.assertEqual(establishment.permit_size, "large")
        self.assertEqual(establishment.latitude, 14.191234)
        self.assertEqual(establishment.longitude, 121.735678)
        self.assertEqual(establishment.remarks, "Existing remarks")
        self.assertEqual(establishment.business_type_id, self.market_stall.id)

    def test_patch_of_business_type_keeps_the_real_type_id_and_permit_data(self):
        establishment = self.create_permitted_establishment()

        response = self.client.patch(
            f"/api/sanitation/establishments/{establishment.id}/",
            {"business_type": self.karaoke.id},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["business_type"], self.karaoke.id)
        self.assertEqual(response.data["business_type_name"], "Karaoke / Video Bar / CSW")
        establishment.refresh_from_db()
        self.assertEqual(establishment.business_type_id, self.karaoke.id)
        self.assertEqual(establishment.permit_number, "LG-2026-007")
        self.assertEqual(establishment.permit_size, "large")
        self.assertEqual(establishment.compliance_status, "for_completion")

    def test_patch_with_null_permit_dates_left_untouched_keeps_them_null(self):
        # Mirrors a production record holding a permit but no recorded dates.
        establishment = self.create_permitted_establishment()
        SanitaryEstablishment.objects.filter(pk=establishment.pk).update(
            permit_issued_date=None, permit_expiry_date=None
        )

        response = self.client.patch(
            f"/api/sanitation/establishments/{establishment.id}/",
            {"owner_name": "Maria Dela Cruz"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        establishment.refresh_from_db()
        self.assertEqual(establishment.owner_name, "Maria Dela Cruz")
        self.assertIsNone(establishment.permit_issued_date)
        self.assertIsNone(establishment.permit_expiry_date)


class BookingManagementApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.user = User.objects.create_user(
            username="tourism_admin",
            password="Tourism@123",
        )
        UserProfile.objects.create(user=cls.user, role=ROLE_TOURISM)
        cls.token = Token.objects.create(user=cls.user)

        ensure_test_reference_tables()
        TouristRecord.objects.bulk_create(
            [TouristRecord(**record) for record in INITIAL_TOURIST_RECORDS],
            ignore_conflicts=True,
        )

    def setUp(self):
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        from unittest.mock import patch
        import datetime
        self.localdate_patcher = patch("django.utils.timezone.localdate", return_value=datetime.date(2026, 4, 1))
        self.mock_localdate = self.localdate_patcher.start()

    def tearDown(self):
        self.localdate_patcher.stop()

    def test_booking_management_returns_summary_and_rows(self):
        response = self.client.get("/api/booking-management/", {"page_size": 20})

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        data = response.json()
        self.assertEqual(data["summary"]["verifiedEntries"], 12)
        self.assertEqual(data["summary"]["pending"], 4)
        self.assertEqual(data["summary"]["arrived"], 7)
        self.assertEqual(data["summary"]["noShow"], 1)
        self.assertEqual(len(data["rows"]), 12)

        first_row = data["rows"][0]
        self.assertIn("country_name", first_row)
        self.assertIn("resort_name", first_row)
        self.assertIn("status_label", first_row)

    def test_booking_management_filters_by_search_and_status(self):
        response = self.client.get(
            "/api/booking-management/",
            {"search": "Ethan", "status": "pending"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        rows = response.json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["survey_id"], "SURV-2026-002")
        self.assertEqual(rows[0]["status"], "pending")

    def test_booking_status_can_be_updated(self):
        self.client.get("/api/booking-management/")

        response = self.client.patch(
            "/api/tourist-records/SURV-2026-002/",
            {"status": BOOKING_STATUS_ARRIVED},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["status"], BOOKING_STATUS_ARRIVED)
        self.assertEqual(
            TouristRecord.objects.get(survey_id="SURV-2026-002").status,
            BOOKING_STATUS_ARRIVED,
        )
        self.assertTrue(
            ActivityLog.objects.filter(
                action=ACTION_UPDATE,
                module=MODULE_TOURISM,
                record_id="SURV-2026-002",
                user=self.user,
            ).exists()
        )


    def _record_payload(self, **overrides):
        region = Region.objects.first()
        province = Province.objects.first() or Province.objects.create(
            id=9001, name="Quezon", region=region
        )
        payload = {
            "first_name": "Maria",
            "last_name": "Reyes",
            "full_name": "Maria Reyes",
            "email": "maria@example.com",
            "consent_confirmed": True,
            "contact_number": "+639171234567",
            "country_id": Country.objects.first().id,
            "region_id": region.id,
            "province_id": province.id,
            "country_of_origin": "",
            "resort_id": Resort.objects.first().resort_id,
            "itinerary_id": Itinerary.objects.first().id,
            "travel_mode_id": TravelMode.objects.first().id,
            "boat_type_id": BoatType.objects.first().id,
            "boat_capacity_fare": "",
            "parking_space": "",
            "visit_purpose_id": VisitPurpose.objects.first().id,
            "arrival_date": "2026-04-01",
            "filipino_count": 3,
            "foreigner_count": 1,
            "total_visitors": 4,
            "total_male": 2,
            "total_female": 2,
            "special_group_count": 0,
            "age_0_7": 0,
            "age_8_59": 4,
            "age_60_above": 0,
            "status": "pending",
        }
        payload.update(overrides)
        return payload

    def test_maubanin_count_above_filipino_or_total_is_rejected(self):
        for value in (4, 5, 99):  # 4 = total head count, still > filipino_count (3)
            response = self.client.post(
                "/api/tourist-records/",
                self._record_payload(maubanin_count=value),
                format="json",
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, value)
            self.assertIn("maubanin_count", response.json())
            self.assertIn(
                "cannot be greater than the Filipino count",
                str(response.json()["maubanin_count"]),
            )
        self.assertFalse(TouristRecord.objects.filter(email="maria@example.com").exists())

    def test_valid_maubanin_count_is_saved_without_changing_totals(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._record_payload(maubanin_count=2),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = TouristRecord.objects.get(email="maria@example.com")
        self.assertEqual(record.maubanin_count, 2)
        self.assertEqual(record.filipino_count, 3)
        self.assertEqual(record.foreigner_count, 1)
        self.assertEqual(record.total_visitors, 4)

    def test_maubanin_count_equal_to_filipino_count_is_allowed(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._record_payload(maubanin_count=3),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)

    def test_omitted_maubanin_count_saves_zero(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._record_payload(),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(
            TouristRecord.objects.get(email="maria@example.com").maubanin_count, 0
        )

    def test_updating_maubanin_count_above_filipino_count_is_rejected(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._record_payload(maubanin_count=1),
            format="json",
        )
        survey_id = response.json()["survey_id"]

        bad = self.client.patch(
            f"/api/tourist-records/{survey_id}/", {"maubanin_count": 4}, format="json"
        )
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

        good = self.client.patch(
            f"/api/tourist-records/{survey_id}/", {"maubanin_count": 3}, format="json"
        )
        self.assertEqual(good.status_code, status.HTTP_200_OK, good.content)
        self.assertEqual(good.json()["maubanin_count"], 3)

    def test_maubanin_count_is_not_added_on_top_of_filipino_in_dashboard(self):
        TouristRecord.objects.all().delete()
        response = self.client.post(
            "/api/tourist-records/",
            self._record_payload(maubanin_count=3, status="arrived"),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)

        dashboard = self.client.get("/api/dashboard/", {"year": "all"})
        self.assertEqual(dashboard.status_code, status.HTTP_200_OK)
        classification = dashboard.json()["classification"]
        # Maubanin is a subset of Filipino: Domestic must equal filipino_count (3), not 3 + 3.
        self.assertEqual(classification["filipino"], 3)
        self.assertEqual(classification["foreign"], 1)


class MobilePublicApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        ensure_test_barangays()
        ensure_test_reference_tables()
        cls.btype, _ = SanitaryBusinessType.objects.get_or_create(
            name="Food Establishment",
            defaults={"inspection_frequency": "quarterly"},
        )
        cls.establishment, _ = SanitaryEstablishment.objects.get_or_create(
            permit_number="LG-2026-001",
            defaults={
                "business_name": "Test Mobile Establishment",
                "owner_name": "Test Owner",
                "business_type": cls.btype,
                "barangay": "Poblacion",
                "address": "123 Main St",
                "has_permit": True,
                "compliance_status": "good_standing",
                "permit_status": "active",
            },
        )
        cls.sanitation_user = User.objects.create_user(
            username="test_mobile_sanitation_inspector",
            password="Password@123",
            email="inspector@test.local",
        )
        UserProfile.objects.create(user=cls.sanitation_user, role=ROLE_SANITATION)
        cls.sanitation_token = Token.objects.create(user=cls.sanitation_user)

    def setUp(self):
        self.client = APIClient()
        from unittest.mock import patch
        import datetime
        self.localdate_patcher = patch("django.utils.timezone.localdate", return_value=datetime.date(2026, 4, 1))
        self.mock_localdate = self.localdate_patcher.start()

    def tearDown(self):
        self.localdate_patcher.stop()

    def test_mobile_bootstrap_is_public(self):
        response = self.client.get("/api/mobile/tourism/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertIn("destinations", data)
        self.assertIn("referenceTables", data)
        self.assertIn("barangays", data)

    def test_mobile_bootstrap_shows_ranked_mauban_destinations_only(self):
        response = self.client.get("/api/mobile/tourism/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        destinations = response.json()["destinations"]
        names = [destination["resort_name"] for destination in destinations]

        self.assertLessEqual(len(destinations), 10)
        self.assertIn("Dona Choleng Camping Resort", names)
        self.assertLess(len(destinations), Resort.objects.count())
        self.assertTrue(all(not name.endswith("(") for name in names))

    def test_mobile_bootstrap_prioritizes_high_visitor_destinations(self):
        ensure_initial_reference_data()
        TouristRecord.objects.create(
            survey_id="MOB-TOP-0001",
            full_name="High Visitor Group",
            contact_number="09170000000",
            country=Country.objects.first(),
            region=Region.objects.first(),
            province=Province.objects.first(),
            arrival_date="2026-04-02",
            resort=Resort.objects.get(resort_name="Orlan Beach Resort"),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=5000,
            filipino_count=5000,
            total_male=2500,
            total_female=2500,
            age_8_59=5000,
            status="arrived",
        )

        response = self.client.get("/api/mobile/tourism/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        first_destination = response.json()["destinations"][0]
        self.assertEqual(first_destination["resort_name"], "Orlan Beach Resort")
        self.assertEqual(first_destination["monthly_arrivals"], 1)

    def test_mobile_tourist_registration_creates_booking_record(self):
        response = self.client.post(
            "/api/mobile/tourism/register-visit/",
            {
                "full_name": "Mobile Tourist",
                "contact_number": "09171234567",
                "arrival_date": "2026-06-01",
                "total_visitors": 2,
                "filipino_count": 2,
                "total_male": 1,
                "total_female": 1,
                "age_8_59": 2,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(
            TouristRecord.objects.filter(full_name="Mobile Tourist").exists()
        )

    def test_mobile_tourist_registration_accepts_full_web_record_fields(self):
        ensure_initial_reference_data()
        self.client.get("/api/mobile/tourism/bootstrap/")
        resort = Resort.objects.first()
        country = Country.objects.first()
        region = Region.objects.first()
        province = Province.objects.first()
        if province is None:
            province = Province.objects.create(id=1, name="Quezon", region=region)
        itinerary = Itinerary.objects.first()
        travel_mode = TravelMode.objects.first()
        boat_type = BoatType.objects.first()
        purpose = VisitPurpose.objects.first()

        response = self.client.post(
            "/api/mobile/tourism/register-visit/",
            {
                "full_name": "Complete Mobile Tourist",
                "email": "complete@example.com",
                "consent_confirmed": True,
                "contact_number": "09171234567",
                "country_id": country.id,
                "region_id": region.id,
                "province_id": province.id,
                "country_of_origin": "Philippines",
                "arrival_date": "2026-06-01",
                "resort_id": resort.resort_id,
                "itinerary_id": itinerary.id,
                "travel_mode_id": travel_mode.id,
                "boat_type_id": boat_type.id,
                "boat_capacity_fare": "20 pax / PHP 150",
                "parking_space": "Municipal parking",
                "visit_purpose_id": purpose.id,
                "total_visitors": 3,
                "filipino_count": 2,
                "maubanin_count": 1,
                "foreigner_count": 0,
                "total_male": 1,
                "total_female": 2,
                "special_group_count": 1,
                "age_0_7": 1,
                "age_8_59": 2,
                "age_60_above": 0,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        record = TouristRecord.objects.get(full_name="Complete Mobile Tourist")
        self.assertEqual(record.country_of_origin, "Philippines")
        self.assertEqual(record.boat_capacity_fare, "20 pax / PHP 150")
        self.assertEqual(record.parking_space, "Municipal parking")
        self.assertEqual(record.special_group_count, 1)
        self.assertEqual(record.age_0_7, 1)

    def test_mobile_feedback_creates_feedback_entry(self):
        self.client.get("/api/mobile/tourism/bootstrap/")
        resort = Resort.objects.first()

        response = self.client.post(
            "/api/mobile/tourism/feedback/",
            {
                "destination_id": resort.resort_id,
                "reviewer": "Mobile Reviewer",
                "rating": 5,
                "message": "Clean and organized.",
                "cleanliness_rating": 5,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(
            FeedbackEntry.objects.filter(reviewer="Mobile Reviewer").exists()
        )

    def test_mobile_sanitation_report_creates_complaint(self):
        response = self.client.post(
            "/api/mobile/sanitation/reports/",
            {
                "complainant_name": "Resident Reporter",
                "contact_number": "09170000000",
                "category": "Improper Garbage Disposal",
                "barangay": "Daungan",
                "location_address": "Walkway beside the plaza",
                "description": "Garbage pile near the walkway.",
                "latitude": 14.186,
                "longitude": 121.73,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        complaint = SanitaryComplaint.objects.get(
            complainant_name="Resident Reporter"
        )
        self.assertEqual(complaint.photo_documentation, "")
        self.assertEqual(complaint.status, "pending")

    def test_mobile_sanitation_report_history_needs_contact_and_complaint_id(self):
        submitted = self.client.post(
            "/api/mobile/sanitation/reports/",
            {
                "complainant_name": "Resident Reporter",
                "contact_number": "09170000000",
                "category": "Contaminated Water Source",
                "barangay": "Daungan",
                "location_address": "Deep well, Purok 2",
                "description": "Water source needs inspection.",
            },
            format="json",
        )

        response = self.client.get(
            "/api/mobile/sanitation/reports/history/",
            {"contact": "09170000000", "reference": submitted.json()["complaint_id"]},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = response.json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "pending")
        self.assertIn("status_label", rows[0])

    def test_mobile_sanitation_permit_verify_by_permit_number(self):
        # Permit numbers are no longer listed publicly; use the known record.
        permit_number = self.establishment.permit_number

        response = self.client.get(
            "/api/mobile/sanitation/permits/verify/",
            {"code": permit_number},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertTrue(data["verified"])
        self.assertEqual(data["permit"]["permit_number"], permit_number)

    def test_mobile_sanitation_bootstrap_is_public(self):
        response = self.client.get("/api/mobile/sanitation/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertIn("establishments", data)
        self.assertIn("businessTypes", data)
        self.assertIn("householdRecords", data)
        self.assertIn("barangays", data)

    def test_mobile_sanitation_inspection_creates_establishment_inspection(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.sanitation_token.key}")
        self.client.get("/api/mobile/sanitation/bootstrap/")
        establishment = SanitaryEstablishment.objects.first()

        response = self.client.post(
            "/api/mobile/sanitation/inspections/",
            {
                "establishment": establishment.id,
                "inspector_name": "Mobile Inspector",
                "inspection_date": "2026-06-03",
                "next_due_date": "2026-07-03",
                "findings": "All inspected.",
                "remarks": "For monitoring.",
                "status_after_inspection": "good_standing",
                "checklist_items": [
                    {
                        "requirement_name": "Proper waste disposal system",
                        "is_complied": True,
                        "notes": "",
                    }
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(
            SanitaryInspection.objects.filter(
                inspector_name="Mobile Inspector",
                establishment=establishment,
            ).exists()
        )

    def test_mobile_household_survey_creates_household_record(self):
        sanitation_user = User.objects.create_user(
            username="test_mobile_survey_sanitation",
            password="Password@123",
            email="mobile_survey_sanitation@test.local",
        )
        UserProfile.objects.create(user=sanitation_user, role=ROLE_SANITATION)
        token = Token.objects.create(user=sanitation_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

        response = self.client.post(
            "/api/mobile/sanitation/household-surveys/",
            {
                "household_head": "Mobile Household",
                "barangay": "Poblacion",
                "address": "Sample Street",
                "male_count": 2,
                "female_count": 3,
                "toilet_type": "none",
                "water_level": "level_1",
                "water_source": "Deep well",
                "waste_disposal": "dumped",
                "remarks": "Needs follow-up inspection.",
                "latitude": 14.186,
                "longitude": 121.73,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        record = HouseholdSanitationRecord.objects.get(
            household_head="Mobile Household"
        )
        self.assertTrue(record.household_code.startswith("HH-"))
        self.assertEqual(record.total_members, 5)
        self.assertEqual(record.status, "violation")


class NotificationViolationTriggerTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Create active admin user
        self.admin_user = User.objects.create_user(
            username="notif_test_admin",
            password="Password@123",
            email="admin@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)

        # Create active sanitation user
        self.sanitation_user = User.objects.create_user(
            username="notif_test_sanitation",
            password="Password@123",
            email="sanitation@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)

        # Authenticate client as sanitation user
        self.token = Token.objects.create(user=self.sanitation_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        # Create test business type and establishment in good_standing
        self.btype = SanitaryBusinessType.objects.create(
            name="Violation Test Bakery",
            inspection_frequency="monthly",
        )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Sweet Treats Bakery",
            owner_name="Baker Bob",
            business_type=self.btype,
            barangay="Poblacion",
            address="100 Rizal St",
            compliance_status="good_standing",
            permit_status="active",
        )

    def test_inspection_violation_triggers_notifications_and_prevents_duplicates(self):
        # 1. Submit a SanitaryInspection with status_after_inspection="violation" via the actual view
        response = self.client.post(
            "/api/sanitation/inspections/",
            {
                "establishment": self.establishment.id,
                "inspector_name": "Inspector Test",
                "inspection_date": "2026-09-11",
                "status_after_inspection": "violation",
                "findings": "Pest infestation and improper food storage",
                "remarks": "Immediate closure order recommended",
                "checklist_items": [
                    {"requirement_name": "Sanitary Permit", "is_complied": True},
                    {"requirement_name": "Pest Control", "is_complied": False},
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        inspection_id = response.json()["id"]

        # a) Assert exactly one Notification exists for the admin user and one for the sanitation user,
        #    both with notification_type="violation_alert"
        admin_notifs = Notification.objects.filter(recipient_user=self.admin_user)
        self.assertEqual(admin_notifs.count(), 1)
        admin_notif = admin_notifs.first()
        self.assertEqual(admin_notif.notification_type, "violation_alert")
        self.assertEqual(admin_notif.severity, "critical")
        self.assertEqual(admin_notif.module, "sanitation")
        self.assertIn("Sweet Treats Bakery", admin_notif.title)

        sanitation_notifs = Notification.objects.filter(recipient_user=self.sanitation_user)
        self.assertEqual(sanitation_notifs.count(), 1)
        sanitation_notif = sanitation_notifs.first()
        self.assertEqual(sanitation_notif.notification_type, "violation_alert")
        self.assertEqual(sanitation_notif.severity, "critical")
        self.assertEqual(sanitation_notif.module, "sanitation")
        self.assertIn("Sweet Treats Bakery", sanitation_notif.title)

        # b) Assert related_object_id matches the establishment's id
        self.assertEqual(admin_notif.related_object_id, str(self.establishment.id))
        self.assertEqual(sanitation_notif.related_object_id, str(self.establishment.id))
        self.assertEqual(admin_notif.related_model, "SanitaryEstablishment")
        self.assertEqual(sanitation_notif.related_model, "SanitaryEstablishment")
        self.assertEqual(admin_notif.action_url, f"/sanitation/establishments/{self.establishment.id}")

        # Refresh establishment to verify its compliance_status is now "violation"
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "violation")

        # c) Calling sync_establishment_after_inspection again with the establishment
        #    already in "violation" status does NOT create additional notifications (no duplicates)
        inspection = SanitaryInspection.objects.get(id=inspection_id)
        from api.services.sanitation import sync_establishment_after_inspection

        sync_establishment_after_inspection(inspection)

        # Assert notification count has NOT increased
        self.assertEqual(Notification.objects.filter(recipient_user=self.admin_user).count(), 1)
        self.assertEqual(Notification.objects.filter(recipient_user=self.sanitation_user).count(), 1)
        self.assertEqual(
            Notification.objects.filter(related_object_id=str(self.establishment.id)).count(),
            2,
        )


class NotificationStaffApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # User A (Admin)
        self.user_a = User.objects.create_user(
            username="notif_api_user_a",
            password="Password@123",
            email="user_a@test.local",
        )
        UserProfile.objects.create(user=self.user_a, role=ROLE_ADMIN)
        self.token_a = Token.objects.create(user=self.user_a)

        # User B (Sanitation)
        self.user_b = User.objects.create_user(
            username="notif_api_user_b",
            password="Password@123",
            email="user_b@test.local",
        )
        UserProfile.objects.create(user=self.user_b, role=ROLE_SANITATION)
        self.token_b = Token.objects.create(user=self.user_b)

        # Create notifications for User A: mix of read/unread and modules
        self.notif_a1 = Notification.objects.create(
            title="User A Sanitation Unread",
            message="Message A1",
            notification_type="violation_alert",
            severity="critical",
            module="sanitation",
            audience_type="user",
            recipient_user=self.user_a,
            is_read=False,
            action_url="/sanitation/establishments/1",
        )
        self.notif_a2 = Notification.objects.create(
            title="User A Tourism Unread",
            message="Message A2",
            notification_type="permit_due",
            severity="warning",
            module="tourism",
            audience_type="user",
            recipient_user=self.user_a,
            is_read=False,
            action_url="/tourism/resorts/1",
        )
        self.notif_a3 = Notification.objects.create(
            title="User A General Read",
            message="Message A3",
            notification_type="system",
            severity="info",
            module="general",
            audience_type="user",
            recipient_user=self.user_a,
            is_read=True,
            read_at=timezone.now(),
            action_url="/general/announcements/1",
        )

        # Create notifications for User B: mix of read/unread and modules
        self.notif_b1 = Notification.objects.create(
            title="User B Sanitation Unread",
            message="Message B1",
            notification_type="violation_alert",
            severity="critical",
            module="sanitation",
            audience_type="user",
            recipient_user=self.user_b,
            is_read=False,
            action_url="/sanitation/establishments/2",
        )
        self.notif_b2 = Notification.objects.create(
            title="User B Tourism Read",
            message="Message B2",
            notification_type="inspection_due",
            severity="warning",
            module="tourism",
            audience_type="user",
            recipient_user=self.user_b,
            is_read=True,
            read_at=timezone.now(),
            action_url="/tourism/resorts/2",
        )

    def test_notification_endpoints_lifecycle(self):
        # Authenticate as User A
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token_a.key}")

        # b) Hits GET /api/notifications/ as user A:
        #    Asserts: only user A's notifications are returned, unread_count matches (2),
        #    and user B's notifications never appear.
        response = self.client.get("/api/notifications/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["unread_count"], 2)
        results = data["results"]
        self.assertEqual(len(results), 3)

        result_ids = {r["id"] for r in results}
        self.assertEqual(result_ids, {self.notif_a1.id, self.notif_a2.id, self.notif_a3.id})
        self.assertNotIn(self.notif_b1.id, result_ids)
        self.assertNotIn(self.notif_b2.id, result_ids)

        # Test filter: ?unread_only=true
        res_unread = self.client.get("/api/notifications/?unread_only=true")
        self.assertEqual(res_unread.status_code, status.HTTP_200_OK)
        data_unread = res_unread.json()
        self.assertEqual(data_unread["unread_count"], 2)  # unread_count unchanged by filter
        self.assertEqual(len(data_unread["results"]), 2)
        self.assertEqual({r["id"] for r in data_unread["results"]}, {self.notif_a1.id, self.notif_a2.id})

        # Test filter: ?module=sanitation
        res_module = self.client.get("/api/notifications/?module=sanitation")
        self.assertEqual(res_module.status_code, status.HTTP_200_OK)
        data_module = res_module.json()
        self.assertEqual(data_module["unread_count"], 2)
        self.assertEqual(len(data_module["results"]), 1)
        self.assertEqual(data_module["results"][0]["id"], self.notif_a1.id)

        # c) Hits PATCH /api/notifications/<id>/read/ on one of user A's notifications
        #    and asserts is_read/read_at are set, and unread_count drops by 1 on a follow-up GET.
        patch_res = self.client.patch(f"/api/notifications/{self.notif_a1.id}/read/")
        self.assertEqual(patch_res.status_code, status.HTTP_200_OK)
        self.notif_a1.refresh_from_db()
        self.assertTrue(self.notif_a1.is_read)
        self.assertIsNotNone(self.notif_a1.read_at)

        # Follow-up GET: unread_count should now be 1
        res_after_read = self.client.get("/api/notifications/")
        self.assertEqual(res_after_read.status_code, status.HTTP_200_OK)
        self.assertEqual(res_after_read.json()["unread_count"], 1)

        # d) Attempts PATCH /api/notifications/<id>/read/ on one of user B's
        #    notification IDs while authenticated as user A, and asserts it returns 404 (not 403)
        cross_res = self.client.patch(f"/api/notifications/{self.notif_b1.id}/read/")
        self.assertEqual(cross_res.status_code, status.HTTP_404_NOT_FOUND)
        self.notif_b1.refresh_from_db()
        self.assertFalse(self.notif_b1.is_read)
        self.assertIsNone(self.notif_b1.read_at)

        # e) Hits POST /api/notifications/mark-all-read/ as user A and
        #    asserts all of user A's notifications are now read and user B's are untouched.
        mark_all_res = self.client.post("/api/notifications/mark-all-read/")
        self.assertEqual(mark_all_res.status_code, status.HTTP_200_OK)
        self.assertEqual(mark_all_res.json()["updated_count"], 1)  # only notif_a2 was unread

        # All of User A's notifications are now read
        self.assertEqual(
            Notification.objects.filter(recipient_user=self.user_a, is_read=False).count(),
            0,
        )
        self.notif_a2.refresh_from_db()
        self.assertTrue(self.notif_a2.is_read)
        self.assertIsNotNone(self.notif_a2.read_at)

        # User B's notifications remain untouched (notif_b1 still unread)
        self.notif_b1.refresh_from_db()
        self.assertFalse(self.notif_b1.is_read)
        self.assertIsNone(self.notif_b1.read_at)
        self.assertEqual(
            Notification.objects.filter(recipient_user=self.user_b, is_read=False).count(),
            1,
        )


class NotificationDueDateScanningTests(TestCase):
    def setUp(self):
        # Create active admin user
        self.admin_user = User.objects.create_user(
            username="notif_scan_admin",
            password="Password@123",
            email="admin_scan@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)

        # Create active sanitation user
        self.sanitation_user = User.objects.create_user(
            username="notif_scan_sanitation",
            password="Password@123",
            email="sanitation_scan@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)

        # Business type
        self.btype = SanitaryBusinessType.objects.create(
            name="Due Date Bakery",
            inspection_frequency="monthly",
        )

        # a) Creates an establishment with permit_expiry_date = today + 10 days, has_permit=True
        self.today = timezone.localdate()
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Sunshine Bakery",
            owner_name="Sunny Day",
            business_type=self.btype,
            barangay="Poblacion",
            address="123 Sunny St",
            compliance_status="good_standing",
            permit_status="active",
            has_permit=True,
            permit_expiry_date=self.today + timedelta(days=10),
        )

    def test_due_date_scanning_lifecycle_and_idempotency(self):
        # b) Calls the command (call_command('evaluate_due_notifications')),
        #    asserts exactly one permit_due notification was created for each of the two users,
        #    with related_due_date matching.
        call_command("evaluate_due_notifications")

        admin_notifs = Notification.objects.filter(
            recipient_user=self.admin_user,
            notification_type="permit_due",
        )
        self.assertEqual(admin_notifs.count(), 1)
        admin_notif = admin_notifs.first()
        self.assertEqual(admin_notif.related_due_date, self.today + timedelta(days=10))
        self.assertEqual(admin_notif.related_object_id, str(self.establishment.id))
        self.assertEqual(admin_notif.severity, "warning")
        self.assertEqual(admin_notif.module, "sanitation")

        sanit_notifs = Notification.objects.filter(
            recipient_user=self.sanitation_user,
            notification_type="permit_due",
        )
        self.assertEqual(sanit_notifs.count(), 1)
        sanit_notif = sanit_notifs.first()
        self.assertEqual(sanit_notif.related_due_date, self.today + timedelta(days=10))
        self.assertEqual(sanit_notif.related_object_id, str(self.establishment.id))

        # c) Calls the command AGAIN immediately and asserts no new notifications were created
        #    (still exactly 1 each) — proves idempotency.
        call_command("evaluate_due_notifications")
        self.assertEqual(
            Notification.objects.filter(
                recipient_user=self.admin_user,
                notification_type="permit_due",
            ).count(),
            1,
        )
        self.assertEqual(
            Notification.objects.filter(
                recipient_user=self.sanitation_user,
                notification_type="permit_due",
            ).count(),
            1,
        )

        # d) Changes the establishment's permit_expiry_date to a new date (simulating a renewal)
        #    and calls the command again — asserts a NEW notification is created (proves idempotency
        #    is keyed to the due date value, not just existence).
        new_expiry_date = self.today + timedelta(days=20)
        self.establishment.permit_expiry_date = new_expiry_date
        self.establishment.save()

        call_command("evaluate_due_notifications")

        admin_notifs_after = Notification.objects.filter(
            recipient_user=self.admin_user,
            notification_type="permit_due",
        )
        self.assertEqual(admin_notifs_after.count(), 2)
        admin_due_dates = set(admin_notifs_after.values_list("related_due_date", flat=True))
        self.assertEqual(admin_due_dates, {self.today + timedelta(days=10), new_expiry_date})

        sanit_notifs_after = Notification.objects.filter(
            recipient_user=self.sanitation_user,
            notification_type="permit_due",
        )
        self.assertEqual(sanit_notifs_after.count(), 2)

        # e) Creates a SanitaryInspection with next_due_date = today + 3 days on an establishment
        #    with compliance_status NOT "violation", calls the command, and asserts an inspection_due
        #    notification was created for each staff user.
        inspection_due_date = self.today + timedelta(days=3)
        SanitaryInspection.objects.create(
            establishment=self.establishment,
            next_due_date=inspection_due_date,
            inspector_name="Inspector Valid",
            inspection_date=self.today,
            status_after_inspection="good_standing",
        )

        call_command("evaluate_due_notifications")

        admin_insp_notifs = Notification.objects.filter(
            recipient_user=self.admin_user,
            notification_type="inspection_due",
            related_object_id=str(self.establishment.id),
        )
        self.assertEqual(admin_insp_notifs.count(), 1)
        self.assertEqual(admin_insp_notifs.first().related_due_date, inspection_due_date)

        sanit_insp_notifs = Notification.objects.filter(
            recipient_user=self.sanitation_user,
            notification_type="inspection_due",
            related_object_id=str(self.establishment.id),
        )
        self.assertEqual(sanit_insp_notifs.count(), 1)
        self.assertEqual(sanit_insp_notifs.first().related_due_date, inspection_due_date)

        # f) Creates a second inspection with next_due_date = today + 3 days but on an
        #    establishment with compliance_status="violation", and asserts NO inspection_due
        #    notification is created for that one (violation already covers it).
        est_violation = SanitaryEstablishment.objects.create(
            business_name="Violating Cafe",
            owner_name="Bad Actor",
            business_type=self.btype,
            barangay="Poblacion",
            address="456 Danger Rd",
            compliance_status="violation",
            permit_status="revoked",
        )
        SanitaryInspection.objects.create(
            establishment=est_violation,
            next_due_date=inspection_due_date,
            inspector_name="Inspector Strict",
            inspection_date=self.today,
            status_after_inspection="violation",
        )

        call_command("evaluate_due_notifications")

        self.assertEqual(
            Notification.objects.filter(
                related_object_id=str(est_violation.id),
                notification_type="inspection_due",
            ).count(),
            0,
        )

        # g) Overdue Permit: establishment with permit_expiry_date = today - 5 days, has_permit=True
        expired_date = self.today - timedelta(days=5)
        est_expired = SanitaryEstablishment.objects.create(
            business_name="Expired Permit Diner",
            owner_name="Late Larry",
            business_type=self.btype,
            barangay="Poblacion",
            address="789 Overdue Ave",
            compliance_status="good_standing",
            permit_status="active",
            has_permit=True,
            permit_expiry_date=expired_date,
        )

        call_command("evaluate_due_notifications")

        admin_expired_notifs = Notification.objects.filter(
            recipient_user=self.admin_user,
            notification_type="permit_due",
            related_object_id=str(est_expired.id),
        )
        self.assertEqual(admin_expired_notifs.count(), 1)
        admin_expired = admin_expired_notifs.first()
        self.assertEqual(admin_expired.severity, "critical")
        self.assertEqual(admin_expired.related_due_date, expired_date)
        self.assertIn("EXPIRED", admin_expired.message)

        sanit_expired_notifs = Notification.objects.filter(
            recipient_user=self.sanitation_user,
            notification_type="permit_due",
            related_object_id=str(est_expired.id),
        )
        self.assertEqual(sanit_expired_notifs.count(), 1)
        sanit_expired = sanit_expired_notifs.first()
        self.assertEqual(sanit_expired.severity, "critical")
        self.assertEqual(sanit_expired.related_due_date, expired_date)

        # h) Overdue Inspection: inspection with next_due_date = today - 2 days on a non-violation establishment
        overdue_insp_date = self.today - timedelta(days=2)
        est_overdue_insp = SanitaryEstablishment.objects.create(
            business_name="Overdue Inspection Cafe",
            owner_name="Forgetful Fred",
            business_type=self.btype,
            barangay="Poblacion",
            address="321 Tardy Rd",
            compliance_status="good_standing",
            permit_status="active",
        )
        SanitaryInspection.objects.create(
            establishment=est_overdue_insp,
            next_due_date=overdue_insp_date,
            inspector_name="Inspector Check",
            inspection_date=self.today - timedelta(days=30),
            status_after_inspection="good_standing",
        )

        call_command("evaluate_due_notifications")

        admin_overdue_insp = Notification.objects.filter(
            recipient_user=self.admin_user,
            notification_type="inspection_due",
            related_object_id=str(est_overdue_insp.id),
        )
        self.assertEqual(admin_overdue_insp.count(), 1)
        admin_oi = admin_overdue_insp.first()
        self.assertEqual(admin_oi.severity, "critical")
        self.assertEqual(admin_oi.related_due_date, overdue_insp_date)
        self.assertIn("OVERDUE", admin_oi.message)

        sanit_overdue_insp = Notification.objects.filter(
            recipient_user=self.sanitation_user,
            notification_type="inspection_due",
            related_object_id=str(est_overdue_insp.id),
        )
        self.assertEqual(sanit_overdue_insp.count(), 1)
        sanit_oi = sanit_overdue_insp.first()
        self.assertEqual(sanit_oi.severity, "critical")
        self.assertEqual(sanit_oi.related_due_date, overdue_insp_date)

        # i) Confirm the original in-window (future due date) cases still have severity="warning" (not critical)
        self.assertEqual(admin_notif.severity, "warning")
        self.assertEqual(sanit_notif.severity, "warning")
        self.assertEqual(admin_insp_notifs.first().severity, "warning")
        self.assertEqual(sanit_insp_notifs.first().severity, "warning")


class NotificationPublicAdvisoryApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        self.admin_user = User.objects.create_user(
            username="notif_pub_admin",
            password="Password@123",
            email="pub_admin@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)
        self.admin_token = Token.objects.create(user=self.admin_user)

        self.tourism_user = User.objects.create_user(
            username="notif_pub_tourism",
            password="Password@123",
            email="pub_tourism@test.local",
        )
        UserProfile.objects.create(user=self.tourism_user, role=ROLE_TOURISM)
        self.tourism_token = Token.objects.create(user=self.tourism_user)

        self.sanitation_user = User.objects.create_user(
            username="notif_pub_sanit",
            password="Password@123",
            email="pub_sanit@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)
        self.sanitation_token = Token.objects.create(user=self.sanitation_user)

        self.establishment_user = User.objects.create_user(
            username="notif_pub_est",
            password="Password@123",
            email="pub_est@test.local",
        )
        UserProfile.objects.create(user=self.establishment_user, role=ROLE_ESTABLISHMENT)
        self.establishment_token = Token.objects.create(user=self.establishment_user)

    def test_establishment_role_cannot_post_advisory(self):
        """a) Asserts a user with role 'establishment' gets 403 on POST."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.establishment_token.key}")
        response = self.client.post(
            "/api/notifications/public/",
            {
                "title": "Establishment Advisory",
                "message": "Should not be allowed",
                "module": "general",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            Notification.objects.filter(title="Establishment Advisory").count(),
            0,
        )

    def test_tourism_role_can_create_advisory(self):
        """b) Asserts a user with role 'tourism' CAN create an advisory."""
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        payload = {
            "title": "Storm Warning Advisory",
            "message": "Heavy rains expected. Tourism activities suspended.",
            "module": "tourism",
            "severity": "warning",
        }
        response = self.client.post(
            "/api/notifications/public/",
            payload,
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.data
        self.assertEqual(data["title"], payload["title"])
        self.assertEqual(data["message"], payload["message"])
        self.assertEqual(data["module"], "tourism")
        self.assertEqual(data["severity"], "warning")
        self.assertEqual(data["notification_type"], "public_advisory")
        self.assertTrue(data["is_active"])

        # Responses must NOT include recipient_user, target_role, is_read, read_at
        self.assertNotIn("recipient_user", data)
        self.assertNotIn("target_role", data)
        self.assertNotIn("is_read", data)
        self.assertNotIn("read_at", data)

        # Confirm in DB
        advisory = Notification.objects.get(title=payload["title"])
        self.assertEqual(advisory.audience_type, "public")
        self.assertEqual(advisory.notification_type, "public_advisory")
        self.assertIsNone(advisory.recipient_user)
        self.assertEqual(advisory.target_role, "")
        self.assertTrue(advisory.is_active)

    def test_unauthenticated_get_only_sees_active_non_expired(self):
        """
        c) Asserts an unauthenticated GET only sees is_active=True, non-expired advisories
        (create one expired, one inactive, one valid — confirm only the valid one appears).
        """
        now = timezone.now()

        valid_advisory = Notification.objects.create(
            title="Valid Advisory",
            message="Valid message",
            notification_type="public_advisory",
            audience_type="public",
            module="general",
            is_active=True,
            expires_at=now + timedelta(days=2),
        )
        expired_advisory = Notification.objects.create(
            title="Expired Advisory",
            message="Expired message",
            notification_type="public_advisory",
            audience_type="public",
            module="general",
            is_active=True,
            expires_at=now - timedelta(days=2),
        )
        inactive_advisory = Notification.objects.create(
            title="Inactive Advisory",
            message="Inactive message",
            notification_type="public_advisory",
            audience_type="public",
            module="general",
            is_active=False,
            expires_at=now + timedelta(days=2),
        )

        self.client.credentials()  # Unauthenticated
        response = self.client.get("/api/notifications/public/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("results", response.data)
        results = response.data["results"]
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["id"], valid_advisory.id)
        self.assertEqual(results[0]["title"], "Valid Advisory")

        # Also confirm an establishment user gets the same public view
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.establishment_token.key}")
        est_response = self.client.get("/api/notifications/public/")
        self.assertEqual(est_response.status_code, status.HTTP_200_OK)
        self.assertIn("results", est_response.data)
        est_results = est_response.data["results"]
        self.assertEqual(len(est_results), 1)
        self.assertEqual(est_results[0]["id"], valid_advisory.id)

    def test_authenticated_admin_get_sees_all_advisories(self):
        """
        d) Asserts an authenticated admin GET sees all three (active, inactive, and expired)
        in the management view.
        """
        now = timezone.now()

        valid = Notification.objects.create(
            title="Active Valid Advisory",
            message="Valid",
            notification_type="public_advisory",
            audience_type="public",
            module="sanitation",
            is_active=True,
            expires_at=now + timedelta(days=5),
        )
        expired = Notification.objects.create(
            title="Expired Advisory",
            message="Expired",
            notification_type="public_advisory",
            audience_type="public",
            module="tourism",
            is_active=True,
            expires_at=now - timedelta(days=1),
        )
        inactive = Notification.objects.create(
            title="Inactive Advisory",
            message="Inactive",
            notification_type="public_advisory",
            audience_type="public",
            module="general",
            is_active=False,
            expires_at=now + timedelta(days=10),
        )

        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        response = self.client.get("/api/notifications/public/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("results", response.data)
        results = response.data["results"]
        self.assertEqual(len(results), 3)
        returned_ids = {item["id"] for item in results}
        self.assertEqual(returned_ids, {valid.id, expired.id, inactive.id})

        # Test module filter for staff
        tourism_resp = self.client.get("/api/notifications/public/?module=tourism")
        self.assertEqual(tourism_resp.status_code, status.HTTP_200_OK)
        self.assertIn("results", tourism_resp.data)
        self.assertEqual(len(tourism_resp.data["results"]), 1)
        self.assertEqual(tourism_resp.data["results"][0]["id"], expired.id)

    def test_patch_deactivates_without_deleting(self):
        """
        e) Asserts PATCH with is_active=False removes it from the public
        unauthenticated view but the row still exists in the DB
        (query it directly to confirm it wasn't deleted).
        """
        advisory = Notification.objects.create(
            title="Water Quality Advisory",
            message="Boil water advisory in effect",
            notification_type="public_advisory",
            audience_type="public",
            module="sanitation",
            is_active=True,
        )

        # Authenticated staff (sanitation) deactivates it
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.sanitation_token.key}")
        patch_resp = self.client.patch(
            f"/api/notifications/public/{advisory.id}/",
            {"is_active": False},
            format="json",
        )
        self.assertEqual(patch_resp.status_code, status.HTTP_200_OK)
        self.assertFalse(patch_resp.data["is_active"])

        # Row still exists in DB
        advisory.refresh_from_db()
        self.assertFalse(advisory.is_active)
        self.assertTrue(Notification.objects.filter(id=advisory.id).exists())

        # Unauthenticated GET no longer sees it
        self.client.credentials()
        get_resp = self.client.get("/api/notifications/public/")
        self.assertEqual(get_resp.status_code, status.HTTP_200_OK)
        self.assertIn("results", get_resp.data)
        ids = [item["id"] for item in get_resp.data["results"]]
        self.assertNotIn(advisory.id, ids)

    def test_patch_non_public_advisory_returns_404(self):
        """
        f) Asserts PATCH on a non-public_advisory notification id (e.g. one
        created by notify_violation) returns 404, not 200 — this
        endpoint must not be able to edit staff notifications.
        """
        staff_notif = Notification.objects.create(
            title="Internal Violation Warning",
            message="Internal alert for inspector",
            notification_type="violation_alert",
            audience_type="user",
            recipient_user=self.admin_user,
            module="sanitation",
            is_active=True,
        )

        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        patch_resp = self.client.patch(
            f"/api/notifications/public/{staff_notif.id}/",
            {"title": "Tampered Title", "is_active": False},
            format="json",
        )
        self.assertEqual(patch_resp.status_code, status.HTTP_404_NOT_FOUND)

        staff_notif.refresh_from_db()
        self.assertEqual(staff_notif.title, "Internal Violation Warning")
        self.assertTrue(staff_notif.is_active)

    def test_user_without_profile_gets_403(self):
        """Asserts a user with no profile gets 403 on POST."""
        user_no_profile = User.objects.create_user(
            username="user_no_profile",
            password="Password@123",
            email="no_profile@test.local",
        )
        token = Token.objects.create(user=user_no_profile)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")
        response = self.client.post(
            "/api/notifications/public/",
            {"title": "Test", "message": "Test", "module": "general"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)


class MobileSanitationAuthTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Create active admin user
        self.admin_user = User.objects.create_user(
            username="test_admin_auth",
            password="Password@123",
            email="admin_auth@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)

        # Create active sanitation user
        self.sanitation_user = User.objects.create_user(
            username="test_sanitation_auth",
            password="Password@123",
            email="sanitation_auth@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)
        self.sanitation_token = Token.objects.create(user=self.sanitation_user)

        # Create active establishment user
        self.establishment_user = User.objects.create_user(
            username="test_establishment_auth",
            password="Password@123",
            email="establishment_auth@test.local",
        )
        UserProfile.objects.create(user=self.establishment_user, role=ROLE_ESTABLISHMENT)
        self.establishment_token = Token.objects.create(user=self.establishment_user)

        # Create test establishment
        self.btype = SanitaryBusinessType.objects.create(
            name="Auth Test Business",
            inspection_frequency="monthly",
        )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Auth Test Establishment",
            owner_name="Test Owner",
            business_type=self.btype,
            barangay="Poblacion",
            address="123 Test St",
            compliance_status="good_standing",
            permit_status="active",
        )

        self.inspection_payload = {
            "establishment": self.establishment.id,
            "inspector_name": "Inspector Auth",
            "inspection_date": "2026-09-11",
            "next_due_date": "2026-10-11",
            "findings": "All compliant.",
            "remarks": "Regular inspection.",
            "status_after_inspection": "good_standing",
            "checklist_items": [
                {"requirement_name": "Sanitary Permit", "is_complied": True, "notes": ""}
            ],
        }

        self.survey_payload = {
            "household_head": "Auth Test Household",
            "barangay": "Poblacion",
            "address": "Zone 1",
            "male_count": 2,
            "female_count": 2,
            "toilet_type": "water_sealed",
            "water_level": "level_3",
            "water_source": "MWSS",
            "waste_disposal": "collected",
            "remarks": "Normal condition.",
            "latitude": "14.186",
            "longitude": "121.73",
        }

    def test_unauthenticated_post_returns_401_with_zero_db_writes(self):
        self.client.credentials()  # No auth
        initial_inspections = SanitaryInspection.objects.count()
        initial_surveys = HouseholdSanitationRecord.objects.count()

        # 1. Inspection endpoint
        resp_insp = self.client.post(
            "/api/mobile/sanitation/inspections/",
            self.inspection_payload,
            format="json",
        )
        self.assertEqual(resp_insp.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(SanitaryInspection.objects.count(), initial_inspections)

        # 2. Household survey endpoint
        resp_surv = self.client.post(
            "/api/mobile/sanitation/household-surveys/",
            self.survey_payload,
            format="json",
        )
        self.assertEqual(resp_surv.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(HouseholdSanitationRecord.objects.count(), initial_surveys)

    def test_establishment_role_post_returns_403_with_zero_db_writes(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.establishment_token.key}")
        initial_inspections = SanitaryInspection.objects.count()
        initial_surveys = HouseholdSanitationRecord.objects.count()

        # 1. Inspection endpoint
        resp_insp = self.client.post(
            "/api/mobile/sanitation/inspections/",
            self.inspection_payload,
            format="json",
        )
        self.assertEqual(resp_insp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn("detail", resp_insp.json())
        self.assertEqual(SanitaryInspection.objects.count(), initial_inspections)

        # 2. Household survey endpoint
        resp_surv = self.client.post(
            "/api/mobile/sanitation/household-surveys/",
            self.survey_payload,
            format="json",
        )
        self.assertEqual(resp_surv.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn("detail", resp_surv.json())
        self.assertEqual(HouseholdSanitationRecord.objects.count(), initial_surveys)

    def test_sanitation_role_post_succeeds_201(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.sanitation_token.key}")
        initial_inspections = SanitaryInspection.objects.count()
        initial_surveys = HouseholdSanitationRecord.objects.count()

        # 1. Inspection endpoint
        resp_insp = self.client.post(
            "/api/mobile/sanitation/inspections/",
            self.inspection_payload,
            format="json",
        )
        self.assertEqual(resp_insp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(SanitaryInspection.objects.count(), initial_inspections + 1)
        self.assertTrue(
            SanitaryInspection.objects.filter(inspector_name="Inspector Auth").exists()
        )

        # 2. Household survey endpoint
        resp_surv = self.client.post(
            "/api/mobile/sanitation/household-surveys/",
            self.survey_payload,
            format="json",
        )
        self.assertEqual(resp_surv.status_code, status.HTTP_201_CREATED)
        self.assertEqual(HouseholdSanitationRecord.objects.count(), initial_surveys + 1)
        self.assertTrue(
            HouseholdSanitationRecord.objects.filter(household_head="Auth Test Household").exists()
        )


class MobileTourismAuthTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Explicit reference models needed for TouristRecord
        self.country, _ = Country.objects.get_or_create(id=991, defaults={"name": "Philippines"})
        self.region, _ = Region.objects.get_or_create(id=991, defaults={"name": "CALABARZON", "code": "04"})
        self.province, _ = Province.objects.get_or_create(id=991, defaults={"name": "Quezon", "region": self.region, "code": "QUE"})
        self.itinerary, _ = Itinerary.objects.get_or_create(id=991, defaults={"name": "Day Tour"})
        self.travel_mode, _ = TravelMode.objects.get_or_create(id=991, defaults={"name": "Private"})
        self.boat_type, _ = BoatType.objects.get_or_create(id=991, defaults={"name": "Motorized"})
        self.visit_purpose, _ = VisitPurpose.objects.get_or_create(id=991, defaults={"name": "Pleasure / Vacation"})
        self.resort, _ = Resort.objects.get_or_create(
            resort_id=991,
            defaults={
                "resort_name": "Test Auth Resort",
                "type": "beach",
                "location": "Mauban",
                "short_description": "Test",
                "access": "boat",
                "latitude": 14.18,
                "longitude": 121.73,
            },
        )

        # Admin user
        self.admin_user = User.objects.create_user(
            username="test_tourism_admin_auth",
            password="Password@123",
            email="admin_tourism_auth@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)
        self.admin_token = Token.objects.create(user=self.admin_user)

        # Tourism user
        self.tourism_user = User.objects.create_user(
            username="test_tourism_staff_auth",
            password="Password@123",
            email="staff_tourism_auth@test.local",
        )
        UserProfile.objects.create(user=self.tourism_user, role=ROLE_TOURISM)
        self.tourism_token = Token.objects.create(user=self.tourism_user)

        # Sanitation user
        self.sanitation_user = User.objects.create_user(
            username="test_sanitation_role_auth",
            password="Password@123",
            email="sanitation_role_auth@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)
        self.sanitation_token = Token.objects.create(user=self.sanitation_user)

        # Establishment user
        self.establishment_user = User.objects.create_user(
            username="test_establishment_role_auth",
            password="Password@123",
            email="establishment_role_auth@test.local",
        )
        UserProfile.objects.create(user=self.establishment_user, role=ROLE_ESTABLISHMENT)
        self.establishment_token = Token.objects.create(user=self.establishment_user)

        # Test Tourist Record with sensitive PII
        self.tourist_name = "Jane Sensitive Tourist"
        self.tourist_email = "jane.sensitive@privatemail.local"
        self.tourist_phone = "09179988776"
        self.survey_id = "MOB-SEC-2026-99"

        self.record = TouristRecord.objects.create(
            survey_id=self.survey_id,
            first_name="Jane",
            last_name="Sensitive Tourist",
            full_name=self.tourist_name,
            email=self.tourist_email,
            contact_number=self.tourist_phone,
            country=self.country,
            region=self.region,
            province=self.province,
            arrival_date=timezone.localdate(),
            resort=self.resort,
            itinerary=self.itinerary,
            travel_mode=self.travel_mode,
            boat_type=self.boat_type,
            visit_purpose=self.visit_purpose,
            total_visitors=2,
            filipino_count=2,
            foreigner_count=0,
            total_male=1,
            total_female=1,
            age_8_59=2,
            status="pending",
        )

        self.check_in_payload = {
            "survey_id": self.survey_id,
            "status": "arrived",
            "total_visitors": 2,
            "filipino_count": 2,
            "foreigner_count": 0,
            "total_male": 1,
            "total_female": 1,
            "age_8_59": 2,
        }

    def test_unauthenticated_requests_return_401_and_contain_no_tourist_pii(self):
        self.client.credentials()  # No auth

        # 1. Lookup endpoint
        resp_lookup = self.client.get(
            f"/api/mobile/tourism/records/lookup/?query={self.survey_id}"
        )
        self.assertEqual(resp_lookup.status_code, status.HTTP_401_UNAUTHORIZED)
        lookup_body = resp_lookup.content.decode("utf-8")
        self.assertNotIn(self.tourist_name, lookup_body)
        self.assertNotIn(self.tourist_email, lookup_body)
        self.assertNotIn(self.tourist_phone, lookup_body)

        # 2. Check-in endpoint
        resp_checkin = self.client.post(
            "/api/mobile/tourism/records/check-in/",
            self.check_in_payload,
            format="json",
        )
        self.assertEqual(resp_checkin.status_code, status.HTTP_401_UNAUTHORIZED)
        checkin_body = resp_checkin.content.decode("utf-8")
        self.assertNotIn(self.tourist_name, checkin_body)
        self.assertNotIn(self.tourist_email, checkin_body)
        self.assertNotIn(self.tourist_phone, checkin_body)

        # Verify zero DB mutation on record status
        self.record.refresh_from_db()
        self.assertEqual(self.record.status, "pending")

        # 3. History endpoint
        resp_history = self.client.get(
            f"/api/mobile/tourism/records/history/?search={self.survey_id}"
        )
        self.assertEqual(resp_history.status_code, status.HTTP_401_UNAUTHORIZED)
        history_body = resp_history.content.decode("utf-8")
        self.assertNotIn(self.tourist_name, history_body)
        self.assertNotIn(self.tourist_email, history_body)
        self.assertNotIn(self.tourist_phone, history_body)

    def test_disallowed_roles_return_403_with_zero_db_mutation(self):
        for role_name, token in [
            ("establishment", self.establishment_token),
            ("sanitation", self.sanitation_token),
        ]:
            self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

            # 1. Lookup
            resp_lookup = self.client.get(
                f"/api/mobile/tourism/records/lookup/?query={self.survey_id}"
            )
            self.assertEqual(
                resp_lookup.status_code,
                status.HTTP_403_FORBIDDEN,
                f"Role {role_name} should be rejected with 403 on lookup",
            )
            self.assertIn("detail", resp_lookup.json())

            # 2. Check-in
            resp_checkin = self.client.post(
                "/api/mobile/tourism/records/check-in/",
                self.check_in_payload,
                format="json",
            )
            self.assertEqual(
                resp_checkin.status_code,
                status.HTTP_403_FORBIDDEN,
                f"Role {role_name} should be rejected with 403 on check-in",
            )
            self.assertIn("detail", resp_checkin.json())
            self.record.refresh_from_db()
            self.assertEqual(self.record.status, "pending")

            # 3. History
            resp_history = self.client.get(
                f"/api/mobile/tourism/records/history/?search={self.survey_id}"
            )
            self.assertEqual(
                resp_history.status_code,
                status.HTTP_403_FORBIDDEN,
                f"Role {role_name} should be rejected with 403 on history",
            )
            self.assertIn("detail", resp_history.json())

    def test_tourism_and_admin_roles_succeed(self):
        # 1. Test tourism role succeeds on lookup, check-in, and history
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")

        resp_lookup = self.client.get(
            f"/api/mobile/tourism/records/lookup/?query={self.survey_id}"
        )
        self.assertEqual(resp_lookup.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_lookup.json()["survey_id"], self.survey_id)
        self.assertEqual(resp_lookup.json()["full_name"], self.tourist_name)

        resp_checkin = self.client.post(
            "/api/mobile/tourism/records/check-in/",
            self.check_in_payload,
            format="json",
        )
        self.assertEqual(resp_checkin.status_code, status.HTTP_200_OK)
        self.record.refresh_from_db()
        self.assertEqual(self.record.status, BOOKING_STATUS_ARRIVED)

        resp_history = self.client.get(
            f"/api/mobile/tourism/records/history/?search={self.survey_id}"
        )
        self.assertEqual(resp_history.status_code, status.HTTP_200_OK)
        records = resp_history.json().get("records", [])
        self.assertTrue(any(r["survey_id"] == self.survey_id for r in records))

        # 2. Test admin role also succeeds
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        resp_admin_lookup = self.client.get(
            f"/api/mobile/tourism/records/lookup/?query={self.survey_id}"
        )
        self.assertEqual(resp_admin_lookup.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_admin_lookup.json()["survey_id"], self.survey_id)


class MobileBootstrapNotificationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        now = timezone.now()

        # 1. Active public advisories
        self.tourism_advisory = Notification.objects.create(
            title="Tourism Advisory Test",
            message="Active public advisory for tourists",
            notification_type=NOTIFICATION_TYPE_PUBLIC_ADVISORY,
            audience_type=NOTIFICATION_AUDIENCE_PUBLIC,
            module=NOTIFICATION_MODULE_TOURISM,
            is_active=True,
        )
        self.sanitation_advisory = Notification.objects.create(
            title="Sanitation Advisory Test",
            message="Active public advisory for sanitation",
            notification_type=NOTIFICATION_TYPE_PUBLIC_ADVISORY,
            audience_type=NOTIFICATION_AUDIENCE_PUBLIC,
            module=NOTIFICATION_MODULE_SANITATION,
            is_active=True,
        )

        # 2. Inactive advisory and expired advisory
        self.inactive_advisory = Notification.objects.create(
            title="Inactive Advisory",
            message="Should not appear in bootstrap",
            notification_type=NOTIFICATION_TYPE_PUBLIC_ADVISORY,
            audience_type=NOTIFICATION_AUDIENCE_PUBLIC,
            module=NOTIFICATION_MODULE_TOURISM,
            is_active=False,
        )
        self.expired_advisory = Notification.objects.create(
            title="Expired Advisory",
            message="Should not appear in bootstrap",
            notification_type=NOTIFICATION_TYPE_PUBLIC_ADVISORY,
            audience_type=NOTIFICATION_AUDIENCE_PUBLIC,
            module=NOTIFICATION_MODULE_TOURISM,
            is_active=True,
            expires_at=now - timedelta(days=1),
        )

    def test_mobile_tourism_bootstrap_dynamic_advisories(self):
        resp = self.client.get("/api/mobile/tourism/bootstrap/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        notifications = resp.json().get("notifications", [])
        notif_ids = [n["id"] for n in notifications]

        # Includes active tourism advisory
        self.assertIn(f"advisory-{self.tourism_advisory.id}", notif_ids)

        # Does NOT include sanitation advisory, inactive advisory, or expired advisory
        self.assertNotIn(f"advisory-{self.sanitation_advisory.id}", notif_ids)
        self.assertNotIn(f"advisory-{self.inactive_advisory.id}", notif_ids)
        self.assertNotIn(f"advisory-{self.expired_advisory.id}", notif_ids)

    def test_mobile_sanitation_bootstrap_dynamic_advisories(self):
        resp = self.client.get("/api/mobile/sanitation/bootstrap/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        notifications = resp.json().get("notifications", [])
        notif_ids = [n["id"] for n in notifications]

        # Includes active sanitation advisory
        self.assertIn(f"advisory-{self.sanitation_advisory.id}", notif_ids)

        # Does NOT include tourism-only advisory, inactive advisory, or expired advisory
        self.assertNotIn(f"advisory-{self.tourism_advisory.id}", notif_ids)
        self.assertNotIn(f"advisory-{self.inactive_advisory.id}", notif_ids)
        self.assertNotIn(f"advisory-{self.expired_advisory.id}", notif_ids)


class NotificationWebhookCronTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = "/api/notifications/evaluate-due/"
        self.secret = "super_secret_cron_key_987"

    @patch("api.views.notifications.config")
    def test_post_without_key_returns_403(self, mock_config):
        mock_config.side_effect = lambda key, default="": self.secret if key == "CRON_SECRET_KEY" else default
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(resp.json()["detail"], "Invalid or missing cron key.")

    @patch("api.views.notifications.config")
    def test_post_with_invalid_key_returns_403(self, mock_config):
        mock_config.side_effect = lambda key, default="": self.secret if key == "CRON_SECRET_KEY" else default
        resp = self.client.post(self.url, HTTP_X_CRON_KEY="wrong_secret")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(resp.json()["detail"], "Invalid or missing cron key.")

    @patch("api.views.notifications.config")
    def test_post_with_valid_key_triggers_evaluation_and_returns_200(self, mock_config):
        mock_config.side_effect = lambda key, default="": self.secret if key == "CRON_SECRET_KEY" else default

        # Test with X-Cron-Key header
        resp1 = self.client.post(self.url, HTTP_X_CRON_KEY=self.secret)
        self.assertEqual(resp1.status_code, status.HTTP_200_OK)
        self.assertEqual(resp1.json()["status"], "success")
        self.assertEqual(resp1.json()["message"], "Due notifications evaluated.")

        # Test with Authorization Bearer header
        resp2 = self.client.post(self.url, HTTP_AUTHORIZATION=f"Bearer {self.secret}")
        self.assertEqual(resp2.status_code, status.HTTP_200_OK)
        self.assertEqual(resp2.json()["status"], "success")
        self.assertEqual(resp2.json()["message"], "Due notifications evaluated.")

    @patch("api.views.notifications.config")
    def test_post_when_unconfigured_returns_503(self, mock_config):
        mock_config.side_effect = lambda key, default="": "" if key == "CRON_SECRET_KEY" else default
        resp = self.client.post(self.url, HTTP_X_CRON_KEY="some_key")
        self.assertEqual(resp.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertEqual(resp.json()["detail"], "Cron trigger is not configured on this server.")


class SanitaryStaffApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        # Admin user
        self.admin = User.objects.create_user(
            username="sanitary_superadmin",
            password="Password@123",
            is_staff=True,
            is_superuser=True,
        )
        UserProfile.objects.create(user=self.admin, role=ROLE_ADMIN)
        self.admin_token, _ = Token.objects.get_or_create(user=self.admin)

        # Existing staff inspector
        self.inspector = User.objects.create_user(
            username="inspector_santos",
            password="Password@123",
            first_name="Maria",
            last_name="Santos",
            email="maria.santos@mauban.gov.ph",
            is_staff=True,
            is_active=True,
        )
        UserProfile.objects.create(user=self.inspector, role=ROLE_SANITATION)
        self.inspector_token, _ = Token.objects.get_or_create(user=self.inspector)

        # Tourist user
        self.tourist = User.objects.create_user(
            username="tourist_guest",
            password="Password@123",
        )
        UserProfile.objects.create(user=self.tourist, role=ROLE_TOURISM)
        self.tourist_token, _ = Token.objects.get_or_create(user=self.tourist)

    def test_get_staff_list(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        response = self.client.get("/api/v1/sanitation/staff/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        usernames = [u["username"] for u in response.json()]
        self.assertIn("inspector_santos", usernames)

    def test_create_staff_success(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        payload = {
            "first_name": "Juan",
            "last_name": "Dela Cruz",
            "username": "inspector_juan",
            "password": "SecurePassword@123",
            "email": "juan.cruz@mauban.gov.ph",
        }
        response = self.client.post("/api/v1/sanitation/staff/", payload)
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        self.assertEqual(data["username"], "inspector_juan")
        self.assertEqual(data["first_name"], "Juan")
        self.assertEqual(data["last_name"], "Dela Cruz")
        self.assertEqual(data["role"], ROLE_SANITATION)
        self.assertTrue(data["is_active"])

        # Check in DB
        created_user = User.objects.get(username="inspector_juan")
        self.assertTrue(created_user.check_password("SecurePassword@123"))
        self.assertTrue(created_user.is_staff)
        self.assertEqual(created_user.profile.role, ROLE_SANITATION)

    def test_create_staff_validation_errors(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        # Missing fields
        resp = self.client.post("/api/v1/sanitation/staff/", {})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

        # Password too short
        resp = self.client.post("/api/v1/sanitation/staff/", {
            "first_name": "Test",
            "last_name": "User",
            "username": "test_short",
            "password": "123",
        })
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("at least 6 characters", resp.json()["detail"])

        # Duplicate username
        resp = self.client.post("/api/v1/sanitation/staff/", {
            "first_name": "Maria",
            "last_name": "Santos",
            "username": "inspector_santos",
            "password": "Password@123",
        })
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("already in use", resp.json()["detail"])

    def test_toggle_staff_status(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        # Deactivate
        resp = self.client.patch(
            f"/api/v1/sanitation/staff/{self.inspector.pk}/",
            {"is_active": False},
            format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertFalse(resp.json()["is_active"])
        self.inspector.refresh_from_db()
        self.assertFalse(self.inspector.is_active)

        # Reactivate
        resp = self.client.patch(
            f"/api/v1/sanitation/staff/{self.inspector.pk}/",
            {"is_active": True},
            format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertTrue(resp.json()["is_active"])
        self.inspector.refresh_from_db()
        self.assertTrue(self.inspector.is_active)

    def test_prevent_self_deactivation(self):
        # Admin cannot deactivate self if admin user was targeted
        admin_profile = self.admin.profile
        admin_profile.role = ROLE_SANITATION
        admin_profile.save()

        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.admin_token.key}")
        resp = self.client.patch(
            f"/api/v1/sanitation/staff/{self.admin.pk}/",
            {"is_active": False},
            format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("cannot deactivate your own", resp.json()["detail"])

    def test_unauthorized_access(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourist_token.key}")
        resp = self.client.get("/api/v1/sanitation/staff/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


from django.core.files.uploadedfile import SimpleUploadedFile
from api.services.upload import (
    MAX_FILE_SIZE_BYTES,
    StorageServiceError,
    UploadValidationError,
    save_image_file,
    validate_image_file,
)


class SecureUploadTests(TestCase):
    def setUp(self):
        ensure_test_barangays()
        self.client = APIClient()
        self.tourism_user = User.objects.create_user(
            username="tourism_staff_upload",
            password="Password@123",
            is_staff=True,
        )
        UserProfile.objects.create(user=self.tourism_user, role=ROLE_TOURISM)
        self.tourism_token, _ = Token.objects.get_or_create(user=self.tourism_user)

        self.inspector_user = User.objects.create_user(
            username="inspector_upload",
            password="Password@123",
            is_staff=True,
        )
        UserProfile.objects.create(user=self.inspector_user, role=ROLE_SANITATION)
        self.inspector_token, _ = Token.objects.get_or_create(user=self.inspector_user)

    def test_validate_image_file_valid_types(self):
        jpg_file = SimpleUploadedFile("photo.jpg", b"fake jpg content", content_type="image/jpeg")
        self.assertEqual(validate_image_file(jpg_file), ".jpg")

        upper_jpg = SimpleUploadedFile("PHOTO.JPG", b"fake jpg content", content_type="image/jpeg")
        self.assertEqual(validate_image_file(upper_jpg), ".jpg")

        jpeg_file = SimpleUploadedFile("photo.JPEG", b"fake jpeg content", content_type="image/jpeg")
        self.assertEqual(validate_image_file(jpeg_file), ".jpg")

        png_file = SimpleUploadedFile("photo.png", b"fake png content", content_type="image/png")
        self.assertEqual(validate_image_file(png_file), ".png")

        webp_file = SimpleUploadedFile("photo.webp", b"fake webp content", content_type="image/webp")
        self.assertEqual(validate_image_file(webp_file), ".webp")

        heic_file = SimpleUploadedFile("photo.heic", b"fake heic content", content_type="image/heic")
        self.assertEqual(validate_image_file(heic_file), ".heic")

        heif_file = SimpleUploadedFile("photo.HEIF", b"fake heif content", content_type="image/heif")
        self.assertEqual(validate_image_file(heif_file), ".heic")

        octet_jpg = SimpleUploadedFile("photo.jpg", b"fake binary", content_type="application/octet-stream")
        self.assertEqual(validate_image_file(octet_jpg), ".jpg")

        octet_no_ext = SimpleUploadedFile("camera_raw", b"fake binary", content_type="application/octet-stream")
        self.assertEqual(validate_image_file(octet_no_ext), ".jpg")

    def test_validate_image_file_oversized(self):
        large_content = b"0" * (MAX_FILE_SIZE_BYTES + 1)
        oversized = SimpleUploadedFile("big.jpg", large_content, content_type="image/jpeg")
        with self.assertRaises(UploadValidationError) as ctx:
            validate_image_file(oversized)
        self.assertEqual(str(ctx.exception), "File size exceeds 5MB limit.")

    def test_validate_image_file_disallowed_type(self):
        pdf_file = SimpleUploadedFile("doc.pdf", b"%PDF-1.4...", content_type="application/pdf")
        with self.assertRaises(UploadValidationError) as ctx:
            validate_image_file(pdf_file)
        self.assertEqual(str(ctx.exception), "Only JPG, PNG, and WebP images are allowed.")

        exe_file = SimpleUploadedFile("malware.exe", b"MZ...", content_type="image/jpeg")
        with self.assertRaises(UploadValidationError) as ctx:
            validate_image_file(exe_file)
        self.assertEqual(str(ctx.exception), "Only JPG, PNG, and WebP images are allowed.")

    def test_save_image_file_generates_uuid_name(self):
        test_file = SimpleUploadedFile("my vacation photo.JPG", b"sample-bytes", content_type="image/jpeg")
        url = save_image_file(test_file, "resorts")
        self.assertIn("resorts/", url)
        self.assertNotIn("my vacation photo", url)
        self.assertTrue(url.endswith(".jpg"))

    def test_save_image_file_storage_error_raises_503_exception(self):
        test_file = SimpleUploadedFile("photo.png", b"sample-bytes", content_type="image/png")
        with patch("api.services.upload.default_storage.save", side_effect=IOError("Storage quota exceeded")):
            with self.assertRaises(StorageServiceError) as ctx:
                save_image_file(test_file, "resorts")
            self.assertEqual(str(ctx.exception), "Image upload failed. Storage service unavailable or full.")

    def test_resort_image_upload_endpoint_success(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        test_file = SimpleUploadedFile("resort_view.png", b"image-content", content_type="image/png")
        response = self.client.post("/api/resorts/upload-image/", {"image": test_file}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn("url", response.json())
        self.assertIn("resorts/", response.json()["url"])

    def test_resort_image_upload_endpoint_oversized_rejected(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        oversized = SimpleUploadedFile("huge.png", b"0" * (MAX_FILE_SIZE_BYTES + 1), content_type="image/png")
        response = self.client.post("/api/resorts/upload-image/", {"image": oversized}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["error"], "File size exceeds 5MB limit.")

    def test_resort_image_upload_endpoint_invalid_type_rejected(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        bad_file = SimpleUploadedFile("document.pdf", b"pdf-data", content_type="application/pdf")
        response = self.client.post("/api/resorts/upload-image/", {"image": bad_file}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["error"], "Only JPG, PNG, and WebP images are allowed.")

    def test_resort_image_upload_endpoint_storage_failure_503(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.tourism_token.key}")
        test_file = SimpleUploadedFile("resort.jpg", b"image-data", content_type="image/jpeg")
        with patch("api.services.upload.default_storage.save", side_effect=Exception("S3 Connection Timeout")):
            response = self.client.post("/api/resorts/upload-image/", {"image": test_file}, format="multipart")
            self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
            self.assertEqual(response.json()["error"], "Image upload failed. Storage service unavailable or full.")

    def test_mobile_sanitation_report_upload_validation(self):
        bad_file = SimpleUploadedFile("evidence.pdf", b"pdf-data", content_type="application/pdf")
        response = self.client.post(
            "/api/mobile/sanitation/reports/",
            {
                "complainant_name": "Upload Tester",
                "contact_number": "09171234567",
                "category": "Improper Garbage Disposal",
                "barangay": "Daungan",
                "location_address": "Creek bank, Purok 1",
                "description": "Garbage dump near creek",
                "photo": bad_file,
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["error"], "Only JPG, PNG, and WebP images are allowed.")

    def test_mobile_sanitation_report_storage_failure_503(self):
        good_file = SimpleUploadedFile("trash.jpg", b"image-data", content_type="image/jpeg")
        with patch("api.services.upload.default_storage.save", side_effect=Exception("S3 Storage Bucket Full")):
            response = self.client.post(
                "/api/mobile/sanitation/reports/",
                {
                    "complainant_name": "Upload Tester",
                    "contact_number": "09171234567",
                    "category": "Improper Garbage Disposal",
                    "barangay": "Daungan",
                "location_address": "Creek bank, Purok 1",
                    "description": "Garbage dump near creek",
                    "photo": good_file,
                },
                format="multipart",
            )
            self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
            self.assertEqual(response.json()["error"], "Image upload failed. Storage service unavailable or full.")


class PublicBoatRenameMigrationTests(TestCase):
    OLD = "Public Boat (P100/ride/head) Sabang Port Only"
    NEW = "Public Boat"

    def setUp(self):
        import importlib

        self.migration = importlib.import_module(
            "api.migrations.0033_rename_public_boat_type"
        )

    def _forward(self):
        from io import StringIO
        from contextlib import redirect_stdout
        from django.apps import apps

        out = StringIO()
        with redirect_stdout(out):
            self.migration.rename_public_boat_forward(apps, None)
        return out.getvalue()

    def _reverse(self):
        from io import StringIO
        from contextlib import redirect_stdout
        from django.apps import apps

        out = StringIO()
        with redirect_stdout(out):
            self.migration.rename_public_boat_reverse(apps, None)
        return out.getvalue()

    def test_renames_only_the_old_public_boat_row_and_keeps_its_id(self):
        BoatType.objects.create(id=901, name=self.OLD)
        BoatType.objects.create(id=902, name="Private Boat (Rates depend on the capacity)")
        BoatType.objects.create(id=903, name="2.0")

        self._forward()

        self.assertEqual(BoatType.objects.get(id=901).name, self.NEW)
        self.assertEqual(
            BoatType.objects.get(id=902).name,
            "Private Boat (Rates depend on the capacity)",
        )
        self.assertEqual(BoatType.objects.get(id=903).name, "2.0")
        self.assertEqual(BoatType.objects.count(), 3)

    def test_skips_with_warning_when_public_boat_already_exists(self):
        BoatType.objects.create(id=901, name=self.OLD)
        BoatType.objects.create(id=902, name=self.NEW)

        output = self._forward()

        self.assertIn("WARNING", output)
        self.assertEqual(BoatType.objects.get(id=901).name, self.OLD)
        self.assertEqual(BoatType.objects.get(id=902).name, self.NEW)
        self.assertEqual(BoatType.objects.filter(name=self.NEW).count(), 1)
        self.assertEqual(BoatType.objects.count(), 2)

    def test_does_nothing_when_old_row_is_absent(self):
        BoatType.objects.create(id=901, name="Speedboat")

        self._forward()

        self.assertEqual(list(BoatType.objects.values_list("name", flat=True)), ["Speedboat"])

    def test_reverse_restores_the_old_name(self):
        BoatType.objects.create(id=901, name=self.OLD)
        self._forward()
        self.assertEqual(BoatType.objects.get(id=901).name, self.NEW)

        self._reverse()

        self.assertEqual(BoatType.objects.get(id=901).name, self.OLD)
        self.assertEqual(BoatType.objects.count(), 1)

    def test_reverse_skips_when_old_name_already_exists(self):
        BoatType.objects.create(id=901, name=self.OLD)
        BoatType.objects.create(id=902, name=self.NEW)

        output = self._reverse()

        self.assertIn("WARNING", output)
        self.assertEqual(BoatType.objects.get(id=901).name, self.OLD)
        self.assertEqual(BoatType.objects.get(id=902).name, self.NEW)



class _ProductionBoatRows:
    """The seven boat rows production has before 0045, plus record helpers."""

    PRIVATE_OLD = "Private Boat (Rates depend on the capacity)"
    FARE = "1-2 pax (One-Way P1500, Two-way P2000)"
    PRODUCTION_ROWS = {
        1: "Public Boat",
        2: PRIVATE_OLD,
        3: "Boat Provided by Resort (As confirmed by both guests and resort)",
        4: "2.0",
        5: "Motorized Banca",
        6: "Speedboat",
        7: "Passenger Boat",
    }
    # Sheet text, old wording and new, and the row each must land on.
    IMPORT_CELLS = {
        "Public Boat": 1,
        "  public boat (sabang)  ": 1,
        "Tourist Boat": 1,
        "": 1,
        "Private Boat": 2,
        PRIVATE_OLD: 2,
        "passenger boat": 2,
        "Boat provided by resort": 2,
        "Boat Provided by Resort (As confirmed by both guests and resort)": 2,
    }

    def setUp(self):
        ensure_test_reference_tables()
        BoatType.objects.all().delete()
        for boat_id, name in self.PRODUCTION_ROWS.items():
            BoatType.objects.create(id=boat_id, name=name)

    def _record(self, survey_id, boat_type_id, fare=""):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Boat Test Group",
            contact_number="09170000000",
            country=Country.objects.first(),
            region=Region.objects.first(),
            province=Province.objects.first(),
            arrival_date="2026-04-02",
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type_id=boat_type_id,
            boat_capacity_fare=fare,
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=4,
            filipino_count=4,
            total_male=2,
            total_female=2,
            age_8_59=4,
            status="arrived",
        )

    def _run(self, function):
        from io import StringIO
        from contextlib import redirect_stdout
        from django.apps import apps

        out = StringIO()
        with redirect_stdout(out):
            function(apps, None)
        return out.getvalue()

    def _names(self):
        return dict(BoatType.objects.values_list("id", "name"))

    def _resolve(self, cells):
        from .services.online_booking import ReferenceResolver, resolve_boat_type

        resolver = ReferenceResolver(commit=False)
        return {cell: resolve_boat_type(resolver, cell) for cell in cells}

    def assert_import_lands_on_rows_1_and_2(self):
        rows_before = BoatType.objects.count()

        resolved = self._resolve(self.IMPORT_CELLS)

        self.assertEqual({cell: item.id for cell, item in resolved.items()}, self.IMPORT_CELLS)
        self.assertTrue(all(not item._state.adding for item in resolved.values()))
        self.assertEqual(BoatType.objects.count(), rows_before)


class BoatCapacityFareFlagTests(_ProductionBoatRows, TestCase):
    """0044 puts the capacity-fare rule on the row; the import maps by id."""

    def setUp(self):
        import importlib

        super().setUp()
        self.flag_migration = importlib.import_module(
            "api.migrations.0044_boattype_requires_capacity_fare"
        )

    def test_flag_migration_sets_only_the_tourist_boat(self):
        self._run(self.flag_migration.flag_tourist_boat)

        self.assertEqual(
            list(BoatType.objects.filter(requires_capacity_fare=True).values_list("id", flat=True)),
            [1],
        )
        self.assertEqual(self._names(), self.PRODUCTION_ROWS)

    def test_import_maps_old_and_new_names_while_the_old_names_are_stored(self):
        self.assert_import_lands_on_rows_1_and_2()

    def test_import_still_creates_a_row_for_unrecognised_text(self):
        item = self._resolve(["3.0"])["3.0"]

        self.assertTrue(item._state.adding)
        self.assertEqual(item.name, "3.0")

    def test_reference_payload_exposes_the_flag(self):
        from .serializers import BoatTypeSerializer

        self._run(self.flag_migration.flag_tourist_boat)

        data = BoatTypeSerializer(BoatType.objects.filter(id__in=[1, 2]), many=True).data
        self.assertEqual(
            [dict(row) for row in data],
            [
                {"id": 1, "name": "Public Boat", "requires_capacity_fare": True, "is_active": True},
                {"id": 2, "name": self.PRIVATE_OLD, "requires_capacity_fare": False, "is_active": True},
            ],
        )


class BoatTypeConsolidationTests(_ProductionBoatRows, TestCase):
    """0045 renames ids 1 and 2 and removes the other rows."""

    def setUp(self):
        import importlib

        super().setUp()
        self.migration = importlib.import_module("api.migrations.0045_consolidate_boat_types")

    def test_forwards_renames_moves_the_passenger_record_and_deletes_the_rest(self):
        public = self._record("SURV-BOAT-001", 1, self.FARE)
        moved = self._record("SURV-BOAT-002", 7, self.FARE)
        private = self._record("SURV-BOAT-003", 2)

        self._run(self.migration.forwards)

        self.assertEqual(self._names(), {1: "Tourist Boat", 2: "Passenger Boat"})
        moved.refresh_from_db()
        self.assertEqual(moved.boat_type_id, 2)
        self.assertEqual(moved.boat_capacity_fare, self.FARE)
        public.refresh_from_db()
        self.assertEqual((public.boat_type_id, public.boat_capacity_fare), (1, self.FARE))
        private.refresh_from_db()
        self.assertEqual(private.boat_type_id, 2)

    def test_forwards_is_a_no_op_when_run_again(self):
        self._record("SURV-BOAT-001", 7, self.FARE)
        self._run(self.migration.forwards)

        self._run(self.migration.forwards)

        self.assertEqual(self._names(), {1: "Tourist Boat", 2: "Passenger Boat"})
        self.assertEqual(TouristRecord.objects.get().boat_type_id, 2)

    def test_forwards_keeps_a_row_that_still_has_records(self):
        self._record("SURV-BOAT-001", 5)

        output = self._run(self.migration.forwards)

        self.assertIn("WARNING", output)
        self.assertEqual(
            self._names(),
            {1: "Tourist Boat", 2: "Passenger Boat", 5: "Motorized Banca"},
        )
        self.assertEqual(TouristRecord.objects.get().boat_type_id, 5)

    def test_forwards_leaves_rows_whose_name_does_not_match(self):
        BoatType.objects.filter(id=4).update(name="Something Else")

        self._run(self.migration.forwards)

        self.assertEqual(
            self._names(),
            {1: "Tourist Boat", 2: "Passenger Boat", 4: "Something Else"},
        )

    def test_backwards_restores_names_and_rows_but_not_the_moved_record(self):
        self._record("SURV-BOAT-001", 7, self.FARE)
        self._run(self.migration.forwards)

        self._run(self.migration.backwards)

        self.assertEqual(self._names(), self.PRODUCTION_ROWS)
        record = TouristRecord.objects.get()
        self.assertEqual((record.boat_type_id, record.boat_capacity_fare), (2, self.FARE))

    def test_import_maps_old_and_new_names_after_the_rename(self):
        self._run(self.migration.forwards)

        self.assert_import_lands_on_rows_1_and_2()

    def test_import_creates_a_row_for_2_0_again_once_it_is_deleted(self):
        self._run(self.migration.forwards)

        item = self._resolve(["2.0"])["2.0"]

        self.assertTrue(item._state.adding)
        self.assertEqual(item.name, "2.0")

    def test_seed_matches_the_consolidated_rows(self):
        self.assertEqual(
            REFERENCE_TABLES["boat_types"],
            [
                {"id": 1, "name": "Tourist Boat", "requires_capacity_fare": True},
                {"id": 2, "name": "Passenger Boat", "requires_capacity_fare": False},
            ],
        )
        self.assertTrue(
            all(record["boat_type_id"] in (1, 2) for record in INITIAL_TOURIST_RECORDS)
        )


class ForeignOriginTests(TestCase):
    """Region and province are required for a Philippine record only, and a
    foreign record without them is counted under its country."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        user = User.objects.create_user(username="origin_admin", password="Origin@123")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        cls.token = Token.objects.create(user=user)
        ensure_test_reference_tables()
        cls.philippines = Country.objects.get(name="Philippines")
        cls.united_states = Country.objects.get(name="United States")
        cls.calabarzon = Region.objects.get(name__icontains="CALABARZON")
        cls.quezon = Province.objects.get(name="Quezon")

    def setUp(self):
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

    def _web_payload(self, **overrides):
        payload = {
            "first_name": "Ana",
            "last_name": "Cruz",
            "full_name": "Ana Cruz",
            "email": "ana@example.com",
            "contact_number": "+639171234567",
            "country_id": self.philippines.id,
            "region_id": self.calabarzon.id,
            "province_id": self.quezon.id,
            "country_of_origin": "",
            "resort_id": Resort.objects.first().resort_id,
            "itinerary_id": Itinerary.objects.first().id,
            "travel_mode_id": TravelMode.objects.first().id,
            "boat_type_id": BoatType.objects.first().id,
            "visit_purpose_id": VisitPurpose.objects.first().id,
            "arrival_date": "2026-04-01",
            "filipino_count": 0,
            "foreigner_count": 2,
            "total_visitors": 2,
            "total_male": 1,
            "total_female": 1,
            "age_8_59": 2,
            "status": "pending",
        }
        payload.update(overrides)
        return payload

    def _record(self, survey_id, country, province=None, region=None, visitors=1):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Origin Group",
            contact_number=f"0917{survey_id[-7:]}",
            country=country,
            region=region,
            province=province,
            arrival_date="2026-04-02",
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=visitors,
            foreigner_count=visitors,
            total_male=visitors,
            age_8_59=visitors,
            status="arrived",
        )

    # --- validation ---------------------------------------------------------

    def test_philippine_record_still_requires_region_and_province(self):
        for missing in ({"region_id": None, "province_id": None}, {"region_id": "", "province_id": ""}):
            with self.subTest(missing=missing):
                response = self.client.post(
                    "/api/tourist-records/", self._web_payload(**missing), format="json"
                )

                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn("region_id", response.json())
                self.assertIn("province_id", response.json())

        payload = self._web_payload()
        del payload["region_id"], payload["province_id"]
        response = self.client.post("/api/tourist-records/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(TouristRecord.objects.count(), 0)

    def test_foreign_record_saves_with_no_region_or_province(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._web_payload(
                country_id=self.united_states.id, region_id=None, province_id=None
            ),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.json())
        record = TouristRecord.objects.get()
        self.assertEqual((record.region_id, record.province_id), (None, None))
        self.assertEqual(record.country_id, self.united_states.id)

    def test_country_stays_required(self):
        response = self.client.post(
            "/api/tourist-records/",
            self._web_payload(country_id=None, region_id=None, province_id=None),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("country_id", response.json())

    def test_rule_follows_the_country_type_not_its_name(self):
        renamed = Country.objects.create(id=950, name="Republika ng Pilipinas", type="local")
        untyped_default = Country.objects.create(id=951, name="Canada")

        for country in (renamed, untyped_default):
            with self.subTest(country=country.name):
                response = self.client.post(
                    "/api/tourist-records/",
                    self._web_payload(country_id=country.id, region_id=None, province_id=None),
                    format="json",
                )
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn("region_id", response.json())

    # --- origin groupings ---------------------------------------------------

    def _seed_origins(self):
        self._record("ORIGIN-0000001", self.philippines, self.quezon, self.calabarzon, visitors=2)
        self._record("ORIGIN-0000002", self.united_states, visitors=5)

    def test_dashboard_top_origin_reports_the_country(self):
        from .services.tourism import build_dashboard_payload

        self._seed_origins()

        metrics = build_dashboard_payload({"year": "2026"})["metrics"]

        self.assertEqual(metrics["topOriginThisMonth"], "United States")

    def test_origin_report_groups_a_record_without_province_under_its_country(self):
        from .services.tourism import build_reports_payload

        self._seed_origins()

        rows = build_reports_payload(
            {"type": "origin", "year": "2026", "include_questions": "false"}
        )["rows"]

        self.assertEqual(
            [(row["name"], row["visitors"]) for row in rows],
            [("United States", 5), ("Quezon", 2)],
        )

    def test_analytics_top_origin_and_top_six_report_the_country(self):
        from .services.tourism import build_tourism_question_answers

        self._seed_origins()

        answers = {item["id"]: item for item in build_tourism_question_answers({"year": "2026"})}
        top_origin = answers["top_origin"]

        self.assertTrue(top_origin["answer"].startswith("United States leads with 5 visitors"))
        self.assertEqual(
            [(item["label"], item["value"]) for item in top_origin["visual"]["items"]],
            [("United States", 5), ("Quezon", 2)],
        )

    # --- mobile -------------------------------------------------------------

    def _mobile_payload(self, **overrides):
        payload = {
            "first_name": "Ana",
            "last_name": "Cruz",
            "contact_number": "+639171234500",
            "arrival_date": "2026-04-01",
            "resort_id": Resort.objects.first().resort_id,
            "filipino_count": 0,
            "foreigner_count": 2,
            "total_visitors": 2,
            "total_male": 1,
            "total_female": 1,
            "age_8_59": 2,
        }
        payload.update(overrides)
        return payload

    def test_mobile_foreign_record_without_region_is_not_given_quezon(self):
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            self._mobile_payload(country_id=self.united_states.id),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = TouristRecord.objects.get()
        self.assertEqual((record.region_id, record.province_id), (None, None))
        self.assertEqual(record.country_of_origin, "")

    def test_mobile_philippine_record_still_defaults_to_calabarzon_quezon(self):
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            self._mobile_payload(country_id=self.philippines.id),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = TouristRecord.objects.get()
        self.assertEqual((record.region_id, record.province_id), (self.calabarzon.id, self.quezon.id))
        self.assertEqual(record.country_of_origin, "Philippines")

    def test_mobile_payload_from_the_current_app_still_saves(self):
        # The installed APK always sends a region and province, even for a
        # foreign country. That must keep working; the values are kept.
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            self._mobile_payload(
                country_id=self.united_states.id,
                region_id=self.calabarzon.id,
                province_id=self.quezon.id,
                country_of_origin="Philippines",
            ),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = TouristRecord.objects.get()
        self.assertEqual((record.region_id, record.province_id), (self.calabarzon.id, self.quezon.id))

    # --- Excel import -------------------------------------------------------

    def _import(self, rows):
        import datetime
        from io import BytesIO
        from openpyxl import Workbook
        from .services.online_booking import COLUMNS, process_online_booking_workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Form Responses 1"
        for offset, row in enumerate(rows):
            values = {
                "arrival_date": datetime.datetime(2026, 4, 2),
                "resort": Resort.objects.first().resort_name,
                "total_visitors": 2,
                "total_male": 1,
                "total_female": 1,
                "itinerary": "Overnight",
                "maubanin_count": 0,
                "special_group_count": 0,
                "foreigner_count": 2,
                "filipino_count": 0,
                "age_0_7": 0,
                "age_8_59": 2,
                "age_60_above": 0,
                "travel_mode": "Private Vehicle",
                "boat_type": "Public Boat",
                "visit_purpose": "Leisure",
                "consent": "Yes",
                **row,
            }
            for key, value in values.items():
                sheet.cell(row=3 + offset, column=COLUMNS[key], value=value)
        file_obj = BytesIO()
        workbook.save(file_obj)
        file_obj.seek(0)
        return process_online_booking_workbook(file_obj, commit=True)

    def test_import_accepts_a_2027_arrival_and_creates_the_record(self):
        import datetime

        result = self._import(
            [{"full_name": "Advance Guest", "contact_number": "09170000004", "country": "Philippines",
              "arrival_date": datetime.datetime(2027, 9, 21)}]
        )

        self.assertEqual(result["valid_count"], 1, result["error_samples"])
        record = TouristRecord.objects.get(full_name="Advance Guest")
        self.assertEqual(record.arrival_date, datetime.date(2027, 9, 21))

    def test_import_window_moves_with_today(self):
        import datetime
        from unittest import mock

        this_year = timezone.localdate().year
        rows = [
            {"full_name": f"Guest {year}", "contact_number": f"0917000{year}", "country": "Philippines",
             "arrival_date": datetime.datetime(year, 6, 1)}
            for year in (this_year - 10, this_year + 2)
        ]
        self.assertEqual(self._import(rows)["valid_count"], 2)

        # The same years are outside the window once today is far enough away.
        later = datetime.date(this_year + 20, 1, 1)
        with mock.patch("api.services.online_booking.timezone.localdate", return_value=later):
            result = self._import(
                [{"full_name": "Old Guest", "contact_number": "09170009999", "country": "Philippines",
                  "arrival_date": datetime.datetime(this_year, 6, 1)}]
            )
        self.assertEqual(result["valid_count"], 0)

    def test_import_rejects_an_absurd_year_with_a_readable_message(self):
        import datetime

        this_year = timezone.localdate().year
        for year in (this_year - 11, this_year + 3, 1926, 2206):
            with self.subTest(year=year):
                result = self._import(
                    [{"full_name": "Typo Guest", "contact_number": "09170000005", "country": "Philippines",
                      "arrival_date": datetime.datetime(year, 6, 1)}]
                )
                self.assertEqual(result["valid_count"], 0)
                self.assertEqual(
                    result["error_samples"][0]["message"],
                    f"arrival_date: Must be an arrival date from {this_year - 10} to {this_year + 2}.",
                )
        self.assertFalse(TouristRecord.objects.filter(full_name="Typo Guest").exists())

    def test_monitoring_import_requires_the_year(self):
        from django.core.management.base import CommandError

        with self.assertRaisesMessage(CommandError, "--year"):
            call_command("import_tourism_excel", "MONITORING_FORM.xlsx")

    def test_import_no_longer_forces_quezon_onto_a_foreign_record(self):
        result = self._import(
            [
                {"full_name": "Foreign Guest", "contact_number": "09170000001", "country": "United States"},
                {"full_name": "New Country Guest", "contact_number": "09170000002", "country": "Canada"},
                {"full_name": "Local Guest", "contact_number": "09170000003", "country": "Philippines"},
            ]
        )

        self.assertEqual(result["valid_count"], 3, result["error_samples"])
        by_guest = {record.full_name: record for record in TouristRecord.objects.all()}
        for guest in ("Foreign Guest", "New Country Guest"):
            with self.subTest(guest=guest):
                record = by_guest[guest]
                self.assertEqual((record.region_id, record.province_id), (None, None))
        self.assertEqual(Country.objects.get(name="Canada").type, "foreign")
        local = by_guest["Local Guest"]
        self.assertEqual((local.region_id, local.province_id), (self.calabarzon.id, self.quezon.id))


class ClearSurv2026013OriginMigrationTests(TestCase):
    """0047 clears the Philippine region and province forced onto SURV-2026-013."""

    def setUp(self):
        import importlib

        ensure_test_reference_tables()
        self.migration = importlib.import_module("api.migrations.0047_clear_surv_2026_013_origin")
        self.western_visayas, _ = Region.objects.get_or_create(
            id=9, defaults={"name": "Region VI - Western Visayas"}
        )
        self.guimaras, _ = Province.objects.get_or_create(
            id=42, defaults={"name": "Guimaras", "region": self.western_visayas}
        )
        self.united_states = Country.objects.get(name="United States")
        self.target = self._record("SURV-2026-013", self.united_states)
        self.other = self._record("SURV-2026-099", self.united_states)

    def _record(self, survey_id, country):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Origin Group",
            contact_number="09170000000",
            country=country,
            region=self.western_visayas,
            province=self.guimaras,
            arrival_date="2026-10-27",
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=4,
            filipino_count=4,
            total_male=2,
            total_female=2,
            age_8_59=4,
            status="pending",
        )

    def _run(self, function):
        from django.apps import apps

        function(apps, None)

    def _origin(self, record):
        record.refresh_from_db()
        return (record.country_id, record.region_id, record.province_id, record.total_visitors)

    def test_forward_clears_only_that_record_and_keeps_its_country_and_counts(self):
        self._run(self.migration.clear_origin)

        self.assertEqual(self._origin(self.target), (self.united_states.id, None, None, 4))
        self.assertEqual(self._origin(self.other), (self.united_states.id, 9, 42, 4))

    def test_forward_does_nothing_once_the_values_have_changed(self):
        TouristRecord.objects.filter(pk="SURV-2026-013").update(province=Province.objects.get(id=1))

        self._run(self.migration.clear_origin)

        self.assertEqual(self._origin(self.target), (self.united_states.id, 9, 1, 4))

    def test_reverse_restores_the_exact_ids(self):
        self._run(self.migration.clear_origin)

        self._run(self.migration.restore_origin)

        self.assertEqual(self._origin(self.target), (self.united_states.id, 9, 42, 4))
        self.assertEqual(self._origin(self.other), (self.united_states.id, 9, 42, 4))


class AnalyticsPercentageTests(TestCase):
    """The Analytics answers divide by the year's visitor total.

    Three records of 6, 3 and 1 visitors (total 10): no single record's head
    count equals the total, so an answer that divides by one record instead
    of the total is wrong whichever record the loop ends on.
    """

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        quezon = Province.objects.get(name="Quezon")
        resorts = list(Resort.objects.order_by("resort_id")[:2])
        purposes = list(VisitPurpose.objects.order_by("id")[:2])
        cls.main_resort, cls.main_purpose = resorts[0], purposes[0]
        rows = [
            # survey_id, visitors, itinerary, resort, purpose, country, province
            ("PCT-0001", 6, "Overnight", resorts[0], purposes[0], "Philippines", quezon),
            ("PCT-0002", 3, "2 Nights", resorts[0], purposes[1], "Philippines", quezon),
            ("PCT-0003", 1, "Same Day", resorts[1], purposes[0], "United States", None),
        ]
        for survey_id, visitors, itinerary, resort, purpose, country, province in rows:
            TouristRecord.objects.create(
                survey_id=survey_id,
                full_name="Percentage Group",
                contact_number=f"0917000{survey_id[-4:]}",
                country=Country.objects.get(name=country),
                region=province.region if province else None,
                province=province,
                arrival_date="2026-04-02",
                resort=resort,
                itinerary=Itinerary.objects.get(name=itinerary),
                travel_mode=TravelMode.objects.first(),
                boat_type=BoatType.objects.first(),
                visit_purpose=purpose,
                total_visitors=visitors,
                filipino_count=visitors,
                total_male=visitors,
                age_8_59=visitors,
                status="arrived",
            )

    def setUp(self):
        from .services.tourism import build_tourism_question_answers

        self.answers = {
            item["id"]: item for item in build_tourism_question_answers({"year": "2026"})
        }

    def test_top_resort_percentage_is_of_the_year_total(self):
        self.assertEqual(
            self.answers["top_resort"]["answer"],
            f"{self.main_resort.resort_name} leads with 9 visitors, equal to 90.0% of the selected total.",
        )

    def test_top_origin_percentage_is_of_the_year_total(self):
        self.assertEqual(
            self.answers["top_origin"]["answer"],
            "Quezon leads with 9 visitors, equal to 90.0% of the selected total.",
        )

    def test_visit_purpose_percentage_is_of_the_year_total(self):
        self.assertEqual(
            self.answers["visit_purpose"]["answer"],
            f"{self.main_purpose.name} leads with 7 visitors, equal to 70.0% of the selected total.",
        )

    def test_average_stay_divides_nights_by_the_year_total(self):
        # 6 x 1 night + 3 x 2 nights + 1 x 0 (same day) = 12 nights / 10 visitors.
        self.assertEqual(self.answers["average_stay"]["visual"]["value"], 1.2)
        self.assertIn("1.2 night(s) per visitor", self.answers["average_stay"]["answer"])


class EntranceFeeRateTests(TestCase):
    """Every fee and revenue figure is head count x PHP 80 (no fee is stored).

    Two arrived records of 3 and 2 visitors on Apr 2, a no-show of 5 on Apr 2
    and a pending record of 4 on Apr 3: only arrived records are counted, so the
    Reports, Arrival Monitoring and the Dashboard all come to 5 visitors, PHP 400.
    """

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        quezon = Province.objects.get(name="Quezon")
        resort = Resort.objects.order_by("resort_id").first()
        rows = [
            ("FEE-0001", 3, "2026-04-02", "arrived"),
            ("FEE-0002", 2, "2026-04-02", "arrived"),
            ("FEE-0003", 5, "2026-04-02", "no_show"),
            ("FEE-0004", 4, "2026-04-03", "pending"),
        ]
        for survey_id, visitors, arrival_date, booking_status in rows:
            TouristRecord.objects.create(
                survey_id=survey_id,
                full_name="Fee Group",
                contact_number=f"0917000{survey_id[-4:]}",
                country=Country.objects.get(name="Philippines"),
                region=quezon.region,
                province=quezon,
                arrival_date=arrival_date,
                resort=resort,
                itinerary=Itinerary.objects.first(),
                travel_mode=TravelMode.objects.first(),
                boat_type=BoatType.objects.first(),
                visit_purpose=VisitPurpose.objects.first(),
                total_visitors=visitors,
                filipino_count=visitors,
                total_male=visitors,
                age_8_59=visitors,
                status=booking_status,
            )

    def test_the_rate_is_80_per_visitor(self):
        from .services.tourism import DISCOUNTED_ENTRANCE_FEE, REGULAR_ENTRANCE_FEE

        self.assertEqual((REGULAR_ENTRANCE_FEE, DISCOUNTED_ENTRANCE_FEE), (80, 64))

    def test_arrival_monitoring_fee_paid_and_fees_collected(self):
        from .services.tourism import build_arrival_monitoring_payload

        payload = build_arrival_monitoring_payload({"date": "2026-04-02"})

        self.assertEqual(
            sorted((row["survey_id"], row["feePaid"]) for row in payload["rows"]),
            [("FEE-0001", 240), ("FEE-0002", 160)],
        )
        self.assertEqual(payload["summary"]["feesCollected"], 400)
        self.assertEqual(payload["dailyTotals"]["feesCollected"], 400)
        self.assertEqual(payload["feePerVisitor"], 80)

    def test_dashboard_total_revenue_collected(self):
        from .services.tourism import build_dashboard_payload

        payload = build_dashboard_payload({"year": "2026"})

        self.assertEqual(payload["metrics"]["totalRevenueCollected"], 400)

    def test_daily_report_revenue_and_average_per_visitor(self):
        from .services.tourism import build_reports_payload

        payload = build_reports_payload(
            {"type": "daily", "year": "2026", "include_questions": "false"}
        )

        self.assertEqual(
            [(row["id"], row["visitors"], row["revenue"], row["avg"]) for row in payload["rows"]],
            [("2026-04-02", 5, 400, 80)],
        )
        self.assertEqual(payload["totals"], {"visitors": 5, "male": 5, "female": 0, "revenue": 400, "avg": 80})


class DiscountedEntranceFeeTests(TestCase):
    """fee = (visitors - discounted) x 80 + discounted x 64, per record and per group.

    Mixed records, so a group that used the wrong discounted count would be off:
      A arrived Apr 2: 4 visitors, 3 discounted, resort 1, purpose 1 -> 272
      B arrived Apr 2: 5 visitors, 0 discounted, resort 1, purpose 2 -> 400
      C arrived Apr 2: 2 visitors, 2 discounted, resort 2, purpose 1 -> 128
      D pending Apr 3: 3 visitors, 1 discounted, resort 2, purpose 2 (counted nowhere)
      E no-show Apr 2: 6 visitors, 6 discounted (counted nowhere)
    """

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        cls.quezon = Province.objects.get(name="Quezon")
        cls.resorts = list(Resort.objects.order_by("resort_id")[:2])
        cls.purposes = list(VisitPurpose.objects.order_by("id")[:2])
        rows = [
            ("DSC-0001", "arrived", "2026-04-02", 4, 3, 0, 0),
            ("DSC-0002", "arrived", "2026-04-02", 5, 0, 0, 1),
            ("DSC-0003", "arrived", "2026-04-02", 2, 2, 1, 0),
            ("DSC-0004", "pending", "2026-04-03", 3, 1, 1, 1),
            ("DSC-0005", "no_show", "2026-04-02", 6, 6, 0, 0),
        ]
        for survey_id, booking_status, arrival_date, visitors, discounted, resort, purpose in rows:
            cls._record(survey_id, booking_status, arrival_date, visitors, discounted, resort, purpose)

    @classmethod
    def _record(cls, survey_id, booking_status, arrival_date, visitors, discounted, resort=0, purpose=0):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Discount Group",
            contact_number=f"0917000{survey_id[-4:]}",
            country=Country.objects.get(name="Philippines"),
            region=cls.quezon.region,
            province=cls.quezon,
            arrival_date=arrival_date,
            resort=cls.resorts[resort],
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=cls.purposes[purpose],
            total_visitors=visitors,
            filipino_count=visitors,
            total_male=visitors,
            age_8_59=visitors,
            discounted_count=discounted,
            status=booking_status,
        )

    # --- the helper ---------------------------------------------------------

    def test_fee_helper(self):
        from .services.tourism import entrance_fee

        self.assertEqual(entrance_fee(5, 0), 400)  # all regular
        self.assertEqual(entrance_fee(3, 3), 192)  # all discounted
        self.assertEqual(entrance_fee(4, 3), 272)  # mixed
        self.assertEqual(entrance_fee(0, 0), 0)  # nobody

    # --- validation ---------------------------------------------------------

    def _web_payload(self, **overrides):
        payload = {
            "first_name": "Ana",
            "last_name": "Cruz",
            "full_name": "Ana Cruz",
            "email": "ana@example.com",
            "contact_number": "+639171234567",
            "country_id": Country.objects.get(name="Philippines").id,
            "region_id": self.quezon.region_id,
            "province_id": self.quezon.id,
            "resort_id": self.resorts[0].resort_id,
            "itinerary_id": Itinerary.objects.first().id,
            "travel_mode_id": TravelMode.objects.first().id,
            "boat_type_id": BoatType.objects.first().id,
            "visit_purpose_id": self.purposes[0].id,
            "arrival_date": "2026-05-01",
            "filipino_count": 3,
            "foreigner_count": 0,
            "total_visitors": 3,
            "total_male": 3,
            "total_female": 0,
            "age_8_59": 3,
            "status": "pending",
        }
        payload.update(overrides)
        return payload

    def _client(self):
        user = User.objects.create_user(username="fee_admin", password="Fee@12345")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        return client

    def test_discounted_count_above_total_visitors_is_rejected(self):
        response = self._client().post(
            "/api/tourist-records/", self._web_payload(discounted_count=4), format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            response.json()["discounted_count"],
            ["Discounted count cannot be greater than the total number of visitors."],
        )

    def test_discounted_count_equal_to_total_visitors_is_allowed(self):
        response = self._client().post(
            "/api/tourist-records/", self._web_payload(discounted_count=3), format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.json())
        self.assertEqual(response.json()["discounted_count"], 3)

    # --- totals -------------------------------------------------------------

    def test_dashboard_total_revenue_collected(self):
        from .services.tourism import build_dashboard_payload

        # A + B + C: 11 visitors, 5 discounted -> 6 x 80 + 5 x 64.
        self.assertEqual(
            build_dashboard_payload({"year": "2026"})["metrics"]["totalRevenueCollected"], 800
        )

    def test_arrival_monitoring_fee_paid_and_fees_collected(self):
        from .services.tourism import build_arrival_monitoring_payload

        payload = build_arrival_monitoring_payload({"date": "2026-04-02"})

        self.assertEqual(
            sorted((row["survey_id"], row["feePaid"]) for row in payload["rows"]),
            [("DSC-0001", 272), ("DSC-0002", 400), ("DSC-0003", 128)],
        )
        self.assertEqual(payload["summary"]["feesCollected"], 800)

    def _report(self, report_type):
        from .services.tourism import build_reports_payload

        return build_reports_payload(
            {"type": report_type, "year": "2026", "include_questions": "false"}
        )

    def test_daily_report_uses_each_days_discounted_count(self):
        payload = self._report("daily")

        self.assertEqual(
            [(row["id"], row["visitors"], row["revenue"], row["avg"]) for row in payload["rows"]],
            [("2026-04-02", 11, 800, 73)],
        )
        self.assertEqual(payload["totals"], {"visitors": 11, "male": 11, "female": 0, "revenue": 800, "avg": 73})

    def test_monthly_report_uses_the_months_discounted_count(self):
        payload = self._report("monthly")

        # April (A + B + C): 11 visitors, 5 discounted -> 6 x 80 + 5 x 64.
        self.assertEqual(
            [(row["id"], row["visitors"], row["revenue"]) for row in payload["rows"]],
            [("2026-04", 11, 800)],
        )

    def test_purpose_report_uses_each_purposes_discounted_count(self):
        payload = self._report("purpose")

        # Purpose 1 (A + C): 6 visitors, 5 discounted. Purpose 2 (B): 5 visitors, none.
        self.assertEqual(
            [(row["name"], row["visitors"], row["revenue"], row["avg"]) for row in payload["rows"]],
            [(self.purposes[0].name, 6, 400, 67), (self.purposes[1].name, 5, 400, 80)],
        )

    def test_resort_report_uses_each_resorts_discounted_count(self):
        payload = self._report("resort")

        # Resort 1 (A + B): 9 visitors, 3 discounted. Resort 2 (C): 2 visitors, 2 discounted.
        self.assertEqual(
            [(row["resort_id"], row["visitors"], row["revenue"], row["avg"]) for row in payload["rows"]],
            [(self.resorts[0].resort_id, 9, 672, 75), (self.resorts[1].resort_id, 2, 128, 64)],
        )
        self.assertEqual(payload["totals"]["revenue"], 800)

    # --- other writers ------------------------------------------------------

    def test_mobile_payload_without_discounted_count_saves_zero(self):
        # Today's APK does not know the field.
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            {
                "first_name": "Ana",
                "last_name": "Cruz",
                "contact_number": "+639171234599",
                "arrival_date": "2026-05-01",
                "resort_id": self.resorts[0].resort_id,
                "filipino_count": 3,
                "foreigner_count": 0,
                "total_visitors": 3,
                "total_male": 2,
                "total_female": 1,
                "age_0_7": 1,
                "age_8_59": 2,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(TouristRecord.objects.get(contact_number__endswith="1234599").discounted_count, 0)

    def test_excel_import_saves_zero(self):
        ForeignOriginTests._import(
            self,
            [
                {
                    "full_name": "Imported Guest",
                    "contact_number": "09170009999",
                    "country": "Philippines",
                    "age_0_7": 1,
                    "age_8_59": 1,
                }
            ],
        )

        self.assertEqual(TouristRecord.objects.get(full_name="Imported Guest").discounted_count, 0)

    def test_migration_backfills_children_plus_seniors(self):
        import importlib
        from django.apps import apps

        migration = importlib.import_module("api.migrations.0048_touristrecord_discounted_count")
        record = self._record("DSC-0006", "arrived", "2026-04-05", 5, 0)
        TouristRecord.objects.filter(pk=record.pk).update(age_0_7=2, age_8_59=2, age_60_above=1)

        migration.backfill_discounted_count(apps, None)

        record.refresh_from_db()
        self.assertEqual(record.discounted_count, 3)


class DiscountedCountSpecialNeedsMigrationTests(TestCase):
    """0049 moves records still on the age-only suggestion to the new one."""

    def setUp(self):
        import importlib

        ensure_test_reference_tables()
        self.migration = importlib.import_module(
            "api.migrations.0049_discounted_count_includes_special_needs"
        )

    def _record(self, survey_id, visitors, age_0_7, age_60_above, special, discounted):
        quezon = Province.objects.get(name="Quezon")
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Special Needs Group",
            contact_number=f"0917000{survey_id[-4:]}",
            country=Country.objects.get(name="Philippines"),
            region=quezon.region,
            province=quezon,
            arrival_date="2026-04-02",
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=visitors,
            filipino_count=visitors,
            total_male=visitors,
            age_0_7=age_0_7,
            age_60_above=age_60_above,
            age_8_59=visitors - age_0_7 - age_60_above,
            special_group_count=special,
            discounted_count=discounted,
            status="arrived",
        )

    def _run(self, function):
        from django.apps import apps

        function(apps, None)

    def _discounted(self):
        return dict(TouristRecord.objects.values_list("survey_id", "discounted_count"))

    def test_forward_adds_special_needs_capped_and_keeps_manual_counts(self):
        # Age-only 2, plus 1 special needs -> 3.
        self._record("SPN-0001", 5, 1, 1, 1, 2)
        # 3 visitors, 2 seniors who are also special needs: 2 + 2 = 4, capped at 3.
        self._record("SPN-0002", 3, 0, 2, 2, 2)
        # Staff set 4 by hand (age-only would be 1): left alone.
        self._record("SPN-0003", 5, 1, 0, 1, 4)
        # Nothing to add: unchanged.
        self._record("SPN-0004", 2, 0, 0, 0, 0)

        self._run(self.migration.include_special_needs)

        self.assertEqual(
            self._discounted(),
            {"SPN-0001": 3, "SPN-0002": 3, "SPN-0003": 4, "SPN-0004": 0},
        )

    def test_reverse_returns_to_the_age_only_value(self):
        self._record("SPN-0001", 5, 1, 1, 1, 2)
        self._record("SPN-0002", 3, 0, 2, 2, 2)
        self._record("SPN-0003", 5, 1, 0, 1, 4)
        self._run(self.migration.include_special_needs)

        self._run(self.migration.back_to_age_only)

        self.assertEqual(self._discounted(), {"SPN-0001": 2, "SPN-0002": 2, "SPN-0003": 4})


class ArrivalMonitoringViewsTests(TestCase):
    """Month ranges, the capped table, and the complete date-ordered export.

    305 arrived records spread over September 2026 (one visitor each), two
    arrived in October, and a pending September record that is never counted.
    """

    SEPTEMBER = {"year": "2026", "from": "2026-09-01", "to": "2026-09-30"}

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        quezon = Province.objects.get(name="Quezon")
        common = dict(
            full_name="Arrival Group",
            contact_number="09170000000",
            country=Country.objects.get(name="Philippines"),
            region=quezon.region,
            province=quezon,
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=1,
            filipino_count=1,
            total_male=1,
            age_8_59=1,
        )
        records = [
            TouristRecord(
                survey_id=f"ARR-{index:04d}",
                arrival_date=f"2026-09-{(index % 30) + 1:02d}",
                status="arrived",
                **common,
            )
            for index in range(305)
        ]
        records += [
            TouristRecord(survey_id="ARR-OCT-1", arrival_date="2026-10-01", status="arrived", **common),
            TouristRecord(survey_id="ARR-OCT-2", arrival_date="2026-10-31", status="arrived", **common),
            TouristRecord(survey_id="ARR-PEND", arrival_date="2026-09-15", status="pending", **common),
        ]
        TouristRecord.objects.bulk_create(records)

    def test_table_is_capped_but_counts_every_record(self):
        from .services.tourism import build_arrival_monitoring_payload

        payload = build_arrival_monitoring_payload(self.SEPTEMBER)

        self.assertEqual(len(payload["rows"]), 300)
        self.assertEqual(payload["rowCount"], 305)
        self.assertEqual(payload["summary"]["totalArrivals"], 305)

    def test_export_returns_every_row_in_date_order(self):
        from .services.tourism import build_arrival_monitoring_export

        payload = build_arrival_monitoring_export(self.SEPTEMBER)
        dates = [row["date"] for row in payload["rows"]]

        self.assertEqual(payload["rowCount"], 305)
        self.assertEqual(len(payload["rows"]), 305)
        self.assertEqual(dates, sorted(dates))
        self.assertEqual((dates[0], dates[-1]), ("2026-09-01", "2026-09-30"))
        self.assertEqual(
            {row["survey_id"] for row in payload["rows"]},
            {f"ARR-{index:04d}" for index in range(305)},
        )

    def test_month_range_selects_only_that_month(self):
        from .services.tourism import build_arrival_monitoring_payload

        october = build_arrival_monitoring_payload(
            {"year": "2026", "from": "2026-10-01", "to": "2026-10-31"}
        )

        self.assertEqual(october["rowCount"], 2)
        self.assertEqual(sorted(row["survey_id"] for row in october["rows"]), ["ARR-OCT-1", "ARR-OCT-2"])

    def test_a_range_with_a_different_year_returns_nothing(self):
        from .services.tourism import build_arrival_monitoring_payload

        payload = build_arrival_monitoring_payload({**self.SEPTEMBER, "year": "2025"})

        self.assertEqual(payload["rowCount"], 0)

    def test_export_endpoint(self):
        user = User.objects.create_user(username="arrival_admin", password="Arrival@123")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

        response = client.get("/api/arrival-monitoring/export/", self.SEPTEMBER)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["rowCount"], 305)
        self.assertEqual(len(response.json()["rows"]), 305)
        self.assertEqual(APIClient().get("/api/arrival-monitoring/export/").status_code, 401)


class NoShowSweepTests(TestCase):
    """The no-show sweep runs at most once a day, plus after any record write."""

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        cls.quezon = Province.objects.get(name="Quezon")

    def setUp(self):
        from django.core.cache import cache

        cache.clear()

    def _record(self, survey_id, arrival_date, booking_status="pending"):
        return TouristRecord(
            survey_id=survey_id,
            full_name="Sweep Group",
            contact_number="09170000000",
            country=Country.objects.get(name="Philippines"),
            region=self.quezon.region,
            province=self.quezon,
            arrival_date=arrival_date,
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=1,
            filipino_count=1,
            total_male=1,
            age_8_59=1,
            status=booking_status,
        )

    def _insert_without_signal(self, *records):
        # bulk_create sends no post_save, so it does not ask for a new sweep.
        TouristRecord.objects.bulk_create(records)

    def _status(self, survey_id):
        return TouristRecord.objects.get(pk=survey_id).status

    def _days(self, offset):
        return timezone.localdate() + timedelta(days=offset)

    def test_past_dated_pending_booking_becomes_no_show(self):
        from .services.no_show import sweep_no_shows

        self._insert_without_signal(
            self._record("SWP-PAST", self._days(-1)),
            self._record("SWP-TODAY", self._days(0)),
            self._record("SWP-FUTURE", self._days(3)),
            self._record("SWP-ARRIVED", self._days(-1), "arrived"),
        )

        sweep_no_shows()

        self.assertEqual(self._status("SWP-PAST"), "no_show")
        self.assertEqual(self._status("SWP-TODAY"), "pending")
        self.assertEqual(self._status("SWP-FUTURE"), "pending")
        self.assertEqual(self._status("SWP-ARRIVED"), "arrived")

    def test_first_request_of_the_day_sweeps_and_a_second_does_not(self):
        from .services.no_show import sweep_no_shows

        self._insert_without_signal(self._record("SWP-1", self._days(-1)))
        sweep_no_shows()
        self.assertEqual(self._status("SWP-1"), "no_show")

        self._insert_without_signal(self._record("SWP-2", self._days(-1)))
        with CaptureQueriesContext(connection) as queries:
            sweep_no_shows()

        self.assertEqual(len(queries), 0)
        self.assertEqual(self._status("SWP-2"), "pending")

    def test_a_new_day_sweeps_again(self):
        from .services import no_show

        no_show.sweep_no_shows()
        self._insert_without_signal(self._record("SWP-1", self._days(0)))

        tomorrow = self._days(1)
        tomorrow_noon = timezone.localtime().replace(
            year=tomorrow.year, month=tomorrow.month, day=tomorrow.day, hour=12
        )
        with patch.object(no_show.timezone, "localdate", return_value=tomorrow), patch.object(
            no_show.timezone, "localtime", return_value=tomorrow_noon
        ):
            no_show.sweep_no_shows()

        self.assertEqual(self._status("SWP-1"), "no_show")

    def test_saving_a_record_today_makes_the_next_request_sweep(self):
        from .services.no_show import sweep_no_shows

        sweep_no_shows()
        self._record("SWP-SAVED", self._days(-2)).save()

        sweep_no_shows()

        self.assertEqual(self._status("SWP-SAVED"), "no_show")

    def test_an_import_makes_the_next_request_sweep(self):
        from .services.no_show import sweep_no_shows

        sweep_no_shows()
        ForeignOriginTests._import(
            self,
            [{"full_name": "Imported Past Guest", "contact_number": "09170001111", "country": "Philippines"}],
        )

        sweep_no_shows()

        self.assertEqual(TouristRecord.objects.get(full_name="Imported Past Guest").status, "no_show")

    def test_a_cache_failure_sweeps_instead_of_skipping(self):
        from .services import no_show

        self._insert_without_signal(self._record("SWP-1", self._days(-1)))

        with patch.object(no_show.cache, "get", side_effect=RuntimeError("cache down")), patch.object(
            no_show.cache, "set", side_effect=RuntimeError("cache down")
        ):
            no_show.sweep_no_shows()

        self.assertEqual(self._status("SWP-1"), "no_show")

    def test_the_marker_expires_at_midnight(self):
        from .services import no_show

        late = timezone.localtime().replace(hour=23, minute=59, second=30, microsecond=0)
        with patch.object(no_show.timezone, "localtime", return_value=late):
            self.assertEqual(no_show._seconds_until_midnight(), 30)

    def test_a_second_get_of_a_page_does_not_write(self):
        user = User.objects.create_user(username="sweep_admin", password="Sweep@123")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

        client.get("/api/dashboard/")
        with CaptureQueriesContext(connection) as queries:
            response = client.get("/api/dashboard/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(
            [q["sql"] for q in queries if q["sql"].lstrip().upper().startswith("UPDATE")]
        )


class QuestionAnswersDeduplicationTests(TestCase):
    """The de-duplicated question answers equal the old, separate computations.

    The oracle below is the code that was replaced: a separate query for each
    leader and its top 6, and separate text and chart functions for the demand
    and validation answers.
    """

    AFFECTED = ("top_resort", "top_origin", "visit_purpose", "high_demand", "validation")

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        cls.quezon = Province.objects.get(name="Quezon")
        cls.resorts = list(Resort.objects.order_by("resort_id")[:2])
        cls.purposes = list(VisitPurpose.objects.order_by("id")[:2])
        cls.us = Country.objects.get(name="United States")

    def _record(self, survey_id, day, visitors, resort=0, purpose=0, booking_status="arrived",
                foreign=False, full_name="Group", contact="09170000000"):
        return TouristRecord(
            survey_id=survey_id,
            full_name=full_name,
            contact_number=contact,
            country=self.us if foreign else Country.objects.get(name="Philippines"),
            region=None if foreign else self.quezon.region,
            province=None if foreign else self.quezon,
            arrival_date=day,
            resort=self.resorts[resort],
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=self.purposes[purpose],
            total_visitors=visitors,
            filipino_count=visitors,
            total_male=visitors,
            age_8_59=visitors,
            status=booking_status,
        )

    # --- the oracle: the replaced code --------------------------------------

    def _old(self, params):
        from .services import tourism as t

        reporting_year = t.get_reporting_year(params)
        # Arrived bookings only (pending is not counted, since 2026-10-11).
        arrived = t.apply_reporting_year(
            TouristRecord.objects.filter(status="arrived"), reporting_year
        )
        records = t.apply_reporting_year(TouristRecord.objects.all(), reporting_year)
        if params.get("from"):
            arrived = arrived.filter(arrival_date__gte=params["from"])
            records = records.filter(arrival_date__gte=params["from"])
        if params.get("to"):
            arrived = arrived.filter(arrival_date__lte=params["to"])
            records = records.filter(arrival_date__lte=params["to"])
        total = t.sum_visitors(arrived)

        def ranking(field, queryset):
            top = t.top_group(queryset, field)
            return (
                t.format_top_answer(top, total, "visitors"),
                [{"label": r["name"], "value": r["total"]} for r in t.top_n_groups(queryset, field, limit=6)],
            )

        def demand_signal():
            latest = arrived.aggregate(latest=Max("arrival_date"))["latest"]
            if not latest:
                return "No recent arrival trend is available yet."
            recent_start = latest - timedelta(days=29)
            previous = arrived.filter(arrival_date__range=(recent_start - timedelta(days=30), recent_start - timedelta(days=1)))
            recent_top = t.top_group(arrived.filter(arrival_date__range=(recent_start, latest)), "resort__resort_name")
            growth = recent_top["total"] - t.sum_visitors(previous.filter(resort__resort_name=recent_top["name"]))
            if not recent_top["name"]:
                return "No recent high-demand resort is available yet."
            return (
                f"{recent_top['name']} shows the strongest recent demand with "
                f"{recent_top['total']} visitors in the latest 30-day window "
                f"({growth:+d} versus the previous 30 days)."
            )

        def demand_visual():
            latest = arrived.aggregate(latest=Max("arrival_date"))["latest"]
            if not latest:
                return {"type": "comparison", "items": []}
            recent_start = latest - timedelta(days=29)
            recent_top = t.top_group(arrived.filter(arrival_date__range=(recent_start, latest)), "resort__resort_name")
            previous_total = t.sum_visitors(arrived.filter(
                arrival_date__range=(recent_start - timedelta(days=30), recent_start - timedelta(days=1)),
                resort__resort_name=recent_top["name"],
            ))
            return {"type": "comparison", "items": [
                {"label": "Previous 30 days", "value": previous_total},
                {"label": "Latest 30 days", "value": recent_top["total"]},
            ]}

        def counts():
            return (
                records.filter(status="pending").count(),
                records.filter(status="no_show").count(),
                records.exclude(contact_number="").values("full_name", "contact_number", "arrival_date", "resort_id")
                .annotate(total=Count("survey_id")).filter(total__gt=1).count(),
                records.filter(Q(full_name="") | Q(contact_number="") | Q(total_visitors=0)
                               | Q(total_male__isnull=True) | Q(total_female__isnull=True)).count(),
            )

        pending, no_show, duplicates, incomplete = counts()
        resort_answer, resort_items = ranking("resort__resort_name", arrived)
        origin_answer, origin_items = ranking("origin", t.with_origin(arrived))
        purpose_answer, purpose_items = ranking("visit_purpose__name", arrived)
        return {
            "top_resort": (resort_answer, {"type": "ranking", "items": resort_items}),
            "top_origin": (origin_answer, {"type": "ranking", "items": origin_items}),
            "visit_purpose": (purpose_answer, {"type": "ranking", "items": purpose_items}),
            "high_demand": (demand_signal(), demand_visual()),
            "validation": (
                f"Needs review: {pending} pending, {no_show} no-show, "
                f"{duplicates} possible duplicates, and {incomplete} incomplete records.",
                {"type": "stack", "items": [
                    {"label": "Pending", "value": pending},
                    {"label": "No-show", "value": no_show},
                    {"label": "Duplicates", "value": duplicates},
                    {"label": "Incomplete", "value": incomplete},
                ]},
            ),
        }

    def _new(self, params):
        from .services.tourism import build_tourism_question_answers

        answers = {item["id"]: item for item in build_tourism_question_answers(params)}
        return {key: (answers[key]["answer"], answers[key]["visual"]) for key in self.AFFECTED}

    def _assert_same(self, params):
        self.assertEqual(self._new(params), self._old(params))

    # --- cases ----------------------------------------------------------------

    def test_same_with_no_records(self):
        self._assert_same({"year": "2026"})

    def test_same_with_mixed_records(self):
        TouristRecord.objects.bulk_create([
            self._record("QA-01", "2026-09-10", 5, resort=0, purpose=0),
            self._record("QA-02", "2026-09-20", 3, resort=1, purpose=1),
            self._record("QA-03", "2026-08-05", 2, resort=0, purpose=0, foreign=True),
            self._record("QA-04", "2026-09-25", 4, resort=1, purpose=1, booking_status="pending"),
            self._record("QA-05", "2026-09-15", 6, resort=0, booking_status="no_show"),
            self._record("QA-06", "2026-09-10", 1, resort=0, purpose=1),  # duplicate of QA-01's key
            self._record("QA-07", "2026-09-01", 2, resort=1, contact=""),  # incomplete
            self._record("QA-08", "2026-07-30", 3, resort=1, purpose=0),  # previous 30-day window
        ])

        for params in (
            {"year": "2026"},
            {"year": "all"},
            {"year": "2026", "from": "2026-09-01", "to": "2026-09-30"},
            {"year": "2026", "from": "2026-12-01", "to": "2026-12-31"},
        ):
            with self.subTest(params=params):
                self._assert_same(params)

    def test_same_with_a_tie(self):
        TouristRecord.objects.bulk_create([
            self._record("QA-11", "2026-09-10", 4, resort=0, purpose=0),
            self._record("QA-12", "2026-09-11", 4, resort=1, purpose=1),
        ])

        self._assert_same({"year": "2026"})

    def test_question_answers_run_17_queries(self):
        from .services.tourism import build_tourism_question_answers

        TouristRecord.objects.bulk_create([self._record("QA-21", "2026-09-10", 2)])
        with CaptureQueriesContext(connection) as queries:
            build_tourism_question_answers({"year": "2026"})

        self.assertEqual(len(queries), 17)


class ReportsMaleFemaleTests(TestCase):
    """Every report tab carries male and female, summed in its existing query."""

    TABS = ("daily", "monthly", "yearly", "resort", "origin", "purpose", "transport", "boat", "no_show")
    # Queries per tab without the question answers; the male and female sums
    # must not add one.
    QUERIES = {
        "daily": 1, "monthly": 1, "yearly": 1, "resort": 2,
        "origin": 1, "purpose": 1, "transport": 1, "boat": 1, "no_show": 1,
    }

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        cls.quezon = Province.objects.get(name="Quezon")
        cls.us = Country.objects.get(name="United States")
        cls.resorts = list(Resort.objects.order_by("resort_id")[:2])
        cls.purposes = list(VisitPurpose.objects.order_by("id")[:2])
        cls.modes = list(TravelMode.objects.order_by("id")[:2])
        cls.boats = list(BoatType.objects.order_by("id")[:2])
        TouristRecord.objects.bulk_create([
            cls._record("MF-01", "2026-08-10", 3, 2, resort=0, purpose=0, mode=0),
            cls._record("MF-02", "2026-09-05", 1, 3, resort=1, purpose=1, mode=1, foreign=True, boat=1),
            # MF-03 shares MF-01's travel mode but not its boat, so the two groupings differ.
            cls._record("MF-03", "2026-09-05", 2, 0, resort=0, purpose=1, mode=0, boat=1),
            cls._record("MF-04", "2026-09-20", 0, 4, resort=1, purpose=0, mode=1, booking_status="no_show"),
            # A pending advance booking in a second year, shaped like SURV-2026-011:
            # not approved, so counted on no tab.
            cls._record("MF-05", "2027-09-21", 2, 2, resort=0, purpose=0, mode=0,
                        booking_status="pending", discounted=1),
            # A pending booking in 2026, large enough to change every tab if counted.
            cls._record("MF-06", "2026-08-10", 5, 5, resort=1, purpose=0, mode=1, booking_status="pending"),
        ])

    @classmethod
    def _record(cls, survey_id, day, male, female, resort, purpose, mode,
                booking_status="arrived", foreign=False, discounted=0, boat=0):
        visitors = male + female
        return TouristRecord(
            survey_id=survey_id,
            full_name="Group",
            contact_number="09170000000",
            country=cls.us if foreign else Country.objects.get(name="Philippines"),
            region=None if foreign else cls.quezon.region,
            province=None if foreign else cls.quezon,
            arrival_date=day,
            resort=cls.resorts[resort],
            itinerary=Itinerary.objects.first(),
            travel_mode=cls.modes[mode],
            boat_type=cls.boats[boat],
            visit_purpose=cls.purposes[purpose],
            total_visitors=visitors,
            foreigner_count=visitors if foreign else 0,
            filipino_count=0 if foreign else visitors,
            total_male=male,
            total_female=female,
            discounted_count=discounted,
            age_8_59=visitors,
            status=booking_status,
        )

    def _report(self, report_type, year="2026"):
        from .services.tourism import build_reports_payload

        return build_reports_payload({"type": report_type, "year": year, "include_questions": "false"})

    def test_each_tab_reports_male_and_female(self):
        r0, r1 = (resort.resort_name for resort in self.resorts)
        p0, p1 = (purpose.name for purpose in self.purposes)
        m0, m1 = (mode.name for mode in self.modes)
        # (name, male, female, visitors) per row, then the totals. No-show
        # counts only MF-04; every other tab counts the arrived MF-01 to MF-03,
        # never the pending MF-05 and MF-06.
        expected = {
            "daily": ([("Aug 10, 2026", 3, 2, 5), ("Sep 05, 2026", 3, 3, 6)], (6, 5, 11)),
            "monthly": ([("August 2026", 3, 2, 5), ("September 2026", 3, 3, 6)], (6, 5, 11)),
            "yearly": ([("2026", 6, 5, 11)], (6, 5, 11)),
            "resort": ([(r0, 5, 2, 7), (r1, 1, 3, 4)], (6, 5, 11)),
            "origin": ([("Quezon", 5, 2, 7), ("United States", 1, 3, 4)], (6, 5, 11)),
            "purpose": ([(p1, 3, 3, 6), (p0, 3, 2, 5)], (6, 5, 11)),
            "transport": ([(m0, 5, 2, 7), (m1, 1, 3, 4)], (6, 5, 11)),
            "boat": ([(self.boats[1].name, 3, 3, 6), (self.boats[0].name, 3, 2, 5)], (6, 5, 11)),
            "no_show": ([(r1, 0, 4, 4)], (0, 4, 4)),
        }
        for report_type, (rows, totals) in expected.items():
            with self.subTest(report_type=report_type):
                payload = self._report(report_type)
                self.assertEqual(
                    [(row["name"], row["male"], row["female"], row["visitors"]) for row in payload["rows"]],
                    rows,
                )
                self.assertEqual(
                    (payload["totals"]["male"], payload["totals"]["female"], payload["totals"]["visitors"]),
                    totals,
                )

    def test_male_plus_female_equals_visitors_on_every_row_and_total(self):
        for report_type in self.TABS:
            with self.subTest(report_type=report_type):
                payload = self._report(report_type)
                self.assertTrue(payload["rows"])
                for row in payload["rows"]:
                    self.assertEqual(row["male"] + row["female"], row["visitors"], row)
                totals = payload["totals"]
                self.assertEqual(totals["male"] + totals["female"], totals["visitors"])
                self.assertEqual(totals["male"], sum(row["male"] for row in payload["rows"]))
                self.assertEqual(totals["female"], sum(row["female"] for row in payload["rows"]))

    def test_yearly_tab_gives_one_row_per_year(self):
        payload = self._report("yearly", year="all")

        # 2026: MF-01 to MF-03, 11 visitors, none discounted -> 11 x 80.
        # 2027: only the pending MF-05, so no row.
        self.assertEqual(
            [(row["id"], row["name"], row["male"], row["female"], row["visitors"], row["revenue"])
             for row in payload["rows"]],
            [(2026, "2026", 6, 5, 11, 880)],
        )
        totals = payload["totals"]
        self.assertEqual(
            (totals["male"], totals["female"], totals["visitors"], totals["revenue"]), (6, 5, 11, 880)
        )
        self.assertEqual(totals["male"] + totals["female"], totals["visitors"])

    def test_yearly_tab_respects_the_year_filter(self):
        for year, expected in (
            ("2026", [(2026, 11)]), ("2027", []), ("2025", []), ("all", [(2026, 11)]),
        ):
            with self.subTest(year=year):
                payload = self._report("yearly", year=year)
                self.assertEqual([(row["id"], row["visitors"]) for row in payload["rows"]], expected)

    def test_2027_has_only_a_pending_booking_so_every_tab_is_empty(self):
        expected = {report_type: [] for report_type in self.TABS}
        for report_type, rows in expected.items():
            with self.subTest(report_type=report_type):
                payload = self._report(report_type, year="2027")
                self.assertEqual(
                    [(r["name"], r["male"], r["female"], r["visitors"], r["revenue"]) for r in payload["rows"]],
                    rows,
                )
                self.assertEqual(payload["filters"]["year"], "2027")

    def test_boat_tab_gives_one_row_per_boat_type_with_its_fee(self):
        payload = self._report("boat")

        # Boat 1: MF-02 + MF-03, 6 visitors. Boat 0: MF-01, 5. None discounted.
        self.assertEqual(
            [(r["name"], r["male"], r["female"], r["visitors"], r["revenue"]) for r in payload["rows"]],
            [(self.boats[1].name, 3, 3, 6, 480), (self.boats[0].name, 3, 2, 5, 400)],
        )
        self.assertEqual(payload["type"], "boat")

    def test_boat_tab_total_equals_the_daily_total(self):
        for year in ("2026", "2027", "all"):
            with self.subTest(year=year):
                boat = self._report("boat", year=year)["totals"]
                daily = self._report("daily", year=year)["totals"]
                self.assertEqual(boat, daily)

    def test_vehicle_tab_is_unchanged_by_the_boat_tab(self):
        m0, m1 = (mode.name for mode in self.modes)

        payload = self._report("transport")

        self.assertEqual(payload["type"], "transport")
        self.assertEqual(
            [(r["id"], r["name"], r["male"], r["female"], r["visitors"], r["revenue"]) for r in payload["rows"]],
            [(m0, m0, 5, 2, 7, 560), (m1, m1, 1, 3, 4, 320)],
        )

    def test_query_count_per_tab(self):
        for report_type in self.TABS:
            with self.subTest(report_type=report_type):
                with CaptureQueriesContext(connection) as queries:
                    self._report(report_type)
                self.assertEqual(len(queries), self.QUERIES[report_type])


class ReportingYearTests(TestCase):
    """The year filters come from the data; a bad year or report type is a 400."""

    REPORT_ENDPOINTS = (
        "/api/reports/",
        "/api/dashboard/",
        "/api/arrival-monitoring/",
        "/api/arrival-monitoring/export/",
        "/api/booking-management/",
    )

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        cls.quezon = Province.objects.get(name="Quezon")

    def _record(self, survey_id, day, visitors=2):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Group",
            contact_number="09170000000",
            country=Country.objects.get(name="Philippines"),
            region=self.quezon.region,
            province=self.quezon,
            arrival_date=day,
            resort=Resort.objects.order_by("resort_id").first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=visitors,
            filipino_count=visitors,
            total_male=visitors,
            age_8_59=visitors,
            status="arrived",
        )

    def _client(self):
        user = User.objects.create_user(username="years_admin", password="Years@12345")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        return client

    def test_no_records_still_lists_the_current_year(self):
        from .services.tourism import current_reporting_year, get_reporting_years

        self.assertEqual(get_reporting_years(), [current_reporting_year()])

    def test_years_come_from_the_data_newest_first(self):
        from .services.tourism import current_reporting_year, get_reporting_years

        self._record("RY-01", "2027-09-21")
        self._record("RY-02", "2024-03-01")
        self._record("RY-03", "2024-11-30")

        expected = sorted({"2027", "2024", current_reporting_year()}, reverse=True)
        with CaptureQueriesContext(connection) as queries:
            self.assertEqual(get_reporting_years(), expected)
        self.assertEqual(len(queries), 1)

    def test_bootstrap_sends_the_reporting_years(self):
        from .services.tourism import get_reporting_years

        self._record("RY-01", "2027-09-21")
        response = self._client().get("/api/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["reportingYears"], get_reporting_years())
        self.assertIn("2027", response.json()["reportingYears"])

    def test_a_missing_year_is_the_current_year(self):
        from .services.tourism import current_reporting_year, get_reporting_year

        for params in ({}, {"year": ""}, {"year": "  "}):
            with self.subTest(params=params):
                self.assertEqual(get_reporting_year(params), current_reporting_year())

    def test_a_junk_year_is_rejected_on_every_report_endpoint(self):
        self._record("RY-01", "2026-09-21")
        client = self._client()
        for year in ("2027x", "abcd", "20", "02026", "２０２６", "-2026", "all-years"):
            for endpoint in self.REPORT_ENDPOINTS:
                with self.subTest(year=year, endpoint=endpoint):
                    response = client.get(endpoint, {"year": year, "date": "all"})
                    self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                    self.assertIn("Unknown year", response.json()["detail"])
                    self.assertNotIn("rows", response.json())

    def test_every_report_tab_rejects_a_junk_year(self):
        from .services.tourism import REPORT_TYPES, ReportParameterError, build_reports_payload

        for report_type in REPORT_TYPES:
            with self.subTest(report_type=report_type):
                with self.assertRaises(ReportParameterError):
                    build_reports_payload({"type": report_type, "year": "2027x", "include_questions": "false"})

    def test_an_unknown_report_type_is_rejected_not_answered_with_resorts(self):
        self._record("RY-01", "2026-09-21")
        client = self._client()

        response = client.get("/api/reports/", {"type": "helicopter", "year": "2026"})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Unknown report type "helicopter"', response.json()["detail"])
        self.assertNotIn("rows", response.json())

    def test_no_report_type_is_still_the_resort_report(self):
        from .services.tourism import build_reports_payload

        self._record("RY-01", "2026-09-21")

        payload = build_reports_payload({"year": "2026", "include_questions": "false"})

        self.assertEqual(payload["type"], "resort")
        self.assertEqual(len(payload["rows"]), 1)


class PublicMobileTourismPrivacyTests(TestCase):
    """The public mobile tourism endpoints carry no visitor's personal data."""

    NAMES = ("Juana Dela Cruz", "Pedro Santos")

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        from api.views.mobile import get_mobile_destination_queryset

        cls.resort = get_mobile_destination_queryset().order_by("resort_id").first()
        quezon = Province.objects.get(name="Quezon")
        common = dict(
            country=Country.objects.get(name="Philippines"),
            region=quezon.region,
            province=quezon,
            resort=cls.resort,
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=2,
            filipino_count=2,
            total_male=1,
            total_female=1,
            age_8_59=2,
            status="arrived",
        )
        TouristRecord.objects.create(
            survey_id="SURV-2026-001", full_name=cls.NAMES[0], contact_number="09171234567",
            arrival_date="2026-09-20", **common,
        )
        TouristRecord.objects.create(
            survey_id="SURV-2026-002", full_name=cls.NAMES[1], contact_number="09179876543",
            arrival_date="2026-09-21", **common,
        )
        cls.reviews = {
            name: FeedbackEntry.objects.create(
                destination=cls.resort, reviewer=name, rating=5, date="2026-09-22",
                title="Visit", message="Clean.",
            )
            for name in ("Juan Dela Cruz", "   ", "123 !!", "Mobile Tourist", "maria")
        }

    def _bootstrap(self, query=None):
        response = APIClient().get("/api/mobile/tourism/bootstrap/", query or {})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response

    def test_the_public_bootstrap_carries_no_survey_id_or_name(self):
        body = self._bootstrap().content.decode()

        self.assertNotIn("SURV-", body)
        for name in self.NAMES:
            self.assertNotIn(name, body)
        self.assertNotIn("approval", body)

    def test_lookup_parameters_change_nothing(self):
        expected = self._bootstrap().json()

        for query in (
            {"contact": "9"},
            {"contact": "09171234567"},
            {"reference": "SURV-2026-001"},
            {"reference": "surv-2026-001"},
            {"contact": "09171234567", "reference": "SURV-2026-001"},
        ):
            with self.subTest(query=query):
                self.assertEqual(self._bootstrap(query).json(), expected)

    def test_notifications_is_always_a_list(self):
        notifications = self._bootstrap().json()["notifications"]

        self.assertIsInstance(notifications, list)
        for item in notifications:
            self.assertEqual(set(item) & {"id", "type", "title", "message"}, {"id", "type", "title", "message"})
            self.assertNotEqual(item["type"], "approval")

    def test_public_reviews_show_initials_never_the_typed_name(self):
        from api.views.mobile import public_reviewer_initials

        cases = {
            "Juan Dela Cruz": "J.D.C.",
            "maria": "M.",
            "  ana   de la paz ": "A.D.L.P.",
            "J.R. Santos-Reyes": "J.S.",
            "": "Mobile Tourist",
            "   ": "Mobile Tourist",
            "123 !!": "Mobile Tourist",
            None: "Mobile Tourist",
            "Mobile Tourist": "Mobile Tourist",
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                self.assertEqual(public_reviewer_initials(name), expected)

        response = APIClient().get(f"/api/mobile/tourism/destinations/{self.resort.resort_id}/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        bylines = sorted(item["reviewer"] for item in response.json()["recentFeedback"])
        self.assertEqual(bylines, sorted(["J.D.C.", "Mobile Tourist", "Mobile Tourist", "Mobile Tourist", "M."]))
        self.assertNotIn("Juan Dela Cruz", response.content.decode())

    def test_staff_feedback_keeps_the_full_name(self):
        user = User.objects.create_user(username="feedback_staff", password="Staff@12345")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

        response = client.get("/api/feedback/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("Juan Dela Cruz", [item["reviewer"] for item in response.json()])
        self.assertEqual(APIClient().get("/api/feedback/").status_code, status.HTTP_401_UNAUTHORIZED)


class PublicFeedbackWhitelistTests(TestCase):
    """The public destination detail shows a review through an explicit list
    of fields (PublicFeedbackSerializer), not everything the staff serializer
    has. A field added later must not reach the public without a test saying so.
    """

    PUBLIC_KEYS = {"id", "reviewer", "rating", "date", "message", "photos", "reply"}
    STAFF_KEYS = {
        "id", "destinationId", "reviewer", "rating", "status", "date",
        "title", "message", "reply", "photos",
        "is_hidden", "hidden_at", "hidden_by_display",
    }
    TYPED_NAME = "Juan Dela Cruz"

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        from api.views.mobile import get_mobile_destination_queryset

        cls.resort = get_mobile_destination_queryset().order_by("resort_id").first()
        cls.entry = FeedbackEntry.objects.create(
            destination=cls.resort,
            reviewer=cls.TYPED_NAME,
            rating=4,
            status=FeedbackEntry.STATUS_POSITIVE,
            date="2026-09-22",
            title="Mobile tourism feedback",
            message="Clean beach." + chr(10) * 2 + "Cleanliness feedback: cleanliness rating 5/5",
            reply="Thank you for visiting.",
            photos=["https://example.test/storage/media/feedback/abc.jpg"],
        )

    def _public(self):
        response = APIClient().get(
            f"/api/mobile/tourism/destinations/{self.resort.resort_id}/"
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response

    def test_the_public_keys_are_exactly_the_whitelist(self):
        items = self._public().json()["recentFeedback"]

        self.assertEqual(len(items), 1)
        self.assertEqual(set(items[0]), self.PUBLIC_KEYS)

    def test_destination_id_status_and_title_are_not_public(self):
        item = self._public().json()["recentFeedback"][0]

        for key in ("destinationId", "status", "title"):
            with self.subTest(key=key):
                self.assertNotIn(key, item)

    def test_the_values_the_app_shows_are_unchanged(self):
        item = self._public().json()["recentFeedback"][0]

        self.assertEqual(item["id"], self.entry.pk)
        self.assertEqual(item["rating"], 4)
        self.assertEqual(item["date"], "2026-09-22")
        self.assertEqual(item["reply"], "Thank you for visiting.")
        self.assertEqual(
            item["photos"], ["https://example.test/storage/media/feedback/abc.jpg"]
        )
        self.assertIn("Clean beach.", item["message"])
        self.assertIn("cleanliness rating 5/5", item["message"])

    def test_the_byline_is_the_initials_and_the_name_is_not_public(self):
        response = self._public()

        self.assertEqual(response.json()["recentFeedback"][0]["reviewer"], "J.D.C.")
        self.assertNotIn(self.TYPED_NAME, response.content.decode())

    def test_the_staff_list_still_has_every_field_and_the_full_name(self):
        user = User.objects.create_user(username="whitelist_staff", password="Staff@12345")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

        response = client.get("/api/feedback/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        item = next(row for row in response.json() if row["id"] == self.entry.pk)
        self.assertEqual(set(item), self.STAFF_KEYS)
        self.assertEqual(item["reviewer"], self.TYPED_NAME)

    def test_the_whitelist_is_an_explicit_list_and_read_only(self):
        from api.views.mobile import PublicFeedbackSerializer

        fields = PublicFeedbackSerializer.Meta.fields
        self.assertIsInstance(fields, list)
        self.assertEqual(set(fields), self.PUBLIC_KEYS)
        self.assertEqual(
            set(PublicFeedbackSerializer.Meta.read_only_fields), self.PUBLIC_KEYS
        )


class FeedbackModerationTests(TestCase):
    """Hide / unhide a review (item #17).

    A hidden review is not shown to the public and does not count in the
    destination's rating; staff still see it, marked, and can show it again.
    """

    PUBLIC_KEYS = {"id", "reviewer", "rating", "date", "message", "photos", "reply"}
    MODERATION_KEYS = {"is_hidden", "hidden_at", "hidden_by_display", "hidden_by"}

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        from api.views.mobile import get_mobile_destination_queryset

        cls.resort = get_mobile_destination_queryset().order_by("resort_id").first()

    # -- helpers -----------------------------------------------------------

    def _user(self, username, role):
        user = User.objects.create_user(username=username, password="Staff@12345")
        UserProfile.objects.create(user=user, role=role)
        return user

    def _client(self, user=None):
        client = APIClient()
        if user is not None:
            client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.get_or_create(user=user)[0].key}")
        return client

    def _review(self, rating=5, **extra):
        values = dict(
            destination=self.resort,
            reviewer="Juan Dela Cruz",
            rating=rating,
            status=FeedbackEntry.STATUS_POSITIVE,
            date="2026-09-22",
            title="Mobile tourism feedback",
            message="Clean beach.",
            reply="Thank you.",
        )
        values.update(extra)
        return FeedbackEntry.objects.create(**values)

    def _rating(self):
        self.resort.refresh_from_db()
        return self.resort.tourism_rating

    def _public(self):
        response = APIClient().get(
            f"/api/mobile/tourism/destinations/{self.resort.resort_id}/"
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.json()["recentFeedback"]

    def _hide(self, client, entry):
        return client.post(f"/api/feedback/{entry.pk}/hide/")

    def _unhide(self, client, entry):
        return client.post(f"/api/feedback/{entry.pk}/unhide/")

    # -- the public side ---------------------------------------------------

    def test_a_hidden_review_is_not_in_the_public_answer(self):
        keep = self._review(rating=5)
        gone = self._review(rating=3)
        staff = self._client(self._user("mod_a", ROLE_TOURISM))
        self.assertEqual(self._hide(staff, gone).status_code, status.HTTP_200_OK)

        ids = [item["id"] for item in self._public()]

        self.assertEqual(ids, [keep.pk])

    def test_six_reviews_one_hidden_still_gives_five(self):
        # The filter must come BEFORE the slice of five: after it, the hidden
        # review would still take one of the five places and four would show.
        entries = [self._review(rating=4, date=f"2026-09-{day:02d}") for day in range(10, 16)]
        staff = self._client(self._user("mod_b", ROLE_TOURISM))
        newest = entries[-1]  # the newest is the one that would be in the five
        self.assertEqual(self._hide(staff, newest).status_code, status.HTTP_200_OK)

        items = self._public()

        self.assertEqual(len(items), 5)
        self.assertNotIn(newest.pk, [item["id"] for item in items])
        self.assertEqual(
            sorted(item["id"] for item in items),
            sorted(entry.pk for entry in entries[:-1]),
        )

    def test_the_public_keys_are_still_the_seven_and_no_moderation_field_shows(self):
        hidden = self._review()
        self._review(rating=4)
        staff = self._client(self._user("mod_c", ROLE_TOURISM))
        self._hide(staff, hidden)

        for item in self._public():
            self.assertEqual(set(item), self.PUBLIC_KEYS)
            self.assertFalse(set(item) & self.MODERATION_KEYS)

    def test_a_public_post_cannot_set_the_moderation_fields(self):
        client = APIClient()
        mod = self._user("mod_d", ROLE_TOURISM)

        response = client.post(
            "/api/mobile/tourism/feedback/",
            {
                "destination_id": self.resort.resort_id,
                "reviewer": "Visitor",
                "rating": 4,
                "message": "Nice.",
                "is_hidden": True,
                "hidden_by": mod.pk,
                "hidden_at": "2026-01-01T00:00:00Z",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        entry = FeedbackEntry.objects.get(pk=response.json()["id"])
        self.assertFalse(entry.is_hidden)
        self.assertIsNone(entry.hidden_at)
        self.assertIsNone(entry.hidden_by)
        # ...and the answer to the sender carries no moderation field.
        self.assertFalse(set(response.json()) & self.MODERATION_KEYS)
        # ...and the review is in the public answer.
        self.assertIn(entry.pk, [item["id"] for item in self._public()])

    # -- the rating --------------------------------------------------------

    def test_the_rating_leaves_out_the_hidden_and_comes_back(self):
        five = self._review(rating=5)
        three = self._review(rating=3)
        staff = self._client(self._user("mod_e", ROLE_TOURISM))
        self.assertEqual(self._rating(), 4.0)

        self._hide(staff, three)
        self.assertEqual(self._rating(), 5.0)

        self._unhide(staff, three)
        self.assertEqual(self._rating(), 4.0)

        self._hide(staff, three)
        self._hide(staff, five)
        self.assertEqual(self._rating(), 0.0)

    def test_hiding_runs_the_save_signal_so_the_stars_change_at_once(self):
        # QuerySet.update() sends no post_save: the rating would stay as it was.
        self._review(rating=5)
        three = self._review(rating=3)
        staff = self._client(self._user("mod_f", ROLE_TOURISM))
        self.assertEqual(self._rating(), 4.0)

        response = self._hide(staff, three)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(self._rating(), 5.0)  # in the same request, no wait

    # -- the actions -------------------------------------------------------

    def test_hide_sets_the_three_fields_from_the_user_and_unhide_clears_them(self):
        entry = self._review()
        user = self._user("mod_g", ROLE_TOURISM)
        staff = self._client(user)
        before = timezone.now()

        response = self._hide(staff, entry)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertTrue(entry.is_hidden)
        self.assertEqual(entry.hidden_by_id, user.pk)
        self.assertGreaterEqual(entry.hidden_at, before)
        body = response.json()
        self.assertTrue(body["is_hidden"])
        self.assertEqual(body["hidden_by_display"], "mod_g")
        self.assertNotIn("hidden_by", body)  # the id of the user is not given

        response = self._unhide(staff, entry)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertFalse(entry.is_hidden)
        self.assertIsNone(entry.hidden_at)
        self.assertIsNone(entry.hidden_by)

    def test_the_who_and_when_come_from_the_server_not_the_request(self):
        entry = self._review()
        user = self._user("mod_h", ROLE_TOURISM)
        other = self._user("someone_else", ROLE_TOURISM)

        response = self._client(user).post(
            f"/api/feedback/{entry.pk}/hide/",
            {"hidden_by": other.pk, "hidden_at": "2001-01-01T00:00:00Z", "is_hidden": False},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertTrue(entry.is_hidden)
        self.assertEqual(entry.hidden_by_id, user.pk)
        self.assertGreater(entry.hidden_at.year, 2001)

    def test_hide_is_idempotent_the_first_who_and_when_stay(self):
        entry = self._review()
        first = self._user("mod_first", ROLE_TOURISM)
        second = self._user("mod_second", ROLE_ADMIN)
        self._hide(self._client(first), entry)
        entry.refresh_from_db()
        when = entry.hidden_at

        response = self._hide(self._client(second), entry)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertEqual(entry.hidden_by_id, first.pk)
        self.assertEqual(entry.hidden_at, when)
        self.assertEqual(response.json()["hidden_by_display"], "mod_first")

    def test_permissions(self):
        entry = self._review()
        paths = [f"/api/feedback/{entry.pk}/hide/", f"/api/feedback/{entry.pk}/unhide/"]

        for path in paths:
            with self.subTest(path=path, who="nobody"):
                self.assertEqual(self._client().post(path).status_code, status.HTTP_401_UNAUTHORIZED)
            for role in (ROLE_TOURIST, ROLE_SANITATION):
                with self.subTest(path=path, who=role):
                    user = self._user(f"mod_{role}_{path[-6:-1]}", role)
                    self.assertEqual(self._client(user).post(path).status_code, status.HTTP_403_FORBIDDEN)

        entry.refresh_from_db()
        self.assertFalse(entry.is_hidden)  # none of those changed it

        for role in (ROLE_TOURISM, ROLE_ADMIN):
            with self.subTest(who=role):
                user = self._user(f"mod_ok_{role}", role)
                self.assertEqual(self._hide(self._client(user), entry).status_code, status.HTTP_200_OK)
                self.assertEqual(self._unhide(self._client(user), entry).status_code, status.HTTP_200_OK)

    def test_an_unknown_review_is_404(self):
        staff = self._client(self._user("mod_i", ROLE_TOURISM))

        self.assertEqual(staff.post("/api/feedback/999999/hide/").status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(staff.post("/api/feedback/999999/unhide/").status_code, status.HTTP_404_NOT_FOUND)

    def test_hiding_changes_nothing_else_of_the_review(self):
        entry = self._review(rating=4, reply="Thank you.", message="Clean beach.")
        staff = self._client(self._user("mod_j", ROLE_TOURISM))

        body = self._hide(staff, entry).json()

        # the answer shows the review as it is, not only the database
        self.assertEqual(body["reply"], "Thank you.")
        self.assertEqual(body["message"], "Clean beach.")
        self.assertEqual(body["rating"], 4)
        entry.refresh_from_db()
        self.assertEqual(entry.reply, "Thank you.")
        self.assertEqual(entry.message, "Clean beach.")
        self.assertEqual(entry.rating, 4)
        self.assertEqual(entry.reviewer, "Juan Dela Cruz")

    def test_patch_cannot_hide_or_unhide(self):
        entry = self._review()
        staff = self._client(self._user("mod_k", ROLE_TOURISM))

        response = staff.patch(
            f"/api/feedback/{entry.pk}/", {"is_hidden": True, "reply": "Edited"}, format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertFalse(entry.is_hidden)
        self.assertEqual(entry.reply, "Edited")  # a reply edit still works

    # -- the staff side: the trap ------------------------------------------

    def test_the_staff_list_still_has_the_hidden_one_marked_with_the_name(self):
        # If the staff list left the hidden review out, the admin could not see
        # what they hid, and so could not show it again.
        entry = self._review()
        user = self._user("mod_l", ROLE_TOURISM)
        staff = self._client(user)
        self._hide(staff, entry)

        listing = staff.get("/api/feedback/")

        self.assertEqual(listing.status_code, status.HTTP_200_OK)
        row = next(item for item in listing.json() if item["id"] == entry.pk)
        self.assertTrue(row["is_hidden"])
        self.assertEqual(row["hidden_by_display"], "mod_l")
        self.assertIsNotNone(row["hidden_at"])
        self.assertEqual(row["reviewer"], "Juan Dela Cruz")  # the full name, as before
        self.assertNotIn("hidden_by", row)

    def test_the_web_bootstrap_still_has_the_hidden_one(self):
        entry = self._review()
        staff = self._client(self._user("mod_m", ROLE_TOURISM))
        self._hide(staff, entry)

        boot = staff.get("/api/bootstrap/")

        self.assertEqual(boot.status_code, status.HTTP_200_OK)
        row = next(item for item in boot.json()["feedbackEntries"] if item["id"] == entry.pk)
        self.assertTrue(row["is_hidden"])

    # -- accountability ----------------------------------------------------

    def test_the_activity_log_says_which_review_and_who(self):
        entry = self._review(title="Mobile tourism feedback")
        user = self._user("mod_n", ROLE_TOURISM)
        staff = self._client(user)
        before = ActivityLog.objects.count()

        self._hide(staff, entry)
        self._hide(staff, entry)  # a repeat changes nothing and is not logged
        self._unhide(staff, entry)

        logs = list(ActivityLog.objects.order_by("id")[before:])
        self.assertEqual(len(logs), 2)
        for log in logs:
            self.assertEqual(log.action, "update")  # no new action
            self.assertEqual(log.user_id, user.pk)
            self.assertEqual(log.record_type, "FeedbackEntry")
            self.assertEqual(log.record_id, str(entry.pk))
            self.assertNotEqual(log.record_label, "Mobile tourism feedback")
            self.assertIn(self.resort.resort_name, log.record_label)
            self.assertIn("2026-09-22", log.record_label)
        self.assertTrue(logs[0].record_label.startswith("Review hidden:"))
        self.assertTrue(logs[1].record_label.startswith("Review shown again:"))

    def test_no_new_activity_action_was_added(self):
        from api.models import ACTIVITY_ACTION_CHOICES

        self.assertEqual([value for value, _ in ACTIVITY_ACTION_CHOICES], ["create", "update", "delete"])

    # -- the migration -----------------------------------------------------

    def test_the_model_has_no_change_without_a_migration(self):
        # makemigrations --check exits non-zero (SystemExit) when a model
        # changed and no migration says so.
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)

    def test_a_review_that_already_existed_is_visible(self):
        # The two live reviews were made before the columns: they take the
        # default and stay visible.
        entry = self._review()

        entry.refresh_from_db()
        self.assertFalse(entry.is_hidden)
        self.assertIsNone(entry.hidden_at)
        self.assertIsNone(entry.hidden_by)
        self.assertIn(entry.pk, [item["id"] for item in self._public()])


class ReferenceOptionSchemaTests(TestCase):
    """The columns behind "registration options" (commit 1: schema only).

    Four choices can be retired (is_active) and an itinerary can say how long
    the stay is (stay_nights). Nothing reads either yet, so nothing that
    exists may change: not the payloads, not the reports.
    """

    ACTIVE_MODELS = (Itinerary, TravelMode, BoatType, VisitPurpose)

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()

    # -- is_active ---------------------------------------------------------

    def test_the_four_have_is_active_and_it_defaults_to_true(self):
        for model in self.ACTIVE_MODELS:
            with self.subTest(model=model.__name__):
                field = model._meta.get_field("is_active")
                self.assertIs(field.default, True)
                row = model.objects.create(id=9000, name="Default check")
                row.refresh_from_db()
                self.assertIs(row.is_active, True)

    def test_a_row_inserted_without_the_column_is_active(self):
        # The old instance, during a deploy, inserts rows without knowing the
        # column. The database default must make them active.
        for model in self.ACTIVE_MODELS:
            with self.subTest(model=model.__name__):
                field = model._meta.get_field("is_active")
                self.assertIs(field.db_default, True)
                table = model._meta.db_table
                columns = "id, name, requires_capacity_fare" if model is BoatType else "id, name"
                values = "9100, 'Old code row', false" if model is BoatType else "9100, 'Old code row'"
                with connection.cursor() as cursor:
                    cursor.execute(f"INSERT INTO {table} ({columns}) VALUES ({values})")
                self.assertIs(model.objects.get(pk=9100).is_active, True)

    def test_the_other_reference_tables_have_no_is_active(self):
        for model in (Country, Region, Province):
            with self.subTest(model=model.__name__):
                names = [f.name for f in model._meta.get_fields()]
                self.assertNotIn("is_active", names)

    def test_only_itinerary_has_stay_nights(self):
        self.assertIn("stay_nights", [f.name for f in Itinerary._meta.get_fields()])
        for model in (TravelMode, BoatType, VisitPurpose, Country, Region, Province):
            with self.subTest(model=model.__name__):
                self.assertNotIn("stay_nights", [f.name for f in model._meta.get_fields()])

    # -- stay_nights -------------------------------------------------------

    def test_stay_nights_is_nullable_and_none_by_default(self):
        field = Itinerary._meta.get_field("stay_nights")
        self.assertTrue(field.null)
        self.assertTrue(field.blank)
        self.assertIsNone(Itinerary(id=9199, name="Unsaved").stay_nights)
        row = Itinerary.objects.create(id=9200, name="Not set yet")
        row.refresh_from_db()
        self.assertIsNone(row.stay_nights)

    def test_stay_nights_keeps_zero_apart_from_not_set(self):
        # 0 means same day; None means "read the name". They are not the same.
        same_day = Itinerary.objects.create(id=9201, name="X", stay_nights=0)
        unset = Itinerary.objects.create(id=9202, name="Y")
        same_day.refresh_from_db()
        unset.refresh_from_db()
        self.assertEqual(same_day.stay_nights, 0)
        self.assertIsNotNone(same_day.stay_nights)
        self.assertIsNone(unset.stay_nights)

    def test_stay_nights_cannot_be_negative(self):
        row = Itinerary(id=9203, name="Negative", stay_nights=-1)
        with self.assertRaises(ValidationError):
            row.full_clean()

    # -- nothing that exists changes ---------------------------------------

    def test_the_mobile_payload_has_the_same_keys_as_before(self):
        # Since commit 2 the WEB payload carries is_active (and stay_nights);
        # the mobile payload must keep exactly the keys it always had.
        from .services.tourism import build_reference_tables_payload

        payload = build_reference_tables_payload(active_only=True)
        for key in ("itineraries", "travelModes", "visitPurposes"):
            with self.subTest(key=key):
                self.assertTrue(payload[key])
                for item in payload[key]:
                    self.assertEqual(set(item), {"id", "name"})
        for item in payload["boatTypes"]:
            self.assertEqual(set(item), {"id", "name", "requires_capacity_fare"})

    def test_an_inactive_row_is_still_in_the_web_payload(self):
        # The web is given every option, retired or not (commit 2 filters the
        # mobile payload only).
        from .services.tourism import build_reference_tables_payload

        row = Itinerary.objects.create(id=9300, name="Retired later", is_active=False)

        ids = [item["id"] for item in build_reference_tables_payload()["itineraries"]]

        self.assertIn(row.id, ids)

    def test_the_reports_read_stay_nights_before_the_name(self):
        from .services.tourism import (
            build_arrival_monitoring_payload,
            build_tourism_question_answers,
        )

        # Opposite on purpose: the name says same day, stay_nights says 5; the
        # name says nights, stay_nights says 0. Since commit 3 stay_nights wins.
        day_name = Itinerary.objects.create(id=9400, name="Weekend 2 Days", stay_nights=5)
        night_name = Itinerary.objects.create(id=9401, name="Camp Out", stay_nights=0)
        quezon = Province.objects.get(name="Quezon")
        common = dict(
            full_name="Group", contact_number="09170000000",
            country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
            arrival_date="2026-04-02", resort=Resort.objects.order_by("resort_id").first(),
            travel_mode=TravelMode.objects.first(), boat_type=BoatType.objects.first(),
            visit_purpose=VisitPurpose.objects.first(), total_visitors=4, filipino_count=4,
            total_male=4, age_8_59=4, status="arrived",
        )
        TouristRecord.objects.create(survey_id="SCH-0001", itinerary=day_name, **common)
        # A different head count, so that a swap of the two would show in the
        # totals (4 and 4 would look the same swapped).
        TouristRecord.objects.create(
            survey_id="SCH-0002", itinerary=night_name,
            **{**common, "total_visitors": 6, "filipino_count": 6, "total_male": 6, "age_8_59": 6},
        )

        rows = {
            row["survey_id"]: row
            for row in build_arrival_monitoring_payload({"date": "2026-04-02"})["rows"]
        }
        self.assertEqual((rows["SCH-0001"]["sameDay"], rows["SCH-0001"]["overnight"]), (0, 4))
        self.assertEqual((rows["SCH-0002"]["sameDay"], rows["SCH-0002"]["overnight"]), (6, 0))

        answers = {a["id"]: a for a in build_tourism_question_answers({"year": "2026"})}
        self.assertEqual(
            answers["stay_type"]["answer"],
            "Same-day visitors: 6; overnight or multi-day visitors: 4.",
        )

    def test_the_model_has_no_change_without_a_migration(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


class ReferenceOptionPayloadTests(TestCase):
    """What each client is offered (registration options, commit 2).

    The mobile app gets only the ACTIVE options, in the keys it always read.
    The web gets every option with is_active, so it can show a retired one and
    name it on old records. The cache holds the full lists and the filtering is
    done after it.
    """

    MOBILE_KEYS = {
        "itineraries": {"id", "name"},
        "travelModes": {"id", "name"},
        "boatTypes": {"id", "name", "requires_capacity_fare"},
        "visitPurposes": {"id", "name"},
    }
    CACHE_KEY = "mobile_reference_tables_v2"

    # Different numbers on purpose (1 active + 2 retired, 1 retired, 1 retired,
    # 3 retired): a mutant that swaps "active" and "retired" cannot give the
    # same answer.
    ITINERARY_ACTIVE, ITINERARY_RETIRED = [9501], [9502, 9503]
    MODE_RETIRED = [9511]
    BOAT_RETIRED = [9521]
    PURPOSE_RETIRED = [9531, 9532, 9533]

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        Itinerary.objects.create(id=9501, name="Added and active", stay_nights=2)
        for i in self.ITINERARY_RETIRED:
            Itinerary.objects.create(id=i, name=f"Retired itinerary {i}", is_active=False)
        for i in self.MODE_RETIRED:
            TravelMode.objects.create(id=i, name=f"Retired mode {i}", is_active=False)
        for i in self.BOAT_RETIRED:
            BoatType.objects.create(
                id=i, name=f"Retired boat {i}", requires_capacity_fare=True, is_active=False
            )
        for i in self.PURPOSE_RETIRED:
            VisitPurpose.objects.create(id=i, name=f"Retired purpose {i}", is_active=False)

    def _retired(self):
        return {
            "itineraries": set(self.ITINERARY_RETIRED),
            "travelModes": set(self.MODE_RETIRED),
            "boatTypes": set(self.BOAT_RETIRED),
            "visitPurposes": set(self.PURPOSE_RETIRED),
        }

    @staticmethod
    def _ids(payload, key):
        return {item["id"] for item in payload[key]}

    def _record(self, survey_id, itinerary):
        quezon = Province.objects.get(name="Quezon")
        return TouristRecord.objects.create(
            survey_id=survey_id, full_name="Group", contact_number="09170000000",
            country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
            arrival_date="2026-04-02", resort=Resort.objects.order_by("resort_id").first(),
            itinerary=itinerary, travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(), visit_purpose=VisitPurpose.objects.first(),
            total_visitors=4, filipino_count=4, total_male=4, age_8_59=4, status="arrived",
        )

    def _admin_client(self):
        user = User.objects.create_user(username="options_admin", password="Staff@12345")
        UserProfile.objects.create(user=user, role=ROLE_ADMIN)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        return client

    # -- who gets what -----------------------------------------------------

    def test_a_retired_option_is_not_in_the_mobile_payload(self):
        from .services.tourism import build_reference_tables_payload

        payload = build_reference_tables_payload(active_only=True)

        for key, retired in self._retired().items():
            with self.subTest(key=key):
                self.assertFalse(self._ids(payload, key) & retired)
        self.assertIn(9501, self._ids(payload, "itineraries"))  # an active one stays

    def test_a_retired_option_is_in_the_web_payload_marked_inactive(self):
        from .services.tourism import build_reference_tables_payload

        payload = build_reference_tables_payload()

        for key, retired in self._retired().items():
            with self.subTest(key=key):
                rows = {item["id"]: item for item in payload[key]}
                for i in retired:
                    self.assertIn(i, rows)
                    self.assertIs(rows[i]["is_active"], False)
        self.assertIs(
            {i["id"]: i for i in payload["itineraries"]}[9501]["is_active"], True
        )

    def test_the_web_endpoint_carries_is_active_and_stay_nights(self):
        response = self._admin_client().get("/api/reference-tables/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        itineraries = {item["id"]: item for item in response.json()["itineraries"]}
        self.assertEqual(itineraries[9501]["stay_nights"], 2)
        self.assertIsNone(itineraries[9502]["stay_nights"])
        self.assertIs(itineraries[9502]["is_active"], False)
        boats = {item["id"]: item for item in response.json()["boatTypes"]}
        self.assertEqual(
            set(boats[9521]), {"id", "name", "requires_capacity_fare", "is_active"}
        )

    def test_the_mobile_payload_keeps_the_old_keys_only(self):
        from .services.tourism import build_reference_tables_payload

        payload = build_reference_tables_payload(active_only=True)

        for key, keys in self.MOBILE_KEYS.items():
            with self.subTest(key=key):
                self.assertTrue(payload[key])
                for item in payload[key]:
                    self.assertEqual(set(item), keys)

    def test_the_mobile_bootstrap_endpoint_has_the_old_shape_and_no_retired(self):
        response = APIClient().get("/api/mobile/tourism/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        tables = response.json()["referenceTables"]
        for key, keys in self.MOBILE_KEYS.items():
            with self.subTest(key=key):
                for item in tables[key]:
                    self.assertEqual(set(item), keys)
                self.assertFalse(self._ids(tables, key) & self._retired()[key])
        self.assertEqual(
            set(tables),
            {"countries", "regions", "provinces", "itineraries", "travelModes",
             "boatTypes", "visitPurposes", "resorts", "boatCapacityFares"},
        )

    def test_other_lists_are_untouched(self):
        from .services.tourism import build_reference_tables_payload

        mobile = build_reference_tables_payload(active_only=True)
        web = build_reference_tables_payload()

        for key in ("countries", "regions", "provinces"):
            with self.subTest(key=key):
                self.assertEqual(mobile[key], web[key])
                self.assertNotIn("is_active", mobile[key][0])

    # -- history -----------------------------------------------------------

    def test_an_old_record_keeps_its_option_name_after_it_is_retired(self):
        from .services.tourism import build_arrival_monitoring_payload

        itinerary = Itinerary.objects.create(id=9600, name="Retired Trail")
        self._record("OPT-0001", itinerary)
        Itinerary.objects.filter(pk=9600).update(is_active=False)

        row = build_arrival_monitoring_payload({"date": "2026-04-02"})["rows"][0]

        self.assertEqual(row["itinerary"], "Retired Trail")
        self.assertEqual(TouristRecord.objects.get(pk="OPT-0001").itinerary.name, "Retired Trail")

    def test_a_retired_option_is_still_accepted_on_a_record(self):
        # A visitor who opened the form before it was retired must not lose
        # their registration to an error: the dropdown is the control.
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            {
                "full_name": "Stale Form", "contact_number": "09171234567",
                "arrival_date": "2026-06-01", "total_visitors": 1, "filipino_count": 1,
                "total_male": 1, "age_8_59": 1, "itinerary_id": 9502,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(TouristRecord.objects.get(full_name="Stale Form").itinerary_id, 9502)

    # -- defaults ----------------------------------------------------------

    def test_a_default_is_never_a_retired_option(self):
        from .views.mobile import default_reference_id

        lowest = Itinerary.objects.order_by("id").first()
        second = Itinerary.objects.order_by("id")[1]
        Itinerary.objects.filter(pk=lowest.pk).update(is_active=False)

        self.assertEqual(default_reference_id(Itinerary), second.pk)

    def test_the_default_is_the_lowest_active_id(self):
        from .views.mobile import default_reference_id

        # Retire everything but two, 9501 and a seeded one: the lower id wins.
        keep = Itinerary.objects.exclude(pk=9501).order_by("id").first().pk
        Itinerary.objects.exclude(pk__in=[keep, 9501]).update(is_active=False)

        self.assertEqual(default_reference_id(Itinerary), min(keep, 9501))

    def test_every_retirable_list_skips_a_retired_first_row(self):
        from .views.mobile import default_reference_id

        for model in (TravelMode, BoatType, VisitPurpose):
            with self.subTest(model=model.__name__):
                first = model.objects.order_by("id").first()
                model.objects.filter(pk=first.pk).update(is_active=False)
                expected = model.objects.filter(is_active=True).order_by("id").first().pk
                self.assertNotEqual(expected, first.pk)
                self.assertEqual(default_reference_id(model), expected)

    def test_with_none_active_there_is_no_guess_and_a_400(self):
        from .views.mobile import default_reference_id

        Itinerary.objects.update(is_active=False)
        before = TouristRecord.objects.count()

        self.assertIsNone(default_reference_id(Itinerary))
        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            {
                "full_name": "No Default", "contact_number": "09171234567",
                "arrival_date": "2026-06-01", "total_visitors": 1, "filipino_count": 1,
                "total_male": 1, "age_8_59": 1,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("itinerary_id", response.json())
        self.assertEqual(TouristRecord.objects.count(), before)

    def test_a_registration_without_ids_gets_an_active_option(self):
        Itinerary.objects.filter(pk=Itinerary.objects.order_by("id").first().pk).update(is_active=False)
        expected = Itinerary.objects.filter(is_active=True).order_by("id").first().pk

        response = APIClient().post(
            "/api/mobile/tourism/register-visit/",
            {
                "full_name": "Defaulted", "contact_number": "09171234567",
                "arrival_date": "2026-06-01", "total_visitors": 1, "filipino_count": 1,
                "total_male": 1, "age_8_59": 1,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(TouristRecord.objects.get(full_name="Defaulted").itinerary_id, expected)

    def test_the_default_of_the_other_models_is_unchanged(self):
        # Resort and the places have no is_active; their default is not part of
        # this change (the Resort default is its own item).
        from .views.mobile import default_reference_id

        self.assertEqual(default_reference_id(Resort), Resort.objects.first().pk)
        self.assertEqual(default_reference_id(Country), Country.objects.first().pk)

    # -- the cache ---------------------------------------------------------

    def test_the_cache_does_not_mix_mobile_and_web(self):
        from django.core.cache import cache
        from .services.tourism import build_reference_tables_payload

        # mobile first, then web: the web must still see the retired ones
        cache.clear()
        build_reference_tables_payload(active_only=True)
        web = build_reference_tables_payload(active_only=False)
        self.assertTrue(self._ids(web, "itineraries") >= set(self.ITINERARY_RETIRED))

        # web first, then mobile: the mobile must still not
        cache.clear()
        build_reference_tables_payload(active_only=False)
        mobile = build_reference_tables_payload(active_only=True)
        self.assertFalse(self._ids(mobile, "itineraries") & set(self.ITINERARY_RETIRED))
        self.assertFalse(self._ids(mobile, "visitPurposes") & set(self.PURPOSE_RETIRED))

    def test_the_cache_holds_the_full_lists_under_v2(self):
        from django.core.cache import cache
        from .services.tourism import build_reference_tables_payload

        build_reference_tables_payload(active_only=True)

        cached = cache.get(self.CACHE_KEY)
        self.assertIsNotNone(cached)
        self.assertTrue(
            {i["id"] for i in cached["itineraries"]} >= set(self.ITINERARY_RETIRED)
        )
        self.assertIsNone(cache.get("mobile_reference_tables_v1"))

    def test_an_entry_of_the_old_shape_under_v1_is_ignored(self):
        from django.core.cache import cache
        from .services.tourism import build_reference_tables_payload

        cache.set(
            "mobile_reference_tables_v1",
            {key: [{"id": 1, "name": "Stale"}] for key in
             ("countries", "regions", "provinces", "itineraries", "travelModes",
              "boatTypes", "visitPurposes")},
            timeout=900,
        )

        web = build_reference_tables_payload()

        self.assertNotIn("Stale", [i["name"] for i in web["itineraries"]])
        self.assertIn("is_active", web["itineraries"][0])

    def test_the_model_has_no_change_without_a_migration(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


# The 12 itineraries the live system has, hard-coded on purpose: the local seed
# has 6, and this test must not depend on what a database holds.
LIVE_ITINERARY_NAMES = [
    "Overnight", "Same Day", "2 Nights", "3 Nights", "4 Nights",
    "5 Nights and above", "6 Nights", "Que", "Island Hopping Circuit",
    "Beach and Resort Day Tour", "Heritage and Lighthouse Stop",
    "Eco Adventure Trail",
]
# A different head count for every record (all primes, none repeated): a mutant
# that moves one record between same day and overnight cannot give the same
# totals.
STAY_FIXTURE_VISITORS = [1, 2, 3, 5, 7, 11, 13, 17, 19, 23, 29, 31]


def _stay_fixture():
    """One arrived record per live itinerary name, on two resorts."""
    ensure_test_reference_tables()
    quezon = Province.objects.get(name="Quezon")
    resorts = list(Resort.objects.order_by("resort_id")[:2])
    for index, (name, visitors) in enumerate(zip(LIVE_ITINERARY_NAMES, STAY_FIXTURE_VISITORS)):
        itinerary = Itinerary.objects.create(id=9701 + index, name=name)
        TouristRecord.objects.create(
            survey_id=f"STY-{index + 1:02d}", full_name="Group", contact_number="09170000000",
            country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
            arrival_date="2026-04-02", resort=resorts[0] if index < 7 else resorts[1],
            itinerary=itinerary, travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(), visit_purpose=VisitPurpose.objects.first(),
            total_visitors=visitors, filipino_count=visitors, total_male=visitors,
            age_8_59=visitors, status="arrived",
        )


def _stay_report_numbers():
    """What the three reports say about stay, on the fixture."""
    from .services.tourism import (
        build_arrival_monitoring_payload,
        build_dashboard_payload,
        build_tourism_question_answers,
    )

    arrival = build_arrival_monitoring_payload({"date": "2026-04-02"})
    rows = {r["survey_id"]: (r["sameDay"], r["overnight"]) for r in arrival["rows"]}
    stay_type = build_dashboard_payload({"year": "2026"})["stayType"]
    answers = {a["id"]: a for a in build_tourism_question_answers({"year": "2026"})}
    return {
        "arrival_rows": rows,
        "dashboard_stay_type": dict(stay_type),
        "stay_type_answer": answers["stay_type"]["answer"],
        "overnight_resort_answer": answers["overnight_resort"]["answer"],
        "average_stay_answer": answers["average_stay"]["answer"],
    }


def _old_name_parse(name):
    """The name reading every report used before itinerary_stay existed, copied
    from the old code as the oracle: (same_day, nights)."""
    text = (name or "").lower()
    if "day" in text or "same" in text:
        return True, 0
    if "2" in text:
        return False, 2
    if "3" in text:
        return False, 3
    if "4" in text:
        return False, 4
    if "5" in text:
        return False, 5
    return False, 1


class ItineraryStayTests(TestCase):
    """Reports read an itinerary's stay_nights, and the NAME only when it is None."""

    # What the old code gave for each of the 12 live names: (same_day, nights).
    # Written out, not computed: if any of these moved, an LGU number moved.
    LIVE_EXPECTED = {
        "Overnight": (False, 1),
        "Same Day": (True, 0),
        "2 Nights": (False, 2),
        "3 Nights": (False, 3),
        "4 Nights": (False, 4),
        "5 Nights and above": (False, 5),
        "6 Nights": (False, 1),  # the old reading missed the 6; the name says 6
        "Que": (False, 1),
        "Island Hopping Circuit": (False, 1),
        "Beach and Resort Day Tour": (True, 0),
        "Heritage and Lighthouse Stop": (False, 1),
        "Eco Adventure Trail": (False, 1),
    }

    # The reports on the fixture with every stay_nights None, taken from the OLD
    # code before it was replaced (the same fixture, run on the old code).
    GOLDEN_WITH_NULLS = {
        "arrival_rows": {
            "STY-01": (0, 1), "STY-02": (2, 0), "STY-03": (0, 3), "STY-04": (0, 5),
            "STY-05": (0, 7), "STY-06": (0, 11), "STY-07": (0, 13), "STY-08": (0, 17),
            "STY-09": (0, 19), "STY-10": (23, 0), "STY-11": (0, 29), "STY-12": (0, 31),
        },
        "dashboard_stay_type": {"dayTour": 25, "overnight": 136},
        "stay_type_answer": "Same-day visitors: 25; overnight or multi-day visitors: 136.",
        "overnight_resort_answer": "Aquazul Hotel and Resort has the highest overnight demand with 96 visitors.",
        "average_stay_answer": "The estimated average stay is 1.3 night(s) per visitor based on itinerary labels.",
    }

    # -- the helper --------------------------------------------------------

    def test_the_twelve_live_names_read_exactly_as_the_old_code_did(self):
        from .services.tourism import itinerary_stay

        self.assertEqual(sorted(self.LIVE_EXPECTED), sorted(LIVE_ITINERARY_NAMES))
        for name, expected in self.LIVE_EXPECTED.items():
            with self.subTest(name=name):
                self.assertEqual(itinerary_stay(name, None), expected)
                self.assertEqual(itinerary_stay(name), expected)  # None is the default
                self.assertEqual(itinerary_stay(name, None), _old_name_parse(name))

    def test_the_name_reading_matches_the_old_code_on_tricky_names(self):
        from .services.tourism import itinerary_stay

        # "2 or 5 nights": the old chain took the 2 first. "Same trip": only
        # "same" makes it same day. None and "" are 1 night.
        for name in ["2 or 5 nights", "2 and 3 nights", "Same trip", "Sameday", "Day 4", "Holiday", "25 nights",
                     "3-4 nights", "", None, "Friday Night", "Weekend 2 Days", "Trip 5", "A 4 B 3"]:
            with self.subTest(name=name):
                self.assertEqual(itinerary_stay(name, None), _old_name_parse(name))
        self.assertEqual(itinerary_stay("2 or 5 nights", None), (False, 2))
        self.assertEqual(itinerary_stay("Same trip", None), (True, 0))

    def test_weekend_2_days_with_two_nights_is_two_nights_not_same_day(self):
        # The bug this field exists for: "day" in the name made it a same-day stay.
        from .services.tourism import itinerary_stay

        self.assertEqual(itinerary_stay("Weekend 2 Days", None), (True, 0))  # the old, wrong, reading
        self.assertEqual(itinerary_stay("Weekend 2 Days", 2), (False, 2))

    def test_stay_nights_wins_over_the_name_both_ways(self):
        from .services.tourism import itinerary_stay

        self.assertEqual(itinerary_stay("Same Day", 3), (False, 3))
        self.assertEqual(itinerary_stay("2 Nights", 0), (True, 0))
        self.assertEqual(itinerary_stay("Overnight", 7), (False, 7))

    def test_zero_is_a_value_and_none_is_not_set(self):
        from .services.tourism import itinerary_stay

        # 0 must NOT fall through to the name: "Overnight" with 0 is same day.
        self.assertEqual(itinerary_stay("Overnight", 0), (True, 0))
        self.assertEqual(itinerary_stay("Overnight", None), (False, 1))

    def test_the_old_helpers_are_gone(self):
        import api.services.tourism as tourism

        self.assertFalse(hasattr(tourism, "get_stay_counts"))
        self.assertFalse(hasattr(tourism, "estimate_stay_nights"))

    # -- the reports, on real data -----------------------------------------

    def test_with_every_stay_nights_none_the_reports_are_what_they_were(self):
        _stay_fixture()
        self.assertFalse(Itinerary.objects.exclude(stay_nights=None).filter(id__gte=9701).exists())

        self.assertEqual(_stay_report_numbers(), self.GOLDEN_WITH_NULLS)

    def test_the_reports_use_stay_nights_where_it_is_set(self):
        _stay_fixture()
        # Three set, nine still None. Eco Adventure Trail (31): None read as
        # 1 night -> same day. Beach and Resort Day Tour (23): same day ->
        # 3 nights. 6 Nights (13): read as 1 -> 6.
        Itinerary.objects.filter(name="Eco Adventure Trail").update(stay_nights=0)
        Itinerary.objects.filter(name="Beach and Resort Day Tour").update(stay_nights=3)
        Itinerary.objects.filter(name="6 Nights").update(stay_nights=6)

        numbers = _stay_report_numbers()

        rows = numbers["arrival_rows"]
        self.assertEqual(rows["STY-12"], (31, 0))
        self.assertEqual(rows["STY-10"], (0, 23))
        self.assertEqual(rows["STY-07"], (0, 13))
        self.assertEqual(rows["STY-02"], (2, 0))   # a None still reads its name
        self.assertEqual(rows["STY-03"], (0, 3))
        self.assertEqual(numbers["dashboard_stay_type"], {"dayTour": 33, "overnight": 128})
        self.assertEqual(
            numbers["stay_type_answer"],
            "Same-day visitors: 33; overnight or multi-day visitors: 128.",
        )
        # nights: 214 before; 6 Nights 13 x 1 -> 13 x 6 (+65); Day Tour 0 -> 23 x 3 (+69);
        # Eco 31 x 1 -> 0 (-31): 317 over 161 visitors.
        self.assertEqual(
            numbers["average_stay_answer"],
            "The estimated average stay is 2.0 night(s) per visitor based on itinerary labels.",
        )
        self.assertEqual(
            numbers["overnight_resort_answer"],
            "Aquazul Hotel and Resort has the highest overnight demand with 88 visitors.",
        )

    def test_a_set_stay_nights_does_not_change_the_name_shown(self):
        _stay_fixture()
        Itinerary.objects.filter(name="Same Day").update(stay_nights=2)

        from .services.tourism import build_arrival_monitoring_payload

        rows = {r["survey_id"]: r for r in build_arrival_monitoring_payload({"date": "2026-04-02"})["rows"]}
        self.assertEqual(rows["STY-02"]["itinerary"], "Same Day")
        self.assertEqual((rows["STY-02"]["sameDay"], rows["STY-02"]["overnight"]), (0, 2))


class ReferenceOptionApiTests(TestCase):
    """Add, rename, retire and bring back the four registration options."""

    KINDS = {
        "itinerary": (Itinerary, "itinerary"),
        "travel-mode": (TravelMode, "travel_mode"),
        "boat-type": (BoatType, "boat_type"),
        "visit-purpose": (VisitPurpose, "visit_purpose"),
    }
    PAYLOAD_KEYS = {
        "itinerary": "itineraries", "travel-mode": "travelModes",
        "boat-type": "boatTypes", "visit-purpose": "visitPurposes",
    }
    CACHE_KEY = "mobile_reference_tables_v2"

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        self.admin = self._client(ROLE_ADMIN, "opt_admin")

    # -- helpers -----------------------------------------------------------

    def _client(self, role=None, username=None):
        client = APIClient()
        if role is not None:
            user = User.objects.create_user(username=username or f"opt_{role}", password="Staff@12345")
            UserProfile.objects.create(user=user, role=role)
            client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
            client.user = user
        return client

    def _url(self, kind, option_id=None, action=None):
        url = f"/api/reference-options/{kind}/"
        if option_id is not None:
            url += f"{option_id}/"
        if action:
            url += f"{action}/"
        return url

    def _record(self, survey_id, **fks):
        quezon = Province.objects.get(name="Quezon")
        values = dict(
            survey_id=survey_id, full_name="Group", contact_number="09170000000",
            country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
            arrival_date="2026-04-02", resort=Resort.objects.order_by("resort_id").first(),
            itinerary=Itinerary.objects.order_by("id").first(),
            travel_mode=TravelMode.objects.order_by("id").first(),
            boat_type=BoatType.objects.order_by("id").first(),
            visit_purpose=VisitPurpose.objects.order_by("id").first(),
            total_visitors=4, filipino_count=4, total_male=4, age_8_59=4, status="arrived",
        )
        values.update(fks)
        return TouristRecord.objects.create(**values)

    def _make(self, kind, option_id, name, uses=0, **extra):
        """An option with exactly `uses` records pointing at it."""
        model, fk = self.KINDS[kind]
        option = model.objects.create(id=option_id, name=name, **extra)
        for i in range(uses):
            self._record(f"OPT-{kind[:3]}-{option_id}-{i}", **{fk: option})
        return option

    def _payload(self, active_only=False):
        from .services.tourism import build_reference_tables_payload

        return build_reference_tables_payload(active_only=active_only)

    def _ids(self, kind, active_only=False):
        return {item["id"] for item in self._payload(active_only)[self.PAYLOAD_KEYS[kind]]}

    # -- who may ----------------------------------------------------------

    def test_only_the_admin_may_use_any_of_the_five_endpoints(self):
        option = self._make("itinerary", 9801, "Guarded", uses=1)
        calls = [
            ("get", self._url("itinerary")),
            ("post", self._url("itinerary")),
            ("patch", self._url("itinerary", option.id)),
            ("post", self._url("itinerary", option.id, "deactivate")),
            ("post", self._url("itinerary", option.id, "reactivate")),
        ]
        for method, url in calls:
            with self.subTest(who="nobody", method=method, url=url):
                self.assertEqual(getattr(self._client(), method)(url, {}, format="json").status_code, 401)
            for role in (ROLE_TOURISM, ROLE_SANITATION, ROLE_TOURIST):
                with self.subTest(who=role, method=method, url=url):
                    client = self._client(role, f"u_{role}_{method}_{url[-12:-1].replace('/', '')}")
                    self.assertEqual(getattr(client, method)(url, {}, format="json").status_code, 403)
        option.refresh_from_db()
        self.assertEqual((option.name, option.is_active), ("Guarded", True))  # none of those changed it

    def test_the_admin_is_let_in(self):
        self.assertEqual(self.admin.get(self._url("itinerary")).status_code, 200)

    def test_an_unknown_kind_is_404_not_500(self):
        option = self._make("itinerary", 9802, "Known")
        for method, url in [
            ("get", "/api/reference-options/boats/"),
            ("post", "/api/reference-options/boats/"),
            ("patch", "/api/reference-options/boats/1/"),
            ("post", "/api/reference-options/boats/1/deactivate/"),
            ("post", "/api/reference-options/boats/1/reactivate/"),
            ("get", "/api/reference-options/Itinerary/"),
        ]:
            with self.subTest(method=method, url=url):
                self.assertEqual(getattr(self.admin, method)(url, {}, format="json").status_code, 404)
        self.assertEqual(self.admin.patch(self._url("itinerary", 99999), {}, format="json").status_code, 404)
        self.assertEqual(self.admin.post(self._url("itinerary", 99999, "deactivate")).status_code, 404)
        self.assertTrue(Itinerary.objects.filter(pk=option.pk).exists())

    def test_there_is_no_delete_and_no_put(self):
        option = self._make("travel-mode", 9803, "Keep Me", uses=1)
        for method, url in [
            ("delete", self._url("travel-mode", option.id)),
            ("delete", self._url("travel-mode")),
            ("put", self._url("travel-mode", option.id)),
            ("delete", self._url("travel-mode", option.id, "deactivate")),
        ]:
            with self.subTest(method=method, url=url):
                self.assertEqual(getattr(self.admin, method)(url).status_code, 405)
        self.assertTrue(TravelMode.objects.filter(pk=option.pk).exists())

    # -- the list ----------------------------------------------------------

    def test_the_list_has_retired_options_and_the_real_record_counts(self):
        # 3, 1, 0 and 2 records: every count different, so options cannot be swapped unseen.
        self._make("visit-purpose", 9810, "Three uses", uses=3)
        self._make("visit-purpose", 9811, "One use", uses=1)
        self._make("visit-purpose", 9812, "No use", uses=0)
        self._make("visit-purpose", 9813, "Retired two", uses=2, is_active=False)

        rows = {r["id"]: r for r in self.admin.get(self._url("visit-purpose")).json()}

        self.assertEqual(
            (rows[9810]["record_count"], rows[9811]["record_count"], rows[9812]["record_count"], rows[9813]["record_count"]),
            (3, 1, 0, 2),
        )
        self.assertIs(rows[9813]["is_active"], False)
        self.assertEqual(set(rows), set(VisitPurpose.objects.values_list("id", flat=True)))

    def test_the_list_carries_the_fields_of_its_kind(self):
        self._make("itinerary", 9814, "Two nights", stay_nights=2)
        self._make("boat-type", 9815, "Ferry", requires_capacity_fare=True)

        itineraries = {r["id"]: r for r in self.admin.get(self._url("itinerary")).json()}
        boats = {r["id"]: r for r in self.admin.get(self._url("boat-type")).json()}
        modes = self.admin.get(self._url("travel-mode")).json()

        self.assertEqual(itineraries[9814]["stay_nights"], 2)
        self.assertIsNone(itineraries[min(itineraries)]["stay_nights"])
        self.assertIs(boats[9815]["requires_capacity_fare"], True)
        self.assertNotIn("stay_nights", modes[0])
        self.assertNotIn("requires_capacity_fare", modes[0])

    # -- adding ------------------------------------------------------------

    def test_add_in_every_kind_and_the_server_gives_the_id(self):
        for kind, (model, _) in self.KINDS.items():
            with self.subTest(kind=kind):
                top = model.objects.order_by("-id").first().id
                response = self.admin.post(self._url(kind), {"name": "  Brand   New  "}, format="json")

                self.assertEqual(response.status_code, 201)
                body = response.json()
                self.assertEqual(body["id"], top + 1)
                self.assertEqual(body["name"], "Brand New")  # trimmed, one space
                self.assertIs(body["is_active"], True)
                self.assertEqual(body["record_count"], 0)
                self.assertTrue(model.objects.filter(pk=top + 1, name="Brand New").exists())

    def test_an_id_sent_by_the_client_is_ignored(self):
        existing = Itinerary.objects.order_by("id").first()
        before = existing.name
        top = Itinerary.objects.order_by("-id").first().id

        taken = self.admin.post(self._url("itinerary"), {"id": existing.id, "name": "Wants an old id"}, format="json")
        far = self.admin.post(self._url("itinerary"), {"id": 99999, "name": "Wants a far id"}, format="json")

        self.assertEqual(taken.status_code, 201)
        self.assertEqual(far.status_code, 201)
        self.assertEqual(taken.json()["id"], top + 1)
        self.assertEqual(far.json()["id"], top + 2)
        self.assertEqual(Itinerary.objects.get(pk=existing.id).name, before)  # the old row was not replaced
        self.assertFalse(Itinerary.objects.filter(pk=99999).exists())

    def test_a_patch_cannot_change_the_id_or_retire_by_the_back_door(self):
        option = self._make("travel-mode", 9820, "Fixed", uses=2)
        url = self._url("travel-mode", option.id)

        # a PATCH that really changes something (a rename), with extras the
        # server must not obey; and one that changes nothing
        renamed = self.admin.patch(
            url,
            {"id": 1, "is_active": False, "record_count": 99, "name": "Fixed Again", "confirm_affected_records": 2},
            format="json",
        )
        nothing = self.admin.patch(url, {"id": 1, "is_active": False, "name": "Fixed Again"}, format="json")

        for response in (renamed, nothing):
            self.assertEqual(response.status_code, 200)
            # what is ANSWERED must be what is saved: still active, same id
            self.assertEqual((response.json()["id"], response.json()["is_active"]), (9820, True))
            self.assertEqual(response.json()["record_count"], 2)  # the server's count
        option.refresh_from_db()
        self.assertEqual((option.id, option.name, option.is_active), (9820, "Fixed Again", True))
        self.assertTrue(TravelMode.objects.filter(pk=1).exists())  # row 1 was not touched

    def test_the_same_name_in_any_dress_is_refused(self):
        for kind, (model, _) in self.KINDS.items():
            self._make(kind, 9830, "Same Day Out")
            for attempt in ["same day out", "Same Day Out ", " Same Day Out", "Same  Day  Out", "SAME   DAY OUT  "]:
                with self.subTest(kind=kind, attempt=attempt):
                    before = model.objects.count()
                    response = self.admin.post(self._url(kind), {"name": attempt}, format="json")
                    self.assertEqual(response.status_code, 400)
                    self.assertIn("name", response.json())
                    self.assertEqual(model.objects.count(), before)

    def test_the_same_name_is_refused_even_against_a_retired_option(self):
        self._make("visit-purpose", 9831, "Old Purpose", is_active=False)

        response = self.admin.post(self._url("visit-purpose"), {"name": "old   purpose"}, format="json")

        self.assertEqual(response.status_code, 400)

    def test_a_rename_onto_another_options_name_is_refused(self):
        self._make("boat-type", 9832, "Alpha Boat")
        other = self._make("boat-type", 9833, "Beta Boat")

        for attempt in ["alpha boat", "Alpha  Boat ", "ALPHA BOAT"]:
            with self.subTest(attempt=attempt):
                response = self.admin.patch(
                    self._url("boat-type", other.id),
                    {"name": attempt, "confirm_affected_records": 0}, format="json",
                )
                self.assertEqual(response.status_code, 400)
        other.refresh_from_db()
        self.assertEqual(other.name, "Beta Boat")

    def test_an_empty_or_huge_name_is_refused(self):
        for value in ["", "   ", None, 5, "x" * 121]:
            with self.subTest(name=value):
                before = Itinerary.objects.count()
                response = self.admin.post(self._url("itinerary"), {"name": value}, format="json")
                self.assertEqual(response.status_code, 400)
                self.assertEqual(Itinerary.objects.count(), before)

    def test_itinerary_stay_nights_and_boat_flag_on_create(self):
        stay = self.admin.post(self._url("itinerary"), {"name": "Three Days", "stay_nights": 3}, format="json")
        boat = self.admin.post(self._url("boat-type"), {"name": "Barge", "requires_capacity_fare": True}, format="json")
        plain = self.admin.post(self._url("boat-type"), {"name": "Raft"}, format="json")

        self.assertEqual(stay.json()["stay_nights"], 3)
        self.assertIs(boat.json()["requires_capacity_fare"], True)
        self.assertIs(plain.json()["requires_capacity_fare"], False)
        for bad in [-1, 61, "3", True, 2.5]:
            with self.subTest(stay_nights=bad):
                self.assertEqual(
                    self.admin.post(self._url("itinerary"), {"name": f"Bad {bad}", "stay_nights": bad}, format="json").status_code,
                    400,
                )
        self.assertEqual(
            self.admin.post(self._url("boat-type"), {"name": "Bad flag", "requires_capacity_fare": "yes"}, format="json").status_code,
            400,
        )

    # -- retiring ----------------------------------------------------------

    def test_the_last_active_option_cannot_be_retired(self):
        for kind, (model, _) in self.KINDS.items():
            with self.subTest(kind=kind):
                keep = model.objects.order_by("id").first()
                model.objects.exclude(pk=keep.pk).update(is_active=False)

                refused = self.admin.post(self._url(kind, keep.id, "deactivate"))

                self.assertEqual(refused.status_code, 409)
                self.assertEqual(refused.json()["code"], "last_active")
                self.assertTrue(model.objects.get(pk=keep.pk).is_active)
                # ...but with a second active one, it is allowed
                other = model.objects.exclude(pk=keep.pk).order_by("id").first()
                self.assertEqual(self.admin.post(self._url(kind, other.id, "reactivate")).status_code, 200)
                self.assertEqual(self.admin.post(self._url(kind, keep.id, "deactivate")).status_code, 200)

    def test_retire_and_bring_back_in_every_kind_and_both_are_idempotent(self):
        for kind, (model, _) in self.KINDS.items():
            with self.subTest(kind=kind):
                option = self._make(kind, 9840, "Seasonal", uses=2)
                logs = ActivityLog.objects.count()

                first = self.admin.post(self._url(kind, option.id, "deactivate"))
                again = self.admin.post(self._url(kind, option.id, "deactivate"))
                self.assertEqual((first.status_code, again.status_code), (200, 200))
                self.assertIs(first.json()["is_active"], False)
                self.assertEqual(ActivityLog.objects.count(), logs + 1)  # the repeat changed nothing: not logged

                back = self.admin.post(self._url(kind, option.id, "reactivate"))
                back_again = self.admin.post(self._url(kind, option.id, "reactivate"))
                self.assertEqual((back.status_code, back_again.status_code), (200, 200))
                self.assertIs(back.json()["is_active"], True)
                self.assertEqual(ActivityLog.objects.count(), logs + 2)

    def test_retiring_does_not_hurt_an_existing_record(self):
        from .services.tourism import build_arrival_monitoring_payload

        itinerary = self._make("itinerary", 9850, "Retiring Trail", uses=0)
        record = self._record("OPT-KEEP-1", itinerary=itinerary)
        self._make("itinerary", 9851, "Other Trail", uses=1)

        response = self.admin.post(self._url("itinerary", itinerary.id, "deactivate"))

        self.assertEqual(response.status_code, 200)
        # (a) the record is whole
        record.refresh_from_db()
        self.assertEqual(record.itinerary_id, 9850)
        self.assertEqual(record.total_visitors, 4)
        # (b) the report still names it
        rows = build_arrival_monitoring_payload({"date": "2026-04-02"})["rows"]
        self.assertEqual({r["survey_id"]: r["itinerary"] for r in rows}["OPT-KEEP-1"], "Retiring Trail")
        # (c) gone from what the mobile app is offered
        self.assertNotIn(9850, self._ids("itinerary", active_only=True))
        self.assertIn(9851, self._ids("itinerary", active_only=True))
        # (d) still in what the web is given, marked
        web = {i["id"]: i for i in self._payload()["itineraries"]}
        self.assertIs(web[9850]["is_active"], False)

    # -- renaming ----------------------------------------------------------

    def test_a_rename_needs_the_exact_count_the_server_counted(self):
        # Three records on this option, one on the other: the right answer is 3.
        mine = self._make("travel-mode", 9860, "Old Ride", uses=3)
        self._make("travel-mode", 9861, "Other Ride", uses=1)
        url = self._url("travel-mode", mine.id)

        missing = self.admin.patch(url, {"name": "New Ride"}, format="json")
        wrong_low = self.admin.patch(url, {"name": "New Ride", "confirm_affected_records": 1}, format="json")
        wrong_high = self.admin.patch(url, {"name": "New Ride", "confirm_affected_records": 4}, format="json")
        as_flag = self.admin.patch(url, {"name": "New Ride", "confirm_affected_records": True}, format="json")
        junk = self.admin.patch(url, {"name": "New Ride", "confirm_affected_records": "three"}, format="json")

        for response in (missing, wrong_low, wrong_high, as_flag, junk):
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.json()["record_count"], 3)  # the real count is given back
            self.assertEqual(response.json()["code"], "confirmation_required")
        mine.refresh_from_db()
        self.assertEqual(mine.name, "Old Ride")  # none of them renamed it

        right = self.admin.patch(url, {"name": "New Ride", "confirm_affected_records": 3}, format="json")
        self.assertEqual(right.status_code, 200)
        self.assertEqual(right.json()["name"], "New Ride")
        mine.refresh_from_db()
        self.assertEqual(mine.name, "New Ride")

    def test_true_is_not_a_count_of_one(self):
        # True == 1 in Python: an option with ONE record must still refuse it.
        option = self._make("visit-purpose", 9864, "One Record", uses=1)

        response = self.admin.patch(
            self._url("visit-purpose", option.id), {"name": "One Record 2", "confirm_affected_records": True}, format="json"
        )

        self.assertEqual(response.status_code, 409)
        option.refresh_from_db()
        self.assertEqual(option.name, "One Record")

    def test_the_count_is_accepted_as_text_digits_too(self):
        mine = self._make("visit-purpose", 9862, "Text Count", uses=2)

        response = self.admin.patch(
            self._url("visit-purpose", mine.id), {"name": "Text Count 2", "confirm_affected_records": "2"}, format="json"
        )

        self.assertEqual(response.status_code, 200)

    def test_renaming_an_option_nobody_uses_still_needs_a_zero(self):
        option = self._make("boat-type", 9863, "Unused Boat", uses=0)
        url = self._url("boat-type", option.id)

        self.assertEqual(self.admin.patch(url, {"name": "Unused 2"}, format="json").status_code, 409)
        self.assertEqual(
            self.admin.patch(url, {"name": "Unused 2", "confirm_affected_records": 0}, format="json").status_code, 200
        )

    def test_an_itinerary_with_no_stay_length_cannot_be_renamed(self):
        option = self._make("itinerary", 9870, "Reads Its Name", uses=2)  # stay_nights is None
        url = self._url("itinerary", option.id)

        refused = self.admin.patch(url, {"name": "Reads Another", "confirm_affected_records": 2}, format="json")

        self.assertEqual(refused.status_code, 409)
        self.assertEqual(refused.json()["code"], "stay_not_set")
        self.assertIn("NAME", refused.json()["detail"])
        option.refresh_from_db()
        self.assertEqual(option.name, "Reads Its Name")

    def test_an_itinerary_with_a_stay_length_can_be_renamed(self):
        option = self._make("itinerary", 9871, "Weekend 2 Days", uses=2, stay_nights=2)

        response = self.admin.patch(
            self._url("itinerary", option.id), {"name": "Two Day Weekend", "confirm_affected_records": 2}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        option.refresh_from_db()
        self.assertEqual((option.name, option.stay_nights), ("Two Day Weekend", 2))

    def test_a_rename_that_also_sets_the_stay_length_is_allowed(self):
        option = self._make("itinerary", 9872, "Weekend 2 Days", uses=1)

        response = self.admin.patch(
            self._url("itinerary", option.id),
            {"name": "Two Day Weekend", "stay_nights": 2, "confirm_affected_records": 1}, format="json",
        )

        self.assertEqual(response.status_code, 200)
        option.refresh_from_db()
        self.assertEqual((option.name, option.stay_nights), ("Two Day Weekend", 2))

    def test_the_other_three_kinds_rename_without_a_stay_length(self):
        for kind in ("travel-mode", "boat-type", "visit-purpose"):
            with self.subTest(kind=kind):
                option = self._make(kind, 9873, "Plain Old", uses=1)
                response = self.admin.patch(
                    self._url(kind, option.id), {"name": "Plain New", "confirm_affected_records": 1}, format="json"
                )
                self.assertEqual(response.status_code, 200)

    def test_changing_the_stay_length_also_needs_the_count(self):
        option = self._make("itinerary", 9874, "Counted Stay", uses=2, stay_nights=1)
        url = self._url("itinerary", option.id)

        self.assertEqual(self.admin.patch(url, {"stay_nights": 3}, format="json").status_code, 409)
        self.assertEqual(
            self.admin.patch(url, {"stay_nights": 3, "confirm_affected_records": 9}, format="json").status_code, 409
        )
        right = self.admin.patch(url, {"stay_nights": 3, "confirm_affected_records": 2}, format="json")
        self.assertEqual(right.status_code, 200)
        self.assertEqual(right.json()["stay_nights"], 3)
        for bad in [-1, 61, "x", True, 1.5]:
            with self.subTest(stay_nights=bad):
                self.assertEqual(
                    self.admin.patch(url, {"stay_nights": bad, "confirm_affected_records": 2}, format="json").status_code,
                    400,
                )

    def test_a_patch_that_changes_nothing_writes_and_logs_nothing(self):
        option = self._make("travel-mode", 9875, "Still", uses=1)
        logs = ActivityLog.objects.count()

        response = self.admin.patch(self._url("travel-mode", option.id), {"name": "Still"}, format="json")
        spaced = self.admin.patch(self._url("travel-mode", option.id), {"name": "  Still "}, format="json")

        self.assertEqual((response.status_code, spaced.status_code), (200, 200))
        self.assertEqual(ActivityLog.objects.count(), logs)

    # -- the cache ---------------------------------------------------------

    def test_every_write_shows_in_the_next_payload(self):
        from django.core.cache import cache

        # POST
        self._payload()  # warms the cache
        self.assertIsNotNone(cache.get(self.CACHE_KEY))
        created = self.admin.post(self._url("visit-purpose"), {"name": "Cache Purpose"}, format="json").json()
        self.assertIsNone(cache.get(self.CACHE_KEY))
        self.assertIn(created["id"], self._ids("visit-purpose"))
        self.assertIn(created["id"], self._ids("visit-purpose", active_only=True))

        # PATCH
        self._payload()
        self.admin.patch(
            self._url("visit-purpose", created["id"]),
            {"name": "Cache Purpose Renamed", "confirm_affected_records": 0}, format="json",
        )
        self.assertIsNone(cache.get(self.CACHE_KEY))
        names = {i["id"]: i["name"] for i in self._payload()["visitPurposes"]}
        self.assertEqual(names[created["id"]], "Cache Purpose Renamed")

        # deactivate
        self._payload(active_only=True)
        self.admin.post(self._url("visit-purpose", created["id"], "deactivate"))
        self.assertIsNone(cache.get(self.CACHE_KEY))
        self.assertNotIn(created["id"], self._ids("visit-purpose", active_only=True))
        self.assertIn(created["id"], self._ids("visit-purpose"))

        # reactivate
        self._payload(active_only=True)
        self.admin.post(self._url("visit-purpose", created["id"], "reactivate"))
        self.assertIsNone(cache.get(self.CACHE_KEY))
        self.assertIn(created["id"], self._ids("visit-purpose", active_only=True))

    def test_a_refused_write_leaves_the_cache_alone(self):
        from django.core.cache import cache

        option = self._make("travel-mode", 9880, "Held", uses=1)
        self._payload()

        self.admin.patch(self._url("travel-mode", option.id), {"name": "Held 2"}, format="json")  # 409
        self.admin.post(self._url("travel-mode"), {"name": ""}, format="json")  # 400

        self.assertIsNotNone(cache.get(self.CACHE_KEY))

    # -- the trace ---------------------------------------------------------

    def test_every_kind_of_write_is_in_the_activity_log(self):
        option = self._make("itinerary", 9890, "Trace Trail", uses=2, stay_nights=1)
        admin_user = self.admin.user

        def newest():
            return ActivityLog.objects.order_by("-id").first()

        before = ActivityLog.objects.count()
        created = self.admin.post(self._url("itinerary"), {"name": "Logged New"}, format="json").json()
        log = newest()
        self.assertEqual((log.user_id, log.module, log.action), (admin_user.pk, "tourism", "create"))
        self.assertEqual((log.record_type, log.record_id), ("Itinerary", str(created["id"])))
        self.assertIn("Logged New", log.record_label)
        self.assertIn("Travel itinerary", log.record_label)

        self.admin.patch(
            self._url("itinerary", option.id), {"name": "Trace Trail 2", "confirm_affected_records": 2}, format="json"
        )
        log = newest()
        self.assertEqual((log.user_id, log.action, log.record_id), (admin_user.pk, "update", str(option.id)))
        self.assertIn("Trace Trail", log.record_label)
        self.assertIn("Trace Trail 2", log.record_label)
        self.assertIn("2 record(s)", log.record_label)

        self.admin.post(self._url("itinerary", option.id, "deactivate"))
        log = newest()
        self.assertEqual((log.user_id, log.action, log.record_id), (admin_user.pk, "update", str(option.id)))
        self.assertIn("deactivated", log.record_label)

        self.admin.post(self._url("itinerary", option.id, "reactivate"))
        log = newest()
        self.assertEqual((log.user_id, log.action), (admin_user.pk, "update"))
        self.assertIn("reactivated", log.record_label)

        self.assertEqual(ActivityLog.objects.count(), before + 4)

    def test_a_refused_write_is_not_logged(self):
        option = self._make("travel-mode", 9891, "Refused", uses=1)
        before = ActivityLog.objects.count()

        self.admin.patch(self._url("travel-mode", option.id), {"name": "Refused 2"}, format="json")  # no count
        self.admin.post(self._url("travel-mode"), {"name": "refused"}, format="json")  # same name
        self._client(ROLE_TOURISM, "tourism_try").post(self._url("travel-mode"), {"name": "Nope"}, format="json")

        self.assertEqual(ActivityLog.objects.count(), before)

    def test_the_model_has_no_change_without_a_migration(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


class ReferenceOptionProtectedNameTests(TestCase):
    """A name the Excel import produces is not renamed (commit 4b).

    The import finds options BY NAME (and finds retired ones too). Rename one
    and the next import creates a new row under the old name, and the reports
    are split between the two. The list of such names lives in ONE place, the
    import's own IMPORT_PRODUCED_NAMES, and the admin API reads it from there.
    """

    KINDS = {
        "itinerary": Itinerary, "travel-mode": TravelMode,
        "boat-type": BoatType, "visit-purpose": VisitPurpose,
    }
    FKS = {"itinerary": "itinerary", "travel-mode": "travel_mode",
           "boat-type": "boat_type", "visit-purpose": "visit_purpose"}

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        user = User.objects.create_user(username="prot_admin", password="Staff@12345")
        UserProfile.objects.create(user=user, role=ROLE_ADMIN)
        self.admin = APIClient()
        self.admin.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

    # -- helpers -----------------------------------------------------------

    @staticmethod
    def _kind_of(model):
        return {Itinerary: "itinerary", TravelMode: "travel-mode",
                BoatType: "boat-type", VisitPurpose: "visit-purpose"}[model]

    def _url(self, kind, option_id=None, action=None):
        url = f"/api/reference-options/{kind}/"
        if option_id is not None:
            url += f"{option_id}/"
        if action:
            url += f"{action}/"
        return url

    def _option_named(self, model, name):
        """The row with this name (any capitals / spacing), or a new one."""
        from .services.online_booking import normalize_key

        for option in model.objects.all():
            if normalize_key(option.name) == normalize_key(name):
                return option
        top = model.objects.order_by("-id").first().id
        return model.objects.create(id=top + 1, name=name)

    def _with_records(self, option, kind, count):
        quezon = Province.objects.get(name="Quezon")
        for i in range(count):
            values = dict(
                survey_id=f"PRT-{kind[:3]}-{option.id}-{i}", full_name="Group",
                contact_number="09170000000", country=Country.objects.get(name="Philippines"),
                region=quezon.region, province=quezon, arrival_date="2026-04-02",
                resort=Resort.objects.order_by("resort_id").first(),
                itinerary=Itinerary.objects.order_by("id").first(),
                travel_mode=TravelMode.objects.order_by("id").first(),
                boat_type=BoatType.objects.order_by("id").first(),
                visit_purpose=VisitPurpose.objects.order_by("id").first(),
                total_visitors=2, filipino_count=2, total_male=2, age_8_59=2, status="arrived",
            )
            values[self.FKS[kind]] = option
            TouristRecord.objects.create(**values)

    def _protected(self):
        from .services import online_booking

        return [
            (self._kind_of(model), name)
            for model, names in online_booking.IMPORT_PRODUCED_NAMES.items()
            for name in names
        ]

    # -- the gate ----------------------------------------------------------

    def test_every_name_the_import_produces_refuses_a_rename(self):
        protected = self._protected()
        self.assertTrue(protected)
        for kind, name in protected:
            with self.subTest(kind=kind, name=name):
                model = self.KINDS[kind]
                option = self._option_named(model, name)
                self._with_records(option, kind, 2)
                stored_name = option.name

                # even with the right count, even with a stay length: the name wins
                response = self.admin.patch(
                    self._url(kind, option.id),
                    {"name": f"{name} Renamed", "confirm_affected_records": 2, "stay_nights": 1}
                    if kind == "itinerary" else {"name": f"{name} Renamed", "confirm_affected_records": 2},
                    format="json",
                )

                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.json()["code"], "rename_protected")
                self.assertIn("Excel import", response.json()["detail"])
                option.refresh_from_db()
                self.assertEqual(option.name, stored_name)  # the database is unchanged
                if kind == "itinerary":
                    self.assertIsNone(option.stay_nights)  # and nothing else was saved

    def test_a_protected_name_is_protected_in_any_dress(self):
        # stored as "private vehicle" (lower case, extra space): still protected
        option = TravelMode.objects.create(id=9901, name="private   vehicle")
        # (an existing row with that name would make this a duplicate; free the seeded one)
        TravelMode.objects.filter(name="Private Vehicle").update(name="Private Vehicle (seed)")

        response = self.admin.patch(
            self._url("travel-mode", option.id), {"name": "Car", "confirm_affected_records": 0}, format="json"
        )

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "rename_protected")
        option.refresh_from_db()
        self.assertEqual(option.name, "private   vehicle")

    def test_a_name_that_is_not_produced_still_renames_with_the_right_count(self):
        option = TravelMode.objects.create(id=9902, name="Public Bus Line")
        self._with_records(option, "travel-mode", 3)

        response = self.admin.patch(
            self._url("travel-mode", option.id), {"name": "Bus Line", "confirm_affected_records": 3}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["name"], "Bus Line")
        self.assertIs(response.json()["rename_protected"], False)
        option.refresh_from_db()
        self.assertEqual(option.name, "Bus Line")

    def test_a_retired_protected_option_is_still_protected_but_can_be_retired(self):
        itinerary = self._option_named(Itinerary, "Same Day")

        retired = self.admin.post(self._url("itinerary", itinerary.id, "deactivate"))
        self.assertEqual(retired.status_code, 200)  # retiring is safe: the import reuses the row
        itinerary.refresh_from_db()
        self.assertIs(itinerary.is_active, False)
        self.assertIs(retired.json()["is_active"], False)

        rename = self.admin.patch(
            self._url("itinerary", itinerary.id),
            {"name": "Day Trip", "stay_nights": 0, "confirm_affected_records": 0}, format="json",
        )

        self.assertEqual(rename.status_code, 409)
        self.assertEqual(rename.json()["code"], "rename_protected")
        itinerary.refresh_from_db()
        self.assertEqual(itinerary.name, "Same Day")

    # -- what the list says ------------------------------------------------

    def test_the_list_marks_exactly_the_protected_options(self):
        protected = {(kind, name.lower()) for kind, name in self._protected()}
        for kind in self.KINDS:
            self._option_named(self.KINDS[kind], "Totally Ordinary Name")  # one that is not protected
        for kind, model in self.KINDS.items():
            with self.subTest(kind=kind):
                rows = self.admin.get(self._url(kind)).json()
                self.assertTrue(rows)
                for row in rows:
                    expected = (kind, " ".join(row["name"].split()).lower()) in protected
                    self.assertIs(row["rename_protected"], expected, row["name"])
                if kind == "boat-type":
                    self.assertFalse(any(r["rename_protected"] for r in rows))  # reached by id, not by name
                self.assertIs(
                    {r["name"]: r for r in rows}["Totally Ordinary Name"]["rename_protected"], False
                )

    def test_a_protected_option_shows_as_protected_after_other_writes(self):
        created = self.admin.post(self._url("travel-mode"), {"name": "Brand New Ride"}, format="json")
        self.assertIs(created.json()["rename_protected"], False)
        # a NEW option cannot take a protected name that already exists
        again = self.admin.post(self._url("travel-mode"), {"name": "public utility vehicle"}, format="json")
        self.assertEqual(again.status_code, 400)

    # -- ONE source --------------------------------------------------------

    def test_the_gate_follows_the_list_the_import_owns(self):
        from .services import online_booking

        ordinary = self._option_named(Itinerary, "Trail X")
        was_protected = self._option_named(Itinerary, "Same Day")
        Itinerary.objects.filter(pk__in=[ordinary.pk, was_protected.pk]).update(stay_nights=2)

        with patch.dict(online_booking.IMPORT_PRODUCED_NAMES, {Itinerary: ("Trail X",)}, clear=True):
            now_protected = self.admin.patch(
                self._url("itinerary", ordinary.id), {"name": "Trail Y", "confirm_affected_records": 0}, format="json"
            )
            no_longer = self.admin.patch(
                self._url("itinerary", was_protected.id), {"name": "Day Trip", "confirm_affected_records": 0}, format="json"
            )
            listing = {r["name"]: r["rename_protected"] for r in self.admin.get(self._url("itinerary")).json()}

        self.assertEqual(now_protected.status_code, 409)
        self.assertEqual(now_protected.json()["code"], "rename_protected")
        self.assertEqual(no_longer.status_code, 200)
        self.assertIs(listing["Trail X"], True)
        self.assertIs(listing["Day Trip"], False)
        ordinary.refresh_from_db()
        was_protected.refresh_from_db()
        self.assertEqual((ordinary.name, was_protected.name), ("Trail X", "Day Trip"))

    def test_the_list_is_exactly_what_the_import_can_produce(self):
        # If the normalizers can say a name that is not in the list (or the list
        # has a name no normalizer says), the two have drifted apart.
        from .services import online_booking as ob

        produced = {
            Itinerary: {ob.normalize_itinerary(v) for v in
                        ["2 nights", "3 night", "4 nights", "5 nights", "day trip", "Overnight", ""]},
            TravelMode: {ob.normalize_travel_mode(v) for v in ["public utility", ""]},
            VisitPurpose: {ob.normalize_purpose(v) for v in ["team building", ""]},
        }
        for model, names in produced.items():
            with self.subTest(model=model.__name__):
                self.assertEqual(names, set(ob.IMPORT_PRODUCED_NAMES[model]))
        # free text is NOT protected: it becomes whatever the sheet says
        self.assertEqual(ob.normalize_itinerary("Que"), "Que")
        self.assertNotIn("Que", ob.IMPORT_PRODUCED_NAMES[Itinerary])
        self.assertEqual(ob.normalize_travel_mode("Jeepney"), "Jeepney")
        self.assertEqual(ob.normalize_purpose("swimming"), "Swimming")

    def test_the_api_and_the_import_compare_names_the_same_way(self):
        from .services.online_booking import normalize_key
        from .views.reference_options import name_key

        for text in ["Same Day", "same day", " Same  Day ", "SAME\tDAY", "Weekend 2 Days", ""]:
            with self.subTest(text=text):
                self.assertEqual(name_key(text), normalize_key(text))

    # -- what the import does with a retired option ------------------------

    def test_the_import_reuses_a_retired_row_instead_of_making_a_new_one(self):
        # This is why retiring a protected option is safe and needs no gate.
        from .services.online_booking import ReferenceResolver, normalize_itinerary

        row = self._option_named(Itinerary, "Same Day")
        Itinerary.objects.filter(pk=row.pk).update(is_active=False)
        before = Itinerary.objects.count()

        resolved = ReferenceResolver(commit=True).named(Itinerary, normalize_itinerary("day trip"))

        self.assertEqual(resolved.pk, row.pk)
        self.assertEqual(Itinerary.objects.count(), before)

    def test_renaming_does_make_the_import_create_a_duplicate(self):
        # The harm the gate prevents, shown by doing it directly (not through
        # the API, which refuses): rename the row, and the import adds a new
        # one under the old name.
        from .services.online_booking import ReferenceResolver, normalize_itinerary

        row = self._option_named(Itinerary, "Same Day")
        Itinerary.objects.filter(pk=row.pk).update(name="Day Trip")
        before = Itinerary.objects.count()

        resolved = ReferenceResolver(commit=True).named(Itinerary, normalize_itinerary("day trip"))

        self.assertNotEqual(resolved.pk, row.pk)
        self.assertEqual(resolved.name, "Same Day")
        self.assertEqual(Itinerary.objects.count(), before + 1)

    def test_the_model_has_no_change_without_a_migration(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


# The 12 itineraries of the live system with the record split of 2026-10-17:
# (name, [visitors of the ARRIVED records], [visitors of the records that are not arrived]).
# Overnight 3 records / 9 visitors, Same Day 5/14, 2 Nights 5/20, 3 Nights 2/8,
# 4 Nights 2/4, 5 Nights and above 1/2, Eco Adventure Trail 1/4, the rest none:
# 19 records, 61 visitors, of which 8 records / 31 visitors arrived.
LIVE_LIKE_SPLIT = [
    ("Overnight", [4], [3, 2]),
    ("Same Day", [5], [4, 2, 2, 1]),
    ("2 Nights", [6, 4], [5, 3, 2]),
    ("3 Nights", [4], [4]),
    ("4 Nights", [3, 1], []),
    ("5 Nights and above", [], [2]),
    ("6 Nights", [], []),
    ("Que", [], []),
    ("Island Hopping Circuit", [], []),
    ("Beach and Resort Day Tour", [], []),
    ("Heritage and Lighthouse Stop", [], []),
    ("Eco Adventure Trail", [4], []),
]


def _live_like_fixture():
    """The 12 itineraries (ids 1..12, live names, stay_nights not set) and the
    live record split, on two resorts."""
    ensure_test_reference_tables()
    quezon = Province.objects.get(name="Quezon")
    resorts = list(Resort.objects.order_by("resort_id")[:2])
    count = 0
    for index, (name, arrived, others) in enumerate(LIVE_LIKE_SPLIT):
        itinerary, _ = Itinerary.objects.update_or_create(
            id=index + 1, defaults={"name": name, "is_active": True, "stay_nights": None}
        )
        for kind, visitors_list in (("arrived", arrived), ("other", others)):
            for visitors in visitors_list:
                count += 1
                status = "arrived" if kind == "arrived" else ("no_show" if count % 3 else "pending")
                TouristRecord.objects.create(
                    survey_id=f"LV-{count:02d}", full_name="Group", contact_number="09170000000",
                    country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
                    arrival_date="2026-04-02", resort=resorts[count % 2], itinerary=itinerary,
                    travel_mode=TravelMode.objects.first(), boat_type=BoatType.objects.first(),
                    visit_purpose=VisitPurpose.objects.first(), total_visitors=visitors,
                    filipino_count=visitors, total_male=visitors, age_8_59=visitors, status=status,
                )
    return count


class ItineraryStayFillTests(TestCase):
    """Migration 0052: the approved stay_nights, by name, with a reverse."""

    APPROVED = {
        1: ("Overnight", 1), 2: ("Same Day", 0), 3: ("2 Nights", 2), 4: ("3 Nights", 3),
        5: ("4 Nights", 4), 6: ("5 Nights and above", 5), 7: ("6 Nights", 6), 8: ("Que", None),
        9: ("Island Hopping Circuit", None), 10: ("Beach and Resort Day Tour", 0),
        11: ("Heritage and Lighthouse Stop", None), 12: ("Eco Adventure Trail", None),
    }

    def setUp(self):
        import importlib

        from django.apps import apps as django_apps

        self.module = importlib.import_module("api.migrations.0052_set_itinerary_stay_nights")
        self.apps = django_apps

    def _fill(self):
        self.module.set_stay_nights(self.apps, None)

    def _unfill(self):
        self.module.clear_stay_nights(self.apps, None)

    def _values(self):
        return {i.id: i.stay_nights for i in Itinerary.objects.filter(id__lte=12).order_by("id")}

    # -- the gate: does any live number move? ------------------------------

    def test_the_fixture_is_the_live_split(self):
        self.assertEqual(_live_like_fixture(), 19)
        self.assertEqual(TouristRecord.objects.count(), 19)
        self.assertEqual(sum(TouristRecord.objects.values_list("total_visitors", flat=True)), 61)
        arrived = TouristRecord.objects.filter(status="arrived")
        self.assertEqual((arrived.count(), sum(arrived.values_list("total_visitors", flat=True))), (8, 31))
        by_name = {
            r["itinerary__name"]: (r["n"], r["v"])
            for r in TouristRecord.objects.values("itinerary__name").annotate(n=Count("survey_id"), v=Sum("total_visitors"))
        }
        self.assertEqual(by_name, {
            "Overnight": (3, 9), "Same Day": (5, 14), "2 Nights": (5, 20), "3 Nights": (2, 8),
            "4 Nights": (2, 4), "5 Nights and above": (1, 2), "Eco Adventure Trail": (1, 4),
        })

    def test_no_number_of_the_live_reports_moves_when_the_fill_runs(self):
        _live_like_fixture()
        before = _stay_report_numbers()

        self._fill()
        after = _stay_report_numbers()

        # the live figures the baseline shows
        self.assertEqual(before["dashboard_stay_type"], {"dayTour": 5, "overnight": 26})
        self.assertEqual(before["stay_type_answer"], "Same-day visitors: 5; overnight or multi-day visitors: 26.")
        self.assertIn("1.8 night(s)", before["average_stay_answer"])
        # and not one of them is different afterwards
        self.assertEqual(after, before)
        self.assertEqual(after["dashboard_stay_type"], {"dayTour": 5, "overnight": 26})
        self.assertIn("1.8 night(s)", after["average_stay_answer"])
        # the fill really ran
        self.assertEqual(self._values()[3], 2)

    def test_the_arrival_rows_of_every_record_stay_the_same(self):
        _live_like_fixture()
        before = _stay_report_numbers()["arrival_rows"]
        self.assertEqual(len(before), 8)  # the 8 arrived records

        self._fill()

        self.assertEqual(_stay_report_numbers()["arrival_rows"], before)

    # -- what the fill sets ------------------------------------------------

    def test_the_fill_sets_exactly_the_approved_values(self):
        _live_like_fixture()

        self._fill()

        self.assertEqual(self._values(), {pk: nights for pk, (name, nights) in self.APPROVED.items()})

    def test_it_matches_by_name_and_skips_a_row_whose_name_does_not_fit_its_id(self):
        _live_like_fixture()
        Itinerary.objects.filter(pk=7).update(name="Seven Nights")  # id 7 but not "6 Nights"
        Itinerary.objects.filter(pk=3).update(name="  2   NIGHTS ")  # same name in another dress: fits

        import io
        from contextlib import redirect_stdout

        out = io.StringIO()
        with redirect_stdout(out):
            self._fill()

        values = self._values()
        self.assertIsNone(values[7])  # skipped, not guessed
        self.assertEqual(values[3], 2)  # matched in spite of capitals and spaces
        self.assertIn("id 7", out.getvalue())
        self.assertIn("Seven Nights", out.getvalue())

    def test_a_name_in_the_wrong_row_is_not_found_by_its_name_either(self):
        _live_like_fixture()
        # "6 Nights" moved to another id: the mapping's id 7 now holds another name
        Itinerary.objects.filter(pk=7).update(name="Que")
        Itinerary.objects.filter(pk=8).update(name="6 Nights")

        self._fill()

        values = self._values()
        self.assertIsNone(values[7])
        self.assertIsNone(values[8])

    def test_a_missing_row_is_skipped_without_a_crash(self):
        _live_like_fixture()
        TouristRecord.objects.filter(itinerary_id=12).delete()
        Itinerary.objects.filter(pk=12).delete()

        self._fill()

        self.assertEqual(self._values()[2], 0)
        self.assertNotIn(12, self._values())

    def test_a_value_an_admin_already_set_is_never_overwritten(self):
        _live_like_fixture()
        Itinerary.objects.filter(pk=3).update(stay_nights=9)  # the approved value is 2

        self._fill()

        self.assertEqual(self._values()[3], 9)
        self.assertEqual(self._values()[4], 3)  # a row that was not set still gets its value

    def test_running_it_twice_changes_nothing_more(self):
        _live_like_fixture()
        self._fill()
        once = self._values()

        self._fill()

        self.assertEqual(self._values(), once)

    def test_the_rows_left_not_set_are_still_not_set(self):
        _live_like_fixture()

        self._fill()

        for pk in (8, 9, 11, 12):
            with self.subTest(pk=pk):
                self.assertIsNone(self._values()[pk])

    def test_it_changes_no_other_column_and_no_other_table(self):
        _live_like_fixture()
        before_itineraries = list(Itinerary.objects.order_by("id").values("id", "name", "is_active"))
        before_modes = list(TravelMode.objects.order_by("id").values("id", "name", "is_active"))
        before_records = list(TouristRecord.objects.order_by("survey_id").values_list("survey_id", "itinerary_id", "status", "total_visitors"))

        self._fill()

        self.assertEqual(list(Itinerary.objects.order_by("id").values("id", "name", "is_active")), before_itineraries)
        self.assertEqual(list(TravelMode.objects.order_by("id").values("id", "name", "is_active")), before_modes)
        self.assertEqual(
            list(TouristRecord.objects.order_by("survey_id").values_list("survey_id", "itinerary_id", "status", "total_visitors")),
            before_records,
        )

    # -- the reverse -------------------------------------------------------

    def test_the_reverse_puts_every_stay_nights_back_to_none(self):
        _live_like_fixture()
        before = _stay_report_numbers()
        self._fill()
        self.assertTrue(any(v is not None for v in self._values().values()))

        self._unfill()

        self.assertEqual(set(self._values().values()), {None})
        self.assertEqual(_stay_report_numbers(), before)

    # -- what the fix is for -----------------------------------------------

    def test_six_nights_is_six_nights_after_the_fill_and_was_one_before(self):
        from .services.tourism import itinerary_stay

        _live_like_fixture()
        itinerary = Itinerary.objects.get(pk=7)
        before = itinerary_stay(itinerary.name, itinerary.stay_nights)

        self._fill()
        itinerary.refresh_from_db()
        after = itinerary_stay(itinerary.name, itinerary.stay_nights)

        self.assertEqual(before, (False, 1))  # the name reading missed the 6
        self.assertEqual(after, (False, 6))

    def test_a_new_six_nights_record_counts_six_nights_in_the_average_stay(self):
        _live_like_fixture()
        quezon = Province.objects.get(name="Quezon")

        def add_record(survey_id):
            return TouristRecord.objects.create(
                survey_id=survey_id, full_name="Group", contact_number="09170000000",
                country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
                arrival_date="2026-04-02", resort=Resort.objects.order_by("resort_id").first(),
                itinerary=Itinerary.objects.get(pk=7), travel_mode=TravelMode.objects.first(),
                boat_type=BoatType.objects.first(), visit_purpose=VisitPurpose.objects.first(),
                total_visitors=3, filipino_count=3, total_male=3, age_8_59=3, status="arrived",
            )

        add_record("LV-NEW-6")
        before = _stay_report_numbers()
        self._fill()
        after = _stay_report_numbers()

        # 31 + 3 = 34 visitors; nights 56 + 3 x 1 = 59 before, 56 + 3 x 6 = 74 after
        self.assertIn("1.7 night(s)", before["average_stay_answer"])  # 59 / 34
        self.assertIn("2.2 night(s)", after["average_stay_answer"])  # 74 / 34
        # still an overnight stay either way
        self.assertEqual(before["dashboard_stay_type"], after["dashboard_stay_type"])

    def test_a_new_record_on_the_beach_and_resort_day_tour_is_same_day_before_and_after(self):
        _live_like_fixture()
        quezon = Province.objects.get(name="Quezon")
        TouristRecord.objects.create(
            survey_id="LV-NEW-10", full_name="Group", contact_number="09170000000",
            country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
            arrival_date="2026-04-02", resort=Resort.objects.order_by("resort_id").first(),
            itinerary=Itinerary.objects.get(pk=10), travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(), visit_purpose=VisitPurpose.objects.first(),
            total_visitors=7, filipino_count=7, total_male=7, age_8_59=7, status="arrived",
        )
        before = _stay_report_numbers()
        self._fill()
        after = _stay_report_numbers()

        self.assertEqual(before["dashboard_stay_type"], {"dayTour": 12, "overnight": 26})
        self.assertEqual(after, before)

    def test_the_rows_left_not_set_read_their_name_with_the_same_result(self):
        from .services.tourism import itinerary_stay

        _live_like_fixture()
        names = {pk: Itinerary.objects.get(pk=pk).name for pk in (8, 9, 11, 12)}
        before = {pk: itinerary_stay(name, Itinerary.objects.get(pk=pk).stay_nights) for pk, name in names.items()}

        self._fill()

        after = {pk: itinerary_stay(name, Itinerary.objects.get(pk=pk).stay_nights) for pk, name in names.items()}
        self.assertEqual(after, before)
        self.assertEqual(set(after.values()), {(False, 1)})

    def test_the_migration_is_data_only(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)
        self.assertEqual(self.module.Migration.dependencies, [("api", "0051_reference_options_active_and_stay")])


class ItineraryStayFillMigrationTests(TransactionTestCase):
    """Migration 0052 through Django's own migrate, forward AND backward."""

    BEFORE = [("api", "0051_reference_options_active_and_stay")]
    AFTER = [("api", "0052_set_itinerary_stay_nights")]

    def tearDown(self):
        # leave the test database at the latest migration for the tests after this
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())

    def _model(self, target):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        return executor.loader.project_state(target).apps.get_model("api", "Itinerary")

    def _values(self, model):
        return {row.id: row.stay_nights for row in model.objects.order_by("id")}

    def test_forward_sets_the_values_and_backward_clears_every_one(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.BEFORE)  # undo 0052 (a no-op on the empty tables)
        model = self._model(self.BEFORE)
        model.objects.all().delete()
        for pk, name in [(2, "Same Day"), (3, "2 Nights"), (7, "6 Nights"), (8, "Que"), (10, "Beach and Resort Day Tour")]:
            model.objects.create(id=pk, name=name)
        self.assertEqual(set(self._values(model).values()), {None})

        MigrationExecutor(connection).migrate(self.AFTER)
        model = self._model(self.AFTER)
        self.assertEqual(self._values(model), {2: 0, 3: 2, 7: 6, 8: None, 10: 0})

        # an admin sets one of the rows that was left not set
        model.objects.filter(pk=8).update(stay_nights=5)

        MigrationExecutor(connection).migrate(self.BEFORE)  # the reverse of 0052
        model = self._model(self.BEFORE)
        self.assertEqual(set(self._values(model).values()), {None})  # all of them, as specified

        MigrationExecutor(connection).migrate(self.AFTER)  # and forward again
        self.assertEqual(self._values(self._model(self.AFTER)), {2: 0, 3: 2, 7: 6, 8: None, 10: 0})


# The four hotlines written in the mobile app, word for word and in its order
# (mobile/lib/screens/tourism_screens.dart, ProfilePage). The app shows each as
# a title and one line of "number"; the desk's line is "number • email".
EMERGENCY_FOUR = [
    # (title, the line the app shows under it, icon)
    ("Mauban MDRRMO / Rescue 24/7", "0998-598-8422 / (042) 784-0123", "rescue"),
    ("Mauban Municipal Police (PNP)", "0998-598-5777 / (042) 784-0200", "police"),
    ("Coast Guard Sub-Station Mauban", "0917-842-7643", "coast_guard"),
    ("Mauban Tourism Information Desk", "(042) 784-0555 \u2022 tourism@mauban.gov.ph", "help_desk"),
]


def _shown_line(contact):
    """The one line the app shows: the numbers, then the email after a bullet."""
    return contact.contact_numbers + (f" \u2022 {contact.email}" if contact.email else "")


class EmergencyContactTests(TestCase):
    """The table of emergency contacts and what the mobile bootstrap sends of it."""

    BOOTSTRAP = "/api/mobile/tourism/bootstrap/"
    PUBLIC_KEYS = {"id", "name", "contact_numbers", "email", "description", "icon"}
    OLD_BOOTSTRAP_KEYS = {
        "referenceTables", "destinations", "featuredDestinations", "barangays", "notifications", "theme",
    }

    def setUp(self):
        import importlib

        from django.apps import apps as django_apps

        ensure_test_reference_tables()
        self.module = importlib.import_module("api.migrations.0054_seed_emergency_contacts")
        self.apps = django_apps
        EmergencyContact.objects.all().delete()

    def _seed_the_four(self):
        EmergencyContact.objects.all().delete()
        self.module.add_contacts(self.apps, None)

    def _payload(self):
        response = APIClient().get(self.BOOTSTRAP)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.json()

    # -- the table ---------------------------------------------------------

    def test_the_fields_are_exactly_these_and_there_is_no_colour_or_hours(self):
        names = {f.name for f in EmergencyContact._meta.get_fields()}
        self.assertEqual(
            names,
            {"id", "name", "contact_numbers", "email", "description",
             "icon", "display_order", "is_active"},
        )
        self.assertFalse({"hours", "category", "color", "colour"} & names)

    def test_the_icons_are_these_twelve(self):
        self.assertEqual(
            [value for value, _ in EmergencyContact.ICON_CHOICES],
            ["rescue", "police", "coast_guard", "fire", "hospital", "ambulance",
             "help_desk", "lgu_office", "power", "water", "port", "other"],
        )

    def test_a_new_contact_is_active_and_of_icon_other_with_the_optional_fields_empty(self):
        contact = EmergencyContact.objects.create(name="Barangay Desk", contact_numbers="0917-000-0000")
        contact.refresh_from_db()
        self.assertIs(contact.is_active, True)
        self.assertEqual(contact.icon, "other")
        self.assertEqual((contact.email, contact.description), ("", ""))
        self.assertEqual(contact.display_order, 0)

    def test_a_row_put_in_without_the_column_is_active(self):
        field = EmergencyContact._meta.get_field("is_active")
        self.assertIs(field.db_default, True)
        table = EmergencyContact._meta.db_table
        with connection.cursor() as cursor:
            cursor.execute(
                f"INSERT INTO {table} (name, contact_numbers, email, description, icon, display_order) "
                "VALUES ('Raw row', '0917-111-1111', '', '', 'other', 0)"
            )
        self.assertIs(EmergencyContact.objects.get(name="Raw row").is_active, True)

    def test_the_order_is_display_order_then_id(self):
        # display orders that run AGAINST the ids, and a tie
        late = EmergencyContact.objects.create(name="Created first", contact_numbers="1", display_order=30)
        early = EmergencyContact.objects.create(name="Created second", contact_numbers="2", display_order=10)
        tie_a = EmergencyContact.objects.create(name="Tie A", contact_numbers="3", display_order=20)
        tie_b = EmergencyContact.objects.create(name="Tie B", contact_numbers="4", display_order=20)

        self.assertEqual(
            [c.pk for c in EmergencyContact.objects.all()],
            [early.pk, tie_a.pk, tie_b.pk, late.pk],
        )

    # -- the four, from the data migration ---------------------------------

    def test_the_four_are_exactly_what_the_app_has_and_in_its_order(self):
        self._seed_the_four()

        rows = list(EmergencyContact.objects.all())

        self.assertEqual(
            [(c.name, _shown_line(c), c.icon) for c in rows],
            EMERGENCY_FOUR,
        )
        self.assertEqual([c.icon for c in rows], ["rescue", "police", "coast_guard", "help_desk"])
        self.assertTrue(all(c.is_active for c in rows))
        orders = [c.display_order for c in rows]
        self.assertEqual(orders, sorted(orders))
        self.assertEqual(len(set(orders)), 4)

    def test_the_desk_keeps_its_number_and_its_email_apart(self):
        self._seed_the_four()
        desk = EmergencyContact.objects.get(icon="help_desk")
        self.assertEqual(desk.contact_numbers, "(042) 784-0555")
        self.assertEqual(desk.email, "tourism@mauban.gov.ph")

    def test_it_matches_the_dart_source_on_disk(self):
        # The app's written list stays as the fallback for good. If someone
        # changes it, this fails until the seed (and a decision) catches up.
        import os
        import re

        from django.conf import settings

        path = os.path.join(settings.BASE_DIR, "..", "mobile", "lib", "screens", "tourism_screens.dart")
        if not os.path.exists(path):
            self.skipTest("the mobile source is not next to the backend here")
        source = open(path, encoding="utf-8").read()
        cards = re.findall(r"_EmergencyContactCard\(\s*icon:[^,]+,\s*title:\s*'([^']*)',\s*number:\s*'([^']*)',", source)

        self.assertEqual(cards, [(title, line) for title, line, _ in EMERGENCY_FOUR])

    def test_running_the_seed_again_adds_nothing(self):
        self._seed_the_four()
        self.module.add_contacts(self.apps, None)
        self.assertEqual(EmergencyContact.objects.count(), 4)

    def test_the_seed_never_overwrites_what_an_admin_changed(self):
        self._seed_the_four()
        police = EmergencyContact.objects.get(icon="police")
        police.contact_numbers = "0999-000-0000"
        police.save()

        self.module.add_contacts(self.apps, None)

        police.refresh_from_db()
        self.assertEqual(police.contact_numbers, "0999-000-0000")
        self.assertEqual(EmergencyContact.objects.count(), 4)

    def test_the_reverse_deletes_those_four_and_nothing_else(self):
        self._seed_the_four()
        mine = EmergencyContact.objects.create(name="Barangay Lupon", contact_numbers="0917-222-2222")
        same_name_other_icon = EmergencyContact.objects.create(
            name="Mauban Municipal Police (PNP)", contact_numbers="x", icon="other"
        )

        self.module.remove_contacts(self.apps, None)

        remaining = set(EmergencyContact.objects.values_list("pk", flat=True))
        self.assertEqual(remaining, {mine.pk, same_name_other_icon.pk})
        self.assertFalse(EmergencyContact.objects.filter(icon__in=["rescue", "coast_guard", "help_desk"]).exists())

    # -- the payload -------------------------------------------------------

    def test_the_bootstrap_sends_the_active_contacts_in_order(self):
        self._seed_the_four()

        sent = self._payload()["emergencyContacts"]

        self.assertEqual(
            [(c["name"], c["contact_numbers"], c["email"], c["icon"]) for c in sent],
            [
                ("Mauban MDRRMO / Rescue 24/7", "0998-598-8422 / (042) 784-0123", "", "rescue"),
                ("Mauban Municipal Police (PNP)", "0998-598-5777 / (042) 784-0200", "", "police"),
                ("Coast Guard Sub-Station Mauban", "0917-842-7643", "", "coast_guard"),
                ("Mauban Tourism Information Desk", "(042) 784-0555", "tourism@mauban.gov.ph", "help_desk"),
            ],
        )

    def test_each_contact_has_exactly_the_public_keys(self):
        self._seed_the_four()
        for contact in self._payload()["emergencyContacts"]:
            self.assertEqual(set(contact), self.PUBLIC_KEYS)  # no display_order, no is_active

    def test_an_inactive_contact_is_not_sent_but_is_still_in_the_database(self):
        self._seed_the_four()
        EmergencyContact.objects.filter(icon="coast_guard").update(is_active=False)

        names = [c["name"] for c in self._payload()["emergencyContacts"]]

        self.assertNotIn("Coast Guard Sub-Station Mauban", names)
        self.assertEqual(len(names), 3)
        self.assertEqual(EmergencyContact.objects.count(), 4)
        self.assertTrue(EmergencyContact.objects.filter(name="Coast Guard Sub-Station Mauban").exists())

    def test_the_payload_follows_display_order(self):
        self._seed_the_four()
        # the desk first, the police last
        EmergencyContact.objects.filter(icon="help_desk").update(display_order=1)
        EmergencyContact.objects.filter(icon="police").update(display_order=99)

        icons = [c["icon"] for c in self._payload()["emergencyContacts"]]

        self.assertEqual(icons, ["help_desk", "rescue", "coast_guard", "police"])

    def test_no_contacts_is_an_empty_list_not_a_missing_key(self):
        self.assertEqual(self._payload()["emergencyContacts"], [])

    def test_the_keys_the_app_already_reads_are_all_still_there_and_unchanged(self):
        # without contacts, then with four: only the new key differs
        without = self._payload()
        self._seed_the_four()
        with_four = self._payload()

        self.assertEqual(set(with_four), self.OLD_BOOTSTRAP_KEYS | {"emergencyContacts"})
        for key in self.OLD_BOOTSTRAP_KEYS:
            with self.subTest(key=key):
                self.assertIn(key, without)
                self.assertEqual(with_four[key], without[key])
        self.assertTrue(with_four["referenceTables"]["itineraries"])
        self.assertTrue(with_four["destinations"])

    def test_the_model_has_no_change_without_a_migration(self):
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


class EmergencyContactMigrationTests(TransactionTestCase):
    """Migration 0054 through Django's own migrate, forward and backward."""

    BEFORE = [("api", "0053_emergency_contact")]
    AFTER = [("api", "0054_seed_emergency_contacts")]

    def tearDown(self):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())

    def _model(self, target):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        return executor.loader.project_state(target).apps.get_model("api", "EmergencyContact")

    def test_forward_puts_in_the_four_and_backward_takes_out_only_them(self):
        MigrationExecutor(connection).migrate(self.BEFORE)  # undo 0054
        model = self._model(self.BEFORE)
        model.objects.all().delete()
        mine = model.objects.create(name="Barangay Lupon", contact_numbers="0917-222-2222")

        MigrationExecutor(connection).migrate(self.AFTER)
        model = self._model(self.AFTER)
        self.assertEqual(
            [(c.name, c.icon) for c in model.objects.order_by("display_order", "id") if c.pk != mine.pk],
            [(name, icon) for name, _line, icon in EMERGENCY_FOUR],
        )
        self.assertEqual(model.objects.count(), 5)

        MigrationExecutor(connection).migrate(self.BEFORE)  # the reverse of 0054
        model = self._model(self.BEFORE)
        self.assertEqual(list(model.objects.values_list("pk", flat=True)), [mine.pk])

        MigrationExecutor(connection).migrate(self.AFTER)  # and forward again
        self.assertEqual(self._model(self.AFTER).objects.count(), 5)


class EmergencyContactApiTests(TestCase):
    """Admin CRUD of the emergency contacts: guards, audit trail, and what the app is sent."""

    LIST = "/api/emergency-contacts/"
    REORDER = "/api/emergency-contacts/reorder/"
    BOOTSTRAP = "/api/mobile/tourism/bootstrap/"

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        EmergencyContact.objects.all().delete()
        self.admin = self._client(ROLE_ADMIN, "ec_admin")
        self.a = EmergencyContact.objects.create(
            name="Alpha Rescue", contact_numbers="111-1111", icon="rescue", display_order=10)
        self.b = EmergencyContact.objects.create(
            name="Bravo Police", contact_numbers="222-2222 / 333-3333", icon="police", display_order=20)
        self.c = EmergencyContact.objects.create(
            name="Charlie Desk", contact_numbers="444-4444", icon="help_desk", display_order=30)

    def _client(self, role=None, username=None):
        client = APIClient()
        if role is not None:
            user = User.objects.create_user(username=username or f"ec_{role}", password="Staff@12345")
            UserProfile.objects.create(user=user, role=role)
            client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
            client.user = user
        return client

    def _url(self, contact=None, action=None):
        url = self.LIST
        if contact is not None:
            url += f"{contact.pk}/"
        if action:
            url += f"{action}/"
        return url

    def _payload_names(self):
        response = APIClient().get(self.BOOTSTRAP)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return [row["name"] for row in response.json()["emergencyContacts"]]

    def _admin_names(self):
        return [row["name"] for row in self.admin.get(self.LIST).json()]

    def _all_calls(self):
        return [
            ("get", self.LIST),
            ("post", self.LIST),
            ("patch", self._url(self.a)),
            ("post", self._url(self.a, "deactivate")),
            ("post", self._url(self.a, "reactivate")),
            ("post", self.REORDER),
        ]

    def _logs(self):
        return ActivityLog.objects.filter(record_type="EmergencyContact")

    # -- guard 1: admin only ----------------------------------------------

    def test_nobody_logged_out_gets_in_on_any_endpoint(self):
        before = list(EmergencyContact.objects.values_list("id", "name", "is_active", "display_order"))
        for method, url in self._all_calls():
            response = getattr(APIClient(), method)(url, {"name": "Intruder"}, format="json")
            self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED, (method, url))
        self.assertEqual(
            before, list(EmergencyContact.objects.values_list("id", "name", "is_active", "display_order")))
        self.assertEqual(self._logs().count(), 0)

    def test_a_logged_in_non_admin_is_refused_on_any_endpoint(self):
        before = list(EmergencyContact.objects.values_list("id", "name", "is_active", "display_order"))
        for role in (ROLE_TOURISM, ROLE_TOURIST):
            client = self._client(role, f"ec_{role}")
            for method, url in self._all_calls():
                response = getattr(client, method)(url, {"name": "Intruder"}, format="json")
                self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN, (role, method, url))
        self.assertEqual(
            before, list(EmergencyContact.objects.values_list("id", "name", "is_active", "display_order")))
        self.assertEqual(self._logs().count(), 0)

    # -- list --------------------------------------------------------------

    def test_the_admin_list_has_every_contact_in_order_with_its_state(self):
        self.b.is_active = False
        self.b.save()
        rows = self.admin.get(self.LIST).json()
        self.assertEqual([r["name"] for r in rows], ["Alpha Rescue", "Bravo Police", "Charlie Desk"])
        self.assertEqual([r["is_active"] for r in rows], [True, False, True])
        self.assertEqual(
            set(rows[0]),
            {"id", "name", "contact_numbers", "email", "description",
             "icon", "display_order", "is_active"},
        )

    # -- create ------------------------------------------------------------

    def test_create_answers_and_saves_and_goes_last_and_active(self):
        response = self.admin.post(self.LIST, {
            "name": "  Delta   Coast  Guard ", "contact_numbers": "555-5555", "email": "d@example.com",
            "description": "Boats", "icon": "coast_guard",
            "is_active": False, "display_order": 1, "id": 9999,
        }, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        body = response.json()
        saved = EmergencyContact.objects.get(pk=body["id"])
        self.assertNotEqual(saved.pk, 9999)
        for source in (body, {f: getattr(saved, f) for f in body}):
            self.assertEqual(source["name"], "Delta Coast Guard")
            self.assertEqual(source["contact_numbers"], "555-5555")
            self.assertEqual(source["email"], "d@example.com")
            self.assertEqual(source["description"], "Boats")
            self.assertEqual(source["icon"], "coast_guard")
            self.assertEqual(source["display_order"], 40)
            self.assertIs(source["is_active"], True)
        self.assertEqual(self._payload_names()[-1], "Delta Coast Guard")

    def test_create_without_a_icon_is_other(self):
        response = self.admin.post(self.LIST, {"name": "Echo", "contact_numbers": "6"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(EmergencyContact.objects.get(name="Echo").icon, "other")

    def test_create_refuses_a_missing_name_or_number_and_saves_nothing(self):
        for body in ({"contact_numbers": "1"}, {"name": "  ", "contact_numbers": "1"}, {"name": "Foxtrot"},
                     {"name": "Foxtrot", "contact_numbers": "1", "email": "not-an-email"}):
            response = self.admin.post(self.LIST, body, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, body)
        self.assertEqual(EmergencyContact.objects.count(), 3)
        self.assertEqual(self._logs().count(), 0)

    # -- guard 2: the same name -------------------------------------------

    def test_create_refuses_the_same_name_by_the_import_rule(self):
        from api.services import online_booking

        for clone in ("alpha rescue", "  ALPHA   Rescue ", "Alpha  rescue"):
            self.assertEqual(online_booking.normalize_key(clone), online_booking.normalize_key("Alpha Rescue"))
            response = self.admin.post(self.LIST, {"name": clone, "contact_numbers": "9"}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, clone)
            self.assertIn("name", response.json())
        self.assertEqual(EmergencyContact.objects.count(), 3)
        self.assertEqual(self._logs().count(), 0)

    def test_rename_to_the_name_of_another_is_refused_and_nothing_changes(self):
        response = self.admin.patch(self._url(self.b), {"name": "  alpha   RESCUE", "description": "new text"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.b.refresh_from_db()
        self.assertEqual(self.b.name, "Bravo Police")
        self.assertEqual(self.b.description, "")
        self.assertEqual(self._logs().count(), 0)

    def test_a_contact_may_be_renamed_to_its_own_name_in_other_capitals(self):
        response = self.admin.patch(self._url(self.a), {"name": "ALPHA  rescue"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["name"], "ALPHA rescue")
        self.a.refresh_from_db()
        self.assertEqual(self.a.name, "ALPHA rescue")

    # -- guard 3: the icon --------------------------------------------

    def test_an_unknown_icon_is_a_400_on_create_and_on_edit(self):
        for bad in ("siren", "", None, 5, "RESCUE", "desk", "Hospital"):
            response = self.admin.post(
                self.LIST, {"name": "Golf", "contact_numbers": "7", "icon": bad}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, bad)
            response = self.admin.patch(self._url(self.a), {"icon": bad}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, bad)
        self.a.refresh_from_db()
        self.assertEqual(self.a.icon, "rescue")
        self.assertFalse(EmergencyContact.objects.filter(name="Golf").exists())

    def test_each_of_the_twelve_icons_is_accepted(self):
        for index, code in enumerate(("rescue", "police", "coast_guard", "fire", "hospital", "ambulance",
                                      "help_desk", "lgu_office", "power", "water", "port", "other")):
            response = self.admin.post(
                self.LIST, {"name": f"Cat {index}", "contact_numbers": "1", "icon": code}, format="json")
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, code)
            self.assertEqual(EmergencyContact.objects.get(name=f"Cat {index}").icon, code)

    # -- edit --------------------------------------------------------------

    def test_editing_a_number_answers_saves_and_reaches_the_next_payload(self):
        before = APIClient().get(self.BOOTSTRAP).json()["emergencyContacts"]
        self.assertEqual(before[1]["contact_numbers"], "222-2222 / 333-3333")
        response = self.admin.patch(self._url(self.b), {"contact_numbers": "777-7777"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["contact_numbers"], "777-7777")
        self.b.refresh_from_db()
        self.assertEqual(self.b.contact_numbers, "777-7777")
        after = APIClient().get(self.BOOTSTRAP).json()["emergencyContacts"]
        self.assertEqual(after[1]["contact_numbers"], "777-7777")
        self.assertNotEqual(before, after)

    def test_edit_ignores_is_active_and_display_order(self):
        response = self.admin.patch(
            self._url(self.a), {"description": "8-5", "is_active": False, "display_order": 99}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.a.refresh_from_db()
        self.assertEqual((self.a.description, self.a.is_active, self.a.display_order), ("8-5", True, 10))
        self.assertIs(response.json()["is_active"], True)
        self.assertEqual(response.json()["display_order"], 10)

    def test_edit_of_an_unknown_contact_is_404(self):
        self.assertEqual(self.admin.patch(self.LIST + "99999/", {"description": "x"}, format="json").status_code, 404)
        self.assertEqual(self.admin.post(self.LIST + "99999/deactivate/").status_code, 404)

    # -- deactivate / reactivate ------------------------------------------

    def test_deactivate_hides_from_the_app_and_reactivate_brings_it_back(self):
        response = self.admin.post(self._url(self.b, "deactivate"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIs(response.json()["is_active"], False)
        self.b.refresh_from_db()
        self.assertFalse(self.b.is_active)
        self.assertEqual(self._payload_names(), ["Alpha Rescue", "Charlie Desk"])
        self.assertEqual(self._admin_names(), ["Alpha Rescue", "Bravo Police", "Charlie Desk"])

        response = self.admin.post(self._url(self.b, "reactivate"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIs(response.json()["is_active"], True)
        self.b.refresh_from_db()
        self.assertTrue(self.b.is_active)
        self.assertEqual(self._payload_names(), ["Alpha Rescue", "Bravo Police", "Charlie Desk"])

    def test_with_none_active_the_payload_is_an_empty_list_and_the_rows_stay(self):
        for contact in (self.a, self.b, self.c):
            self.assertEqual(self.admin.post(self._url(contact, "deactivate")).status_code, 200)
        response = APIClient().get(self.BOOTSTRAP)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["emergencyContacts"], [])
        self.assertEqual(EmergencyContact.objects.filter(is_active=False).count(), 3)

    def test_deactivating_twice_logs_once(self):
        self.admin.post(self._url(self.a, "deactivate"))
        self.admin.post(self._url(self.a, "deactivate"))
        self.assertEqual(self._logs().count(), 1)

    # -- no hard delete ----------------------------------------------------

    def test_delete_is_not_allowed_anywhere_and_nothing_is_removed(self):
        for url in (self._url(self.a), self.LIST, self._url(self.a, "deactivate"), self.REORDER):
            response = self.admin.delete(url)
            self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED, url)
        self.assertEqual(EmergencyContact.objects.count(), 3)

    # -- guard 4: reorder --------------------------------------------------

    def test_reorder_changes_the_admin_list_and_the_payload(self):
        ids = [self.c.pk, self.a.pk, self.b.pk]
        response = self.admin.post(self.REORDER, {"ids": ids}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual([r["id"] for r in response.json()], ids)
        self.assertEqual(
            list(EmergencyContact.objects.order_by("display_order", "id").values_list("id", flat=True)), ids)
        self.assertEqual(self._payload_names(), ["Charlie Desk", "Alpha Rescue", "Bravo Police"])
        self.assertEqual(self._admin_names(), ["Charlie Desk", "Alpha Rescue", "Bravo Police"])

    def test_reorder_with_a_bad_list_is_a_400_and_changes_nothing(self):
        before = list(EmergencyContact.objects.values_list("id", "display_order"))
        for ids in (
            [self.a.pk, self.b.pk],                       # missing one
            [self.c.pk, self.b.pk, self.a.pk, 99999],     # an extra id
            [self.c.pk, self.b.pk, 99999],                # a nonexistent id in place of one
            [self.c.pk, self.c.pk, self.a.pk, self.b.pk], # repeated
            [],
            "abc",
            [self.c.pk, True, self.a.pk],
        ):
            response = self.admin.post(self.REORDER, {"ids": ids}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, ids)
            self.assertEqual(before, list(EmergencyContact.objects.values_list("id", "display_order")), ids)
        self.assertEqual(self._logs().count(), 0)

    def test_reorder_does_not_take_true_for_the_id_one(self):
        EmergencyContact.objects.all().delete()
        for pk, name, order in ((1, "One", 10), (2, "Two", 20), (3, "Three", 30)):
            EmergencyContact.objects.create(id=pk, name=name, contact_numbers="1", display_order=order)
        response = self.admin.post(self.REORDER, {"ids": [3, True, 2]}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            list(EmergencyContact.objects.order_by("id").values_list("id", "display_order")),
            [(1, 10), (2, 20), (3, 30)])
        self.assertEqual(self._logs().count(), 0)

    # -- guard 5: the ActivityLog -----------------------------------------

    def test_create_is_logged_with_who_which_and_what(self):
        body = self.admin.post(self.LIST, {"name": "Hotel Desk", "contact_numbers": "8"}, format="json").json()
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.action, ACTION_CREATE)
        self.assertEqual(log.record_id, str(body["id"]))
        self.assertIn("added", log.record_label)
        self.assertIn("Hotel Desk", log.record_label)

    def test_edit_is_logged_with_who_which_and_what(self):
        self.admin.patch(self._url(self.b), {"contact_numbers": "777-7777"}, format="json")
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.b.pk))
        self.assertIn("Bravo Police", log.record_label)
        self.assertIn("edited", log.record_label)
        self.assertIn("222-2222 / 333-3333", log.record_label)
        self.assertIn("777-7777", log.record_label)

    def test_deactivate_is_logged_with_who_which_and_what(self):
        self.admin.post(self._url(self.c, "deactivate"))
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.c.pk))
        self.assertIn("Charlie Desk", log.record_label)
        self.assertIn("deactivated", log.record_label)

    def test_reactivate_is_logged_with_who_which_and_what(self):
        self.c.is_active = False
        self.c.save()
        self.admin.post(self._url(self.c, "reactivate"))
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.c.pk))
        self.assertIn("Charlie Desk", log.record_label)
        self.assertIn("brought back", log.record_label)

    def test_reorder_is_logged_with_who_and_the_new_order(self):
        self.admin.post(self.REORDER, {"ids": [self.b.pk, self.c.pk, self.a.pk]}, format="json")
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertIn("reordered", log.record_label)
        label = log.record_label
        self.assertTrue(label.index("Bravo Police") < label.index("Charlie Desk") < label.index("Alpha Rescue"))

    # -- guard 6: no cache -------------------------------------------------

    def test_the_payload_contacts_are_not_cached(self):
        # Read once (a cache would keep this answer), then change, then read.
        self.assertEqual(self._payload_names(), ["Alpha Rescue", "Bravo Police", "Charlie Desk"])
        self.admin.post(self._url(self.a, "deactivate"))
        self.assertEqual(self._payload_names(), ["Bravo Police", "Charlie Desk"])


def _repo_file(*parts):
    from pathlib import Path

    return Path(__file__).resolve().parent.parent.parent.joinpath(*parts)


def web_fare_tiers():
    """The nine tiers as BookingManagement.js has them NOW, read from the real
    file (not copied here), so that drift between the two cannot hide."""
    import re

    source = _repo_file("frontend", "src", "tourism", "pages", "BookingManagement.js").read_text(encoding="utf-8")
    block = source[source.index("const boatCapacityFareOptions = ["):]
    block = block[: block.index("];")]
    pairs = re.findall(r'\{ id: "([^"]*)", name: "([^"]*)" \}', block)
    assert all(a == b for a, b in pairs), "a web tier has an id different from its name"
    return [name for _id, name in pairs]


def dart_fare_tiers():
    """The tiers as mobile/lib/utils/helpers.dart has them NOW (the leading
    empty string is the 'Select capacity and fare' choice, not a tier)."""
    import re

    source = _repo_file("mobile", "lib", "utils", "helpers.dart").read_text(encoding="utf-8")
    block = source[source.index("const boatCapacityFareOptions = ["):]
    block = block[: block.index("];")]
    return [text for text in re.findall(r"'([^']*)'", block) if text]


class BoatCapacityFareModelTests(TestCase):
    """The table for the Boat Capacity and Fare tiers. Nothing reads it yet."""

    def test_the_fields_are_exactly_these(self):
        fields = {f.name: f for f in BoatCapacityFare._meta.get_fields()}
        self.assertEqual(set(fields), {"id", "name", "display_order", "is_active"})
        self.assertEqual(type(fields["id"]).__name__, "BigAutoField")
        self.assertEqual(fields["name"].max_length, 160)
        self.assertEqual(fields["display_order"].default, 0)
        self.assertIs(fields["is_active"].default, True)
        self.assertIs(fields["is_active"].db_default, True)
        self.assertEqual(BoatCapacityFare._meta.ordering, ["display_order", "id"])
        self.assertTrue(issubclass(BoatCapacityFare, ActiveOption))

    def test_the_name_is_as_long_as_the_column_that_holds_it_on_a_record(self):
        # a tier must always fit in TouristRecord.boat_capacity_fare
        self.assertEqual(
            BoatCapacityFare._meta.get_field("name").max_length,
            TouristRecord._meta.get_field("boat_capacity_fare").max_length,
        )

    def test_nothing_points_at_it(self):
        self.assertEqual(
            [f.name for f in BoatCapacityFare._meta.get_fields() if f.is_relation], []
        )
        self.assertEqual(
            [f.name for f in TouristRecord._meta.get_fields()
             if f.is_relation and getattr(f.related_model, "__name__", "") == "BoatCapacityFare"],
            [],
        )

    def test_a_row_put_in_without_the_column_is_active(self):
        table = BoatCapacityFare._meta.db_table
        with connection.cursor() as cursor:
            cursor.execute(f"INSERT INTO {table} (name, display_order) VALUES ('Raw tier', 0)")
        self.assertIs(BoatCapacityFare.objects.get(name="Raw tier").is_active, True)

    def test_the_order_is_display_order_then_id(self):
        BoatCapacityFare.objects.all().delete()
        late = BoatCapacityFare.objects.create(name="Created first", display_order=30)
        early = BoatCapacityFare.objects.create(name="Created second", display_order=10)
        tie_a = BoatCapacityFare.objects.create(name="Tie A", display_order=20)
        tie_b = BoatCapacityFare.objects.create(name="Tie B", display_order=20)
        self.assertEqual(
            list(BoatCapacityFare.objects.all()), [early, tie_a, tie_b, late]
        )

    def test_the_migration_that_makes_the_table_only_makes_the_table(self):
        import importlib
        from django.db import migrations

        make = importlib.import_module("api.migrations.0055_boat_capacity_fare")
        self.assertEqual(len(make.Migration.operations), 1)
        operation = make.Migration.operations[0]
        self.assertIsInstance(operation, migrations.CreateModel)
        self.assertEqual(operation.name, "BoatCapacityFare")
        self.assertEqual(make.Migration.dependencies, [("api", "0054_seed_emergency_contacts")])

        seed = importlib.import_module("api.migrations.0056_seed_boat_capacity_fares")
        self.assertEqual(len(seed.Migration.operations), 1)
        self.assertIsInstance(seed.Migration.operations[0], migrations.RunPython)
        self.assertEqual(seed.Migration.dependencies, [("api", "0055_boat_capacity_fare")])


class BoatCapacityFareSeedTests(TransactionTestCase):
    """Migrations 0055/0056 through Django's own migrate, forward and backward."""

    BEFORE = [("api", "0055_boat_capacity_fare")]
    AFTER = [("api", "0056_seed_boat_capacity_fares")]
    SEED_LOGGER = "api.migrations.0056_seed_boat_capacity_fares"

    def tearDown(self):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())

    def _model(self, target):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        return executor.loader.project_state(target).apps.get_model("api", "BoatCapacityFare")

    def _fresh_before(self):
        MigrationExecutor(connection).migrate(self.BEFORE)  # undo 0056 (a real reverse)
        model = self._model(self.BEFORE)
        model.objects.all().delete()
        return model

    def test_the_two_files_the_tiers_are_read_from_still_agree_with_each_other(self):
        web, dart = web_fare_tiers(), dart_fare_tiers()
        self.assertEqual(len(web), 9)
        self.assertEqual(len(dart), 9)
        self.assertEqual(web, dart)

    def test_the_real_migrate_puts_in_nine_active_rows_equal_letter_for_letter_to_both_files(self):
        self._fresh_before()
        MigrationExecutor(connection).migrate(self.AFTER)
        rows = list(self._model(self.AFTER).objects.order_by("display_order", "id"))

        self.assertEqual(len(rows), 9)
        self.assertTrue(all(row.is_active for row in rows))
        self.assertEqual([row.display_order for row in rows], [10, 20, 30, 40, 50, 60, 70, 80, 90])
        names = [row.name for row in rows]
        self.assertEqual(names, web_fare_tiers())
        self.assertEqual(names, dart_fare_tiers())
        # the commas were NOT tidied: records hold exactly this text
        self.assertEqual(names[0], "1-2 pax (One-Way P1500, Two-way P2000)")
        self.assertEqual(names[1], "3-4 pax (One-Way P2000, Two-way P2,500)")
        self.assertFalse(any("P1,500" in n or "P2,000" in n for n in names))

    def test_running_the_real_migration_a_second_time_still_leaves_nine(self):
        self._fresh_before()
        MigrationExecutor(connection).migrate(self.AFTER)
        self.assertEqual(self._model(self.AFTER).objects.count(), 9)

        # Forget that 0056 ran (keep its rows) and let Django run it again.
        executor = MigrationExecutor(connection)
        executor.recorder.record_unapplied("api", "0056_seed_boat_capacity_fares")
        with self.assertLogs(self.SEED_LOGGER, level="WARNING") as logs:
            MigrationExecutor(connection).migrate(self.AFTER)

        model = self._model(self.AFTER)
        self.assertEqual(model.objects.count(), 9)
        self.assertEqual(len(logs.records), 9)  # every one of the nine was skipped, and said so
        self.assertEqual(
            [r.name for r in model.objects.order_by("display_order", "id")], web_fare_tiers()
        )

    def test_a_tier_with_the_same_name_in_other_capitals_and_spaces_is_skipped_and_logged(self):
        model = self._fresh_before()
        mine = model.objects.create(name="  1-2 PAX (one-way p1500,   two-way p2000)  ", display_order=0)

        with self.assertLogs(self.SEED_LOGGER, level="WARNING") as logs:
            MigrationExecutor(connection).migrate(self.AFTER)

        model = self._model(self.AFTER)
        self.assertEqual(model.objects.count(), 9)  # 8 new + the one that was already there
        self.assertEqual(len(logs.records), 1)
        self.assertIn("1-2 pax (One-Way P1500, Two-way P2000)", logs.records[0].getMessage())
        kept = model.objects.get(pk=mine.pk)
        self.assertEqual((kept.name, kept.display_order), (mine.name, 0))  # untouched
        self.assertFalse(model.objects.filter(name="1-2 pax (One-Way P1500, Two-way P2000)").exists())

    def test_the_real_reverse_takes_out_the_nine_and_keeps_rows_added_by_hand(self):
        self._fresh_before()
        MigrationExecutor(connection).migrate(self.AFTER)
        model = self._model(self.AFTER)
        added = model.objects.create(name="21-25 pax (One-Way P6,000, Two-way P6,500)", display_order=100)
        # a seeded tier that an admin re-wrote is no longer "one of the nine"
        reworded = model.objects.create(name="1-2 pax (One-Way P1,500, Two-way P2,000)", display_order=5)
        self.assertEqual(model.objects.count(), 11)

        MigrationExecutor(connection).migrate(self.BEFORE)  # the reverse of 0056
        model = self._model(self.BEFORE)
        self.assertEqual(
            sorted(model.objects.values_list("pk", flat=True)), sorted([added.pk, reworded.pk])
        )
        self.assertEqual(model.objects.count(), 2)

        MigrationExecutor(connection).migrate(self.AFTER)  # and forward again
        self.assertEqual(self._model(self.AFTER).objects.count(), 11)

    def test_going_back_past_0055_removes_the_table_and_touches_no_other(self):
        before = set(connection.introspection.table_names())
        MigrationExecutor(connection).migrate([("api", "0054_seed_emergency_contacts")])
        gone = before - set(connection.introspection.table_names())
        # Going back past 0055 also takes back every later migration, so the tables
        # that migrations AFTER 0055 created (Mobile App Content, 0060) go with it.
        later = {"api_mobilecontentitem", "api_mobilecontentblock"}
        self.assertEqual(gone - later, {BoatCapacityFare._meta.db_table})

        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())  # forward again, to the latest
        self.assertEqual(set(connection.introspection.table_names()), before)


class BoatCapacityFarePayloadTests(TestCase):
    """referenceTables.boatCapacityFares (boat fare, commit 2).

    The web gets every tier with its id, order and state; the mobile app gets
    the active tiers' NAMES, in order, and no id. The list is read on every
    request, outside the 15-minute cache of the reference tables.
    """

    MOBILE = "/api/mobile/tourism/bootstrap/"
    WEB_BOOTSTRAP = "/api/bootstrap/"
    WEB_TABLES = "/api/reference-tables/"
    CACHE_KEY = "mobile_reference_tables_v2"
    OLD_LISTS = ("itineraries", "travelModes", "boatTypes", "visitPurposes")

    def setUp(self):
        import importlib

        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        self.seed = importlib.import_module("api.migrations.0056_seed_boat_capacity_fares")
        user = User.objects.create_user(username="bcf_payload_staff", password="Staff@12345")
        UserProfile.objects.create(user=user, role=ROLE_TOURISM)
        self.staff = APIClient()
        self.staff.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

    # -- helpers -----------------------------------------------------------

    def _nine(self):
        BoatCapacityFare.objects.all().delete()
        for index, name in enumerate(self.seed.TIERS):
            BoatCapacityFare.objects.create(name=name, display_order=10 * (index + 1))

    def _mobile(self):
        response = APIClient().get(self.MOBILE)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.json()["referenceTables"]["boatCapacityFares"]

    def _web(self):
        response = self.staff.get(self.WEB_TABLES)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.json()["boatCapacityFares"]

    # -- the shape ----------------------------------------------------------

    def test_the_mobile_list_is_the_active_names_in_order_and_nothing_else(self):
        self._nine()
        mobile = self._mobile()
        self.assertEqual(mobile, [{"name": name} for name in self.seed.TIERS])
        for item in mobile:
            self.assertEqual(set(item), {"name"})
            self.assertNotIn("id", item)

    def test_the_web_list_is_every_tier_with_id_name_order_and_state(self):
        self._nine()
        BoatCapacityFare.objects.filter(display_order=30).update(is_active=False)
        web = self._web()
        self.assertEqual(len(web), 9)
        for item in web:
            self.assertEqual(set(item), {"id", "name", "display_order", "is_active"})
        self.assertEqual([item["name"] for item in web], self.seed.TIERS)
        self.assertEqual([item["display_order"] for item in web], [10, 20, 30, 40, 50, 60, 70, 80, 90])
        self.assertEqual([item["is_active"] for item in web], [True, True, False] + [True] * 6)
        self.assertEqual(
            [item["id"] for item in web],
            list(BoatCapacityFare.objects.order_by("display_order").values_list("id", flat=True)),
        )

    def test_both_web_endpoints_and_the_mobile_one_carry_the_key(self):
        self._nine()
        self.assertEqual(len(self._web()), 9)
        bootstrap = self.staff.get(self.WEB_BOOTSTRAP)
        self.assertEqual(bootstrap.status_code, status.HTTP_200_OK)
        self.assertEqual(len(bootstrap.json()["referenceTables"]["boatCapacityFares"]), 9)
        self.assertEqual(len(self._mobile()), 9)

    def test_no_id_reaches_the_mobile_app_whatever_the_table_holds(self):
        BoatCapacityFare.objects.all().delete()
        BoatCapacityFare.objects.create(id=777, name="Only tier", display_order=1)
        BoatCapacityFare.objects.create(id=778, name="Retired tier", display_order=2, is_active=False)
        mobile = self._mobile()
        self.assertEqual(mobile, [{"name": "Only tier"}])
        self.assertNotIn("777", str(mobile))
        self.assertNotIn("id", mobile[0])

    def test_an_empty_table_gives_an_empty_list_for_both_not_an_error_or_a_missing_key(self):
        BoatCapacityFare.objects.all().delete()
        self.assertEqual(self._mobile(), [])
        self.assertEqual(self._web(), [])

    def test_with_every_tier_retired_the_app_gets_an_empty_list_and_the_web_still_has_them(self):
        self._nine()
        BoatCapacityFare.objects.update(is_active=False)
        self.assertEqual(self._mobile(), [])
        web = self._web()
        self.assertEqual(len(web), 9)
        self.assertEqual({item["is_active"] for item in web}, {False})

    # -- the server wins ----------------------------------------------------

    def _different_list(self):
        """Not the seeded nine: a tenth tier put FIRST, two tiers swapped in
        order, and one retired. Nothing a page could guess from the nine."""
        self._nine()
        BoatCapacityFare.objects.create(name="21-25 pax (One-Way P6,000, Two-way P6,500)", display_order=5)
        third = BoatCapacityFare.objects.get(display_order=30)
        fourth = BoatCapacityFare.objects.get(display_order=40)
        third.display_order, fourth.display_order = 40, 30
        third.save()
        fourth.save()
        BoatCapacityFare.objects.filter(display_order=70).update(is_active=False)

    def test_the_servers_list_comes_out_not_the_nine(self):
        self._different_list()
        tiers = self.seed.TIERS
        expected_web = [
            "21-25 pax (One-Way P6,000, Two-way P6,500)",  # the tenth, first
            tiers[0], tiers[1],
            tiers[3], tiers[2],  # swapped
            tiers[4], tiers[5],
            tiers[6],  # retired: in the web list, not in the app's
            tiers[7], tiers[8],
        ]
        expected_mobile = [name for name in expected_web if name != tiers[6]]

        self.assertNotEqual(expected_mobile, tiers)
        self.assertEqual([item["name"] for item in self._web()], expected_web)
        self.assertEqual([item["name"] for item in self._mobile()], expected_mobile)
        self.assertEqual(len(expected_web), 10)
        self.assertEqual(len(expected_mobile), 9)  # as many as the seed, and a different nine

    # -- no 15-minute delay ---------------------------------------------------

    def test_a_change_shows_in_the_very_next_bootstrap_even_with_the_cache_warm(self):
        from django.core.cache import cache

        self._nine()
        # warm everything that is cached, as a first visit would
        self.assertEqual(len(self._mobile()), 9)
        self.assertEqual(len(self._web()), 9)
        self.assertIsNotNone(cache.get(self.CACHE_KEY))

        BoatCapacityFare.objects.filter(display_order=50).update(is_active=False)

        mobile = [item["name"] for item in self._mobile()]
        self.assertNotIn(self.seed.TIERS[4], mobile)
        self.assertEqual(len(mobile), 8)
        web = {item["name"]: item for item in self._web()}
        self.assertIn(self.seed.TIERS[4], web)
        self.assertIs(web[self.seed.TIERS[4]]["is_active"], False)

        # and a rename, a reorder and a new tier, each in the very next answer
        BoatCapacityFare.objects.filter(display_order=10).update(name="Renamed first tier", display_order=95)
        BoatCapacityFare.objects.create(name="Brand new tier", display_order=1)
        mobile = [item["name"] for item in self._mobile()]
        self.assertEqual(mobile[0], "Brand new tier")
        self.assertEqual(mobile[-1], "Renamed first tier")
        self.assertNotIn(self.seed.TIERS[0], mobile)

    def test_the_list_is_not_in_the_cache_and_the_cache_key_is_still_v2(self):
        from django.core.cache import cache

        from .services.tourism import MOBILE_REFERENCE_TABLES_CACHE_KEY

        self._nine()
        self._mobile()
        self._web()
        self.assertEqual(MOBILE_REFERENCE_TABLES_CACHE_KEY, "mobile_reference_tables_v2")
        cached = cache.get(self.CACHE_KEY)
        self.assertIsNotNone(cached)
        self.assertNotIn("boatCapacityFares", cached)

    # -- nothing else moved ----------------------------------------------------

    def test_the_four_existing_lists_do_not_depend_on_the_new_one(self):
        from .services.tourism import build_reference_tables_payload

        self._nine()
        with_nine_mobile = build_reference_tables_payload(active_only=True)
        with_nine_web = build_reference_tables_payload()
        BoatCapacityFare.objects.all().delete()
        BoatCapacityFare.objects.create(name="Something else entirely", display_order=1)
        other_mobile = build_reference_tables_payload(active_only=True)
        other_web = build_reference_tables_payload()

        for key in self.OLD_LISTS:
            with self.subTest(key=key):
                self.assertEqual(with_nine_mobile[key], other_mobile[key])
                self.assertEqual(with_nine_web[key], other_web[key])
        self.assertEqual(
            set(with_nine_mobile),
            {"countries", "regions", "provinces", "itineraries", "travelModes", "boatTypes",
             "visitPurposes", "resorts", "boatCapacityFares"},
        )
        # the shapes the app has always read
        for key, keys in {
            "itineraries": {"id", "name"}, "travelModes": {"id", "name"},
            "boatTypes": {"id", "name", "requires_capacity_fare"}, "visitPurposes": {"id", "name"},
        }.items():
            for item in with_nine_mobile[key]:
                self.assertEqual(set(item), keys)

    def test_the_cached_lists_are_not_touched_by_adding_the_key(self):
        from django.core.cache import cache

        self._nine()
        self._web()
        cached_before = {key: list(value) for key, value in cache.get(self.CACHE_KEY).items()}
        self._mobile()
        self._web()
        cached_after = cache.get(self.CACHE_KEY)
        self.assertEqual(set(cached_before), set(cached_after))
        self.assertEqual(cached_before, {key: list(value) for key, value in cached_after.items()})

    def test_adding_the_key_changes_no_stored_data(self):
        self._nine()
        before = (TouristRecord.objects.count(), BoatCapacityFare.objects.count())
        self._mobile()
        self._web()
        self.assertEqual(before, (TouristRecord.objects.count(), BoatCapacityFare.objects.count()))


class _BoatFareApiBase:
    """Shared set-up: four tiers (NOT the seeded nine, different names and
    orders), records that hold some of them, and an admin and a non-admin."""

    LIST = "/api/boat-capacity-fares/"
    REORDER = "/api/boat-capacity-fares/reorder/"
    A_NAME = "Alpha 1-2 pax (One-Way P1000, Two-way P1500)"
    B_NAME = "Bravo 3-4 pax (One-Way P2000, Two-way P2,500)"
    C_NAME = "Charlie 5-6 pax (One-Way P3,000, Two-way P3,500)"
    D_NAME = "Delta 7-8 pax (One-Way P4,000, Two-way P4,500)"
    ORPHAN = "Zulu 9-10 pax (One-Way P9,000, Two-way P9,500)"  # held by a record, in no tier

    def _make_client(self, role, username):
        user = User.objects.create_user(username=username, password="Staff@12345")
        UserProfile.objects.create(user=user, role=role)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        client.user = user
        return client

    def _record(self, survey_id, fare):
        return TouristRecord.objects.create(
            survey_id=survey_id,
            full_name="Boat Fare Test Group",
            contact_number="09170000000",
            country=Country.objects.first(),
            region=Region.objects.first(),
            province=Province.objects.first(),
            arrival_date="2026-04-02",
            resort=Resort.objects.first(),
            itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(),
            boat_type=BoatType.objects.first(),
            boat_capacity_fare=fare,
            visit_purpose=VisitPurpose.objects.first(),
            total_visitors=4,
            filipino_count=4,
            total_male=2,
            total_female=2,
            age_8_59=4,
            status="arrived",
        )

    def setUpData(self):
        from django.core.cache import cache

        cache.clear()
        ensure_test_reference_tables()
        TouristRecord.objects.all().delete()
        BoatCapacityFare.objects.all().delete()
        # created in a different order from their display_order
        self.c = BoatCapacityFare.objects.create(name=self.C_NAME, display_order=30)
        self.a = BoatCapacityFare.objects.create(name=self.A_NAME, display_order=10)
        self.d = BoatCapacityFare.objects.create(name=self.D_NAME, display_order=40)
        self.b = BoatCapacityFare.objects.create(name=self.B_NAME, display_order=20)
        # A is held by 2 records, C by 1, B and D by none; one record holds a text in no tier; one holds none
        self._record("BF-001", self.A_NAME)
        self._record("BF-002", self.A_NAME)
        self._record("BF-003", self.C_NAME)
        self._record("BF-004", self.ORPHAN)
        self._record("BF-005", "")
        self.admin = self._make_client(ROLE_ADMIN, "bf_admin")

    def _url(self, fare=None, action=None):
        url = self.LIST
        if fare is not None:
            url += f"{fare.pk}/"
        if action:
            url += f"{action}/"
        return url

    def _stored(self):
        """Every tier row as it is in the database NOW."""
        return list(BoatCapacityFare.objects.order_by("pk").values_list("pk", "name", "display_order", "is_active"))

    def _held(self):
        """Every record's tier text as it is in the database NOW."""
        return list(TouristRecord.objects.order_by("pk").values_list("pk", "boat_capacity_fare"))

    def _logs(self):
        return ActivityLog.objects.filter(record_type="BoatCapacityFare")


class BoatCapacityFareApiTests(_BoatFareApiBase, TestCase):
    """The admin API of the Boat Capacity and Fare tiers (boat fare, commit 3)."""

    def setUp(self):
        self.setUpData()

    # -- who may use it: 401 and 403 on EACH endpoint ---------------------

    def _anonymous_is_refused(self, method, url):
        before, held, logs = self._stored(), self._held(), self._logs().count()
        response = getattr(APIClient(), method)(url, {"name": "Intruder", "ids": []}, format="json")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED, (method, url))
        self.assertEqual((self._stored(), self._held(), self._logs().count()), (before, held, logs))

    def _non_admin_is_refused(self, method, url):
        before, held, logs = self._stored(), self._held(), self._logs().count()
        for role in (ROLE_TOURISM, ROLE_TOURIST):
            client = self._make_client(role, f"bf_{role}_{method}_{abs(hash(url)) % 1000}")
            response = getattr(client, method)(url, {"name": "Intruder", "ids": []}, format="json")
            self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN, (role, method, url))
        self.assertEqual((self._stored(), self._held(), self._logs().count()), (before, held, logs))

    def test_401_on_the_list(self):
        self._anonymous_is_refused("get", self.LIST)

    def test_403_on_the_list(self):
        self._non_admin_is_refused("get", self.LIST)

    def test_401_on_create(self):
        self._anonymous_is_refused("post", self.LIST)

    def test_403_on_create(self):
        self._non_admin_is_refused("post", self.LIST)

    def test_401_on_rename(self):
        self._anonymous_is_refused("patch", self._url(self.b))

    def test_403_on_rename(self):
        self._non_admin_is_refused("patch", self._url(self.b))

    def test_401_on_deactivate(self):
        self._anonymous_is_refused("post", self._url(self.b, "deactivate"))

    def test_403_on_deactivate(self):
        self._non_admin_is_refused("post", self._url(self.b, "deactivate"))

    def test_401_on_reactivate(self):
        self._anonymous_is_refused("post", self._url(self.b, "reactivate"))

    def test_403_on_reactivate(self):
        self._non_admin_is_refused("post", self._url(self.b, "reactivate"))

    def test_401_on_reorder(self):
        self._anonymous_is_refused("post", self.REORDER)

    def test_403_on_reorder(self):
        self._non_admin_is_refused("post", self.REORDER)

    # -- routes ----------------------------------------------------------------

    def test_the_literal_reorder_route_is_registered_before_the_id_route(self):
        from django.urls import resolve

        from api.views import boat_capacity_fares as views

        self.assertIs(resolve(self.REORDER).func, views.boat_capacity_fare_reorder)
        self.assertEqual(resolve(self.REORDER).kwargs, {})
        self.assertIs(resolve(self._url(self.a)).func, views.boat_capacity_fare_detail)
        self.assertEqual(resolve(self._url(self.a)).kwargs, {"fare_id": self.a.pk})
        from api import urls as api_urls

        routes = [str(p.pattern) for p in api_urls.urlpatterns if "boat-capacity-fares" in str(p.pattern)]
        self.assertLess(routes.index("boat-capacity-fares/reorder/"), routes.index("boat-capacity-fares/<int:fare_id>/"))

    # -- no DELETE ----------------------------------------------------------------

    def test_delete_is_405_on_the_list_the_item_and_every_action(self):
        before, held = self._stored(), self._held()
        for url in (self.LIST, self._url(self.a), self._url(self.a, "deactivate"),
                    self._url(self.a, "reactivate"), self.REORDER):
            self.assertEqual(self.admin.delete(url).status_code, status.HTTP_405_METHOD_NOT_ALLOWED, url)
        self.assertEqual((self._stored(), self._held(), self._logs().count()), (before, held, 0))

    # -- reading ------------------------------------------------------------------

    def test_the_list_has_every_tier_in_order_with_its_state_and_record_count(self):
        self.d.is_active = False
        self.d.save()
        rows = self.admin.get(self.LIST).json()
        self.assertEqual([r["name"] for r in rows], [self.A_NAME, self.B_NAME, self.C_NAME, self.D_NAME])
        self.assertEqual([r["display_order"] for r in rows], [10, 20, 30, 40])
        self.assertEqual([r["is_active"] for r in rows], [True, True, True, False])
        for row in rows:
            self.assertEqual(set(row), {"id", "name", "display_order", "is_active", "record_count"})

    def test_record_count_is_exact_text_and_zero_for_a_tier_nobody_holds(self):
        rows = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        # tier 1 (A) is held by 2 records and tier 3 (C) by 1; B and D by none
        self.assertEqual(rows, {self.A_NAME: 2, self.B_NAME: 0, self.C_NAME: 1, self.D_NAME: 0})
        # a text in no tier, an empty text, and a near-miss do not count
        self._record("BF-006", self.A_NAME.upper())
        self._record("BF-007", " " + self.B_NAME)
        rows = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        self.assertEqual(rows[self.A_NAME], 2)
        self.assertEqual(rows[self.B_NAME], 0)

    def test_the_list_is_two_queries_for_any_number_of_tiers_not_one_per_tier(self):
        from django.db import connection as conn
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(conn) as few:
            self.assertEqual(self.admin.get(self.LIST).status_code, 200)
        for index in range(25):
            BoatCapacityFare.objects.create(name=f"Extra tier {index}", display_order=100 + index)
            self._record(f"BF-X{index}", f"Extra tier {index}")
        with CaptureQueriesContext(conn) as many:
            rows = self.admin.get(self.LIST).json()
        self.assertEqual(len(rows), 29)
        self.assertEqual(len(many), len(few))

    # -- create --------------------------------------------------------------------

    def test_create_answers_and_stores_a_tier_last_and_active(self):
        response = self.admin.post(
            self.LIST,
            {"name": "  Echo   11-12 pax   (One-Way P5,000)  ", "id": 9999, "is_active": False, "display_order": 1},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        body = response.json()
        stored = BoatCapacityFare.objects.get(pk=body["id"])
        self.assertNotEqual(stored.pk, 9999)
        for source in (body, {"id": stored.pk, "name": stored.name, "display_order": stored.display_order,
                              "is_active": stored.is_active, "record_count": 0}):
            self.assertEqual(source["name"], "Echo 11-12 pax (One-Way P5,000)")
            self.assertEqual(source["display_order"], 50)
            self.assertIs(source["is_active"], True)
            self.assertEqual(source["record_count"], 0)
        self.assertEqual(BoatCapacityFare.objects.count(), 5)

    def test_a_new_tier_whose_text_some_record_already_holds_counts_those_records(self):
        response = self.admin.post(self.LIST, {"name": self.ORPHAN}, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["record_count"], 1)
        listed = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        self.assertEqual(listed[self.ORPHAN], 1)

    def test_create_refuses_a_bad_name_and_stores_and_logs_nothing(self):
        before = self._stored()
        for bad in (None, "", "   ", 5, ["x"], {"a": 1}, "x" * 161):
            response = self.admin.post(self.LIST, {"name": bad}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, repr(bad)[:30])
            self.assertEqual(self._stored(), before)
        self.assertEqual(self.admin.post(self.LIST, {}, format="json").status_code, 400)
        self.assertEqual(self._stored(), before)
        self.assertEqual(self._logs().count(), 0)

    def test_create_accepts_a_name_of_exactly_160_letters(self):
        name = "n" * 160
        response = self.admin.post(self.LIST, {"name": name}, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(BoatCapacityFare.objects.get(pk=response.json()["id"]).name, name)

    def test_create_refuses_the_same_name_by_the_import_rule(self):
        from api.services import online_booking

        before = self._stored()
        for clone in (self.A_NAME.lower(), "  " + self.A_NAME.replace(" ", "   ") + " ", self.A_NAME.upper()):
            self.assertEqual(online_booking.normalize_key(clone), online_booking.normalize_key(self.A_NAME))
            response = self.admin.post(self.LIST, {"name": clone}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, clone)
            self.assertIn("name", response.json())
            self.assertEqual(self._stored(), before)
        self.assertEqual(self._logs().count(), 0)

    # -- rename ----------------------------------------------------------------------

    def test_rename_of_a_tier_nobody_holds_answers_and_stores_and_leaves_the_rest_alone(self):
        before = {pk: (name, order, active) for pk, name, order, active in self._stored()}
        response = self.admin.patch(
            self._url(self.b),
            {"name": "  Bravo   3-4 pax (One-Way P2,500, Two-way P3,000) ", "is_active": False, "display_order": 99},
            format="json",
        )
        new = "Bravo 3-4 pax (One-Way P2,500, Two-way P3,000)"
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = response.json()
        self.assertEqual((body["id"], body["name"], body["display_order"], body["is_active"]), (self.b.pk, new, 20, True))
        self.assertEqual(body["record_count"], 0)

        now = {pk: (name, order, active) for pk, name, order, active in self._stored()}
        self.assertEqual(now[self.b.pk], (new, 20, True))  # the name changed; is_active and order were ignored
        for pk in (self.a.pk, self.c.pk, self.d.pk):
            self.assertEqual(now[pk], before[pk])
        self.assertEqual(self._held(), self._held())

    def test_rename_of_a_held_tier_without_a_confirmation_is_409_confirm_required_and_nothing_changes(self):
        for tier, held in ((self.a, 2), (self.c, 1)):
            before, records = self._stored(), self._held()
            response = self.admin.patch(self._url(tier), {"name": "Brand new wording"}, format="json")
            self.assertEqual(response.status_code, status.HTTP_409_CONFLICT, tier.name)
            body = response.json()
            self.assertEqual(body["code"], "confirm_required")
            self.assertNotEqual(body["code"], "text_in_use")
            self.assertEqual(body["record_count"], held)  # the SERVER's number
            self.assertIn(f"{held} record(s) hold the old text", body["detail"])  # the number, in words
            self.assertIn("they keep the old text", body["detail"])
            self.assertEqual(self._stored(), before)  # the tier is exactly as it was: the NAME did not change
            self.assertEqual(self._held(), records)  # and so are the records
        self.assertEqual(self._logs().count(), 0)

    def test_rename_of_an_unheld_tier_goes_through_with_no_confirmation(self):
        response = self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["name"], "Bravo renamed")
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, "Bravo renamed")

    def test_a_confirmation_that_is_not_needed_does_not_get_in_the_way(self):
        for sent in (0, 5, "7", None):
            tier = BoatCapacityFare.objects.create(name=f"Spare tier {sent}", display_order=90)
            response = self.admin.patch(
                self._url(tier), {"name": f"Spare renamed {sent}", "confirm_affected_records": sent}, format="json"
            )
            self.assertEqual(response.status_code, status.HTTP_200_OK, sent)
            self.assertEqual(BoatCapacityFare.objects.get(pk=tier.pk).name, f"Spare renamed {sent}")

    def test_a_confirmation_that_is_too_high_is_409_and_nothing_changes(self):
        before, records = self._stored(), self._held()
        response = self.admin.patch(
            self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 3}, format="json"
        )  # A is held by 2
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.json()["code"], "confirm_required")
        self.assertEqual(response.json()["record_count"], 2)  # the real number, not the one sent
        self.assertEqual(self._stored(), before)
        self.assertEqual(self._held(), records)
        self.assertEqual(self._logs().count(), 0)

    def test_a_confirmation_that_is_too_low_is_409_and_nothing_changes(self):
        before, records = self._stored(), self._held()
        response = self.admin.patch(
            self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 1}, format="json"
        )
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.json()["code"], "confirm_required")
        self.assertEqual(response.json()["record_count"], 2)
        self.assertEqual(self._stored(), before)
        self.assertEqual(self._held(), records)
        self.assertEqual(self._logs().count(), 0)

    def test_a_confirmation_that_is_not_a_whole_number_is_409_and_nothing_changes(self):
        before = self._stored()
        for bad in (True, False, None, "abc", "", -2, 2.5, [2], {"n": 2}, 0):
            response = self.admin.patch(
                self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": bad}, format="json"
            )
            self.assertEqual(response.status_code, status.HTTP_409_CONFLICT, repr(bad))
            self.assertEqual(response.json()["code"], "confirm_required", repr(bad))
            self.assertEqual(self._stored(), before, repr(bad))
        self.assertEqual(self._logs().count(), 0)

    def test_a_confirmation_with_the_right_number_renames_the_tier_and_leaves_every_record_on_the_old_text(self):
        # Three records hold tier A (two from the fixture and one more).
        self._record("BF-A3", self.A_NAME)
        holders = list(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).order_by("pk").values_list("pk", flat=True))
        self.assertEqual(len(holders), 3)
        every_record_before = self._held()
        tiers_before = {pk: (name, order, active) for pk, name, order, active in self._stored()}

        response = self.admin.patch(
            self._url(self.a), {"name": "  Alpha   brand new   wording ", "confirm_affected_records": 3}, format="json"
        )

        # the answer
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = response.json()
        self.assertEqual((body["id"], body["name"], body["display_order"], body["is_active"]),
                         (self.a.pk, "Alpha brand new wording", 10, True))
        self.assertEqual(body["record_count"], 0)  # nobody holds the NEW text
        # the tier table, read again: the name is NEW, nothing else about the tier moved
        now = {pk: (name, order, active) for pk, name, order, active in self._stored()}
        self.assertEqual(now[self.a.pk], ("Alpha brand new wording", 10, True))
        for pk in (self.b.pk, self.c.pk, self.d.pk):
            self.assertEqual(now[pk], tiers_before[pk])
        # the three records, read again from the database: the OLD text, letter for letter
        for pk in holders:
            self.assertEqual(TouristRecord.objects.get(pk=pk).boat_capacity_fare, self.A_NAME)
        # and no record at all was changed
        self.assertEqual(self._held(), every_record_before)

    def test_after_the_rename_the_new_text_is_held_by_nobody_and_the_log_says_how_many_were_left(self):
        self._record("BF-A3", self.A_NAME)
        self.admin.patch(self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 3}, format="json")
        listed = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        self.assertEqual(listed["Alpha renamed"], 0)
        self.assertNotIn(self.A_NAME, listed)  # the old text is no longer a tier...
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).count(), 3)  # ...but 3 records hold it
        log = self._logs().get()
        self.assertIn("3 record(s) keep the old text", log.record_label)

    def test_the_log_of_a_rename_has_the_old_name_the_new_name_and_the_number_left_behind(self):
        self.admin.patch(self._url(self.c), {"name": "Charlie renamed", "confirm_affected_records": 1}, format="json")
        log = self._the_log()
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.c.pk))
        self.assertIn("renamed", log.record_label)
        self.assertIn(self.C_NAME, log.record_label)        # the old name
        self.assertIn("Charlie renamed", log.record_label)  # the new name
        self.assertIn("1 record(s) keep the old text", log.record_label)  # how many were left behind
        self.assertIn("no record was changed", log.record_label)

    def test_a_rename_of_an_unheld_tier_is_logged_with_a_count_of_zero(self):
        self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
        log = self._the_log()
        self.assertIn(self.B_NAME, log.record_label)
        self.assertIn("Bravo renamed", log.record_label)
        self.assertIn("0 record(s) keep the old text", log.record_label)

    # The names of the fields of the API, which a person at the LGU has no use for.
    API_WORDS = ("confirm_affected_records", "confirm_required", "record_count", "text_in_use",
                 "display_order", "is_active", "api", "json", "payload", "field")

    def _refusals(self):
        """Every way the 409 can come: no number, a wrong one (high and low), a bad one."""
        out = []
        for payload in (
            {"name": "Alpha renamed"},
            {"name": "Alpha renamed", "confirm_affected_records": 99},
            {"name": "Alpha renamed", "confirm_affected_records": 1},
            {"name": "Alpha renamed", "confirm_affected_records": "abc"},
            {"name": "Alpha renamed", "confirm_affected_records": True},
        ):
            out.append((payload, self.admin.patch(self._url(self.a), payload, format="json")))
        return out

    def test_the_words_of_the_409_are_for_a_person_not_the_names_of_api_fields(self):
        for payload, response in self._refusals():
            self.assertEqual(response.status_code, status.HTTP_409_CONFLICT, payload)
            detail = response.json()["detail"]
            for word in self.API_WORDS:
                self.assertNotIn(word, detail.lower(), (word, payload))
            self.assertNotIn("_", detail, payload)  # no snake_case name of a field at all

    def test_the_409_says_how_many_records_hold_the_text_that_they_keep_it_and_that_a_confirmation_is_needed(self):
        response = self.admin.patch(self._url(self.a), {"name": "Alpha renamed"}, format="json")
        detail = response.json()["detail"]
        self.assertEqual(
            detail,
            "2 record(s) hold the old text of this tier. Renaming the tier does not change them: they keep "
            "the old text, and it will show as (not in the list) in their own dropdown. "
            "The rename needs a confirmation.",
        )
        # the number in the words is the server's, whatever was sent
        wrong = self.admin.patch(
            self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 99}, format="json"
        ).json()["detail"]
        self.assertIn("2 record(s) hold the old text", wrong)
        self.assertNotIn("99", wrong)

    def test_the_code_and_the_count_the_front_end_reads_are_still_there(self):
        for payload, response in self._refusals():
            body = response.json()
            self.assertEqual(response.status_code, status.HTTP_409_CONFLICT, payload)
            self.assertEqual(body["code"], "confirm_required", payload)  # what the front end reads
            self.assertEqual(body["record_count"], 2, payload)  # the real number
            self.assertEqual(set(body), {"detail", "code", "record_count"}, payload)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.a.pk).name, self.A_NAME)  # nothing changed

    def test_the_race_refusal_uses_the_same_words_with_the_new_number(self):
        from unittest import mock

        from api.views import boat_capacity_fares as views

        real_count = views._count_holders
        arrived = []

        def count_then_a_booking_arrives(name):
            count = real_count(name)
            if name == self.A_NAME and not arrived:
                arrived.append(self._record("BF-LATE", self.A_NAME))
            return count

        with mock.patch.object(views, "_count_holders", side_effect=count_then_a_booking_arrives):
            response = self.admin.patch(
                self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 2}, format="json"
            )
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.json()["code"], "confirm_required")
        self.assertEqual(response.json()["record_count"], 3)
        self.assertIn("3 record(s) hold the old text", response.json()["detail"])
        self.assertNotIn("confirm_affected_records", response.json()["detail"])

    def test_the_old_ban_is_gone_text_in_use_is_never_answered(self):
        for payload in ({"name": "x1"}, {"name": "x2", "confirm_affected_records": 99}, {"name": "x3", "confirm_affected_records": 2}):
            response = self.admin.patch(self._url(self.a), payload, format="json")
            self.assertNotIn("text_in_use", response.content.decode())

    def test_no_write_to_the_records_table_is_ever_made_by_any_of_the_writes(self):
        import re

        from django.db import connection as conn
        from django.test.utils import CaptureQueriesContext

        self._record("BF-A3", self.A_NAME)
        with CaptureQueriesContext(conn) as queries:
            self.admin.post(self.LIST, {"name": "Echo new"}, format="json")
            self.admin.patch(self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 3}, format="json")
            self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
            self.admin.post(self._url(self.c, "deactivate"))
            self.admin.post(self._url(self.c, "reactivate"))
            self.admin.post(self.REORDER, {"ids": [self.d.pk, self.c.pk, self.b.pk, self.a.pk]}, format="json")
        # the table each statement WRITES to (the one named right after UPDATE / INSERT INTO / DELETE FROM);
        # a subquery that only READS api_touristrecord inside an UPDATE of the tier table is not a write to it
        target = re.compile(r'^\s*(?:UPDATE|INSERT\s+INTO|DELETE\s+FROM|REPLACE\s+INTO)\s+"?([A-Za-z0-9_]+)"?', re.I)
        written = [m.group(1) for q in queries.captured_queries if (m := target.match(q["sql"]))]
        self.assertIn("api_boatcapacityfare", written)  # the tier table WAS written
        self.assertIn("api_activitylog", written)  # and so was the log
        self.assertNotIn("api_touristrecord", written)  # the records never
        self.assertGreater(len(queries.captured_queries), 10)  # the writes did run
        reads = [q["sql"] for q in queries.captured_queries if "api_touristrecord" in q["sql"]]
        self.assertTrue(reads)  # records were counted, only counted
    def test_the_old_text_stays_valid_on_its_records_so_a_refused_rename_leaves_them_listed(self):
        self.admin.patch(self._url(self.a), {"name": "Brand new wording"}, format="json")
        listed = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        self.assertEqual(listed[self.A_NAME], 2)
        self.assertNotIn("Brand new wording", listed)

    def test_rename_to_the_name_of_another_tier_is_refused_and_nothing_changes(self):
        before = self._stored()
        response = self.admin.patch(self._url(self.b), {"name": self.D_NAME.upper()}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self._stored(), before)
        self.assertEqual(self._logs().count(), 0)

    def test_rename_to_its_own_name_changes_nothing_and_logs_nothing(self):
        before = self._stored()
        response = self.admin.patch(self._url(self.b), {"name": "  " + self.B_NAME + " "}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["name"], self.B_NAME)
        self.assertEqual(self._stored(), before)
        self.assertEqual(self._logs().count(), 0)

    def test_a_change_of_capitals_alone_is_a_rename_and_is_stored(self):
        response = self.admin.patch(self._url(self.b), {"name": self.B_NAME.upper()}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["name"], self.B_NAME.upper())
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, self.B_NAME.upper())

    def test_a_tier_may_be_renamed_to_a_text_that_records_already_hold(self):
        response = self.admin.patch(self._url(self.b), {"name": self.ORPHAN}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["record_count"], 1)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, self.ORPHAN)

    def test_rename_refuses_a_bad_name_and_an_unknown_tier(self):
        before = self._stored()
        for bad in (None, "", "   ", 5, ["x"], "x" * 161):
            response = self.admin.patch(self._url(self.b), {"name": bad}, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, repr(bad)[:30])
        self.assertEqual(self.admin.patch(self._url(self.b), {"is_active": False}, format="json").status_code, 400)
        self.assertEqual(self._stored(), before)
        self.assertEqual(self.admin.patch(self.LIST + "99999/", {"name": "x"}, format="json").status_code, 404)
        self.assertEqual(self._logs().count(), 0)

    # -- deactivate / reactivate ---------------------------------------------------

    def test_deactivate_and_reactivate_answer_and_store(self):
        response = self.admin.post(self._url(self.a, "deactivate"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIs(response.json()["is_active"], False)
        self.assertEqual(response.json()["record_count"], 2)  # retiring does not touch the records
        self.assertIs(BoatCapacityFare.objects.get(pk=self.a.pk).is_active, False)
        self.assertEqual(self._held(), self._held())
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).count(), 2)

        response = self.admin.post(self._url(self.a, "reactivate"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIs(response.json()["is_active"], True)
        self.assertIs(BoatCapacityFare.objects.get(pk=self.a.pk).is_active, True)

    def test_the_last_active_tier_may_be_deactivated_so_none_can_be_active(self):
        for tier in (self.a, self.b, self.c, self.d):
            response = self.admin.post(self._url(tier, "deactivate"))
            self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(BoatCapacityFare.objects.filter(is_active=True).count(), 0)
        self.assertEqual(BoatCapacityFare.objects.count(), 4)  # still all there
        self.assertEqual(self._logs().count(), 4)

    def test_doing_it_twice_changes_and_logs_once(self):
        self.admin.post(self._url(self.b, "deactivate"))
        self.admin.post(self._url(self.b, "deactivate"))
        self.assertEqual(self._logs().count(), 1)
        self.admin.post(self._url(self.d, "reactivate"))  # already active
        self.assertEqual(self._logs().count(), 1)

    def test_deactivating_or_reactivating_an_unknown_tier_is_404(self):
        self.assertEqual(self.admin.post(self.LIST + "99999/deactivate/").status_code, 404)
        self.assertEqual(self.admin.post(self.LIST + "99999/reactivate/").status_code, 404)
        self.assertEqual(self._logs().count(), 0)

    # -- reorder -----------------------------------------------------------------------

    def test_reorder_answers_and_stores_the_new_order(self):
        ids = [self.d.pk, self.b.pk, self.a.pk, self.c.pk]
        response = self.admin.post(self.REORDER, {"ids": ids}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = response.json()
        self.assertEqual([r["id"] for r in body], ids)
        self.assertEqual([r["display_order"] for r in body], [10, 20, 30, 40])
        self.assertEqual({r["name"]: r["record_count"] for r in body},
                         {self.A_NAME: 2, self.B_NAME: 0, self.C_NAME: 1, self.D_NAME: 0})

        stored = list(BoatCapacityFare.objects.order_by("display_order", "id").values_list("pk", "display_order"))
        self.assertEqual(stored, [(ids[0], 10), (ids[1], 20), (ids[2], 30), (ids[3], 40)])
        self.assertEqual(BoatCapacityFare.objects.count(), 4)

    def test_reorder_with_a_bad_list_is_400_and_the_database_does_not_move(self):
        a, b, c, d = (t.pk for t in (self.a, self.b, self.c, self.d))
        bad_bodies = [
            {"ids": [a, b, c]},                    # one missing
            {"ids": [a, b]},                       # two missing
            {"ids": [a, b, c, d, 99999]},          # an extra
            {"ids": [a, b, c, 99999]},             # a stranger in place of one
            {"ids": [d, d, a, b]},                 # repeated
            {"ids": [a, a, a, a]},                 # all the same
            {"ids": [a, b, c, d, a]},              # every tier, and one of them twice
            {"ids": [a, b, c, d, d]},              # likewise, at the end
            {"ids": [a, a, b, c, d]},              # likewise, at the start
            {"ids": []},                           # empty
            {"ids": "abc"},                        # not a list
            {"ids": None},
            {"ids": {"0": a}},
            {"ids": [True, b, c, d]},              # a boolean
            {"ids": [a, b, c, "4"]},               # text
            {"ids": [a, b, c, 4.5]},               # a number with a fraction
            {},                                    # no ids
            [a, b, c, d],                          # a bare list, not {"ids": ...}
        ]
        for body in bad_bodies:
            before, records = self._stored(), self._held()
            response = self.admin.post(self.REORDER, body, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, body)
            self.assertEqual(self._stored(), before, body)  # every row, every column, as before
            self.assertEqual(self._held(), records, body)
            self.assertEqual(self._logs().count(), 0, body)

    def test_reorder_with_an_empty_list_is_400_even_when_there_is_no_tier_at_all(self):
        BoatCapacityFare.objects.all().delete()
        response = self.admin.post(self.REORDER, {"ids": []}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)  # not a 500
        self.assertEqual(self._stored(), [])
        self.assertEqual(self._logs().count(), 0)

    def test_a_text_in_other_capitals_or_with_a_stray_space_does_not_hold_the_tier(self):
        self._record("BF-NEAR1", self.B_NAME.upper())
        self._record("BF-NEAR2", " " + self.B_NAME)
        before_records = self._held()
        listed = {r["name"]: r["record_count"] for r in self.admin.get(self.LIST).json()}
        self.assertEqual(listed[self.B_NAME], 0)

        response = self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)  # nobody holds EXACTLY this text
        self.assertEqual(response.json()["record_count"], 0)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, "Bravo renamed")
        self.assertEqual(self._held(), before_records)  # the near-miss records are untouched

    def test_reorder_with_ids_equal_to_one_in_a_table_where_true_would_match(self):
        BoatCapacityFare.objects.all().delete()
        for pk, name, order in ((1, "One", 10), (2, "Two", 20), (3, "Three", 30)):
            BoatCapacityFare.objects.create(id=pk, name=name, display_order=order)
        before = self._stored()
        response = self.admin.post(self.REORDER, {"ids": [3, True, 2]}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self._stored(), before)

    # -- the ActivityLog: who, which tier, what ----------------------------------------

    def _the_log(self):
        log = self._logs().get()
        self.assertEqual(log.user, self.admin.user)
        self.assertEqual(log.module, MODULE_TOURISM)
        return log

    def test_create_leaves_a_trace(self):
        body = self.admin.post(self.LIST, {"name": "Foxtrot 13-14 pax"}, format="json").json()
        log = self._the_log()
        self.assertEqual(log.action, ACTION_CREATE)
        self.assertEqual(log.record_id, str(body["id"]))
        self.assertIn("added", log.record_label)
        self.assertIn("Foxtrot 13-14 pax", log.record_label)

    def test_rename_leaves_a_trace_with_the_old_and_the_new_text(self):
        self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
        log = self._the_log()
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.b.pk))
        self.assertIn("renamed", log.record_label)
        self.assertIn(self.B_NAME, log.record_label)
        self.assertIn("Bravo renamed", log.record_label)

    def test_deactivate_leaves_a_trace(self):
        self.admin.post(self._url(self.c, "deactivate"))
        log = self._the_log()
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.c.pk))
        self.assertIn("deactivated", log.record_label)
        self.assertIn(self.C_NAME, log.record_label)

    def test_reactivate_leaves_a_trace(self):
        BoatCapacityFare.objects.filter(pk=self.c.pk).update(is_active=False)
        self.admin.post(self._url(self.c, "reactivate"))
        log = self._the_log()
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertEqual(log.record_id, str(self.c.pk))
        self.assertIn("brought back", log.record_label)
        self.assertIn(self.C_NAME, log.record_label)

    def test_reorder_leaves_a_trace_with_the_new_order(self):
        self.admin.post(self.REORDER, {"ids": [self.d.pk, self.c.pk, self.b.pk, self.a.pk]}, format="json")
        log = self._the_log()
        self.assertEqual(log.action, ACTION_UPDATE)
        self.assertIn("reordered", log.record_label)
        label = log.record_label
        self.assertTrue(label.index("Delta") < label.index("Charlie") < label.index("Bravo") < label.index("Alpha"))

    def test_a_failed_write_leaves_no_trace_of_any_kind(self):
        self.admin.post(self.LIST, {"name": ""}, format="json")                              # bad name
        self.admin.post(self.LIST, {"name": self.A_NAME}, format="json")                     # duplicate
        self.admin.patch(self._url(self.a), {"name": "x"}, format="json")                    # held text
        self.admin.patch(self._url(self.b), {"name": self.D_NAME}, format="json")            # duplicate
        self.admin.post(self.REORDER, {"ids": [self.a.pk]}, format="json")                  # incomplete
        self.admin.post(self.LIST + "99999/deactivate/")                                     # unknown
        self.assertEqual(self._logs().count(), 0)


class BoatCapacityFareRenameGuardTests(_BoatFareApiBase, TransactionTestCase):
    """The rename guard and the lock.

    A TransactionTestCase, so that "inside an atomic block" is something the test
    can see (a TestCase is wrapped in one the whole time). The real PostgreSQL
    lock cannot be run here (the tests use SQLite, which has one writer at a
    time), so what is proved is the order and the block it is in, and that the
    check and the write are one statement that refuses if a record is there.
    """

    VIEWS = "api.views.boat_capacity_fares"

    def setUp(self):
        self.setUpData()

    def test_a_record_that_arrives_between_the_count_and_the_write_is_still_refused_for_an_unheld_tier(self):
        from unittest import mock

        from api.views import boat_capacity_fares as views

        real_count = views._count_holders
        arrived = []

        def count_then_a_booking_arrives(name):
            count = real_count(name)  # the count: nobody holds it, so no confirmation is asked
            if name == self.B_NAME and not arrived:
                arrived.append(self._record("BF-LATE", self.B_NAME))  # ...and now someone does
            return count

        with mock.patch.object(views, "_count_holders", side_effect=count_then_a_booking_arrives):
            response = self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")

        self.assertEqual(len(arrived), 1)  # the booking really did arrive in the gap
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.json()["code"], "confirm_required")
        self.assertEqual(response.json()["record_count"], 1)  # the number as it is NOW
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, self.B_NAME)  # not renamed
        self.assertEqual(self._logs().count(), 0)
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.B_NAME).count(), 1)

    def test_a_record_that_arrives_after_the_admin_confirmed_the_number_makes_the_number_wrong_and_is_refused(self):
        from unittest import mock

        from api.views import boat_capacity_fares as views

        real_count = views._count_holders
        arrived = []

        def count_then_a_booking_arrives(name):
            count = real_count(name)  # 2: the number the admin saw and confirmed
            if name == self.A_NAME and not arrived:
                arrived.append(self._record("BF-LATE", self.A_NAME))  # a third record arrives in the gap
            return count

        before_records = self._held()
        with mock.patch.object(views, "_count_holders", side_effect=count_then_a_booking_arrives):
            response = self.admin.patch(
                self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 2}, format="json"
            )

        self.assertEqual(len(arrived), 1)
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.json()["code"], "confirm_required")
        self.assertEqual(response.json()["record_count"], 3)  # the admin must now confirm 3, not 2
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.a.pk).name, self.A_NAME)  # nothing renamed
        self.assertEqual(self._logs().count(), 0)
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).count(), 3)
        self.assertEqual(
            [row for row in self._held() if row[0] != "BF-LATE"], before_records
        )  # the old records are as they were

    def test_confirming_the_new_number_after_that_refusal_then_goes_through(self):
        from unittest import mock

        from api.views import boat_capacity_fares as views

        real_count = views._count_holders
        arrived = []

        def count_then_a_booking_arrives(name):
            count = real_count(name)
            if name == self.A_NAME and not arrived:
                arrived.append(self._record("BF-LATE", self.A_NAME))
            return count

        with mock.patch.object(views, "_count_holders", side_effect=count_then_a_booking_arrives):
            first = self.admin.patch(self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": 2}, format="json")
        self.assertEqual(first.status_code, 409)
        second = self.admin.patch(
            self._url(self.a), {"name": "Alpha renamed", "confirm_affected_records": first.json()["record_count"]}, format="json"
        )
        self.assertEqual(second.status_code, status.HTTP_200_OK)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.a.pk).name, "Alpha renamed")
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).count(), 3)  # all three left on the old text
        self.assertIn("3 record(s) keep the old text", self._logs().get().record_label)
    def test_the_lock_the_check_and_the_write_are_in_one_atomic_block_in_that_order(self):
        from unittest import mock

        from django.db import connection as conn

        from api.views import boat_capacity_fares as views

        seen = []

        def spy(name, real):
            def wrapper(*args, **kwargs):
                seen.append((name, conn.in_atomic_block))
                return real(*args, **kwargs)

            return wrapper

        self.assertFalse(conn.in_atomic_block)  # here, outside the view, there is no block
        with mock.patch.object(views, "_lock", spy("lock", views._lock)), \
                mock.patch.object(views, "_count_holders", spy("count", views._count_holders)), \
                mock.patch.object(views, "_rename_if_still_held_by", spy("rename", views._rename_if_still_held_by)):
            response = self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        order = [name for name, _inside in seen]
        self.assertEqual(order[:3], ["lock", "count", "rename"])  # the lock FIRST, then the check, then the write
        self.assertTrue(all(inside for _name, inside in seen), seen)  # and every one of them inside the block

    def test_the_log_is_written_inside_the_same_block_as_the_write(self):
        from unittest import mock

        from django.db import connection as conn

        from api.views import boat_capacity_fares as views

        inside = []
        real_log = views.log_activity

        def spy(*args, **kwargs):
            inside.append(conn.in_atomic_block)
            return real_log(*args, **kwargs)

        with mock.patch.object(views, "log_activity", spy):
            self.admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
            self.admin.post(self._url(self.c, "deactivate"))
            self.admin.post(self.REORDER, {"ids": [self.d.pk, self.c.pk, self.b.pk, self.a.pk]}, format="json")
            self.admin.post(self.LIST, {"name": "Golf new"}, format="json")
        self.assertEqual(inside, [True, True, True, True])

    def test_a_log_that_fails_takes_the_write_with_it(self):
        from unittest import mock

        from api.views import boat_capacity_fares as views

        with mock.patch.object(views, "log_activity", side_effect=RuntimeError("the log is down")):
            admin = self.admin
            admin.raise_request_exception = False
            response = admin.patch(self._url(self.b), {"name": "Bravo renamed"}, format="json")
        self.assertEqual(response.status_code, 500)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, self.B_NAME)  # no write without its trace

    def test_the_rename_statement_alone_matches_only_while_the_count_is_the_expected_one(self):
        from api.views import boat_capacity_fares as views

        # B is held by nobody: expected 0 matches, expected 1 or 5 does not
        self.assertEqual(views._rename_if_still_held_by(self.b.pk, self.B_NAME, "Bravo x", 1), 0)
        self.assertEqual(views._rename_if_still_held_by(self.b.pk, self.B_NAME, "Bravo x", 5), 0)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.b.pk).name, self.B_NAME)
        self.assertEqual(views._rename_if_still_held_by(self.b.pk, self.B_NAME, "Bravo renamed", 0), 1)
        # A is held by 2: expected 2 matches, 0, 1 and 3 do not
        for wrong in (0, 1, 3):
            self.assertEqual(views._rename_if_still_held_by(self.a.pk, self.A_NAME, "Alpha x", wrong), 0)
        self.assertEqual(BoatCapacityFare.objects.get(pk=self.a.pk).name, self.A_NAME)
        # the old name no longer matches (the tier is not called that any more)
        self.assertEqual(views._rename_if_still_held_by(self.b.pk, self.B_NAME, "again", 0), 0)
        names = dict(BoatCapacityFare.objects.values_list("pk", "name"))
        self.assertEqual(names[self.b.pk], "Bravo renamed")
        self.assertEqual(names[self.a.pk], self.A_NAME)
        self.assertEqual(views._rename_if_still_held_by(self.a.pk, self.A_NAME, "Alpha renamed", 2), 1)
        self.assertEqual(TouristRecord.objects.filter(boat_capacity_fare=self.A_NAME).count(), 2)  # the records stayed
    def test_on_postgresql_the_writers_take_the_advisory_lock(self):
        from unittest import mock

        from django.db import connection as conn

        from api.views import boat_capacity_fares as views

        cursor = mock.MagicMock()
        cursor.__enter__.return_value = cursor
        with mock.patch.object(conn, "vendor", "postgresql"), mock.patch.object(conn, "cursor", return_value=cursor):
            views._lock()
        cursor.execute.assert_called_once_with("SELECT pg_advisory_xact_lock(%s)", [views.LOCK_ID])
        # and it is its own lock, not one the other admin APIs hold
        self.assertNotIn(views.LOCK_ID, {7100001, 7100002, 7100003, 7100004, 7100005})

    def test_on_sqlite_the_lock_does_nothing(self):
        from unittest import mock

        from django.db import connection as conn

        from api.views import boat_capacity_fares as views

        with mock.patch.object(conn, "cursor") as cursor:
            views._lock()
        cursor.assert_not_called()


class PeakSeasonShareTests(TestCase):
    """Key Insights percentages divide by the total the sentence names."""

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        quezon = Province.objects.get(name="Quezon")
        cls.resorts = list(Resort.objects.order_by("resort_id")[:2])
        purposes = list(VisitPurpose.objects.order_by("id")[:2])
        common = dict(
            full_name="Group", contact_number="09170000000", itinerary=Itinerary.objects.first(),
            travel_mode=TravelMode.objects.first(), boat_type=BoatType.objects.first(),
        )

        def record(survey_id, day, visitors, resort=0, booking_status="arrived", foreign=False):
            # The one foreign record also has the second purpose and resort, so
            # every leader (resort, origin, purpose) is 13 of 20, not all of it.
            return TouristRecord(
                survey_id=survey_id, arrival_date=day, resort=cls.resorts[resort], total_visitors=visitors,
                country=Country.objects.get(name="United States" if foreign else "Philippines"),
                region=None if foreign else quezon.region, province=None if foreign else quezon,
                visit_purpose=purposes[1 if foreign else 0],
                foreigner_count=visitors if foreign else 0, filipino_count=0 if foreign else visitors,
                total_male=visitors, age_8_59=visitors, status=booking_status, **common,
            )

        # 2026: July 5, August 12 (7 + 5), September 3 -> 20 visitors. The no-show
        # in October and the pending booking in September are counted nowhere;
        # counted, the pending 30 would make September the peak.
        TouristRecord.objects.bulk_create([
            record("PK-01", "2026-07-10", 5),
            record("PK-02", "2026-08-03", 7, resort=1, foreign=True),
            record("PK-03", "2026-08-20", 5),
            record("PK-04", "2026-09-15", 3),
            record("PK-05", "2026-10-02", 40, booking_status="no_show"),
            record("PK-06", "2026-09-01", 30, booking_status="pending"),
        ])

    def _answers(self, **params):
        from .services.tourism import build_tourism_question_answers

        return {a["id"]: a for a in build_tourism_question_answers({"year": "2026", **params})}

    @staticmethod
    def _stated(answer):
        import re

        match = re.search(r"leads with (\d+) visitors, equal to ([\d.]+)% of the selected total", answer)
        return int(match.group(1)), float(match.group(2))

    def test_peak_month_is_its_share_of_the_selected_total(self):
        peak = self._answers()["peak_month"]

        self.assertEqual(peak["answer"], "August 2026 leads with 12 visitors, equal to 60.0% of the selected total.")
        self.assertEqual(
            (peak["visual"]["value"], peak["visual"]["total"], peak["visual"]["percentage"]), (12, 20, 60.0)
        )

    def test_gauge_and_sentence_show_the_same_percentage(self):
        for params in ({}, {"year": "all"}, {"from": "2026-08-01", "to": "2026-09-30"}):
            with self.subTest(params=params):
                peak = self._answers(**params)["peak_month"]
                count, stated = self._stated(peak["answer"])
                self.assertEqual(stated, peak["visual"]["percentage"])
                self.assertEqual(count, peak["visual"]["value"])

        # August to September: August 12 of 15 visitors.
        self.assertEqual(self._answers(**{"from": "2026-08-01", "to": "2026-09-30"})["peak_month"]["visual"]["percentage"], 80.0)

    def test_every_share_divides_by_the_selected_total(self):
        answers = self._answers()
        total = 20  # all 2026 visitors except the no-show
        expected = {"top_resort": (13, 65.0), "top_origin": (13, 65.0), "visit_purpose": (13, 65.0), "peak_month": (12, 60.0)}
        for key, (count, share) in expected.items():
            with self.subTest(key=key):
                self.assertEqual(self._stated(answers[key]["answer"]), (count, share))
                self.assertEqual(share, round(count / total * 100, 1))

    def test_no_data_is_zero_percent(self):
        peak = self._answers(year="2025")["peak_month"]

        self.assertEqual(peak["answer"], "No matching records are available yet.")
        self.assertEqual(peak["visual"]["percentage"], 0)


class VisitorDemographicsVisualTests(TestCase):
    """Visitor Demographics draws the two numbers its answer states; the other cards are as before."""

    @classmethod
    def setUpTestData(cls):
        PeakSeasonShareTests.setUpTestData.__func__(cls)

    def _answers(self, **params):
        from .services.tourism import build_tourism_question_answers

        return {a["id"]: a for a in build_tourism_question_answers({"year": "2026", **params})}

    def test_doughnut_items_equal_the_numbers_in_its_answer(self):
        import re

        for params in ({}, {"year": "all"}, {"year": "2025"}):
            with self.subTest(params=params):
                demographics = self._answers(**params)["classification"]
                stated = re.fullmatch(
                    r"Domestic \(Filipino\): (\d+), Foreign \(International\): (\d+)\.", demographics["answer"]
                )
                self.assertEqual(demographics["visual"]["type"], "split")
                self.assertEqual(
                    demographics["visual"]["items"],
                    [
                        {"label": "Domestic (Filipino)", "value": int(stated.group(1))},
                        {"label": "Foreign (International)", "value": int(stated.group(2))},
                    ],
                )
        self.assertEqual([i["value"] for i in self._answers()["classification"]["visual"]["items"]], [13, 7])

    def test_the_other_ten_answers_and_visuals_are_unchanged(self):
        # The payload these cards had before the Demographics fix, on this fixture.
        expected = {
            "top_resort": ("Dona Choleng Camping Resort leads with 13 visitors, equal to 65.0% of the selected total.",
                           {"type": "ranking", "items": [{"label": "Dona Choleng Camping Resort", "value": 13}, {"label": "Aquazul Hotel and Resort", "value": 7}]}),
            "month_compare": ("September 2026 has 3 visitors, which is 9 lower than August 2026 (12).",
                              {"type": "comparison", "items": [{"label": "August 2026", "value": 12}, {"label": "September 2026", "value": 3}]}),
            "peak_month": ("August 2026 leads with 12 visitors, equal to 60.0% of the selected total.",
                           {"type": "share", "label": "August 2026", "value": 12, "total": 20, "percentage": 60.0}),
            "stay_type": ("Same-day visitors: 0; overnight or multi-day visitors: 20.",
                          {"type": "split", "items": [{"label": "Same Day", "value": 0}, {"label": "Overnight / multi-day", "value": 20}]}),
            "overnight_resort": ("Dona Choleng Camping Resort has the highest overnight demand with 13 visitors.",
                                 {"type": "ranking", "items": [{"label": "Dona Choleng Camping Resort", "value": 13}, {"label": "Aquazul Hotel and Resort", "value": 7}]}),
            "average_stay": ("The estimated average stay is 1.0 night(s) per visitor based on itinerary labels.",
                             {"type": "metric", "label": "Average stay", "value": 1.0, "unit": "night(s)"}),
            "top_origin": ("Quezon leads with 13 visitors, equal to 65.0% of the selected total.",
                           {"type": "ranking", "items": [{"label": "Quezon", "value": 13}, {"label": "United States", "value": 7}]}),
            "visit_purpose": ("Leisure leads with 13 visitors, equal to 65.0% of the selected total.",
                              {"type": "ranking", "items": [{"label": "Leisure", "value": 13}, {"label": "Vacation", "value": 7}]}),
            "high_demand": ("Dona Choleng Camping Resort shows the strongest recent demand with 8 visitors in the latest 30-day window (+8 versus the previous 30 days).",
                            {"type": "comparison", "items": [{"label": "Previous 30 days", "value": 0}, {"label": "Latest 30 days", "value": 8}]}),
            "validation": ("Needs review: 1 pending, 1 no-show, 0 possible duplicates, and 0 incomplete records.",
                           {"type": "stack", "items": [{"label": "Pending", "value": 1}, {"label": "No-show", "value": 1}, {"label": "Duplicates", "value": 0}, {"label": "Incomplete", "value": 0}]}),
        }
        answers = self._answers()
        self.assertEqual(list(answers), [
            "top_resort", "month_compare", "peak_month", "classification", "stay_type", "overnight_resort",
            "average_stay", "top_origin", "visit_purpose", "high_demand", "validation",
        ])
        for key, (answer, visual) in expected.items():
            with self.subTest(key=key):
                self.assertEqual((answers[key]["answer"], answers[key]["visual"]), (answer, visual))


class ReportsCountArrivedOnlyTests(TestCase):
    """Reports count arrived bookings only; pending is counted nowhere in Reports.

    One arrived booking, one pending and one no-show, each with a size of its
    own, so any wrong inclusion shows: arrived 3 (2026-05-04), pending 7
    (2026-05-04, same resort), no-show 11 (2026-05-04).
    """

    VISITOR_TABS = ("daily", "monthly", "yearly", "resort", "origin", "purpose", "transport", "boat")

    @classmethod
    def setUpTestData(cls):
        ensure_test_reference_tables()
        quezon = Province.objects.get(name="Quezon")
        resort = Resort.objects.order_by("resort_id").first()
        for survey_id, visitors, booking_status in (("AO-01", 3, "arrived"), ("AO-02", 7, "pending"), ("AO-03", 11, "no_show")):
            TouristRecord.objects.create(
                survey_id=survey_id, full_name="Arrival Group", contact_number=f"0917111{survey_id[-2:]}00",
                country=Country.objects.get(name="Philippines"), region=quezon.region, province=quezon,
                arrival_date="2026-05-04", resort=resort, itinerary=Itinerary.objects.first(),
                travel_mode=TravelMode.objects.first(), boat_type=BoatType.objects.first(),
                visit_purpose=VisitPurpose.objects.first(), total_visitors=visitors, filipino_count=visitors,
                total_male=visitors, total_female=0, age_8_59=visitors, status=booking_status,
            )

    def _report(self, report_type, year="2026"):
        from .services.tourism import build_reports_payload

        return build_reports_payload({"type": report_type, "year": year, "include_questions": "false"})

    def test_a_pending_booking_appears_in_no_visitor_tab(self):
        for report_type in self.VISITOR_TABS:
            with self.subTest(report_type=report_type):
                payload = self._report(report_type)
                self.assertEqual([row["visitors"] for row in payload["rows"]], [3])
                self.assertEqual(payload["totals"]["visitors"], 3)

    def test_an_arrived_booking_appears_on_every_visitor_tab(self):
        TouristRecord.objects.filter(survey_id="AO-02").update(status="arrived")
        for report_type in self.VISITOR_TABS:
            with self.subTest(report_type=report_type):
                payload = self._report(report_type)
                self.assertEqual(payload["totals"]["visitors"], 10)
                self.assertEqual(payload["totals"]["revenue"], 800)

    def test_the_no_show_tab_still_counts_no_shows(self):
        payload = self._report("no_show")

        self.assertEqual([row["visitors"] for row in payload["rows"]], [11])
        self.assertEqual(payload["totals"]["visitors"], 11)

    def test_key_insights_count_arrived_only_and_the_review_card_counts_pending(self):
        from .services.tourism import build_tourism_question_answers

        answers = {item["id"]: item for item in build_tourism_question_answers({"year": "2026"})}

        self.assertEqual(answers["classification"]["answer"], "Domestic (Filipino): 3, Foreign (International): 0.")
        self.assertEqual(answers["peak_month"]["visual"]["total"], 3)
        self.assertIn("1 pending, 1 no-show", answers["validation"]["answer"])

    def test_reports_match_arrival_monitoring_for_the_same_year(self):
        from .services.tourism import build_arrival_monitoring_payload

        for year in ("2026", "2027", "all"):
            with self.subTest(year=year):
                reports = self._report("daily", year=year)["totals"]
                monitoring = build_arrival_monitoring_payload({"year": year, "date": "all"})["summary"]
                self.assertEqual(
                    (reports["visitors"], reports["male"], reports["female"], reports["revenue"]),
                    (monitoring["totalArrivals"], monitoring["totalMale"], monitoring["totalFemale"], monitoring["feesCollected"]),
                )


class MobileFeedbackPhotoUploadTests(TestCase):
    """Feedback with photos, posted the way the mobile app's _multipartPost sends it."""

    URL = "/api/mobile/tourism/feedback/"
    JPEG = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00" + b"\x00" * 64

    def setUp(self):
        ensure_test_reference_tables()
        self.client = APIClient()
        self.resort = Resort.objects.first()

        from unittest.mock import MagicMock

        self.storage = MagicMock()
        self.storage.save.side_effect = lambda name, _file: name
        self.storage.url.side_effect = lambda name: f"https://storage.test/{name}"
        patcher = patch("api.services.upload.default_storage", self.storage)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _fields(self):
        # Same string fields as ApiService.submitFeedback
        return {
            "destination_id": str(self.resort.resort_id),
            "reviewer": "Mobile Tester",
            "rating": "4",
            "message": "Nice place",
            "cleanliness_rating": "5",
            "sanitation_comment": "Clean",
        }

    def _photo(self, index=0, name=None, content=None, content_type="image/jpeg"):
        return SimpleUploadedFile(
            name or f"report_1700000000000_{index}.jpg",
            content if content is not None else self.JPEG,
            content_type=content_type,
        )

    def _post(self, photos):
        data = self._fields()
        data["photo"] = photos  # one 'photo' part per image, like the app
        return self.client.post(self.URL, data, format="multipart")

    def test_one_photo_is_uploaded_and_urls_are_stored(self):
        response = self._post([self._photo(0)])

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        photos = response.json()["photos"]
        self.assertEqual(len(photos), 1)
        self.assertRegex(photos[0], r"^https://storage\.test/feedback/[0-9a-f]{32}\.jpg$")
        entry = FeedbackEntry.objects.get(reviewer="Mobile Tester")
        self.assertEqual(entry.photos, photos)
        self.assertNotIn("report_1700000000000", photos[0])  # storage name is a UUID, not the client's

    def test_two_photos_are_uploaded(self):
        png = self._photo(
            1,
            "report_1700000000000_1.png",
            content=b"\x89PNG\r\n\x1a\n" + b"\x00" * 32,
            content_type="image/png",
        )
        response = self._post([self._photo(0), png])

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        photos = response.json()["photos"]
        self.assertEqual(len(photos), 2)
        self.assertEqual(len(set(photos)), 2)
        self.assertTrue(photos[1].endswith(".png"))
        self.assertEqual(self.storage.save.call_count, 2)

    def test_five_photos_allowed_six_rejected(self):
        ok = self._post([self._photo(i) for i in range(5)])
        self.assertEqual(ok.status_code, status.HTTP_201_CREATED, ok.content)
        self.assertEqual(len(ok.json()["photos"]), 5)

        self.storage.save.reset_mock()
        too_many = self._post([self._photo(i) for i in range(6)])
        self.assertEqual(too_many.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("maximum of 5 photos", too_many.json()["detail"])
        self.storage.save.assert_not_called()

    def test_feedback_without_photos_still_works(self):
        as_json = self.client.post(self.URL, self._fields(), format="json")
        self.assertEqual(as_json.status_code, status.HTTP_201_CREATED, as_json.content)
        self.assertEqual(as_json.json()["photos"], [])

        as_multipart = self.client.post(self.URL, self._fields(), format="multipart")
        self.assertEqual(as_multipart.status_code, status.HTTP_201_CREATED, as_multipart.content)
        self.assertEqual(as_multipart.json()["photos"], [])
        self.storage.save.assert_not_called()

    def test_json_body_with_photo_url_list_still_works(self):
        payload = self._fields()
        payload["photos"] = ["https://storage.test/feedback/existing.jpg"]
        response = self.client.post(self.URL, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(response.json()["photos"], ["https://storage.test/feedback/existing.jpg"])

    def test_non_image_file_is_rejected_with_clear_400(self):
        response = self._post([self._photo(name="notes.txt", content=b"hello", content_type="text/plain")])

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["detail"], "Only JPG, PNG, and WebP images are allowed.")
        self.assertFalse(FeedbackEntry.objects.filter(reviewer="Mobile Tester").exists())
        self.storage.save.assert_not_called()

    def test_file_that_is_not_really_an_image_is_rejected(self):
        response = self._post([self._photo(name="photo.jpg", content=b"%PDF-1.4 not an image at all")])

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["detail"], "The uploaded file is not a valid image.")
        self.storage.save.assert_not_called()

    def test_generic_octet_stream_is_not_accepted_as_an_image(self):
        octet = "application/octet-stream"
        # no image extension -> rejected even though the bytes look like a JPEG
        no_ext = self._post([self._photo(name="camera_raw", content_type=octet)])
        self.assertEqual(no_ext.status_code, status.HTTP_400_BAD_REQUEST)
        # image extension but arbitrary bytes -> rejected
        junk = self._post([self._photo(name="photo.jpg", content=b"MZ\x90\x00 arbitrary binary", content_type=octet)])
        self.assertEqual(junk.status_code, status.HTTP_400_BAD_REQUEST)
        self.storage.save.assert_not_called()

        # a real JPEG named .jpg but sent as octet-stream (some Android clients) is still fine
        real = self._post([self._photo(name="photo.jpg", content_type=octet)])
        self.assertEqual(real.status_code, status.HTTP_201_CREATED, real.content)

    def test_oversized_photo_is_rejected(self):
        big = self._photo(content=b"\xff\xd8\xff" + b"0" * (5 * 1024 * 1024))
        response = self._post([big])

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["detail"], "File size exceeds 5MB limit.")
        self.storage.save.assert_not_called()

    def test_one_bad_photo_stores_nothing(self):
        bad = self._photo(1, name="notes.txt", content=b"x", content_type="text/plain")
        response = self._post([self._photo(0), bad])

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.storage.save.assert_not_called()
        self.assertFalse(FeedbackEntry.objects.filter(reviewer="Mobile Tester").exists())

    def test_valid_photo_larger_than_django_memory_threshold_is_accepted(self):
        # Uploads over 2.5 MB are spooled to a temp file by Django. Real phone photos are often
        # 2.5-5 MB, so this must work (copying request.data used to crash on such files).
        photo = self._photo(content=self.JPEG + b"\x00" * (3 * 1024 * 1024))
        response = self._post([photo])

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(len(response.json()["photos"]), 1)
        self.storage.save.assert_called_once()

    def test_storage_failure_returns_503_and_no_entry(self):
        self.storage.save.side_effect = IOError("Storage quota exceeded")
        response = self._post([self._photo(0)])

        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertIn("Image upload failed", response.json()["detail"])
        self.assertFalse(FeedbackEntry.objects.filter(reviewer="Mobile Tester").exists())


class AmbulantFoodVendorBusinessTypeTests(TestCase):
    """Migration 0034: the client-confirmed Ambulant Food Vendor type, with no requirements yet."""

    NAME = "Ambulant Food Vendor"

    def setUp(self):
        import importlib

        self.migration = importlib.import_module(
            "api.migrations.0034_add_ambulant_food_vendor_business_type"
        )

    def _run_migration(self):
        from django.apps import apps

        self.migration.add_ambulant_food_vendor(apps, None)

    def _ambulant(self):
        return SanitaryBusinessType.objects.get(name=self.NAME)

    def test_type_exists_with_monthly_inspections(self):
        ambulant = self._ambulant()
        self.assertEqual(ambulant.inspection_frequency, "monthly")
        self.assertEqual(ambulant.description, "")

    def test_type_has_no_requirements_for_sp_or_large(self):
        ambulant = self._ambulant()
        self.assertEqual(ambulant.requirements.count(), 0)
        for permit_size in ("sp", "large"):
            self.assertFalse(ambulant.requirements.filter(permit_size=permit_size).exists())

    def test_running_again_does_not_duplicate_or_change_the_row(self):
        before = self._ambulant()
        self._run_migration()
        self.assertEqual(SanitaryBusinessType.objects.filter(name__iexact=self.NAME).count(), 1)
        self.assertEqual(self._ambulant().pk, before.pk)

    def test_existing_row_in_other_casing_is_left_untouched(self):
        SanitaryBusinessType.objects.filter(name=self.NAME).delete()
        existing = SanitaryBusinessType.objects.create(
            name="ambulant food vendor", inspection_frequency="quarterly"
        )

        self._run_migration()

        self.assertEqual(SanitaryBusinessType.objects.filter(name__iexact=self.NAME).count(), 1)
        existing.refresh_from_db()
        self.assertEqual(existing.name, "ambulant food vendor")
        self.assertEqual(existing.inspection_frequency, "quarterly")

    def test_seed_requirement_sync_adds_no_requirements_for_sp_or_large(self):
        from .seed_data import SANITARY_REQUIREMENTS
        from .seeders import sync_sanitary_business_types_and_requirements

        self.assertFalse(
            any(group["business_type"] == self.NAME for group in SANITARY_REQUIREMENTS)
        )

        sync_sanitary_business_types_and_requirements()

        ambulant = self._ambulant()
        self.assertEqual(ambulant.inspection_frequency, "monthly")
        self.assertEqual(ambulant.requirements.count(), 0)

    def test_bootstrap_offers_the_type_with_no_requirements(self):
        client = APIClient()
        user = User.objects.create_user(username="ambulant_staff", password="Password@123")
        UserProfile.objects.create(user=user, role=ROLE_SANITATION)
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")

        response = client.get("/api/sanitation/bootstrap/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        ambulant = [t for t in response.json()["businessTypes"] if t["name"] == self.NAME]
        self.assertEqual(len(ambulant), 1)
        self.assertEqual(ambulant[0]["inspection_frequency"], "monthly")
        self.assertEqual(ambulant[0]["requirements"], [])

    def test_establishment_registered_as_sp_or_large_gets_no_requirements(self):
        client = APIClient()
        user = User.objects.create_user(username="ambulant_register", password="Password@123")
        UserProfile.objects.create(user=user, role=ROLE_SANITATION)
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        ambulant = self._ambulant()

        for permit_size in ("sp", "large"):
            response = client.post(
                "/api/sanitation/establishments/",
                {
                    "business_name": f"Ambulant Vendor {permit_size}",
                    "owner_name": "Ana Reyes",
                    "business_type": ambulant.id,
                    "permit_size": permit_size,
                    "barangay": "Daungan",
                    "address": "1 Test St",
                    "contact_number": "09170000000",
                    "latitude": 14.188,
                    "longitude": 121.732,
                    "has_permit": False,
                    "permit_number": "",
                    "permit_issued_date": None,
                    "permit_expiry_date": None,
                    "compliance_status": "no_permit",
                    "permit_status": "no_permit",
                },
                format="json",
            )
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
            created = SanitaryEstablishment.objects.get(pk=response.data["id"])
            self.assertEqual(created.business_type_id, ambulant.id)
            self.assertEqual(created.permit_size, permit_size)

        self.assertEqual(ambulant.requirements.count(), 0)


class InspectionDraftIsolationTests(TestCase):
    """A draft is work in progress, so it must not touch live records.

    Saving a draft must leave the establishment's compliance and permit status
    alone and raise no violation notification. Finalizing that same draft
    applies the result exactly once.
    """

    def setUp(self):
        self.client = APIClient()

        self.admin_user = User.objects.create_user(
            username="draft_test_admin",
            password="Password@123",
            email="draft-admin@test.local",
        )
        UserProfile.objects.create(user=self.admin_user, role=ROLE_ADMIN)

        self.sanitation_user = User.objects.create_user(
            username="draft_test_sanitation",
            password="Password@123",
            email="draft-sanitation@test.local",
        )
        UserProfile.objects.create(user=self.sanitation_user, role=ROLE_SANITATION)

        self.token = Token.objects.create(user=self.sanitation_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        self.btype = SanitaryBusinessType.objects.create(
            name="Draft Test Cafe",
            inspection_frequency="monthly",
        )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Draft Street Cafe",
            owner_name="Owner Ana",
            business_type=self.btype,
            barangay="Poblacion",
            address="12 Quezon St",
            compliance_status="good_standing",
            permit_status="active",
        )

    def violation_payload(self, is_draft):
        return {
            "establishment": self.establishment.id,
            "inspector_name": "Inspector Test",
            "inspection_date": "2026-09-11",
            "status_after_inspection": "violation",
            "findings": "Draft in progress",
            "is_draft": is_draft,
            "checklist_items": [
                {"requirement_name": "Pest Control", "is_complied": False},
            ],
        }

    def notification_count(self):
        return Notification.objects.filter(
            recipient_user__in=[self.admin_user, self.sanitation_user]
        ).count()

    def assertEstablishmentUntouched(self):
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "good_standing")
        self.assertEqual(self.establishment.permit_status, "active")
        self.assertEqual(self.notification_count(), 0)

    def test_saving_a_draft_leaves_the_establishment_alone(self):
        response = self.client.post(
            "/api/sanitation/inspections/",
            self.violation_payload(is_draft=True),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(response.json()["is_draft"])
        self.assertEstablishmentUntouched()

    def test_updating_a_draft_leaves_the_establishment_alone(self):
        created = self.client.post(
            "/api/sanitation/inspections/",
            self.violation_payload(is_draft=True),
            format="json",
        )
        inspection_id = created.json()["id"]

        payload = self.violation_payload(is_draft=True)
        payload["findings"] = "Still drafting"
        response = self.client.put(
            f"/api/sanitation/inspections/{inspection_id}/",
            payload,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEstablishmentUntouched()

    def test_mobile_draft_leaves_the_establishment_alone(self):
        response = self.client.post(
            "/api/mobile/sanitation/inspections/",
            self.violation_payload(is_draft=True),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEstablishmentUntouched()

    def test_finalizing_a_draft_applies_it_once(self):
        created = self.client.post(
            "/api/sanitation/inspections/",
            self.violation_payload(is_draft=True),
            format="json",
        )
        inspection_id = created.json()["id"]
        self.assertEstablishmentUntouched()

        response = self.client.put(
            f"/api/sanitation/inspections/{inspection_id}/",
            self.violation_payload(is_draft=False),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "violation")
        self.assertEqual(self.establishment.permit_status, "active")
        self.assertEqual(
            Notification.objects.filter(recipient_user=self.admin_user).count(), 1
        )
        self.assertEqual(
            Notification.objects.filter(recipient_user=self.sanitation_user).count(), 1
        )

    def test_a_final_inspection_still_applies_immediately(self):
        response = self.client.post(
            "/api/sanitation/inspections/",
            self.violation_payload(is_draft=False),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "violation")
        self.assertEqual(self.establishment.permit_status, "active")
        self.assertEqual(
            Notification.objects.filter(recipient_user=self.admin_user).count(), 1
        )


class InspectionNextDueDateTests(TestCase):
    """The next due date follows the business type's inspection frequency.

    annual -> +1 year, quarterly -> +3 months, monthly -> +1 month. An
    unrecognised frequency yields no suggestion rather than a silent monthly
    one, and a date sent by the client always wins.
    """

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="due_date_sanitation",
            password="Password@123",
            email="due@test.local",
        )
        UserProfile.objects.create(user=self.user, role=ROLE_SANITATION)
        self.token = Token.objects.create(user=self.user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

    def establishment_with(self, frequency, name):
        btype = SanitaryBusinessType.objects.create(
            name=f"Due Date {name}",
            inspection_frequency=frequency,
        )
        return SanitaryEstablishment.objects.create(
            business_name=f"Due Date {name} Shop",
            owner_name="Owner",
            business_type=btype,
            barangay="Poblacion",
            address="1 Main St",
            compliance_status="good_standing",
            permit_status="active",
        )

    def post_inspection(self, establishment, **extra):
        payload = {
            "establishment": establishment.id,
            "inspector_name": "Inspector Test",
            "inspection_date": "2026-03-15",
            "status_after_inspection": "good_standing",
        }
        payload.update(extra)
        return self.client.post(
            "/api/sanitation/inspections/", payload, format="json"
        )

    def test_annual_adds_one_year(self):
        establishment = self.establishment_with("annual", "Annual")
        response = self.post_inspection(establishment)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["next_due_date"], "2027-03-15")

    def test_quarterly_adds_three_months(self):
        establishment = self.establishment_with("quarterly", "Quarterly")
        response = self.post_inspection(establishment)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["next_due_date"], "2026-06-15")

    def test_monthly_adds_one_month(self):
        establishment = self.establishment_with("monthly", "Monthly")
        response = self.post_inspection(establishment)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["next_due_date"], "2026-04-15")

    def test_unknown_frequency_suggests_nothing(self):
        establishment = self.establishment_with("monthly", "Unknown")
        SanitaryBusinessType.objects.filter(
            pk=establishment.business_type_id
        ).update(inspection_frequency="fortnightly")

        response = self.post_inspection(establishment)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIsNone(response.json()["next_due_date"])

    def test_explicit_next_due_date_wins(self):
        establishment = self.establishment_with("annual", "Explicit")
        response = self.post_inspection(establishment, next_due_date="2026-05-01")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["next_due_date"], "2026-05-01")

    def test_a_draft_is_not_given_a_due_date(self):
        establishment = self.establishment_with("annual", "Draft")
        response = self.post_inspection(establishment, is_draft=True)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIsNone(response.json()["next_due_date"])

    def test_month_end_rolls_back_to_a_real_date(self):
        establishment = self.establishment_with("monthly", "MonthEnd")
        response = self.post_inspection(
            establishment, inspection_date="2026-01-31"
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["next_due_date"], "2026-02-28")


class SanitationInspectorListTests(TestCase):
    """The inspector picker needs names only, and only to sanitation staff.

    It must expose no contact details, must list active sanitation and admin
    accounts, and must be closed to every other role.
    """

    ENDPOINT = "/api/sanitation/inspectors/"

    def setUp(self):
        self.client = APIClient()

        self.sanitation_user = self.make_user(
            "insp_list_sanitation", ROLE_SANITATION, first="Ana", last="Reyes"
        )
        self.admin_user = self.make_user(
            "insp_list_admin", ROLE_ADMIN, first="Ben", last="Cruz"
        )
        self.nameless_user = self.make_user("insp_list_nameless", ROLE_SANITATION)
        self.inactive_user = self.make_user(
            "insp_list_inactive", ROLE_SANITATION, first="Old", last="Staff"
        )
        self.inactive_user.is_active = False
        self.inactive_user.save(update_fields=["is_active"])

        self.tourism_user = self.make_user("insp_list_tourism", ROLE_TOURISM)
        self.tourist_user = self.make_user("insp_list_tourist", ROLE_TOURIST)
        self.establishment_user = self.make_user(
            "insp_list_establishment", ROLE_ESTABLISHMENT
        )

    def make_user(self, username, role, first="", last=""):
        user = User.objects.create_user(
            username=username,
            password="Password@123",
            email=f"{username}@test.local",
            first_name=first,
            last_name=last,
        )
        UserProfile.objects.create(user=user, role=role)
        return user

    def authenticate(self, user):
        token, _ = Token.objects.get_or_create(user=user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

    def test_sanitation_staff_can_list_inspectors(self):
        self.authenticate(self.sanitation_user)
        response = self.client.get(self.ENDPOINT)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        names = [row["name"] for row in response.json()]
        self.assertIn("Ana Reyes", names)
        self.assertIn("Ben Cruz", names)

    def test_only_id_and_name_are_exposed(self):
        self.authenticate(self.sanitation_user)
        response = self.client.get(self.ENDPOINT)

        for row in response.json():
            self.assertEqual(set(row.keys()), {"id", "name"})

        body = response.content.decode()
        self.assertNotIn("@test.local", body)
        self.assertNotIn("date_joined", body)

    def test_an_account_without_a_name_falls_back_to_its_username(self):
        self.authenticate(self.sanitation_user)
        response = self.client.get(self.ENDPOINT)

        names = [row["name"] for row in response.json()]
        self.assertIn("insp_list_nameless", names)

    def test_inactive_accounts_are_excluded(self):
        self.authenticate(self.sanitation_user)
        response = self.client.get(self.ENDPOINT)

        names = [row["name"] for row in response.json()]
        self.assertNotIn("Old Staff", names)

    def test_other_roles_are_refused(self):
        for user in [
            self.tourism_user,
            self.tourist_user,
            self.establishment_user,
        ]:
            with self.subTest(user=user.username):
                self.authenticate(user)
                response = self.client.get(self.ENDPOINT)
                self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_anonymous_is_refused(self):
        self.client.credentials()
        response = self.client.get(self.ENDPOINT)

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_the_endpoint_is_read_only(self):
        self.authenticate(self.sanitation_user)
        response = self.client.post(self.ENDPOINT, {"name": "New"}, format="json")

        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)


class InspectionStatusRequiredTests(TestCase):
    """A finished inspection must say what it found.

    Silently defaulting a missing status to good_standing recorded a
    compliance judgement nobody made, so a final inspection without one is now
    rejected. Drafts are still work in progress and may omit it.
    """

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="status_required_sanitation",
            password="Password@123",
            email="status@test.local",
        )
        UserProfile.objects.create(user=self.user, role=ROLE_SANITATION)
        self.token = Token.objects.create(user=self.user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        self.btype = SanitaryBusinessType.objects.create(
            name="Status Required Cafe",
            inspection_frequency="monthly",
        )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Status Required Shop",
            owner_name="Owner",
            business_type=self.btype,
            barangay="Poblacion",
            address="1 Main St",
            compliance_status="upcoming",
            permit_status="renewal_due",
        )

    def payload(self, **extra):
        data = {
            "establishment": self.establishment.id,
            "inspector_name": "Inspector Test",
            "inspection_date": "2026-03-15",
        }
        data.update(extra)
        return data

    def post(self, endpoint, **extra):
        return self.client.post(endpoint, self.payload(**extra), format="json")

    def test_web_final_without_status_is_rejected(self):
        response = self.post("/api/sanitation/inspections/", is_draft=False)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("status_after_inspection", response.json())

        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "upcoming")
        self.assertEqual(SanitaryInspection.objects.count(), 0)

    def test_mobile_final_without_status_is_rejected(self):
        response = self.post("/api/mobile/sanitation/inspections/", is_draft=False)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(SanitaryInspection.objects.count(), 0)

    def test_a_final_inspection_defaults_to_nothing_when_the_key_is_absent(self):
        # No is_draft key at all: the serializer still treats it as final.
        response = self.client.post(
            "/api/sanitation/inspections/", self.payload(), format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_a_draft_may_omit_the_status(self):
        response = self.post("/api/sanitation/inspections/", is_draft=True)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(response.json()["is_draft"])

    def test_a_final_with_a_status_is_accepted(self):
        response = self.post(
            "/api/sanitation/inspections/",
            is_draft=False,
            status_after_inspection="violation",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.compliance_status, "violation")

    def test_the_mobile_client_still_works_because_it_always_sends_one(self):
        # Mirrors the payload the distributed APK sends.
        response = self.client.post(
            "/api/mobile/sanitation/inspections/",
            self.payload(
                next_due_date="2026-04-15",
                findings="",
                remarks="",
                status_after_inspection="good_standing",
                is_draft=False,
                checklist_items=[],
            ),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    def test_finalizing_a_draft_without_a_status_is_rejected(self):
        created = self.post("/api/sanitation/inspections/", is_draft=True)
        inspection_id = created.json()["id"]

        response = self.client.put(
            f"/api/sanitation/inspections/{inspection_id}/",
            self.payload(is_draft=False),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class NotYetInspectedStatusTests(TestCase):
    """An establishment nobody has inspected must not claim Good Standing.

    The new `not_yet_inspected` status is its own bucket: it is never counted
    as compliant, and the first finalized inspection replaces it with whatever
    the inspector actually found.
    """

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="nyi_sanitation",
            password="Password@123",
            email="nyi@test.local",
        )
        UserProfile.objects.create(user=self.user, role=ROLE_SANITATION)
        self.token = Token.objects.create(user=self.user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        self.btype = SanitaryBusinessType.objects.create(
            name="Not Yet Inspected Cafe",
            inspection_frequency="monthly",
        )

    def make_establishment(self, name, **extra):
        return SanitaryEstablishment.objects.create(
            business_name=name,
            owner_name="Owner",
            business_type=self.btype,
            barangay="Poblacion",
            address="1 Main St",
            **extra,
        )

    def test_a_new_establishment_starts_not_yet_inspected(self):
        establishment = self.make_establishment("Brand New Shop")

        self.assertEqual(establishment.compliance_status, "not_yet_inspected")
        self.assertEqual(
            establishment.get_compliance_status_display(), "Not Yet Inspected"
        )

    def test_the_status_is_a_valid_choice(self):
        field = SanitaryEstablishment._meta.get_field("compliance_status")
        values = [value for value, _ in field.choices]

        self.assertIn("not_yet_inspected", values)
        self.assertEqual(field.default, "not_yet_inspected")

    def test_an_explicit_status_still_wins(self):
        establishment = self.make_establishment(
            "Explicit Shop", compliance_status="violation"
        )

        self.assertEqual(establishment.compliance_status, "violation")

    def test_it_is_not_mapped_to_a_permit_status(self):
        from api.services.sanitation import PERMIT_STATUS_BY_COMPLIANCE

        self.assertNotIn("not_yet_inspected", PERMIT_STATUS_BY_COMPLIANCE)

    def test_the_status_counts_give_it_its_own_bucket(self):
        from api.services.sanitation import get_establishment_status_counts

        self.make_establishment("Never Seen One")
        self.make_establishment("Never Seen Two")
        self.make_establishment("Compliant One", compliance_status="good_standing")

        counts = get_establishment_status_counts(
            SanitaryEstablishment.objects.all()
        )

        self.assertEqual(counts["not_yet_inspected"], 2)
        self.assertEqual(counts["good"], 1)
        self.assertEqual(counts["total"], 3)

    def test_the_compliance_rate_ignores_never_inspected_records(self):
        from api.services.sanitation import build_sanitation_question_answers

        self.make_establishment("Compliant", compliance_status="good_standing")
        self.make_establishment("Violator", compliance_status="violation")
        self.make_establishment("Never Inspected")

        answers = build_sanitation_question_answers(
            SanitaryEstablishment.objects.all()
        )
        rate = next(
            item for item in answers if item["id"] == "compliance_rate"
        )

        # 1 of the 2 inspected establishments, not 1 of 3.
        self.assertIn("50.0", rate["answer"])

    def test_the_first_final_inspection_replaces_the_status(self):
        establishment = self.make_establishment("Freshly Inspected")
        self.assertEqual(establishment.compliance_status, "not_yet_inspected")

        response = self.client.post(
            "/api/sanitation/inspections/",
            {
                "establishment": establishment.id,
                "inspector_name": "Inspector Test",
                "inspection_date": "2026-03-15",
                "status_after_inspection": "for_completion",
                "is_draft": False,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        establishment.refresh_from_db()
        self.assertEqual(establishment.compliance_status, "for_completion")

    def test_a_draft_leaves_it_not_yet_inspected(self):
        establishment = self.make_establishment("Draft Only")

        response = self.client.post(
            "/api/sanitation/inspections/",
            {
                "establishment": establishment.id,
                "inspector_name": "Inspector Test",
                "inspection_date": "2026-03-15",
                "is_draft": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        establishment.refresh_from_db()
        self.assertEqual(establishment.compliance_status, "not_yet_inspected")


class NotYetInspectedMigrationTests(TransactionTestCase):
    """Migration 0035 changes the column definition, not the data.

    A default only applies to rows created afterwards, so every establishment
    that already carries a status must come through the migration unchanged.
    """

    migrate_from = [("api", "0034_add_ambulant_food_vendor_business_type")]
    migrate_to = [("api", "0035_add_not_yet_inspected_compliance_status")]

    def test_existing_rows_keep_their_status_across_the_migration(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_from)
        executor.loader.build_graph()

        old_apps = executor.loader.project_state(self.migrate_from).apps
        BusinessType = old_apps.get_model("api", "SanitaryBusinessType")
        Establishment = old_apps.get_model("api", "SanitaryEstablishment")

        btype = BusinessType.objects.create(
            name="Migration Proof Type", inspection_frequency="monthly"
        )
        existing = {
            "good_standing": None,
            "upcoming": None,
            "for_completion": None,
            "violation": None,
            "no_permit": None,
        }
        for value in existing:
            record = Establishment.objects.create(
                business_name=f"Migration Proof {value}",
                owner_name="Owner",
                business_type=btype,
                barangay="Poblacion",
                address="1 Main St",
                compliance_status=value,
            )
            existing[value] = record.pk

        before = dict(
            Establishment.objects.values_list("pk", "compliance_status")
        )

        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_to)
        executor.loader.build_graph()

        new_apps = executor.loader.project_state(self.migrate_to).apps
        MigratedEstablishment = new_apps.get_model("api", "SanitaryEstablishment")
        after = dict(
            MigratedEstablishment.objects.values_list("pk", "compliance_status")
        )

        # Not one row changed.
        self.assertEqual(before, after)
        for value, pk in existing.items():
            self.assertEqual(after[pk], value)

        # Only rows created after the migration pick up the new default.
        fresh = MigratedEstablishment.objects.create(
            business_name="Created After Migration",
            owner_name="Owner",
            business_type_id=btype.pk,
            barangay="Poblacion",
            address="1 Main St",
        )
        self.assertEqual(fresh.compliance_status, "not_yet_inspected")


class ClientInspectionFrequencyMigrationTests(TestCase):
    """Migration 0036: inspection frequencies per the client's form, keyed by name."""

    # Production values before this migration (public bootstrap, read-only).
    PRODUCTION_BEFORE = {
        "Water Refilling Station": "monthly",
        "Agro-industrial Establishment (Poultry / Piggery Farm)": "quarterly",
        "Sub-contractor": "annual",
        "Restaurant / Food Establishment": "monthly",
        "Massage / Physical Therapy": "quarterly",
        "Public Market Stall": "monthly",
        "Food Establishment": "monthly",
        "Commercial Non Food": "monthly",
        "Drug Store": "annual",
        "Resort / Picnic Ground": "quarterly",
        "Boatman": "annual",
        "Funeral Parlor": "annual",
        "Burial Ground": "annual",
        "Private Laboratory & Clinic": "annual",
        "Karaoke / Video Bar / CSW": "monthly",
        "Ambulant Food Vendor": "monthly",
    }

    EXPECTED = {
        "Restaurant / Food Establishment": "quarterly",
        "Public Market Stall": "quarterly",
        "Food Establishment": "quarterly",
        "Sub-contractor": "quarterly",
        "Boatman": "quarterly",
        "Agro-industrial Establishment (Poultry / Piggery Farm)": "quarterly",
        "Resort / Picnic Ground": "annual",
        "Karaoke / Video Bar / CSW": "annual",
        "Funeral Parlor": "annual",
        "Burial Ground": "annual",
        "Private Laboratory & Clinic": "annual",
        "Massage / Physical Therapy": "annual",
        "Water Refilling Station": "monthly",
        "Ambulant Food Vendor": "monthly",
        # "Depends" on the client's form: unchanged.
        "Commercial Non Food": "monthly",
        "Drug Store": "annual",
    }

    def setUp(self):
        import importlib

        self.migration = importlib.import_module(
            "api.migrations.0036_set_client_inspection_frequencies"
        )
        for name, frequency in self.PRODUCTION_BEFORE.items():
            SanitaryBusinessType.objects.update_or_create(
                name=name,
                defaults={"inspection_frequency": frequency, "description": f"{name} notes"},
            )

    def _run_migration(self):
        from django.apps import apps

        self.migration.set_client_inspection_frequencies(apps, None)

    def _frequencies(self):
        return dict(SanitaryBusinessType.objects.values_list("name", "inspection_frequency"))

    def test_sets_the_client_frequencies(self):
        self._run_migration()
        frequencies = self._frequencies()
        for name, frequency in self.EXPECTED.items():
            self.assertEqual(frequencies[name], frequency, name)

    def test_running_again_changes_nothing(self):
        self._run_migration()
        first = self._frequencies()
        self._run_migration()
        self.assertEqual(self._frequencies(), first)

    def test_depends_types_and_unlisted_types_are_untouched(self):
        SanitaryBusinessType.objects.create(name="Future Type", inspection_frequency="quarterly")
        SanitaryBusinessType.objects.filter(name="Drug Store").update(inspection_frequency="quarterly")

        self._run_migration()

        frequencies = self._frequencies()
        self.assertEqual(frequencies["Future Type"], "quarterly")
        self.assertEqual(frequencies["Drug Store"], "quarterly")
        self.assertEqual(frequencies["Commercial Non Food"], "monthly")

    def test_only_inspection_frequency_changes(self):
        before = {
            row["name"]: row
            for row in SanitaryBusinessType.objects.values("id", "name", "description")
        }
        self._run_migration()
        after = {
            row["name"]: row
            for row in SanitaryBusinessType.objects.values("id", "name", "description")
        }
        self.assertEqual(before, after)

    def test_missing_names_are_skipped_without_failing(self):
        SanitaryBusinessType.objects.filter(name__in=["Boatman", "Funeral Parlor"]).delete()

        self._run_migration()

        frequencies = self._frequencies()
        self.assertNotIn("Boatman", frequencies)
        self.assertNotIn("Funeral Parlor", frequencies)
        self.assertEqual(frequencies["Sub-contractor"], "quarterly")

    def test_reverse_is_a_noop(self):
        from django.db.migrations import RunPython

        operation = self.migration.Migration.operations[0]
        self.assertIs(operation.reverse_code, RunPython.noop)

    def test_seed_data_agrees_with_the_client_frequencies(self):
        # The optional seeder (USE_SEED_DATA) update_or_creates frequencies from
        # seed_data; if it disagreed it would silently undo this migration.
        from .seed_data import SANITARY_BUSINESS_TYPES

        for row in SANITARY_BUSINESS_TYPES:
            expected = self.migration.CLIENT_INSPECTION_FREQUENCIES.get(row["name"])
            if expected is not None:
                self.assertEqual(row["inspection_frequency"], expected, row["name"])


class EstablishmentPermitNumberTests(TestCase):
    """Recording an existing sanitary permit number when registering an establishment."""

    LIST_URL = "/api/sanitation/establishments/"

    def setUp(self):
        self.client = APIClient()
        staff = User.objects.create_user(username="permit_staff", password="Password@123")
        UserProfile.objects.create(user=staff, role=ROLE_SANITATION)
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=staff).key}")
        self.btype = SanitaryBusinessType.objects.create(
            name="Permit Test Store", inspection_frequency="annual"
        )

    def _payload(self, **overrides):
        payload = {
            "business_name": "Permit Test Store",
            "owner_name": "Ana Reyes",
            "business_type": self.btype.id,
            "barangay": "Daungan",
            "address": "5 Pier Rd",
            "compliance_status": "not_yet_inspected",
            "has_permit": False,
            "permit_number": "",
            "permit_issued_date": None,
            "permit_expiry_date": None,
            "permit_status": "no_permit",
        }
        payload.update(overrides)
        return payload

    def _with_permit(self, number="SP-2026-777", **overrides):
        return self._payload(
            has_permit=True,
            permit_number=number,
            permit_issued_date="2026-03-01",
            permit_expiry_date="2027-03-01",
            permit_status="active",
            **overrides,
        )

    def test_registering_with_a_permit_stores_it_trimmed_and_not_yet_inspected(self):
        response = self.client.post(self.LIST_URL, self._with_permit("  SP-2026-777  "), format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = SanitaryEstablishment.objects.get(pk=response.json()["id"])
        self.assertEqual(record.permit_number, "SP-2026-777")
        self.assertTrue(record.has_permit)
        self.assertEqual(record.permit_status, "active")
        self.assertEqual(str(record.permit_expiry_date), "2027-03-01")
        self.assertEqual(record.compliance_status, "not_yet_inspected")

    def test_registering_without_a_permit_is_unchanged(self):
        response = self.client.post(self.LIST_URL, self._payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        record = SanitaryEstablishment.objects.get(pk=response.json()["id"])
        self.assertFalse(record.has_permit)
        self.assertEqual(record.permit_number, "")
        self.assertIsNone(record.permit_issued_date)
        self.assertEqual(record.permit_status, "no_permit")

    def test_many_records_may_have_no_permit_number(self):
        for name in ("First Store", "Second Store"):
            response = self.client.post(self.LIST_URL, self._payload(business_name=name), format="json")
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)

    def test_duplicate_permit_number_is_rejected_ignoring_case_and_spaces(self):
        self.client.post(self.LIST_URL, self._with_permit("SP-2026-777"), format="json")
        count_before = SanitaryEstablishment.objects.count()

        response = self.client.post(
            self.LIST_URL,
            self._with_permit("  sp-2026-777 ", business_name="Copy Store"),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("already recorded", " ".join(response.json()["permit_number"]))
        self.assertEqual(SanitaryEstablishment.objects.count(), count_before)

    def test_editing_a_record_to_another_records_permit_number_is_rejected(self):
        self.client.post(self.LIST_URL, self._with_permit("SP-2026-777"), format="json")
        other = self.client.post(
            self.LIST_URL, self._payload(business_name="Other Store"), format="json"
        ).json()

        response = self.client.patch(
            f"{self.LIST_URL}{other['id']}/", {"permit_number": "SP-2026-777"}, format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_a_record_keeps_its_own_permit_number_when_edited(self):
        created = self.client.post(self.LIST_URL, self._with_permit(), format="json").json()

        response = self.client.patch(
            f"{self.LIST_URL}{created['id']}/",
            {"permit_number": "SP-2026-777", "remarks": "Checked"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)


class PublicPermitVerifyExposureTests(TestCase):
    """Exploit: permits/verify/ must only confirm a permit, by permit number."""

    URL = "/api/mobile/sanitation/permits/verify/"
    ALLOWED = {
        "verified": None,
        "establishment": {"business_name", "business_type_name", "barangay", "permit_number"},
        "permit": {
            "permit_number",
            "permit_status",
            "permit_status_label",
            "permit_issued_date",
            "permit_expiry_date",
        },
    }

    def setUp(self):
        btype = SanitaryBusinessType.objects.create(name="Verify Type", inspection_frequency="annual")
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Verify Bakery",
            owner_name="Hidden Owner Person",
            business_type=btype,
            barangay="Daungan",
            address="77 Private Lane",
            contact_number="09179998877",
            has_permit=True,
            permit_number="SP-2026-900",
            permit_issued_date="2026-01-05",
            permit_expiry_date="2027-01-05",
            permit_status="active",
            latitude=14.1911,
            longitude=121.7311,
        )

    def _verify(self, code):
        return APIClient().get(self.URL, {"code": code})

    def test_lookup_by_business_name_fails(self):
        self.assertEqual(self._verify("Verify Bakery").status_code, status.HTTP_404_NOT_FOUND)

    def test_lookup_by_establishment_id_fails(self):
        response = self._verify(str(self.establishment.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_permit_number_matches_trimmed_and_case_insensitive(self):
        response = self._verify("  sp-2026-900 ")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.json()["verified"])
        self.assertEqual(response.json()["permit"]["permit_number"], "SP-2026-900")

    def test_response_keys_are_exactly_the_allowed_set(self):
        data = self._verify("SP-2026-900").json()
        self.assertEqual(set(data), set(self.ALLOWED))
        self.assertEqual(set(data["establishment"]), self.ALLOWED["establishment"])
        self.assertEqual(set(data["permit"]), self.ALLOWED["permit"])

    def test_no_owner_contact_address_or_coordinates(self):
        body = self._verify("SP-2026-900").content.decode("utf-8")
        for secret in ("Hidden Owner Person", "09179998877", "77 Private Lane", "14.1911", "121.7311"):
            self.assertFalse(secret in body, f"{secret!r} leaked")

    def test_not_found_response_reveals_nothing(self):
        response = self._verify("SP-0000-000")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(set(response.json()), {"verified", "detail"})


class PublicReportHistoryExposureTests(TestCase):
    """Exploit: report history must return only the caller's own report."""

    URL = "/api/mobile/sanitation/reports/history/"

    def setUp(self):
        self.mine = SanitaryComplaint.objects.create(
            complaint_id="SAN-HIST-0001",
            complainant_name="Reporter One",
            contact_number="0917 123 4567",
            category="Improper waste disposal",
            barangay="Poblacion",
            reported_date="2026-09-01",
            description="Mine",
        )
        self.other = SanitaryComplaint.objects.create(
            complaint_id="SAN-HIST-0002",
            complainant_name="Reporter Two",
            contact_number="09179990000",
            category="Unsafe water source",
            barangay="Daungan",
            reported_date="2026-09-02",
            description="Someone else's report",
        )

    def _history(self, **params):
        return APIClient().get(self.URL, params)

    def test_partial_contact_returns_nothing(self):
        response = self._history(contact="0917")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertNotIn("rows", response.json())

    def test_partial_contact_with_a_reference_returns_nothing(self):
        response = self._history(contact="0917", reference="SAN-HIST-0001")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["rows"], [])

    def test_wrong_reference_returns_nothing(self):
        response = self._history(contact="09171234567", reference="SAN-HIST-0002")
        self.assertEqual(response.json()["rows"], [])

    def test_reference_alone_returns_nothing(self):
        response = self._history(reference="SAN-HIST-0001")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("contact number and the complaint ID", response.json()["detail"])

    def test_the_correct_pair_returns_only_that_report(self):
        for contact in ("09171234567", "+63 917 123 4567", "0917-123-4567"):
            response = self._history(contact=contact, reference="san-hist-0001")
            rows = response.json()["rows"]
            self.assertEqual([row["complaint_id"] for row in rows], ["SAN-HIST-0001"], contact)
            self.assertEqual(response.json()["summary"]["total"], 1)


class SanitationStaffBootstrapTests(TestCase):
    """The staff-only home for the records the public bootstrap used to expose."""

    URL = "/api/mobile/sanitation/staff-bootstrap/"

    def setUp(self):
        btype = SanitaryBusinessType.objects.create(name="Staff Type", inspection_frequency="annual")
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Staff Only Store",
            owner_name="Owner Staffview",
            business_type=btype,
            barangay="Daungan",
            address="1 Staff St",
            contact_number="09170001234",
            permit_number="SP-2026-321",
        )
        SanitaryComplaint.objects.create(
            complaint_id="SAN-STAFF-0001",
            complainant_name="Complainant",
            contact_number="09170005555",
            category="Improper waste disposal",
            barangay="Daungan",
            reported_date="2026-09-01",
            description="Staff-only complaint",
        )
        HouseholdSanitationRecord.objects.create(
            household_code="HH-STAFF-0001",
            household_head="Household Staffview",
            barangay="Daungan",
        )

    def _client_for(self, role):
        user = User.objects.create_user(username=f"staffboot_{role}", password="Password@123")
        UserProfile.objects.create(user=user, role=role)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        return client

    def test_anonymous_is_rejected(self):
        self.assertEqual(APIClient().get(self.URL).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_non_sanitation_roles_are_forbidden(self):
        for role in (ROLE_TOURISM, ROLE_TOURIST, ROLE_ESTABLISHMENT):
            response = self._client_for(role).get(self.URL)
            self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN, role)

    def test_sanitation_and_admin_get_the_staff_records(self):
        for role in (ROLE_SANITATION, ROLE_ADMIN):
            response = self._client_for(role).get(self.URL)
            self.assertEqual(response.status_code, status.HTTP_200_OK, role)
            data = response.json()
            self.assertEqual(
                set(data),
                {"establishments", "inspections", "complaintData", "householdRecords", "notifications"},
            )
            self.assertEqual(data["establishments"][0]["owner_name"], "Owner Staffview")
            self.assertEqual(data["establishments"][0]["permit_number"], "SP-2026-321")
            self.assertEqual(set(data["complaintData"]), {"summary", "rows"})
            self.assertEqual(data["complaintData"]["rows"][0]["complaint_id"], "SAN-STAFF-0001")
            self.assertEqual(data["householdRecords"][0]["household_head"], "Household Staffview")

    def test_shapes_match_the_existing_mobile_serializers(self):
        data = self._client_for(ROLE_SANITATION).get(self.URL).json()
        self.assertEqual(
            set(data["establishments"][0]),
            {
                "id", "business_name", "owner_name", "business_type", "business_type_name",
                "inspection_frequency", "permit_size", "barangay", "address", "contact_number",
                "has_permit", "permit_number", "permit_issued_date", "permit_expiry_date",
                "compliance_status", "compliance_status_label", "permit_status",
                "permit_status_label", "latitude", "longitude", "open_complaints",
            },
        )
        self.assertEqual(set(data["complaintData"]["summary"]), {"total", "pending", "open"})


class PublicSanitationBootstrapExposureTests(TestCase):
    """Exploit: the anonymous sanitation bootstrap must expose no staff records."""

    URL = "/api/mobile/sanitation/bootstrap/"
    SECRETS = (
        "Owner Publicleak",
        "09170007777",
        "SP-2026-LEAK",
        "9 Leak Street",
        "Complaint Publicleak",
        "09170008888",
        "Household Publicleak",
        "Inspector Publicleak",
        "14.1987",
        "121.7412",
        "14.1765",
        "121.7234",
    )

    def setUp(self):
        from django.core.cache import cache

        cache.clear()  # business types and barangays are cached for 15 minutes
        self.btype = SanitaryBusinessType.objects.create(name="Leak Type", inspection_frequency="annual")
        establishment = SanitaryEstablishment.objects.create(
            business_name="Leak Store",
            owner_name="Owner Publicleak",
            business_type=self.btype,
            barangay="Daungan",
            address="9 Leak Street",
            contact_number="09170007777",
            has_permit=True,
            permit_number="SP-2026-LEAK",
            permit_expiry_date=timezone.localdate() + timedelta(days=10),
            latitude=14.1987,
            longitude=121.7412,
        )
        SanitaryInspection.objects.create(
            establishment=establishment,
            inspector_name="Inspector Publicleak",
            inspection_date=timezone.localdate(),
        )
        SanitaryComplaint.objects.create(
            complaint_id="SAN-LEAK-0001",
            complainant_name="Complaint Publicleak",
            contact_number="09170008888",
            category="Improper waste disposal",
            barangay="Daungan",
            reported_date=timezone.localdate(),
            description="Leak test",
        )
        HouseholdSanitationRecord.objects.create(
            household_code="HH-LEAK-0001",
            household_head="Household Publicleak",
            barangay="Daungan",
            latitude=14.1765,
            longitude=121.7234,
        )
        self.advisory = Notification.objects.create(
            title="Boil water advisory",
            message="Public advisory",
            notification_type=NOTIFICATION_TYPE_PUBLIC_ADVISORY,
            audience_type=NOTIFICATION_AUDIENCE_PUBLIC,
            module=NOTIFICATION_MODULE_SANITATION,
            is_active=True,
        )

    def test_no_staff_record_values_are_exposed(self):
        body = APIClient().get(self.URL).content.decode("utf-8")
        for secret in self.SECRETS + ("Leak Store", "SAN-LEAK-0001", "HH-LEAK-0001"):
            self.assertFalse(secret in body, f"{secret!r} leaked")
        # No permit-expiry notice (it named the establishment and its expiry).
        self.assertFalse("permit-" in body, "permit-expiry notice leaked")
        self.assertFalse("expires on" in body, "permit-expiry notice leaked")

    def test_record_keys_are_present_but_empty_for_old_apks(self):
        from .services.sanitation import COMMUNITY_REPORT_CATEGORY_PRIORITIES

        data = APIClient().get(self.URL).json()
        self.assertEqual(
            set(data),
            {
                "businessTypes",
                "establishments",
                "inspections",
                "complaintData",
                "householdRecords",
                "barangays",
                "reportCategories",
                "notifications",
            },
        )
        self.assertIsInstance(data["reportCategories"], list)
        self.assertEqual(data["reportCategories"], list(COMMUNITY_REPORT_CATEGORY_PRIORITIES))
        self.assertEqual(data["establishments"], [])
        self.assertEqual(data["inspections"], [])
        self.assertEqual(data["complaintData"], {"summary": {}, "rows": []})
        self.assertEqual(data["householdRecords"], [])

    def test_public_data_is_kept(self):
        data = APIClient().get(self.URL).json()
        self.assertIn("Leak Type", [item["name"] for item in data["businessTypes"]])
        self.assertIsInstance(data["barangays"], list)
        self.assertEqual(
            [item["id"] for item in data["notifications"]], [f"advisory-{self.advisory.id}"]
        )


class ThrottleSettingsTests(TestCase):
    """General throttle settings."""

    def test_login_is_not_throttled(self):
        client = APIClient()
        statuses = {
            client.post(
                "/api/auth/login/", {"username": "nobody", "password": "wrong"}, format="json"
            ).status_code
            for _ in range(12)
        }
        self.assertNotIn(status.HTTP_429_TOO_MANY_REQUESTS, statuses)

    def test_throttle_uses_a_cache_shared_across_processes(self):
        from django.conf import settings

        self.assertEqual(
            settings.CACHES["throttle"]["BACKEND"],
            "django.core.cache.backends.db.DatabaseCache",
        )


@override_settings(TRACKING_CODE_KEY="test-tracking-code-key")
class ThrottleCacheTableMigrationTests(TransactionTestCase):
    """Render runs `migrate` but not build.sh, so a migration must create the
    throttle cache table; without it every throttled request would fail with 500."""

    TABLE = "api_throttle_cache"
    BEFORE = [("api", "0036_set_client_inspection_frequencies")]

    def tearDown(self):
        # Leave the table in place for other tests whatever happened here.
        call_command("createcachetable", self.TABLE, verbosity=0)

    def _tables(self):
        return connection.introspection.table_names()

    def test_migrate_creates_the_table_and_throttled_requests_never_return_500(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.BEFORE)
        with connection.cursor() as cursor:
            cursor.execute(f"DROP TABLE IF EXISTS {self.TABLE}")
        self.assertNotIn(self.TABLE, self._tables())

        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())

        self.assertIn(self.TABLE, self._tables())

        client = APIClient(raise_request_exception=False)
        statuses = [
            client.post(
                "/api/mobile/sanitation/establishment-status/",
                {"code": "MBN-2222-2222"},
                format="json",
            ).status_code
            for _ in range(21)
        ]
        self.assertTrue(all(code < 500 for code in statuses), statuses)
        self.assertEqual(statuses[-1], status.HTTP_429_TOO_MANY_REQUESTS)


class CommunityReportIdentityTests(TestCase):
    """Client decision: anonymous community reports are not accepted."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        ensure_test_barangays()

    def _payload(self, **overrides):
        payload = {
            "complainant_name": "Juana Reporter",
            "contact_number": "0917 123 4567",
            "category": "Severe Sewage Overflow",
            "priority": "low",
            "barangay": "Daungan",
            "location_address": "Corner of Rizal St.",
            "description": "The septic tank is leaking at the corner.",
            "latitude": 14.19,
            "longitude": 121.73,
        }
        payload.update(overrides)
        return payload

    def _post(self, **overrides):
        return APIClient().post(self.URL, self._payload(**overrides), format="json")

    def test_missing_name_is_rejected(self):
        response = self._post(complainant_name="  ")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.json()["detail"], "Please enter your name.")
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_missing_contact_is_rejected(self):
        response = self._post(contact_number="")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            response.json()["detail"], "Please enter a valid mobile number (e.g. 09171234567)."
        )
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_invalid_contact_is_rejected(self):
        for contact in ("12345", "0817 123 4567", "0917123456", "091712345678", "abc"):
            response = self._post(contact_number=contact)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, contact)
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_anonymous_flag_with_blank_name_is_rejected(self):
        response = self._post(complainant_name="", is_anonymous=True, anonymous=True)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_anonymous_flag_is_ignored_when_identity_is_given(self):
        response = self._post(is_anonymous=True)
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(
            SanitaryComplaint.objects.get().complainant_name, "Juana Reporter"
        )

    def test_valid_report_is_saved_with_its_category_urgency(self):
        for contact in ("09171234567", "+63 917 123 4567"):
            SanitaryComplaint.objects.all().delete()
            response = self._post(contact_number=contact)
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
            complaint = SanitaryComplaint.objects.get()
            self.assertEqual(complaint.complainant_name, "Juana Reporter")
            self.assertEqual(complaint.contact_number, "09171234567")
            # Urgency comes from the category, never from what the client sends.
            self.assertEqual(complaint.priority, "high")
            self.assertEqual(response.json()["priority"], "high")


class CommunityReportUrgencyTests(TestCase):
    """The server, not the reporter, decides urgency from the category."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        ensure_test_barangays()

    def _post(self, **overrides):
        payload = {
            "complainant_name": "Juana Reporter",
            "contact_number": "09171234567",
            "category": "Improper Garbage Disposal",
            "barangay": "Daungan",
            "location_address": "Purok 3",
            "description": "Garbage is piling up.",
        }
        payload.update(overrides)
        return APIClient().post(self.URL, payload, format="json")

    def test_client_urgency_is_ignored_for_a_standard_category(self):
        response = self._post(priority="high", urgency="urgent")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(SanitaryComplaint.objects.get().priority, "medium")

    def test_each_group_gets_its_urgency(self):
        expected = {
            "Contaminated Water Source": "high",
            "Hazardous / Medical Waste": "high",
            "Severe Sewage Overflow": "high",
            "Pest & Rodents Infestation": "medium",
            "Other Sanitation Concern": "low",
        }
        for index, (category, priority) in enumerate(expected.items()):
            response = self._post(
                category=category, priority="low" if priority != "low" else "high",
                contact_number=f"091700000{index:02d}",
            )
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, category)
            self.assertEqual(response.json()["priority"], priority, category)

    def test_unknown_category_is_rejected(self):
        response = self._post(category="Something Else")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("category", response.json()["detail"].lower())
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_missing_category_is_rejected(self):
        response = self._post(category="")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(SanitaryComplaint.objects.exists())


class CommunityReportRateLimitTests(TestCase):
    """Server-side limits on public community reports (the device counter can be reset)."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def _post(self, client=None, remote_addr="127.0.0.1", **overrides):
        payload = {
            "complainant_name": "Juana Reporter",
            "contact_number": "09171234567",
            "category": "Improper Garbage Disposal",
            "barangay": "Daungan",
            "location_address": "Purok 3",
            "description": "Garbage is piling up.",
        }
        payload.update(overrides)
        return (client or APIClient()).post(
            self.URL, payload, format="json", REMOTE_ADDR=remote_addr
        )

    def test_sixth_report_from_the_same_contact_in_a_day_is_refused(self):
        for index in range(5):
            response = self._post()
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, index)

        response = self._post(contact_number="+63 917 123 4567")
        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertEqual(
            response.json()["detail"],
            "This contact number has reached 5 reports today. Please try again tomorrow.",
        )
        self.assertEqual(SanitaryComplaint.objects.count(), 5)

    def test_a_different_contact_still_succeeds(self):
        for _ in range(5):
            self._post()
        response = self._post(contact_number="09179998888")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    def test_rejected_reports_do_not_use_up_the_contact_quota(self):
        for _ in range(3):
            self.assertEqual(
                self._post(category="Not A Category").status_code,
                status.HTTP_400_BAD_REQUEST,
            )
        for index in range(5):
            self.assertEqual(self._post().status_code, status.HTTP_201_CREATED, index)

    def test_more_than_twenty_reports_an_hour_from_one_address_is_refused(self):
        for index in range(20):
            response = self._post(contact_number=f"091700{index:05d}")
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, index)

        response = self._post(contact_number="09179990000")
        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertEqual(
            response.json()["detail"],
            "Too many reports from this device this hour. Please try again later.",
        )

        other_address = self._post(contact_number="09179990001", remote_addr="10.0.0.2")
        self.assertEqual(other_address.status_code, status.HTTP_201_CREATED)


class CommunityReportAddressTests(TestCase):
    """The reporter's typed location is its own required field."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        ensure_test_barangays()

    def _post(self, **overrides):
        payload = {
            "complainant_name": "Juana Reporter",
            "contact_number": "09171234567",
            "category": "Improper Garbage Disposal",
            "barangay": "Daungan",
            "location_address": "  Corner of Rizal St.  ",
            "description": "Garbage is piling up.",
        }
        payload.update(overrides)
        return APIClient().post(self.URL, payload, format="json")

    def test_missing_address_is_rejected(self):
        for value in ("", "   "):
            response = self._post(location_address=value)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, repr(value))
            detail = response.json()["detail"]
            self.assertEqual(detail, "Please enter the location or address of the problem.")
        # The field left out entirely.
        response = APIClient().post(
            self.URL,
            {
                "complainant_name": "Juana Reporter",
                "contact_number": "09171234567",
                "category": "Improper Garbage Disposal",
                "barangay": "Daungan",
                "description": "Garbage is piling up.",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_address_is_saved_to_its_field_and_description_is_unchanged(self):
        response = self._post()
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        complaint = SanitaryComplaint.objects.get()
        self.assertEqual(complaint.location_address, "Corner of Rizal St.")
        self.assertEqual(complaint.description, "Garbage is piling up.")
        self.assertNotIn("location_address", response.json())  # public receipt only


class CommunityReportRateLimitMessageTests(TestCase):
    """429s carry only the bilingual text; the wait is in Retry-After."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def _post(self, contact, remote_addr="127.0.0.1"):
        return APIClient().post(
            self.URL,
            {
                "complainant_name": "Juana Reporter",
                "contact_number": contact,
                "category": "Improper Garbage Disposal",
                "barangay": "Daungan",
                "location_address": "Purok 3",
                "description": "Garbage is piling up.",
            },
            format="json",
            REMOTE_ADDR=remote_addr,
        )

    def test_contact_limit_detail_is_exactly_the_bilingual_message(self):
        from .throttles import CommunityReportContactRateThrottle

        for _ in range(5):
            self._post("09171234567")
        response = self._post("09171234567")

        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertEqual(response.json()["detail"], CommunityReportContactRateThrottle.message)
        self.assertGreater(int(response["Retry-After"]), 0)

    def test_ip_limit_detail_is_exactly_the_bilingual_message(self):
        from .throttles import CommunityReportIpRateThrottle

        for index in range(20):
            self._post(f"091700{index:05d}", remote_addr="10.9.9.9")
        response = self._post("09179990000", remote_addr="10.9.9.9")

        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertEqual(response.json()["detail"], CommunityReportIpRateThrottle.message)
        self.assertGreater(int(response["Retry-After"]), 0)


def _community_report_payload(**overrides):
    payload = {
        "complainant_name": "Juana Reporter",
        "contact_number": "09171234567",
        "category": "Improper Garbage Disposal",
        "barangay": "Daungan",
        "location_address": "Purok 3",
        "description": "Garbage is piling up.",
    }
    payload.update(overrides)
    return payload


class CommunityReportIdempotencyTests(TestCase):
    """Resending the same form (e.g. after a timeout) never creates a second report."""

    URL = "/api/mobile/sanitation/reports/"
    SID = "7f1c2a9e-4b3d-4e8f-9a1b-2c3d4e5f6a7b"

    def setUp(self):
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def _post(self, **overrides):
        return APIClient().post(self.URL, _community_report_payload(**overrides), format="json")

    def test_same_id_twice_creates_one_report_and_returns_it(self):
        first = self._post(client_submission_id=self.SID)
        second = self._post(client_submission_id=self.SID)

        self.assertEqual(first.status_code, status.HTTP_201_CREATED, first.content)
        self.assertEqual(second.status_code, status.HTTP_200_OK, second.content)
        self.assertEqual(second.json(), first.json())
        self.assertEqual(SanitaryComplaint.objects.count(), 1)
        self.assertEqual(SanitaryComplaint.objects.get().client_submission_id, self.SID)

    def test_a_new_id_creates_a_new_report(self):
        self._post(client_submission_id=self.SID)
        response = self._post(client_submission_id="another-form-fill-0001")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(SanitaryComplaint.objects.count(), 2)

    def test_missing_id_is_still_accepted(self):
        for _ in range(2):
            self.assertEqual(self._post().status_code, status.HTTP_201_CREATED)
        self.assertEqual(SanitaryComplaint.objects.count(), 2)
        self.assertTrue(
            all(value is None for value in SanitaryComplaint.objects.values_list("client_submission_id", flat=True))
        )

    def test_resending_is_not_blocked_when_the_contact_is_at_its_limit(self):
        for index in range(5):
            response = self._post(client_submission_id=f"fill-{index}")
            self.assertEqual(response.status_code, status.HTTP_201_CREATED, index)
        self.assertEqual(self._post(client_submission_id="fill-5").status_code, 429)

        resend = self._post(client_submission_id="fill-4")
        self.assertEqual(resend.status_code, status.HTTP_200_OK, resend.content)
        self.assertEqual(SanitaryComplaint.objects.count(), 5)

    def test_resending_does_not_use_up_the_ip_limit(self):
        self._post(client_submission_id=self.SID)
        for _ in range(25):
            self.assertEqual(
                self._post(client_submission_id=self.SID).status_code, status.HTTP_200_OK
            )
        response = self._post(client_submission_id="fresh", contact_number="09179998888")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)


class _LargePhotoMixin:
    """Photos bigger than Django's in-memory upload limit (2.5 MB).

    Such a photo is spooled to a temporary file, and a request with one used to
    fail with 500 "cannot pickle 'BufferedRandom' instances" (APK 1.0.3 on a
    real phone, 2026-09-27) because the view deep-copied request.data.
    """

    LARGE = 3 * 1024 * 1024
    JPEG_HEAD = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    PNG_HEAD = b"\x89PNG\r\n\x1a\n"

    def setUp(self):
        from unittest.mock import MagicMock

        self.storage = MagicMock()
        self.storage.save.side_effect = lambda name, _file: name
        self.storage.url.side_effect = lambda name: f"https://storage.test/{name}"
        patcher = patch("api.services.upload.default_storage", self.storage)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _photo(self, head, size, name, content_type):
        from django.core.files.uploadedfile import SimpleUploadedFile

        return SimpleUploadedFile(name, head + b"\x00" * (size - len(head)), content_type=content_type)


class CommunityReportLargePhotoTests(_LargePhotoMixin, TestCase):
    URL = "/api/mobile/sanitation/reports/"
    SID = "0d6f3b0e-2f4a-4c1d-8e5b-9a7c6b5d4e3f"

    def setUp(self):
        super().setUp()
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def _post(self, photo, **overrides):
        # Same shape as ApiService._multipartPost: string fields plus one
        # 'photo' part per image.
        data = _community_report_payload(client_submission_id=self.SID, **overrides)
        data["photo"] = [photo]
        return APIClient().post(self.URL, data, format="multipart")

    def _assert_saved_with_photo(self, response, filename_suffix):
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        complaint = SanitaryComplaint.objects.get()
        self.assertEqual(complaint.client_submission_id, self.SID)
        self.assertTrue(complaint.photo_documentation.endswith(filename_suffix))
        self.assertEqual(self.storage.save.call_count, 1)

    def test_a_large_jpeg_is_saved(self):
        photo = self._photo(self.JPEG_HEAD, self.LARGE, "report_1700000000000_0.jpg", "image/jpeg")
        self._assert_saved_with_photo(self._post(photo), ".jpg")

    def test_a_large_png_is_saved(self):
        photo = self._photo(self.PNG_HEAD, self.LARGE, "report_1700000000000_0.png", "image/png")
        self._assert_saved_with_photo(self._post(photo), ".png")

    def test_a_small_photo_is_still_saved(self):
        photo = self._photo(self.JPEG_HEAD, 200 * 1024, "report_1700000000000_0.jpg", "image/jpeg")
        self._assert_saved_with_photo(self._post(photo), ".jpg")

    def test_resending_a_large_photo_returns_the_saved_report(self):
        make = lambda: self._photo(self.JPEG_HEAD, self.LARGE, "report_1700000000000_0.jpg", "image/jpeg")
        first = self._post(make())
        second = self._post(make())

        self.assertEqual(first.status_code, status.HTTP_201_CREATED, first.content)
        self.assertEqual(second.status_code, status.HTTP_200_OK, second.content)
        self.assertEqual(second.json()["complaint_id"], first.json()["complaint_id"])
        self.assertEqual(SanitaryComplaint.objects.count(), 1)
        self.assertEqual(self.storage.save.call_count, 1)

    def test_an_unexpected_error_does_not_leak_exception_text(self):
        photo = self._photo(self.JPEG_HEAD, 200 * 1024, "report_1700000000000_0.jpg", "image/jpeg")
        with patch("api.views.mobile.generate_complaint_id", side_effect=RuntimeError("secret internals")):
            response = self._post(photo)

        self.assertEqual(response.status_code, status.HTTP_500_INTERNAL_SERVER_ERROR)
        self.assertNotIn("secret internals", response.content.decode())
        self.assertEqual(SanitaryComplaint.objects.count(), 0)


class InspectionLargePhotoTests(_LargePhotoMixin, TestCase):
    """The staff inspection endpoint copied request.data the same way."""

    def setUp(self):
        super().setUp()
        user = User.objects.create_user(username="large_photo_inspector", password="Password@123")
        UserProfile.objects.create(user=user, role=ROLE_SANITATION)
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=user).key}")
        btype = SanitaryBusinessType.objects.create(name="Large Photo Cafe", inspection_frequency="monthly")
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Large Photo Shop",
            owner_name="Owner",
            business_type=btype,
            barangay="Daungan",
            address="1 Main St",
        )

    def test_a_large_inspection_photo_is_saved(self):
        photo = self._photo(self.JPEG_HEAD, self.LARGE, "inspection.jpg", "image/jpeg")
        response = self.client.post(
            "/api/mobile/sanitation/inspections/",
            {
                "establishment": str(self.establishment.id),
                "inspector_name": "Inspector Test",
                "inspection_date": "2026-09-27",
                "status_after_inspection": "good_standing",
                "is_draft": "false",
                "photo": photo,
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        inspection = SanitaryInspection.objects.get()
        self.assertTrue(inspection.photo_documentation.endswith(".jpg"))


class CommunityReportConcurrentSubmitTests(TransactionTestCase):
    """Two requests with the same id arriving together still make one report."""

    URL = "/api/mobile/sanitation/reports/"
    SID = "concurrent-fill-0001"

    def setUp(self):
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def test_concurrent_duplicates_create_exactly_one_row(self):
        import threading

        from django.db import connections

        barrier = threading.Barrier(2)
        results = []
        errors = []

        # Both requests pass the "already submitted?" check before either
        # saves, then both insert: the losing insert must turn into a 200 for
        # the winner's report, not a second row or a 500.
        #
        # The second insert is started 0.5 s after the first. SQLite's
        # in-memory test database answers truly simultaneous writes with
        # "database table is locked" instead of waiting; PostgreSQL makes the
        # second insert wait for the first and then raises IntegrityError,
        # which is exactly the path this staggering exercises.
        import time

        from api.views import mobile as mobile_views

        real_lookup = mobile_views.find_existing_community_report

        def lookup_after_both_arrive(submission_id):
            found = real_lookup(submission_id)
            if found is None:
                try:
                    order = barrier.wait(timeout=10)
                except threading.BrokenBarrierError:
                    order = 0
                if order == 1:
                    time.sleep(0.5)
            return found

        def send():
            try:
                client = APIClient(raise_request_exception=False)
                response = client.post(
                    self.URL,
                    _community_report_payload(client_submission_id=self.SID),
                    format="json",
                )
                results.append((response.status_code, response.json().get("complaint_id")))
            except Exception as exc:  # pragma: no cover - reported below
                errors.append(repr(exc))
            finally:
                connections.close_all()

        with patch.object(
            mobile_views, "find_existing_community_report", side_effect=lookup_after_both_arrive
        ):
            threads = [threading.Thread(target=send) for _ in range(2)]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join(timeout=30)

        self.assertEqual(errors, [])
        self.assertEqual(sorted(code for code, _ in results), [200, 201], results)
        self.assertEqual(len({complaint_id for _, complaint_id in results}), 1, results)
        self.assertEqual(SanitaryComplaint.objects.filter(client_submission_id=self.SID).count(), 1)
        self.assertEqual(SanitaryComplaint.objects.count(), 1)


class CommunityReportDuplicateInsideTransactionTests(TestCase):
    """The duplicate path must work when the request runs inside an outer
    transaction (ATOMIC_REQUESTS, or a caller's atomic block).

    On PostgreSQL a failed INSERT aborts the whole transaction unless it ran
    in its own savepoint; the lookup that follows would then fail with
    "current transaction is aborted" and the request would return 500.
    """

    URL = "/api/mobile/sanitation/reports/"
    SID = "outer-atomic-0001"

    def setUp(self):
        ensure_test_barangays()
        from django.core.cache import caches

        caches["throttle"].clear()

    def test_forced_duplicate_inside_an_outer_atomic_block_returns_200(self):
        import itertools

        from django.db import transaction

        from api.views import mobile as mobile_views

        # Preserve the savepoint regression using a submission with its original
        # payload proof. Historical rows without that proof now fail closed.
        created = APIClient().post(
            self.URL, _community_report_payload(client_submission_id=self.SID), format="json"
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        existing = SanitaryComplaint.objects.get(client_submission_id=self.SID)

        # The IP throttle, the contact throttle and the view's own check all
        # "miss" the saved report (as when two requests race), so the insert
        # runs and fails on the unique id. The lookup after that is real.
        real_lookup = mobile_views.find_existing_community_report
        calls = itertools.count()

        def racing_lookup(submission_id):
            return None if next(calls) < 3 else real_lookup(submission_id)

        with patch.object(mobile_views, "find_existing_community_report", side_effect=racing_lookup):
            with transaction.atomic():
                response = APIClient(raise_request_exception=False).post(
                    self.URL,
                    _community_report_payload(client_submission_id=self.SID),
                    format="json",
                )
                # The outer transaction is still usable afterwards.
                rows = SanitaryComplaint.objects.filter(client_submission_id=self.SID).count()

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content[:300])
        self.assertEqual(response.json()["complaint_id"], existing.complaint_id)
        self.assertEqual(rows, 1)


class CommunityReportBarangayTests(TestCase):
    """Reports must name an active official Mauban barangay."""

    URL = "/api/mobile/sanitation/reports/"

    def setUp(self):
        from django.core.cache import caches

        caches["throttle"].clear()
        ensure_test_barangays()

    def _post(self, barangay):
        return APIClient().post(
            self.URL, _community_report_payload(barangay=barangay), format="json"
        )

    def test_a_name_that_is_not_a_mauban_barangay_is_rejected(self):
        for name in ("Poblacion", "Cagsiay", "", "Unspecified"):
            response = self._post(name)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, name)
            detail = response.json()["detail"].lower()
            self.assertIn("barangay", detail)
        self.assertFalse(SanitaryComplaint.objects.exists())

    def test_an_inactive_barangay_is_rejected(self):
        Barangay.objects.filter(name="Daungan").update(is_active=False)
        self.assertEqual(self._post("Daungan").status_code, status.HTTP_400_BAD_REQUEST)

    def test_case_and_spacing_are_ignored_and_the_official_spelling_is_saved(self):
        response = self._post("  cagsiay   ii ")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        self.assertEqual(SanitaryComplaint.objects.get().barangay, "Cagsiay II")
        self.assertNotIn("barangay", response.json())  # public receipt only

    def test_the_official_list_is_the_forty_mauban_barangays(self):
        from .seed_data import MAUBAN_BARANGAYS

        self.assertEqual(len(MAUBAN_BARANGAYS), 40)
        self.assertNotIn("Poblacion", MAUBAN_BARANGAYS)


# --- Owner's tracking code (Establishment Portal, step 1) -------------------

TRACKING_TEST_KEY = "test-tracking-code-key"
TRACKING_CODE_PATTERN = r"^MBN-[2-9A-HJKMNP-Z]{4}-[2-9A-HJKMNP-Z]{4}$"


def _tracking_hash(code, key=TRACKING_TEST_KEY):
    """The expected stored digest, computed independently of the service."""
    import hashlib
    import hmac

    body = code.replace("MBN-", "").replace("-", "")
    return hmac.new(key.encode(), body.encode(), hashlib.sha256).hexdigest()


class TrackingCodeFormatTests(TestCase):
    def test_generated_codes_use_the_unambiguous_alphabet(self):
        from .services.tracking_codes import ALPHABET, new_tracking_code

        self.assertEqual(len(ALPHABET), 31)
        for character in "0O1IL":
            self.assertNotIn(character, ALPHABET)
        codes = {new_tracking_code() for _ in range(300)}
        self.assertGreater(len(codes), 295)
        for code in codes:
            self.assertRegex(code, TRACKING_CODE_PATTERN)

    def test_codes_are_accepted_in_any_case_with_or_without_dashes_or_spaces(self):
        from .services.tracking_codes import normalize_tracking_code

        for typed in (
            "MBN-2ABC-3DEF",
            "mbn-2abc-3def",
            "MBN2ABC3DEF",
            " mbn 2abc 3def ",
            "Mbn - 2aBc - 3dEf",
            "2ABC3DEF",
            "2abc-3def",
        ):
            self.assertEqual(normalize_tracking_code(typed), "2ABC3DEF", typed)

    def test_malformed_codes_normalize_to_nothing(self):
        from .services.tracking_codes import normalize_tracking_code

        for typed in (None, "", "   ", "MBN-2ABC-3DE", "MBN-2ABC-3DEFG", "MBN-2ABC-3DE0",
                      "MBN-2ABC-3DEO", "MBN-2ABC-3DE1", "MBN-2ABC-3DEI", "MBN-2ABC-3DEL",
                      "XYZ-2ABC-3DEF", "MBN-2ABC-3DE!"):
            self.assertEqual(normalize_tracking_code(typed), "", typed)


@override_settings(TRACKING_CODE_KEY=TRACKING_TEST_KEY)
class TrackingCodeGenerateTests(TestCase):
    """POST /api/sanitation/establishments/<id>/tracking-code/ (staff only)."""

    def setUp(self):
        self.btype = SanitaryBusinessType.objects.create(
            name="Tracking Code Carinderia", inspection_frequency="monthly"
        )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Aling Nena Carinderia",
            owner_name="Nena Santos",
            business_type=self.btype,
            barangay="Daungan",
            address="Purok 1",
            contact_number="09171234567",
            permit_number="SP-2026-0101",
        )
        self.staff = self._user("tracking_staff", ROLE_SANITATION, first_name="Maria", last_name="Santos")
        self.admin = self._user("tracking_admin", ROLE_ADMIN)

    def _user(self, username, role, **names):
        user = User.objects.create_user(username=username, password="Password@123", **names)
        UserProfile.objects.create(user=user, role=role)
        return user

    def _client(self, user=None):
        client = APIClient()
        if user is not None:
            token, _ = Token.objects.get_or_create(user=user)
            client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")
        return client

    def _url(self, establishment_id=None):
        return f"/api/sanitation/establishments/{establishment_id or self.establishment.id}/tracking-code/"

    def _generate(self, user=None):
        return self._client(user or self.staff).post(self._url())

    def test_staff_get_a_code_and_the_slip_details_once(self):
        response = self._generate()

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.assertEqual(response["Cache-Control"], "no-store")
        body = response.json()
        self.assertRegex(body["tracking_code"], TRACKING_CODE_PATTERN)
        self.assertEqual(
            body["establishment"],
            {
                "id": self.establishment.id,
                "business_name": "Aling Nena Carinderia",
                "permit_number": "SP-2026-0101",
                "business_type_name": "Tracking Code Carinderia",
                "barangay": "Daungan",
            },
        )
        self.assertTrue(body["issued_at"])
        self.assertEqual(body["issued_by"], "Maria Santos")

    def test_only_the_hash_is_stored(self):
        code = self._generate().json()["tracking_code"]

        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.tracking_code_hash, _tracking_hash(code))
        self.assertEqual(len(self.establishment.tracking_code_hash), 64)
        body = code.replace("MBN-", "").replace("-", "")
        row = SanitaryEstablishment.objects.filter(pk=self.establishment.pk).values().get()
        for value in row.values():
            self.assertNotIn(body, str(value))
            self.assertNotIn(code, str(value))

    def test_who_and_when_are_recorded(self):
        before = timezone.now()
        self._generate()

        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.tracking_code_issued_by, self.staff)
        self.assertGreaterEqual(self.establishment.tracking_code_issued_at, before)

    def test_a_new_code_replaces_the_old_one(self):
        first = self._generate().json()["tracking_code"]
        second = self._generate(self.admin).json()["tracking_code"]

        self.assertNotEqual(first, second)
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.tracking_code_hash, _tracking_hash(second))
        self.assertNotEqual(self.establishment.tracking_code_hash, _tracking_hash(first))
        self.assertEqual(self.establishment.tracking_code_issued_by, self.admin)

    def test_the_activity_log_records_the_issue_but_not_the_code(self):
        code = self._generate().json()["tracking_code"]

        log = ActivityLog.objects.get(record_type="SanitaryEstablishment", action=ACTION_UPDATE)
        self.assertEqual(log.user, self.staff)
        self.assertEqual(log.record_id, str(self.establishment.id))
        self.assertIn("tracking code", log.record_label.lower())
        body = code.replace("MBN-", "").replace("-", "")
        self.assertNotIn(body, log.record_label)

    def test_printing_a_slip_does_not_reorder_the_records(self):
        updated_at = SanitaryEstablishment.objects.get(pk=self.establishment.pk).updated_at
        self._generate()
        self.assertEqual(SanitaryEstablishment.objects.get(pk=self.establishment.pk).updated_at, updated_at)

    def test_non_staff_cannot_generate(self):
        tourism = self._user("tracking_tourism", ROLE_TOURISM)
        owner = self._user("tracking_owner", ROLE_ESTABLISHMENT)
        tourist = self._user("tracking_tourist", ROLE_TOURIST)

        self.assertIn(self._client().post(self._url()).status_code, (401, 403))
        for user in (tourism, owner, tourist):
            self.assertEqual(self._generate(user).status_code, status.HTTP_403_FORBIDDEN, user.username)
        self.establishment.refresh_from_db()
        self.assertIsNone(self.establishment.tracking_code_hash)
        self.assertIsNone(self.establishment.tracking_code_issued_at)

    def test_only_post_is_allowed(self):
        response = self._client(self.staff).get(self._url())
        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

    def test_an_unknown_establishment_is_404(self):
        response = self._client(self.staff).post(self._url(establishment_id=999999))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_records_show_when_a_slip_was_issued_but_never_the_hash(self):
        self._generate()
        self.establishment.refresh_from_db()
        digest = self.establishment.tracking_code_hash
        client = self._client(self.staff)

        detail = client.get(f"/api/sanitation/establishments/{self.establishment.id}/").json()
        self.assertTrue(detail["tracking_code_issued_at"])
        self.assertEqual(detail["tracking_code_issued_by_name"], "Maria Santos")
        for url in (
            f"/api/sanitation/establishments/{self.establishment.id}/",
            "/api/sanitation/establishments/",
            "/api/sanitation/bootstrap/",
            "/api/mobile/sanitation/staff-bootstrap/",
            "/api/mobile/sanitation/bootstrap/",
        ):
            text = client.get(url).content.decode()
            self.assertNotIn(digest, text, url)
            self.assertNotIn("tracking_code_hash", text, url)

    def test_a_record_without_a_slip_says_so(self):
        detail = self._client(self.staff).get(f"/api/sanitation/establishments/{self.establishment.id}/").json()
        self.assertIsNone(detail["tracking_code_issued_at"])
        self.assertEqual(detail["tracking_code_issued_by_name"], "")

    def test_the_code_fields_cannot_be_written_through_the_record(self):
        self._generate()
        self.establishment.refresh_from_db()
        digest = self.establishment.tracking_code_hash
        issued_at = self.establishment.tracking_code_issued_at

        response = self._client(self.admin).patch(
            f"/api/sanitation/establishments/{self.establishment.id}/",
            {
                "tracking_code_hash": "0" * 64,
                "tracking_code_issued_at": "2020-01-01T00:00:00Z",
                "tracking_code_issued_by": self.admin.id,
                "remarks": "edited",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.remarks, "edited")
        self.assertEqual(self.establishment.tracking_code_hash, digest)
        self.assertEqual(self.establishment.tracking_code_issued_at, issued_at)
        self.assertEqual(self.establishment.tracking_code_issued_by, self.staff)

    def test_a_hash_collision_draws_a_new_code(self):
        other = SanitaryEstablishment.objects.create(
            business_name="Other", owner_name="O", business_type=self.btype,
            barangay="Daungan", address="Purok 2",
        )
        other.tracking_code_hash = _tracking_hash("MBN-2222-2222")
        other.save()

        with patch(
            "api.services.tracking_codes.new_tracking_code",
            side_effect=["MBN-2222-2222", "MBN-3333-3333"],
        ):
            response = self._generate()

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.assertEqual(response.json()["tracking_code"], "MBN-3333-3333")
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.tracking_code_hash, _tracking_hash("MBN-3333-3333"))

    @override_settings(TRACKING_CODE_KEY="", DEBUG=False)
    def test_without_a_key_in_production_it_is_503_and_nothing_is_stored(self):
        response = self._generate()

        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertIn("not configured", response.json()["detail"])
        self.establishment.refresh_from_db()
        self.assertIsNone(self.establishment.tracking_code_hash)
        # The rest of the records still work.
        detail = self._client(self.staff).get(f"/api/sanitation/establishments/{self.establishment.id}/")
        self.assertEqual(detail.status_code, status.HTTP_200_OK)

    @override_settings(TRACKING_CODE_KEY="", DEBUG=True)
    def test_in_debug_the_secret_key_stands_in(self):
        from django.conf import settings

        code = self._generate().json()["tracking_code"]
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.tracking_code_hash, _tracking_hash(code, settings.SECRET_KEY))

    @override_settings(TRACKING_CODE_KEY="", DEBUG=False)
    def test_a_missing_key_is_warned_about_at_startup(self):
        from .services.tracking_codes import warn_if_tracking_code_key_missing

        with self.assertLogs("api.services.tracking_codes", level="WARNING") as logs:
            warn_if_tracking_code_key_missing()
        self.assertIn("TRACKING_CODE_KEY", logs.output[0])

    @override_settings(DEBUG=False)
    def test_no_warning_when_the_key_is_set(self):
        from .services.tracking_codes import warn_if_tracking_code_key_missing

        with self.assertNoLogs("api.services.tracking_codes", level="WARNING"):
            warn_if_tracking_code_key_missing()


@override_settings(TRACKING_CODE_KEY=TRACKING_TEST_KEY)
class EstablishmentStatusTests(TestCase):
    """POST /api/mobile/sanitation/establishment-status/ (public, by tracking code)."""

    URL = "/api/mobile/sanitation/establishment-status/"
    CODE = "MBN-7KQ4-XP2M"
    KEYS = {
        "business_name",
        "business_type",
        "barangay",
        "permit_number",
        "permit_status",
        "permit_status_label",
        "permit_expiry_date",
        "days_left",
        "is_expired",
        "renewal_notice",
        "expired_notice",
        "suspended_notice",
        "requirements",
        "requirements_note",
    }
    NOT_FOUND = "Tracking code not found. Check the code on your Owner's Slip."

    def setUp(self):
        from django.core.cache import caches

        from .models import SanitaryRequirement

        caches["throttle"].clear()
        self.today = timezone.localdate()
        self.btype = SanitaryBusinessType.objects.create(
            name="Status Carinderia", inspection_frequency="monthly"
        )
        for size, name in (
            ("sp", "Health Certificate"),
            ("sp", "Water Potability Test"),
            ("large", "Pest Control Contract"),
        ):
            SanitaryRequirement.objects.create(
                business_type=self.btype, permit_size=size, requirement_name=name
            )
        self.establishment = SanitaryEstablishment.objects.create(
            business_name="Aling Nena Carinderia",
            owner_name="Nena Santos",
            business_type=self.btype,
            barangay="Daungan",
            address="Purok 1, Rizal St.",
            contact_number="09171234567",
            permit_number="SP-2026-0101",
            permit_issued_date=self.today - timedelta(days=200),
            permit_expiry_date=self.today + timedelta(days=165),
            permit_status="active",
            compliance_status="good_standing",
            remarks="Inspector notes: grease trap",
            latitude=14.19,
            longitude=121.73,
        )
        self._set_code(self.establishment, self.CODE)

    def _set_code(self, establishment, code):
        establishment.tracking_code_hash = _tracking_hash(code)
        establishment.save()

    def _post(self, code, ip="10.0.0.1"):
        return APIClient().post(self.URL, {"code": code}, format="json", REMOTE_ADDR=ip)

    def _status(self, **changes):
        SanitaryEstablishment.objects.filter(pk=self.establishment.pk).update(**changes)
        response = self._post(self.CODE)
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        return response.json()

    # -- what is returned ------------------------------------------------------

    def test_the_answer_has_exactly_the_public_keys(self):
        from .models import SanitaryComplaint

        SanitaryInspection.objects.create(
            establishment=self.establishment,
            inspector_name="Inspector Juan",
            inspection_date=self.today,
            status_after_inspection="good_standing",
            remarks="Secret inspection note",
        )
        SanitaryComplaint.objects.create(
            complaint_id="CMP-2026-9001", establishment=self.establishment,
            complainant_name="Complainer Pedro", contact_number="09998887777",
            category="Improper Garbage Disposal", barangay="Daungan",
            description="Complaint text", reported_date=self.today,
        )
        other = SanitaryEstablishment.objects.create(
            business_name="Other Shop", owner_name="Other Owner", business_type=self.btype,
            barangay="Daungan", address="Purok 9",
        )
        self._set_code(other, "MBN-2222-3333")

        response = self._post(self.CODE)

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.assertEqual(response["Cache-Control"], "no-store")
        body = response.json()
        self.assertEqual(set(body), self.KEYS)
        self.assertEqual(body["business_name"], "Aling Nena Carinderia")
        self.assertEqual(body["business_type"], "Status Carinderia")
        self.assertEqual(body["barangay"], "Daungan")
        self.assertEqual(body["permit_number"], "SP-2026-0101")
        self.assertEqual(body["permit_status"], "active")
        self.assertEqual(body["permit_status_label"], "Active")
        text = response.content.decode()
        for private in (
            "Nena Santos", "09171234567", "Purok 1", "14.19", "121.73", "grease trap",
            "Inspector Juan", "Secret inspection note", "Complainer Pedro", "09998887777",
            "Complaint text", "CMP-2026-9001", "Other Shop", "Other Owner",
            self.establishment.tracking_code_hash,
        ):
            self.assertNotIn(private, text, private)
        for key in ("owner", "contact", "address", "latitude", "longitude", "remarks",
                    "inspection", "inspector", "complaint", "id", "user"):
            self.assertNotIn(key, body)

    def test_the_code_may_be_typed_in_any_case_with_or_without_dashes_or_spaces(self):
        for typed in ("MBN-7KQ4-XP2M", "mbn-7kq4-xp2m", "MBN7KQ4XP2M", " mbn 7kq4 xp2m ",
                      "7kq4-xp2m", "7KQ4XP2M", "Mbn - 7Kq4 - xP2m"):
            response = self._post(typed)
            self.assertEqual(response.status_code, status.HTTP_200_OK, typed)
            self.assertEqual(response.json()["business_name"], "Aling Nena Carinderia")

    def test_form_encoded_bodies_work_too(self):
        response = APIClient().post(self.URL, {"code": "mbn-7kq4-xp2m"}, REMOTE_ADDR="10.0.0.1")
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)

    # -- not found ---------------------------------------------------------------

    def test_wrong_empty_malformed_and_replaced_codes_get_the_same_404(self):
        self._set_code(self.establishment, "MBN-9WWW-8XXX")  # the old code is replaced

        bodies = set()
        for typed in (self.CODE, "MBN-2222-2222", "", "   ", "MBN-0000-OOOO", "hello",
                      "MBN-7KQ4-XP2", "SP-2026-0101", None):
            response = self._post(typed)
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND, typed)
            self.assertEqual(response["Cache-Control"], "no-store")
            bodies.add(response.content)
        missing = APIClient().post(self.URL, {}, format="json", REMOTE_ADDR="10.0.0.1")
        self.assertEqual(missing.status_code, status.HTTP_404_NOT_FOUND)
        bodies.add(missing.content)

        self.assertEqual(len(bodies), 1)
        self.assertEqual(json.loads(bodies.pop()), {"detail": self.NOT_FOUND})

    def test_a_bad_format_still_computes_the_hmac_and_queries(self):
        with patch("api.services.tracking_codes.hash_tracking_code",
                   wraps=__import__("api.services.tracking_codes", fromlist=["x"]).hash_tracking_code) as spy:
            self._post("not a code at all")
            self._post("")
        self.assertEqual(spy.call_count, 2)

    def test_the_lookup_is_one_query(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as queries:
            self._post("MBN-2222-2222")
        lookups = [q["sql"] for q in queries.captured_queries if "api_sanitaryestablishment" in q["sql"]]
        self.assertEqual(len(lookups), 1, lookups)
        self.assertIn("tracking_code_hash", lookups[0])

    def test_only_post_is_allowed(self):
        response = APIClient().get(self.URL, {"code": self.CODE})
        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

    # -- limits and configuration --------------------------------------------------

    def test_the_21st_request_in_an_hour_from_one_address_is_refused(self):
        statuses = [self._post("MBN-2222-2222").status_code for _ in range(20)]
        self.assertEqual(set(statuses), {status.HTTP_404_NOT_FOUND})

        refused = self._post(self.CODE)
        self.assertEqual(refused.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertEqual(refused["Cache-Control"], "no-store")
        self.assertTrue(int(refused["Retry-After"]) > 0)
        detail = refused.json()["detail"]
        self.assertEqual(detail, "Too many attempts from this device. Please try again in an hour.")
        self.assertNotIn("Expected available", detail)

        # Another address is not affected.
        self.assertEqual(self._post(self.CODE, ip="10.0.0.2").status_code, status.HTTP_200_OK)

    def test_successful_lookups_do_not_count(self):
        # Owners behind one mobile-carrier address must not lock each other out.
        statuses = [self._post(self.CODE).status_code for _ in range(25)]
        self.assertEqual(statuses, [status.HTTP_200_OK] * 25)

        for _ in range(19):
            self.assertEqual(self._post("MBN-2222-2222").status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self._post(self.CODE).status_code, status.HTTP_200_OK)
        self.assertEqual(self._post("MBN-2222-2222").status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self._post(self.CODE).status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_once_refused_every_request_from_that_address_is_refused(self):
        for _ in range(20):
            self._post("MBN-2222-2222")
        for typed in (self.CODE, "MBN-2222-2222", ""):
            self.assertEqual(self._post(typed).status_code, status.HTTP_429_TOO_MANY_REQUESTS, typed)

    @override_settings(TRACKING_CODE_KEY="", DEBUG=False)
    def test_without_a_key_in_production_it_is_503(self):
        response = self._post(self.CODE)
        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertIn("not configured", response.json()["detail"])
        self.assertEqual(response["Cache-Control"], "no-store")

    # -- expiry ----------------------------------------------------------------------

    def test_a_permit_well_before_expiry(self):
        body = self._status()
        self.assertEqual(body["permit_expiry_date"], (self.today + timedelta(days=165)).isoformat())
        self.assertEqual(body["days_left"], 165)
        self.assertFalse(body["is_expired"])
        self.assertIsNone(body["renewal_notice"])
        self.assertIsNone(body["expired_notice"])

    def test_sixty_days_left_shows_the_renewal_notice(self):
        body = self._status(permit_expiry_date=self.today + timedelta(days=60))
        self.assertEqual(body["days_left"], 60)
        self.assertEqual(
            body["renewal_notice"],
            "Your sanitary permit expires in 60 days. Please renew at the Sanitary Office.",
        )
        self.assertIsNone(body["expired_notice"])

    def test_sixty_one_days_left_shows_no_notice(self):
        body = self._status(permit_expiry_date=self.today + timedelta(days=61))
        self.assertIsNone(body["renewal_notice"])

    def test_expiring_today_is_not_expired_yet(self):
        body = self._status(permit_expiry_date=self.today)
        self.assertEqual(body["days_left"], 0)
        self.assertFalse(body["is_expired"])
        self.assertEqual(body["permit_status"], "active")
        self.assertEqual(
            body["renewal_notice"],
            "Your sanitary permit expires today. Please renew at the Sanitary Office.",
        )
        self.assertIsNone(body["expired_notice"])

    def test_a_past_date_is_expired_whatever_the_stored_status(self):
        body = self._status(permit_expiry_date=self.today - timedelta(days=1), permit_status="active")

        self.assertEqual(body["days_left"], -1)
        self.assertTrue(body["is_expired"])
        self.assertEqual(body["permit_status"], "expired")
        self.assertEqual(body["permit_status_label"], "Expired")
        self.assertIsNone(body["renewal_notice"])
        self.assertEqual(
            body["expired_notice"],
            "Your sanitary permit has expired. Please renew at the Sanitary Office.",
        )
        self.establishment.refresh_from_db()
        self.assertEqual(self.establishment.permit_status, "active")  # stored field unchanged

    def test_no_expiry_date(self):
        body = self._status(permit_expiry_date=None)
        self.assertIsNone(body["permit_expiry_date"])
        self.assertIsNone(body["days_left"])
        self.assertFalse(body["is_expired"])
        self.assertIsNone(body["renewal_notice"])
        self.assertIsNone(body["expired_notice"])

    def test_days_left_use_the_manila_date(self):
        from datetime import datetime, timezone as dt_timezone

        # 2026-03-01 17:00 UTC is already 2026-03-02 in Manila.
        SanitaryEstablishment.objects.filter(pk=self.establishment.pk).update(
            permit_expiry_date="2026-03-12"
        )
        with patch("django.utils.timezone.now",
                   return_value=datetime(2026, 3, 1, 17, 0, tzinfo=dt_timezone.utc)):
            body = self._post(self.CODE).json()
        self.assertEqual(body["days_left"], 10)

    # -- suspended / no permit ---------------------------------------------------------

    def test_suspended_adds_only_the_contact_notice(self):
        body = self._status(permit_status="suspended", compliance_status="violation")

        self.assertEqual(body["permit_status"], "suspended")
        self.assertEqual(body["permit_status_label"], "Suspended")
        self.assertEqual(body["suspended_notice"], "Please contact the Sanitary Office.")
        self.assertNotIn("violation", response_text := json.dumps(body))
        self.assertNotIn("grease", response_text)

    def test_not_suspended_has_no_suspended_notice(self):
        self.assertIsNone(self._status()["suspended_notice"])

    def test_no_permit_number_is_null(self):
        body = self._status(permit_number="", has_permit=False, permit_status="no_permit",
                            permit_expiry_date=None, permit_issued_date=None)
        self.assertIsNone(body["permit_number"])
        self.assertEqual(body["permit_status"], "no_permit")
        self.assertEqual(body["permit_status_label"], "No Permit")
        self.assertIsNone(body["days_left"])

    # -- requirements checklist ------------------------------------------------------------

    def _renewal(self, stage="requirements_review", submitted=(), renewal_id="REN-1", days=0):
        from .models import SanitaryPermitRenewal

        renewal = SanitaryPermitRenewal.objects.create(
            renewal_id=renewal_id,
            establishment=self.establishment,
            permit_number="SP-2026-0101",
            expiration_date=self.today + timedelta(days=30),
            stage=stage,
            submitted_requirements=list(submitted),
        )
        if days:
            SanitaryPermitRenewal.objects.filter(pk=renewal.pk).update(
                created_at=timezone.now() - timedelta(days=days)
            )
        return renewal

    def test_the_newest_unreleased_renewal_marks_submitted_and_missing(self):
        self._renewal(renewal_id="REN-OLD", submitted=["Health Certificate", "Water Potability Test"], days=40)
        self._renewal(renewal_id="REN-NEW", submitted=["health certificate "])
        self._renewal(renewal_id="REN-DONE", stage="released", submitted=[])

        body = self._status()

        self.assertEqual(
            body["requirements"],
            [
                {"name": "Health Certificate", "submitted": True},
                {"name": "Water Potability Test", "submitted": False},
            ],
        )
        self.assertIsNone(body["requirements_note"])

    def test_a_large_permit_uses_the_large_requirements(self):
        self._renewal(submitted=["Pest Control Contract"])
        body = self._status(permit_size="large")
        self.assertEqual(body["requirements"], [{"name": "Pest Control Contract", "submitted": True}])

    def test_a_submitted_item_that_is_no_longer_configured_is_still_listed(self):
        self._renewal(submitted=["Old Barangay Clearance"])
        names = [item["name"] for item in self._status()["requirements"]]
        self.assertEqual(names, ["Health Certificate", "Water Potability Test", "Old Barangay Clearance"])

    def test_without_an_open_renewal_the_list_is_neutral(self):
        self._renewal(stage="released", submitted=["Health Certificate"])

        body = self._status()

        self.assertEqual(
            body["requirements"],
            [
                {"name": "Health Certificate", "submitted": None},
                {"name": "Water Potability Test", "submitted": None},
            ],
        )
        self.assertEqual(body["requirements_note"], "Bring at renewal")

    def test_a_type_without_requirements_says_so(self):
        from .models import SanitaryRequirement

        SanitaryRequirement.objects.all().delete()
        body = self._status()
        self.assertEqual(body["requirements"], [])
        self.assertEqual(body["requirements_note"], "No requirements set yet")

        self._renewal()
        body = self._status()
        self.assertEqual(body["requirements"], [])
        self.assertEqual(body["requirements_note"], "No requirements set yet")

    def test_a_size_without_its_own_list_falls_back_to_the_type_list(self):
        from .models import SanitaryRequirement

        SanitaryRequirement.objects.filter(permit_size="large").delete()
        body = self._status(permit_size="large")
        self.assertEqual(
            [item["name"] for item in body["requirements"]],
            ["Health Certificate", "Water Potability Test"],
        )


class TourismThemeSettingTests(TestCase):
    def setUp(self):
        from .models import TourismThemeSetting

        self.Model = TourismThemeSetting

    def test_load_creates_the_row_once_and_returns_the_same_row(self):
        self.assertEqual(self.Model.objects.count(), 0)
        first = self.Model.load()
        second = self.Model.load()
        self.assertEqual(first.pk, 1)
        self.assertEqual(second.pk, first.pk)
        self.assertEqual(self.Model.objects.count(), 1)

    def test_saving_a_second_instance_does_not_create_a_second_row(self):
        self.Model.load()
        self.Model(primary_color="#123456").save()
        self.assertEqual(self.Model.objects.count(), 1)
        self.assertEqual(self.Model.load().primary_color, "#123456")

    def test_defaults_are_the_original_green(self):
        setting = self.Model.load()
        self.assertEqual(setting.primary_color, "#2FA34A")
        self.assertTrue(setting.mobile_follows_web)
        self.assertEqual(setting.mobile_primary_color, "")
        self.assertEqual(
            setting.saved_colors,
            [{"hex": "#2FA34A", "label": "Original (Green)"}],
        )

    def test_lowercase_hex_is_stored_uppercase(self):
        setting = self.Model.load()
        setting.primary_color = "#ff8800"
        setting.mobile_primary_color = "#00aabb"
        setting.saved_colors = [{"hex": "#abcdef", "label": "Test"}]
        setting.save()
        setting.refresh_from_db()
        self.assertEqual(setting.primary_color, "#FF8800")
        self.assertEqual(setting.mobile_primary_color, "#00AABB")
        self.assertEqual(setting.saved_colors, [{"hex": "#ABCDEF", "label": "Test"}])

    def test_primary_color_must_be_a_six_digit_hex(self):
        from django.core.exceptions import ValidationError

        setting = self.Model.load()
        for bad in ["2FA34A", "#2FA34", "#2FA34AA", "#GGGGGG", "", "red"]:
            setting.primary_color = bad
            with self.assertRaises(ValidationError, msg=bad):
                setting.full_clean()
        setting.primary_color = "#2fa34a"
        setting.full_clean()

    def test_mobile_primary_color_accepts_empty_and_rejects_malformed(self):
        from django.core.exceptions import ValidationError

        setting = self.Model.load()
        setting.mobile_primary_color = ""
        setting.full_clean()
        setting.mobile_primary_color = "#12345"
        with self.assertRaises(ValidationError):
            setting.full_clean()
        setting.mobile_primary_color = "#A1B2C3"
        setting.full_clean()


class TourismThemeApiTests(TestCase):
    URL = "/api/tourism-theme/"
    READ_KEYS = {"primary_color", "mobile_follows_web", "mobile_primary_color", "saved_colors", "updated_at"}

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.anon = APIClient()

    def _client_for(self, username, role, is_staff=True):
        user = User.objects.create_user(username=username, password="pass-12345", is_staff=is_staff)
        UserProfile.objects.create(user=user, role=role)
        token = Token.objects.create(user=user)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")
        return user, client

    def _admin(self):
        return self._client_for("theme_admin", ROLE_ADMIN)

    def _put(self, client, **payload):
        body = {"primary_color": "#FF8800", **payload}
        return client.put(self.URL, body, format="json")

    # --- read -------------------------------------------------------------

    def test_unauthenticated_get_returns_the_public_keys(self):
        response = self.anon.get(self.URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(set(response.json().keys()), self.READ_KEYS)
        self.assertEqual(response.json()["primary_color"], "#2FA34A")

    def test_read_never_exposes_updated_by_or_a_username(self):
        admin, client = self._admin()
        self.assertEqual(self._put(client).status_code, status.HTTP_200_OK)
        for response in (self.anon.get(self.URL), client.get(self.URL)):
            body = response.json()
            self.assertNotIn("updated_by", body)
            self.assertNotIn(admin.username, response.content.decode())

    # --- write gate -------------------------------------------------------

    def test_anonymous_put_is_forbidden(self):
        self.assertEqual(self._put(self.anon).status_code, status.HTTP_403_FORBIDDEN)

    def test_tourism_staff_cannot_write_even_with_is_staff(self):
        _, client = self._client_for("tourism_staff", ROLE_TOURISM, is_staff=True)
        self.assertEqual(self._put(client).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.anon.get(self.URL).json()["primary_color"], "#2FA34A")

    def test_sanitation_staff_cannot_write(self):
        _, client = self._client_for("sanitation_staff", ROLE_SANITATION, is_staff=True)
        self.assertEqual(self._put(client).status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_write_and_updated_by_is_set(self):
        from .models import TourismThemeSetting

        admin, client = self._admin()
        response = self._put(client, mobile_follows_web=False, mobile_primary_color="#00aabb")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["primary_color"], "#FF8800")
        setting = TourismThemeSetting.load()
        self.assertEqual(setting.updated_by, admin)
        self.assertFalse(setting.mobile_follows_web)
        self.assertEqual(setting.mobile_primary_color, "#00AABB")

    # --- validation -------------------------------------------------------

    def test_primary_color_is_required_and_must_be_hex(self):
        _, client = self._admin()
        self.assertEqual(client.put(self.URL, {}, format="json").status_code, status.HTTP_400_BAD_REQUEST)
        for bad in ["FF8800", "#FF880", "#FF88000", "#GG8800", "", "orange"]:
            response = self._put(client, primary_color=bad)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, bad)
            self.assertIn("primary_color", response.json())

    def test_mobile_primary_color_accepts_empty_and_rejects_malformed(self):
        _, client = self._admin()
        self.assertEqual(self._put(client, mobile_primary_color="").status_code, status.HTTP_200_OK)
        response = self._put(client, mobile_primary_color="#12345")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("mobile_primary_color", response.json())

    def test_saved_colors_rejects_malformed_payloads(self):
        _, client = self._admin()
        green = {"hex": "#2FA34A", "label": "Original (Green)"}
        bad_payloads = {
            "not a list": {"hex": "#112233", "label": "x"},
            "too many": [green] + [{"hex": f"#1122{i:02d}", "label": f"c{i}"} for i in range(12)],
            "extra key": [green, {"hex": "#112233", "label": "x", "note": "y"}],
            "missing label": [green, {"hex": "#112233"}],
            "entry not an object": [green, "#112233"],
            "bad hex": [green, {"hex": "#11223", "label": "x"}],
            "hex not a string": [green, {"hex": 112233, "label": "x"}],
            "empty label": [green, {"hex": "#112233", "label": ""}],
            "blank label": [green, {"hex": "#112233", "label": "   "}],
            "label too long": [green, {"hex": "#112233", "label": "x" * 41}],
            "label not a string": [green, {"hex": "#112233", "label": 7}],
        }
        for name, saved in bad_payloads.items():
            response = self._put(client, saved_colors=saved)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, name)
            self.assertIn("saved_colors", response.json(), name)
        ok = self._put(client, saved_colors=[green, {"hex": "#112233", "label": "x" * 40}])
        self.assertEqual(ok.status_code, status.HTTP_200_OK)

    def test_original_green_is_restored_when_omitted(self):
        _, client = self._admin()
        response = self._put(client, saved_colors=[{"hex": "#ff8800", "label": "Orange"}])
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.json()["saved_colors"],
            [{"hex": "#2FA34A", "label": "Original (Green)"}, {"hex": "#FF8800", "label": "Orange"}],
        )
        self.assertEqual(self.anon.get(self.URL).json()["saved_colors"][0]["hex"], "#2FA34A")

    def test_green_given_in_lowercase_is_not_duplicated(self):
        _, client = self._admin()
        response = self._put(client, saved_colors=[{"hex": "#2fa34a", "label": "Green"}])
        self.assertEqual(response.json()["saved_colors"], [{"hex": "#2FA34A", "label": "Green"}])

    def test_twelve_colours_without_green_is_rejected(self):
        _, client = self._admin()
        saved = [{"hex": f"#1122{i:02d}", "label": f"c{i}"} for i in range(12)]
        response = self._put(client, saved_colors=saved)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("saved_colors", response.json())

    # --- no cache ---------------------------------------------------------
    #
    # The theme is not cached. A per-process cache served the old colour from
    # every gunicorn worker the save did not reach, for its whole timeout. To
    # stand in for "another worker", these tests change the row WITHOUT going
    # through the view, so nothing could have invalidated anything: a cache
    # would still answer with the old colour.

    BOOTSTRAPS = (
        "/api/tourism-theme/",
        "/api/mobile/tourism/bootstrap/",
        "/api/bootstrap/",
    )

    @staticmethod
    def _theme_of(response, path):
        body = response.json()
        return body["theme"] if "theme" in body else body

    def _set_behind_the_back(self, **fields):
        from .models import TourismThemeSetting

        TourismThemeSetting.load()  # make sure the row exists
        TourismThemeSetting.objects.filter(pk=1).update(**fields)

    def test_a_write_shows_at_once(self):
        self.assertEqual(self.anon.get(self.URL).json()["primary_color"], "#2FA34A")
        _, client = self._admin()
        self.assertEqual(self._put(client, primary_color="#123ABC").status_code, status.HTTP_200_OK)
        self.assertEqual(self.anon.get(self.URL).json()["primary_color"], "#123ABC")

    def test_a_change_made_by_another_process_shows_at_once_on_every_endpoint(self):
        _, client = self._admin()
        # Warm anything that might remember the colour.
        for path in self.BOOTSTRAPS:
            self.assertEqual(client.get(path).status_code, status.HTTP_200_OK)

        self._set_behind_the_back(primary_color="#ABCDEF", mobile_follows_web=False, mobile_primary_color="#0155FE")

        for path in self.BOOTSTRAPS:
            theme = self._theme_of(client.get(path), path)
            self.assertEqual(theme["primary_color"], "#ABCDEF", path)
            self.assertEqual(theme["mobile_primary_color"], "#0155FE", path)
            self.assertIs(theme["mobile_follows_web"], False, path)

        # ...and again, in the other direction, with nothing in between.
        self._set_behind_the_back(primary_color="#2FA34A", mobile_follows_web=True)
        for path in self.BOOTSTRAPS:
            theme = self._theme_of(client.get(path), path)
            self.assertEqual(theme["primary_color"], "#2FA34A", path)
            self.assertIs(theme["mobile_follows_web"], True, path)

    def test_the_theme_is_one_query_a_read(self):
        from .services.tourism import read_tourism_theme

        read_tourism_theme()  # the row exists
        with self.assertNumQueries(1):
            read_tourism_theme()

    # --- bootstrap wiring -------------------------------------------------

    def test_both_bootstrap_endpoints_include_the_theme(self):
        _, client = self._admin()
        self.assertEqual(self._put(client, primary_color="#AA5500").status_code, status.HTTP_200_OK)

        web = client.get("/api/bootstrap/")
        self.assertEqual(web.status_code, status.HTTP_200_OK)
        self.assertEqual(web.json()["theme"]["primary_color"], "#AA5500")
        self.assertEqual(set(web.json()["theme"].keys()), self.READ_KEYS)

        mobile = self.anon.get("/api/mobile/tourism/bootstrap/")
        self.assertEqual(mobile.status_code, status.HTTP_200_OK)
        self.assertEqual(mobile.json()["theme"]["primary_color"], "#AA5500")
        self.assertNotIn("updated_by", mobile.json()["theme"])

import re
from pathlib import Path


class ActivityLogDigestConstraintTests(_BoatFareApiBase, TestCase):
    """The colleague's unique constraint on ActivityLog (migration 0059,
    uniq_sanitation_public_digest) covers record_id ONLY for rows with module
    "sanitation" and record_type "SanitationPublicSubmission". The logs of the boat
    fare tiers and of the hotlines are module "tourism": the same action on the same
    thing, repeated, must never be refused by it."""

    NAME = "uniq_sanitation_public_digest"

    def setUp(self):
        self.setUpData()

    def test_the_constraint_exists_here_and_does_refuse_what_it_is_meant_to(self):
        from django.db import IntegrityError, transaction

        with connection.cursor() as cursor:
            constraints = connection.introspection.get_constraints(cursor, ActivityLog._meta.db_table)
        self.assertIn(self.NAME, constraints)  # so the tests below are not vacuous
        ActivityLog.objects.create(module="sanitation", action="create", record_type="SanitationPublicSubmission", record_id="7", record_label="x")
        with self.assertRaises(IntegrityError), transaction.atomic():
            ActivityLog.objects.create(module="sanitation", action="create", record_type="SanitationPublicSubmission", record_id="7", record_label="x")

    def test_a_tier_deactivated_and_brought_back_twice_the_same_day_is_logged_four_times(self):
        before = self._logs().count()
        for _ in range(2):
            self.assertEqual(self.admin.post(self._url(self.b, "deactivate"), {}, format="json").status_code, 200)
            self.assertEqual(self.admin.post(self._url(self.b, "reactivate"), {}, format="json").status_code, 200)
        rows = list(self._logs().order_by("pk").values_list("record_id", "record_label")[before:])
        self.assertEqual(len(rows), 4)
        self.assertEqual(len({record_id for record_id, _ in rows}), 1)  # the same record_id four times
        self.assertEqual(rows[0], rows[2])  # and identical rows, which the constraint does not look at
        self.assertEqual(rows[1], rows[3])

    def test_a_tier_renamed_and_renamed_back_is_logged_each_time(self):
        before = self._logs().count()
        original = self.b.name
        for _ in range(2):
            self.assertEqual(self.admin.patch(self._url(self.b), {"name": "Bravo temporary"}, format="json").status_code, 200)
            self.assertEqual(self.admin.patch(self._url(self.b), {"name": original}, format="json").status_code, 200)
        self.assertEqual(self._logs().count() - before, 4)

    def test_a_hotline_deactivated_and_brought_back_twice_the_same_day_is_logged_four_times(self):
        contact = EmergencyContact.objects.create(name="Echo Desk", contact_numbers="555-5555", icon="help_desk", display_order=90)
        base = "/api/emergency-contacts/"
        before = ActivityLog.objects.filter(record_type="EmergencyContact").count()
        for _ in range(2):
            self.assertEqual(self.admin.post(f"{base}{contact.pk}/deactivate/", {}, format="json").status_code, 200)
            self.assertEqual(self.admin.post(f"{base}{contact.pk}/reactivate/", {}, format="json").status_code, 200)
        rows = ActivityLog.objects.filter(record_type="EmergencyContact").order_by("pk")
        self.assertEqual(rows.count() - before, 4)
        self.assertEqual(set(rows.values_list("module", flat=True)), {"tourism"})

    def test_our_logs_are_module_tourism_and_so_outside_the_constraint(self):
        self.admin.post(self._url(self.b, "deactivate"), {}, format="json")
        self.assertEqual(set(self._logs().values_list("module", "record_type")), {("tourism", "BoatCapacityFare")})


# -- Mobile App Content, commit 1: the two tables and their seed ----------------------

MOBILE_CONTENT_DART = Path(__file__).resolve().parents[2] / "mobile" / "lib" / "screens" / "tourism_screens.dart"
# What the seed's three placeholders stand for in the Dart (the port guide is the only place).
PORT_FEE_REGULAR = "₱80"
PORT_FEE_DISCOUNTED = "₱64"
# Dart icon -> the name stored (the table in docs of the spec; the app draws the picture).
MOBILE_CONTENT_ICONS = {
    "beach_access": "beach",
    "workspace_premium": "craft",
    "water_drop": "waterfall",
    "restaurant_menu": "food",
    "directions_boat_outlined": "boat",
    "festival_outlined": "festival",
    "receipt_long_outlined": "receipt",
    "delete_outline": "waste",
    "event_available_outlined": "event",
    "eco_outlined": "eco",
}


def _mobile_content_module():
    import importlib

    return importlib.import_module("api.migrations.0061_seed_mobile_content")


class MobileContentMigrationTests(TransactionTestCase):
    """Migrations 0060 and 0061 through Django's own migrate, forward AND backward."""

    # BEFORE is the migration just before ours: the tests never reverse anyone else's.
    BEFORE = [("api", "0059_activitylog_public_digest_uniqueness")]
    MIDDLE = [("api", "0060_mobile_content")]
    AFTER = [("api", "0061_seed_mobile_content")]
    NEW_TABLES = {"api_mobilecontentitem", "api_mobilecontentblock"}

    def tearDown(self):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())

    def _models(self, target):
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        state = executor.loader.project_state(target).apps
        return state.get_model("api", "MobileContentItem"), state.get_model("api", "MobileContentBlock")

    def _schema(self):
        """{table: [(column, type), ...]} of every table that is not one of the two new ones."""
        out = {}
        with connection.cursor() as cursor:
            for table in connection.introspection.table_names(cursor):
                if table in self.NEW_TABLES:
                    continue
                description = connection.introspection.get_table_description(cursor, table)
                out[table] = sorted((col.name, str(col.type_code)) for col in description)
        return out

    def _tables(self):
        with connection.cursor() as cursor:
            return set(connection.introspection.table_names(cursor))

    def test_a_the_real_migrate_gives_ten_items_and_five_blocks_and_changes_no_other_table(self):
        MigrationExecutor(connection).migrate(self.BEFORE)
        self.assertFalse(self.NEW_TABLES & self._tables())
        schema_before = self._schema()

        MigrationExecutor(connection).migrate(self.AFTER)
        self.assertTrue(self.NEW_TABLES <= self._tables())
        Item, Block = self._models(self.AFTER)
        self.assertEqual(Item.objects.count(), 10)
        self.assertEqual(Block.objects.count(), 5)
        self.assertEqual(
            sorted(Block.objects.values_list("key", flat=True)),
            ["fee_office", "port_guide", "travel_guide", "welcome", "why_visit"],
        )
        # g) no existing table got a new, changed or removed column
        self.assertEqual(self._schema(), schema_before)

    def test_d_a_second_run_adds_nothing_and_changes_nothing(self):
        from django.apps import apps as django_apps

        MigrationExecutor(connection).migrate(self.BEFORE)  # a clean start: the seed is run by this migrate
        MigrationExecutor(connection).migrate(self.AFTER)
        Item, Block = self._models(self.AFTER)
        self.assertEqual((Item.objects.count(), Block.objects.count()), (10, 5))
        # an admin has edited one of each: the second run must not undo it or add a twin
        Item.objects.filter(title="Environmental Fee").update(short_text="edited", full_text="a long text")
        Block.objects.filter(key="welcome").update(subtitle="edited too")
        snapshot = (
            list(Item.objects.order_by("id").values()),
            list(Block.objects.order_by("id").values()),
        )
        module = _mobile_content_module()
        module.add_content(django_apps, None)
        module.add_content(django_apps, None)
        self.assertEqual(Item.objects.count(), 10)
        self.assertEqual(Block.objects.count(), 5)
        self.assertEqual(
            (list(Item.objects.order_by("id").values()), list(Block.objects.order_by("id").values())),
            snapshot,
        )

    def test_e_the_reverse_removes_what_was_seeded_and_keeps_what_was_added_by_hand(self):
        MigrationExecutor(connection).migrate(self.BEFORE)  # a clean start: the seed is run by this migrate
        MigrationExecutor(connection).migrate(self.AFTER)
        Item, Block = self._models(self.AFTER)
        self.assertEqual((Item.objects.count(), Block.objects.count()), (10, 5))
        Item.objects.create(section="travel_guide", title="Added by hand", short_text="mine", icon="info", display_order=99)
        Block.objects.create(key="added_by_hand", title="mine")

        MigrationExecutor(connection).migrate(self.MIDDLE)  # the reverse of 0061 only
        self.assertTrue(self.NEW_TABLES <= self._tables())  # the tables are still there
        Item, Block = self._models(self.MIDDLE)
        self.assertEqual(list(Item.objects.values_list("title", flat=True)), ["Added by hand"])
        self.assertEqual(list(Block.objects.values_list("key", flat=True)), ["added_by_hand"])

        MigrationExecutor(connection).migrate(self.AFTER)  # and forward again: the seed returns beside it
        Item, Block = self._models(self.AFTER)
        self.assertEqual(Item.objects.count(), 11)
        self.assertEqual(Block.objects.count(), 6)

    def test_g_the_migrations_only_create_and_seed(self):
        import importlib
        from io import StringIO

        from django.db import migrations as dj_migrations

        m60 = importlib.import_module("api.migrations.0060_mobile_content").Migration
        m61 = importlib.import_module("api.migrations.0061_seed_mobile_content").Migration
        self.assertEqual(m60.dependencies, [("api", "0059_activitylog_public_digest_uniqueness")])
        self.assertEqual(m61.dependencies, [("api", "0060_mobile_content")])
        self.assertTrue(all(isinstance(op, dj_migrations.CreateModel) for op in m60.operations))
        self.assertEqual(sorted(op.name for op in m60.operations), ["MobileContentBlock", "MobileContentItem"])
        self.assertEqual([type(op) for op in m61.operations], [dj_migrations.RunPython])
        out = StringIO()
        call_command("sqlmigrate", "api", "0060", stdout=out, verbosity=0)
        sql = out.getvalue().upper()
        self.assertIn("CREATE TABLE", sql)
        self.assertNotIn("ALTER", sql)
        self.assertNotIn("DROP", sql)
        call_command("makemigrations", "--check", "--dry-run", verbosity=0)


class MobileContentSeedTests(TestCase):
    """The seed against the Dart file the app is built from: letter for letter."""

    def setUp(self):
        from django.apps import apps as django_apps

        from api.models import MobileContentBlock, MobileContentItem

        self.Item, self.Block = MobileContentItem, MobileContentBlock
        self.Item.objects.all().delete()
        self.Block.objects.all().delete()
        _mobile_content_module().add_content(django_apps, None)
        self.dart = MOBILE_CONTENT_DART.read_text(encoding="utf-8").replace("\r\n", "\n")

    # -- helpers ---------------------------------------------------------------------

    def _literal(self, text):
        """The text as the Dart writes it: single quotes, nothing to escape."""
        self.assertNotIn("'", text)
        self.assertNotIn("\\", text)
        self.assertNotIn("$", text)
        return "'" + text + "'"

    def _block(self, key):
        return self.Block.objects.get(key=key)

    def _fee_filled(self, text):
        """The port guide with its placeholders replaced by what the Dart has written."""
        office = self._block("fee_office").body
        return (
            text.replace("{fee_regular}", PORT_FEE_REGULAR)
            .replace("{fee_discounted}", PORT_FEE_DISCOUNTED)
            .replace("{fee_office}", office)
        )

    # -- b) letter for letter ---------------------------------------------------------

    CONSTRUCTORS = {
        ("why_visit", "tile"): ("_buildPillarItem", "sub"),
        ("travel_guide", "tile"): ("_buildGuideTile", "desc"),
        ("travel_guide", "banner"): ("InfoBanner", "text"),
    }

    def _dart_pattern(self, item):
        name, field = self.CONSTRUCTORS[(item.section, item.style)]
        return re.compile(
            re.escape(name)
            + r"\(\s*icon: Icons\.(\w+),\s*title: "
            + re.escape(self._literal(item.title))
            + r",\s*"
            + field
            + r": "
            + re.escape(self._literal(self._fee_filled(item.short_text)))
            + r","
        )

    def test_b_every_item_is_in_the_dart_letter_for_letter_with_its_icon(self):
        self.assertEqual(self.Item.objects.count(), 10)
        for item in self.Item.objects.all():
            match = self._dart_pattern(item).search(self.dart)
            self.assertIsNotNone(match, f"{item.title!r} is not in the Dart as stored")
            self.assertEqual(MOBILE_CONTENT_ICONS[match.group(1)], item.icon, item.title)

    def test_b_the_dart_has_no_item_that_is_not_stored(self):
        self.assertEqual(len(re.findall(r"_buildPillarItem\(\s*icon: Icons\.", self.dart)), 4)
        self.assertEqual(len(re.findall(r"_buildGuideTile\(\s*icon: Icons\.", self.dart)), 4)
        self.assertEqual(len(re.findall(r"InfoBanner\(\s*icon: Icons\.", self.dart)), 2)

    def test_b_the_welcome_banner_why_visit_and_the_heading_are_in_the_dart(self):
        welcome = self._block("welcome")
        self.assertEqual(welcome.title.count("{name}"), 1)
        guest = welcome.title.replace("{name}", "Explorer")
        member = welcome.title.replace("{name}", "${widget.profile.displayName}")
        self.assertIn(f"? {self._literal(guest)}", self.dart)
        # (the member form has Dart's own ${...} interpolation, so it is written out here)
        self.assertIn(": '" + member + "'", self.dart)
        self.assertEqual(welcome.body, "")
        after_greeting = self.dart[self.dart.index("'" + member + "'"):]
        self.assertTrue(
            re.match(r".*?Text\(\s*" + re.escape(self._literal(welcome.subtitle)), after_greeting, re.S),
            "the welcome text is not the one that follows the greeting",
        )
        why = self._block("why_visit")
        region = self.dart[self.dart.index("class _MaubanTownHighlightsCard"):self.dart.index("Widget _buildPillarItem")]
        self.assertTrue(
            re.search(
                r"Text\(\s*" + re.escape(self._literal(why.title)) + r".*?Text\(\s*" + re.escape(self._literal(why.subtitle))
                + r".*?Text\(\s*" + re.escape(self._literal(why.body)),
                region,
                re.S,
            ),
            "Why Visit: title, subtitle and paragraph are not in this order in the Dart",
        )
        heading = self._block("travel_guide")
        region = self.dart[self.dart.index("class _MaubanTownGuideSection"):self.dart.index("class DestinationListPage")]
        self.assertTrue(re.search(r"Text\(\s*" + re.escape(self._literal(heading.title)), region))
        self.assertEqual((heading.subtitle, heading.body), ("", ""))

    def test_b_the_port_guide_is_the_dart_lines_once_the_placeholders_are_filled(self):
        port = self._block("port_guide")
        region = self.dart[self.dart.index("Widget _buildPortGuideCard"):self.dart.index("class ProfilePage")]
        self.assertTrue(re.search(r"Text\(\s*" + re.escape(self._literal(port.title)), region))
        lines = port.body.split("\n")
        self.assertEqual(len(lines), 4)
        for line in lines:
            filled = "• " + self._fee_filled(line)
            self.assertTrue(
                re.search(r"Text\(\s*" + re.escape(self._literal(filled)) + r",", region),
                f"not in the Dart port guide: {filled!r}",
            )
        # and the place of the Tourism Office is the one the Dart writes
        self.assertIn(self._block("fee_office").body, self.dart)
        self.assertEqual((self._block("fee_office").title, self._block("fee_office").subtitle), ("", ""))

    # -- c) style, section and order are where the Dart shows them ---------------------

    def test_c_the_two_notices_are_banners_and_the_four_guides_and_four_pillars_are_tiles(self):
        banners = list(self.Item.objects.filter(style="banner").order_by("display_order").values_list("title", flat=True))
        self.assertEqual(banners, ["Upcoming: Fiesta de Mauban & Buntal Festival", "Clean Coast & Eco-Tourism Policy"])
        self.assertEqual(self.Item.objects.filter(section="travel_guide", style="banner").count(), 2)
        self.assertEqual(self.Item.objects.filter(section="travel_guide", style="tile").count(), 4)
        self.assertEqual(self.Item.objects.filter(section="why_visit", style="tile").count(), 4)
        self.assertEqual(self.Item.objects.filter(section="why_visit", style="banner").count(), 0)
        # the Dart agrees: each stored item is built by the constructor of its style
        # (b checks the match; here no OTHER constructor writes the same title)
        for item in self.Item.objects.all():
            expected = self.CONSTRUCTORS[(item.section, item.style)][0]
            for name in {"_buildPillarItem", "_buildGuideTile", "InfoBanner"} - {expected}:
                other = re.search(
                    re.escape(name) + r"\(\s*icon: Icons\.\w+,\s*title: " + re.escape(self._literal(item.title)), self.dart
                )
                self.assertIsNone(other, f"{item.title!r} is drawn by {name} in the Dart, not {expected}")

    def test_c_the_order_is_the_order_the_app_shows_today(self):
        # tiles and pillars: the order they are written in the Dart
        for section, style in (("why_visit", "tile"), ("travel_guide", "tile"), ("travel_guide", "banner")):
            items = list(self.Item.objects.filter(section=section, style=style).order_by("display_order", "id"))
            positions = [self.dart.index("title: " + self._literal(i.title)) for i in items]
            self.assertEqual(positions, sorted(positions), (section, style))
        orders = [i.display_order for i in self.Item.objects.filter(section="travel_guide").order_by("display_order")]
        self.assertEqual(len(set(orders)), 6)
        # the Home page builds the guides first and the notices after them
        self.assertLess(self.dart.index("_MaubanTownGuideSection()"), self.dart.index("const InfoBanner("))
        tiles = self.Item.objects.filter(section="travel_guide", style="tile").order_by("-display_order").first()
        banner = self.Item.objects.filter(section="travel_guide", style="banner").order_by("display_order").first()
        self.assertLess(tiles.display_order, banner.display_order)

    def test_c_everything_is_active_and_nothing_can_be_tapped_yet(self):
        self.assertFalse(self.Item.objects.filter(is_active=False).exists())
        self.assertEqual(set(self.Item.objects.values_list("full_text", flat=True)), {""})

    # -- f) the placeholders ----------------------------------------------------------

    def _all_texts(self):
        out = []
        for item in self.Item.objects.all():
            out += [item.title, item.short_text, item.full_text]
        for block in self.Block.objects.all():
            out += [block.title, block.subtitle, block.body]
        return out

    def test_f_the_only_placeholders_are_the_four_and_each_is_where_it_belongs(self):
        found = set()
        for text in self._all_texts():
            found |= set(re.findall(r"\{[^{}]*\}", text))
            self.assertEqual(text.count("{"), text.count("}"), text)
            self.assertEqual(text.count("{"), len(re.findall(r"\{[^{}]*\}", text)), text)
        self.assertEqual(found, {"{name}", "{fee_regular}", "{fee_discounted}", "{fee_office}"})
        # {name} is in the greeting only; the three fee ones are in the port guide only
        for block in self.Block.objects.all():
            for field in ("title", "subtitle", "body"):
                text = getattr(block, field)
                if block.key == "welcome" and field == "title":
                    self.assertEqual(set(re.findall(r"\{[^{}]*\}", text)), {"{name}"})
                elif block.key == "port_guide" and field == "body":
                    self.assertEqual(
                        set(re.findall(r"\{[^{}]*\}", text)), {"{fee_regular}", "{fee_discounted}", "{fee_office}"}
                    )
                else:
                    self.assertEqual(re.findall(r"\{[^{}]*\}", text), [], (block.key, field))
        # the items: only the two tiles that name the Tourism Office have a placeholder, and it is {fee_office}
        with_office = []
        for item in self.Item.objects.all():
            for text in (item.title, item.short_text, item.full_text):
                holders = re.findall(r"\{[^{}]*\}", text)
                if holders:
                    self.assertEqual(holders, ["{fee_office}"], item.title)
                    with_office.append(item.title)
        self.assertEqual(sorted(with_office), ["Cagbalete Island Boat Guide", "Environmental Fee"])

    def test_f_the_place_of_the_tourism_office_has_one_source_and_three_readers(self):
        """"2nd Floor" is written ONCE, in the block fee_office; the three texts that name
        the place say {fee_office}."""
        self.assertEqual(self._block("fee_office").body, "Tourism Office, 2nd Floor, Mauban Municipal Hall")
        readers = []
        for item in self.Item.objects.all():
            for field in ("title", "short_text", "full_text"):
                readers += [(item.title, field)] * getattr(item, field).count("{fee_office}")
        for block in self.Block.objects.exclude(key="fee_office"):
            for field in ("title", "subtitle", "body"):
                readers += [(block.key, field)] * getattr(block, field).count("{fee_office}")
        self.assertEqual(
            sorted(readers),
            [
                ("Cagbalete Island Boat Guide", "short_text"),
                ("Environmental Fee", "short_text"),
                ("port_guide", "body"),
            ],
        )
        # and no literal place, in whole or in part, anywhere else
        for item in self.Item.objects.all():
            for text in (item.title, item.short_text, item.full_text):
                self.assertNotIn("2nd Floor", text, item.title)
                self.assertNotIn("Municipal Hall", text, item.title)
        for block in self.Block.objects.exclude(key="fee_office"):
            for text in (block.title, block.subtitle, block.body):
                self.assertNotIn("2nd Floor", text, block.key)
                self.assertNotIn("Municipal Hall", text, block.key)

    def test_f_no_amount_of_money_is_stored_anywhere(self):
        for text in self._all_texts():
            self.assertIsNone(re.search(r"(₱|PHP|Php)\s?\d", text), text)
        port = self._block("port_guide").body
        self.assertIn("Environmental Fee: {fee_regular} per visitor", port)
        self.assertIn("pay {fee_discounted}.", port)
        self.assertIn("Pay at the {fee_office}.", port)

    # -- the shape of the two tables -----------------------------------------------------

    def test_the_default_of_style_and_is_active_is_in_the_database_itself(self):
        """A row put in by code that does not know these columns is a tile and is active."""
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO api_mobilecontentitem (section, icon, title, short_text, full_text, display_order) "
                "VALUES ('travel_guide', 'info', 'Raw row', 'x', '', 5)"
            )
        row = self.Item.objects.get(title="Raw row")
        self.assertEqual((row.style, row.is_active), ("tile", True))

    def test_the_icon_names_are_the_eleven_of_the_spec_and_the_seed_uses_ten(self):
        self.assertEqual(
            set(self.Item.ICON_NAMES),
            set(MOBILE_CONTENT_ICONS.values()) | {"info"},
        )
        self.assertEqual(
            set(self.Item.objects.values_list("icon", flat=True)),
            set(MOBILE_CONTENT_ICONS.values()),
        )

    def test_the_seed_file_is_what_the_script_makes_from_the_dart(self):
        import subprocess
        import sys

        script = MOBILE_CONTENT_DART.parents[3] / "backend" / "tools" / "extract_mobile_content_seed.py"
        result = subprocess.run([sys.executable, "-I", str(script), "--check"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
